"""조작 유형별 탐지율 표 — 한글/Word 표 형식(무채색). docx + png 동시 생성."""
import json, sys
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import rcParams

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

rcParams["font.family"] = "Malgun Gothic"
rcParams["axes.unicode_minus"] = False
rcParams["figure.dpi"] = 180
rcParams["savefig.bbox"] = "tight"

E = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
V = json.loads(Path(sys.argv[2]).read_text(encoding="utf-8"))
S = json.loads(Path(sys.argv[3]).read_text(encoding="utf-8"))
G = json.loads(Path(sys.argv[4]).read_text(encoding="utf-8"))
OUT_PNG = Path(sys.argv[5])
OUT_DOCX = Path(sys.argv[6])


def ev(name):
    return E["edits"]["600"][name] + E["edits"]["1000"][name], 60


def va(name):
    return V["result"][name]["hit"], V["result"][name]["n"]


def gp(gap):
    return G["result"]["600"][str(gap)] + G["result"]["2000"][str(gap)], 60


ORIG = "우리 회사는 내년 3월 15일에 새 스마트워치 오로라를 전국 직영 매장 스무 곳에서 출시한다. 가격은 29만 9천 원으로 정했다."

ROWS = [
    ("문장 하나를 통째로 지움", "우리 회사는 내년 3월 15일에 새 스마트워치 오로라를 전국 직영 매장 스무 곳에서 출시한다.", ev("문장 1개 삭제")),
    ("존댓말로 고쳐 씀", "우리 회사는 … 스무 곳에서 출시합니다. 가격은 29만 9천 원으로 정했습니다.", ev("존댓말 변환")),
    ("조사만 전부 바꿈", "우리 회사가 내년 3월 15일에 새 스마트워치 오로라는 전국 직영 매장 …", va("조사만 전부 바꾸기")),
    ("문장 순서를 바꿈", "가격은 29만 9천 원으로 정했다. 우리 회사는 내년 3월 15일에 …", ev("문장 순서 섞기")),
    ("불릿 형식으로 바꿈", "- 우리 회사는 내년 3월 15일에 … 출시한다   - 가격은 29만 9천 원으로 정했다", ev("불릿 형식")),
    ("글자를 5% 지움", "우리 회사는 내년 3월 5일에 새 스마트워 오로라를 전국 직 매장 스무 곳에서 출시한다.", va("글자 5% 삭제")),
    ("오타를 3~5% 섞음", "우리 회사는 내년 3월 15일에 새 스마트워치 오로라률 전국 지경 매장 스무 곳에서 출시한다.",
     (V["result"]["글자 3% 오타"]["hit"] + V["result"]["글자 5% 오타"]["hit"], 60)),
    ("숫자를 전부 다른 값으로", "우리 회사는 내년 7월 28일에 … 가격은 61만 4천 원으로 정했다.", ev("숫자 전부 변경")),
    ("40자마다 단어 하나 바꿈", "우리 회사는 내년 3월 15일에 새 스마트워치 오로라를 전국 직영 매장 스무 곳에서 ○○○.", gp(40)),
    ("오타를 10% 섞음", "우리 회샤는 내년 3월 15일에 새 스마트워치 오로랄를 전국 직엽 매장 슈무 곳에서 출시한다.", va("글자 10% 오타")),
    ("30자마다 단어 하나 바꿈", "우리 회사는 내년 3월 15일에 새 ○○○○ 오로라를 전국 직영 ○○ 스무 곳에서 출시한다.", gp(30)),
    ("20자마다 단어 하나 바꿈", "우리 회사는 내년 ○○○○ 새 스마트워치 ○○○를 전국 ○○ 매장 스무 ○○○ 출시한다.", gp(20)),
    ("어절 순서를 섞어 다시 씀", "스무 곳에서 우리 회사는 오로라를 내년 3월 15일에 전국 직영 매장 출시한다.", va("문장 안 어절 순서 섞기")),
    ("40자 미만만 입력", "오로라 3월 15일 출시", (S["result"]["산문"]["30"] + S["result"]["표"]["30"], 60)),
    ("자기 말로 다시 씀·요약·번역", "새 스마트워치를 내년 봄에 전국 매장에서 선보일 계획이다.", None),
]

HEAD = ["조작 방식", "그렇게 조작한 입력", "탐지율", "시험 건수"]


def cells(name, ex, res):
    if res is None:
        return [name, ex, "측정 안 함", "―"]
    return [name, ex, f"{100.0 * res[0] / res[1]:.0f}%", f"{res[0]}/{res[1]}"]


TABLE = [HEAD] + [cells(*r) for r in ROWS]
CAPTION = ("※ 예시는 이해를 돕기 위한 가상 문장이며, 실제 측정은 기밀문서 20종에서 600~2,000자를 뽑아 조작마다 변형해 진행했다. "
           "○는 다른 단어로 바꿔 쓴 자리.")

# ---------- docx ----------
doc = Document()
st = doc.styles["Normal"]
st.font.name = "맑은 고딕"
st.font.size = Pt(9)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")

p = doc.add_paragraph()
run = p.add_run("[표] 조작 유형별 탐지율 (기밀문서 20종 등록, 2026-09-19 측정)")
run.bold = True
run.font.size = Pt(11)
doc.add_paragraph("기밀문서 원문: " + ORIG)

t = doc.add_table(rows=len(TABLE), cols=4)
t.style = "Table Grid"
t.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, row in enumerate(TABLE):
    for j, val in enumerate(row):
        cell = t.cell(i, j)
        cell.text = ""
        par = cell.paragraphs[0]
        r = par.add_run(val)
        r.font.size = Pt(9)
        r.font.name = "맑은 고딕"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        if i == 0:
            r.bold = True
        if j >= 2:
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph(CAPTION)
doc.save(OUT_DOCX)

# ---------- png (같은 표를 무채색으로) ----------
COL_X = [0.02, 0.26, 0.845, 0.925]
COL_W = [0.24, 0.585, 0.08, 0.055]
H = 0.0505
TOP = 0.845
fig = plt.figure(figsize=(12.6, 8.0))
ax = fig.add_axes([0, 0, 1, 1])
ax.set_xlim(0, 1)
ax.set_ylim(0, 1)
ax.axis("off")

fig.text(0.02, 0.965, "[표] 조작 유형별 탐지율 (기밀문서 20종 등록, 2026-09-19 측정)", fontsize=12, fontweight="bold", color="black")
fig.text(0.02, 0.925, "기밀문서 원문: " + ORIG, fontsize=9.5, color="black")

left, right = COL_X[0], COL_X[-1] + COL_W[-1]
n = len(TABLE)
for i in range(n + 1):
    y = TOP - i * H
    ax.plot([left, right], [y, y], color="black", lw=0.8, solid_capstyle="butt")
for x in [COL_X[0]] + [COL_X[k] for k in range(1, 4)] + [right]:
    ax.plot([x, x], [TOP, TOP - n * H], color="black", lw=0.8, solid_capstyle="butt")

for i, row in enumerate(TABLE):
    y = TOP - i * H - H / 2
    for j, val in enumerate(row):
        size = 9.3 if j == 1 else 9.6
        if j >= 2:
            fig.text(COL_X[j] + COL_W[j] / 2, y - 0.008, val, fontsize=size, color="black", ha="center")
        else:
            fig.text(COL_X[j] + 0.008, y - 0.008, val, fontsize=size, color="black")

fig.text(0.02, TOP - n * H - 0.028, CAPTION, fontsize=8.8, color="black")
fig.savefig(OUT_PNG)
print("saved", OUT_PNG, OUT_DOCX)
