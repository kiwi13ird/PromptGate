"""
공통 문서 조립기 (v2 — 결재란/문서정보 박스/표 음영 추가).
서브에이전트는 이 모듈을 직접 수정하지 말고, content/<doc_id>.json 콘텐츠 스펙만 작성한 뒤 CLI로 호출한다.

사용법:
    python docgen.py content/01_사업전략_방향_보고서.json

콘텐츠 스펙(JSON) 형식 (docx/pdf):
{
  "doc_id": "01_사업전략_방향_보고서",
  "format": "docx" | "pdf",
  "title": "문서 제목",
  "header_text": "대외비 - 경영기획실",
  "doc_meta": {
    "doc_no": "NW-STRAT-2026-014",
    "date": "2026-08-10",
    "author": "전략기획팀장 이수민",
    "distribution": "경영기획실 및 임원 한정"
  },
  "approval_line": [
    {"role": "기안", "name": "이수민"},
    {"role": "검토", "name": "강민석"},
    {"role": "승인", "name": "김도현"}
  ],
  "sections": [
    {
      "heading": "1. 절 제목",
      "paragraphs": ["문단1", "문단2"],
      "table": {"headers": [...], "rows": [[...], [...]]}   // 없으면 생략
    }
  ]
}

doc_meta / approval_line은 없으면 생략 가능(하위 호환).

콘텐츠 스펙(JSON) 형식 (xlsx — 순수 표/대장류 문서용):
{
  "doc_id": "06_급여대장",
  "format": "xlsx",
  "title": "문서 제목",
  "header_text": "대외비 - 인사팀",
  "sheet_name": "급여대장",
  "doc_meta": {...},   // docx/pdf와 동일, 생략 가능
  "table": {"headers": [...], "rows": [[...], [...]]}
}
"""
import json
import sys
from pathlib import Path

KOREAN_FONT = "맑은 고딕"
KOREAN_FONT_TTF = r"C:\Windows\Fonts\malgun.ttf"
KOREAN_FONT_BOLD_TTF = r"C:\Windows\Fonts\malgunbd.ttf"
COMPANY_NAME = "㈜네오웍스"

OUT_DIR = Path(__file__).parent


def build_docx(spec: dict, out_path: Path) -> None:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    def set_cjk_font(run, size=11, bold=False, color=None):
        run.font.name = KOREAN_FONT
        run.font.size = Pt(size)
        run.bold = bold
        run._element.rPr.rFonts.set(qn("w:eastAsia"), KOREAN_FONT)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)

    def shade_cell(cell, fill="D9D9D9"):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        tcPr.append(shd)

    def set_bottom_border(paragraph, size=18, color="1F3864"):
        pPr = paragraph._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), str(size))
        bottom.set(qn("w:space"), "4")
        bottom.set(qn("w:color"), color)
        pBdr.append(bottom)
        pPr.append(pBdr)

    def add_page_number_field(paragraph):
        run = paragraph.add_run()
        set_cjk_font(run, size=9)
        fld1 = OxmlElement("w:fldChar")
        fld1.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = "PAGE"
        fld2 = OxmlElement("w:fldChar")
        fld2.set(qn("w:fldCharType"), "end")
        run._r.append(fld1)
        run._r.append(instr)
        run._r.append(fld2)

    doc = Document()

    # 머리말 (보안등급 표기) — FR-PRE-003 검증 대상
    header = doc.sections[0].header
    header_para = header.paragraphs[0]
    header_run = header_para.add_run(spec.get("header_text", ""))
    set_cjk_font(header_run, size=9, bold=True, color="C00000")

    # 바닥글 (회사명 + 페이지 번호)
    footer = doc.sections[0].footer
    footer_para = footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer_para.add_run(f"{COMPANY_NAME}  -  ")
    set_cjk_font(fr, size=9)
    add_page_number_field(footer_para)

    # 결재란 (기안/검토/승인) — 우측 상단
    approval = spec.get("approval_line")
    if approval:
        n = len(approval)
        atable = doc.add_table(rows=2, cols=n)
        atable.alignment = WD_TABLE_ALIGNMENT.RIGHT
        atable.style = "Table Grid"
        for i, item in enumerate(approval):
            rc = atable.rows[0].cells[i].paragraphs[0].add_run(item["role"])
            set_cjk_font(rc, size=9, bold=True)
            shade_cell(atable.rows[0].cells[i])
            atable.rows[0].cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            nc = atable.rows[1].cells[i].paragraphs[0].add_run(item["name"])
            set_cjk_font(nc, size=9)
            atable.rows[1].cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            for cell in (atable.rows[0].cells[i], atable.rows[1].cells[i]):
                cell.width = Pt(60)
        doc.add_paragraph()

    # 제목
    title_p = doc.add_paragraph()
    title_run = title_p.add_run(spec["title"])
    set_cjk_font(title_run, size=18, bold=True, color="1F3864")
    set_bottom_border(title_p)

    # 문서정보 박스 (문서번호/작성일/작성자/배포범위)
    meta = spec.get("doc_meta")
    if meta:
        meta_p = doc.add_paragraph()
        meta_text = (
            f"문서번호: {meta.get('doc_no', '-')}    "
            f"작성일: {meta.get('date', '-')}    "
            f"작성자: {meta.get('author', '-')}    "
            f"배포범위: {meta.get('distribution', '-')}"
        )
        mr = meta_p.add_run(meta_text)
        set_cjk_font(mr, size=9, color="595959")
        doc.add_paragraph()

    for section in spec.get("sections", []):
        if section.get("heading"):
            h = doc.add_paragraph()
            hr = h.add_run(section["heading"])
            set_cjk_font(hr, size=13, bold=True, color="1F3864")

        for para_text in section.get("paragraphs", []):
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.3
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run(para_text)
            set_cjk_font(r, size=11)

        table_spec = section.get("table")
        if table_spec:
            headers = table_spec["headers"]
            rows = table_spec["rows"]
            table = doc.add_table(rows=1 + len(rows), cols=len(headers))
            table.style = "Table Grid"
            for i, h in enumerate(headers):
                cell = table.rows[0].cells[i]
                cell_run = cell.paragraphs[0].add_run(h)
                set_cjk_font(cell_run, size=10, bold=True, color="FFFFFF")
                shade_cell(cell, fill="1F3864")
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r_idx, row in enumerate(rows, start=1):
                if r_idx % 2 == 0:
                    for c_idx in range(len(headers)):
                        shade_cell(table.rows[r_idx].cells[c_idx], fill="F2F2F2")
                for c_idx, val in enumerate(row):
                    cell_run = table.rows[r_idx].cells[c_idx].paragraphs[0].add_run(str(val))
                    set_cjk_font(cell_run, size=10)
            doc.add_paragraph()

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


def build_pdf(spec: dict, out_path: Path) -> None:
    from xml.sax.saxutils import escape as _xml_escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        HRFlowable,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    pdfmetrics.registerFont(TTFont("Malgun", KOREAN_FONT_TTF))
    pdfmetrics.registerFont(TTFont("Malgun-Bold", KOREAN_FONT_BOLD_TTF))

    def esc(text) -> str:
        # reportlab Paragraph interprets its text as mini-XML, so a bare "&"
        # (e.g. "R&D") must be escaped or it renders as a broken entity like "R&D;".
        return _xml_escape(str(text))

    header_text = spec.get("header_text", "")

    def draw_header_footer(canvas, doc_):
        canvas.saveState()
        canvas.setFont("Malgun-Bold", 8)
        canvas.setFillColor(colors.HexColor("#C00000"))
        canvas.drawString(20 * mm, 287 * mm, header_text)
        canvas.setFillColor(colors.black)
        canvas.setFont("Malgun", 8)
        canvas.drawCentredString(105 * mm, 12 * mm, f"{COMPANY_NAME}  -  {doc_.page}")
        canvas.restoreState()

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(
        str(out_path),
        pagesize=A4,
        topMargin=25 * mm,
        bottomMargin=20 * mm,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
    )

    title_style = ParagraphStyle("Title", fontName="Malgun-Bold", fontSize=19, textColor=colors.HexColor("#1F3864"), spaceAfter=4)
    meta_style = ParagraphStyle("Meta", fontName="Malgun", fontSize=8.5, textColor=colors.HexColor("#595959"), spaceBefore=6, spaceAfter=14)
    heading_style = ParagraphStyle("Heading", fontName="Malgun-Bold", fontSize=12.5, textColor=colors.HexColor("#1F3864"), spaceBefore=12, spaceAfter=6)
    body_style = ParagraphStyle("Body", fontName="Malgun", fontSize=10.5, leading=17, spaceAfter=7)

    story = []

    # 결재란 (우측 상단)
    approval = spec.get("approval_line")
    if approval:
        role_row = [a["role"] for a in approval]
        name_row = [a["name"] for a in approval]
        atable = Table([role_row, name_row], colWidths=[20 * mm] * len(approval), hAlign="RIGHT")
        atable.setStyle(
            TableStyle(
                [
                    ("FONTNAME", (0, 0), (-1, -1), "Malgun"),
                    ("FONTNAME", (0, 0), (-1, 0), "Malgun-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9D9D9")),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        story.append(atable)
        story.append(Spacer(1, 10))

    story.append(Paragraph(esc(spec["title"]), title_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#1F3864")))

    meta = spec.get("doc_meta")
    if meta:
        meta_text = (
            f"문서번호: {esc(meta.get('doc_no', '-'))} &nbsp;&nbsp;&nbsp; "
            f"작성일: {esc(meta.get('date', '-'))} &nbsp;&nbsp;&nbsp; "
            f"작성자: {esc(meta.get('author', '-'))} &nbsp;&nbsp;&nbsp; "
            f"배포범위: {esc(meta.get('distribution', '-'))}"
        )
        story.append(Paragraph(meta_text, meta_style))
    else:
        story.append(Spacer(1, 10))

    for section in spec.get("sections", []):
        if section.get("heading"):
            story.append(Paragraph(esc(section["heading"]), heading_style))
        for para_text in section.get("paragraphs", []):
            story.append(Paragraph(esc(para_text), body_style))

        table_spec = section.get("table")
        if table_spec:
            headers = table_spec["headers"]
            rows = table_spec["rows"]
            data = [headers] + [[str(v) for v in row] for row in rows]
            row_bg = [
                ("BACKGROUND", (0, r), (-1, r), colors.HexColor("#F2F2F2"))
                for r in range(2, len(data), 2)
            ]
            t = Table(data, hAlign="LEFT")
            t.setStyle(
                TableStyle(
                    [
                        ("FONTNAME", (0, 0), (-1, -1), "Malgun"),
                        ("FONTNAME", (0, 0), (-1, 0), "Malgun-Bold"),
                        ("FONTSIZE", (0, 0), (-1, -1), 9),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3864")),
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 5),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                        ("TOPPADDING", (0, 0), (-1, -1), 4),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                        *row_bg,
                    ]
                )
            )
            story.append(Spacer(1, 4))
            story.append(t)
            story.append(Spacer(1, 10))

    doc.build(story, onFirstPage=draw_header_footer, onLaterPages=draw_header_footer)


def check_font_coverage(spec: dict) -> None:
    """Malgun Gothic covers Korean + common Hanja but NOT all Japanese-only
    kanji (e.g. 会/区/郎 use different glyph forms than Korean Hanja). No
    Japanese font is installed on this machine, so any such character is
    silently dropped by reportlab. Warn loudly instead of failing silently.
    """
    try:
        from fontTools.ttLib import TTFont as _TTFont
    except ImportError:
        return  # fonttools not installed; skip the check rather than block generation

    cmap = _TTFont(KOREAN_FONT_TTF).getBestCmap()

    def walk(value, path):
        if isinstance(value, str):
            missing = sorted({c for c in value if 0x4E00 <= ord(c) <= 0x9FFF and ord(c) not in cmap})
            if missing:
                print(f"WARNING: unsupported glyph(s) {missing} at {path} -- will render blank/garbled. Use romaji or rephrase.")
        elif isinstance(value, dict):
            for k, v in value.items():
                walk(v, f"{path}.{k}")
        elif isinstance(value, list):
            for i, v in enumerate(value):
                walk(v, f"{path}[{i}]")

    walk(spec, "spec")


def build_xlsx(spec: dict, out_path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = spec.get("sheet_name", "Sheet1")[:31]
    ws.sheet_view.rightToLeft = False

    table_spec = spec["table"]
    headers = table_spec["headers"]
    rows = table_spec["rows"]
    n_cols = len(headers)

    navy = PatternFill("solid", fgColor="1F3864")
    stripe = PatternFill("solid", fgColor="F2F2F2")
    white_bold = Font(name=KOREAN_FONT, size=10, bold=True, color="FFFFFF")
    title_font = Font(name=KOREAN_FONT, size=15, bold=True, color="1F3864")
    header_font = Font(name=KOREAN_FONT, size=9, bold=True, color="C00000")
    meta_font = Font(name=KOREAN_FONT, size=9, color="595959")
    body_font = Font(name=KOREAN_FONT, size=10)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    c = ws.cell(row=1, column=1, value=spec.get("header_text", ""))
    c.font = header_font

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n_cols)
    c = ws.cell(row=2, column=1, value=spec["title"])
    c.font = title_font

    meta = spec.get("doc_meta")
    header_row = 4
    if meta:
        meta_text = (
            f"문서번호: {meta.get('doc_no', '-')}    작성일: {meta.get('date', '-')}    "
            f"작성자: {meta.get('author', '-')}    배포범위: {meta.get('distribution', '-')}"
        )
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=n_cols)
        c = ws.cell(row=3, column=1, value=meta_text)
        c.font = meta_font
        header_row = 5

    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=h)
        cell.font = white_bold
        cell.fill = navy
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r_idx, row in enumerate(rows, start=1):
        excel_row = header_row + r_idx
        for c_idx, val in enumerate(row, start=1):
            cell = ws.cell(row=excel_row, column=c_idx, value=val)
            cell.font = body_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if r_idx % 2 == 0:
                cell.fill = stripe

    for col_idx, h in enumerate(headers, start=1):
        max_len = max([len(str(h))] + [len(str(r[col_idx - 1])) for r in rows]) if rows else len(str(h))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 4, 10), 40)

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_path))


def main():
    if len(sys.argv) != 2:
        print("usage: python docgen.py <content_spec.json>")
        sys.exit(1)

    spec_path = Path(sys.argv[1])
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    check_font_coverage(spec)

    fmt = spec["format"]
    out_name = f"{spec['doc_id']}.{fmt}"
    out_path = OUT_DIR / out_name

    if fmt == "docx":
        build_docx(spec, out_path)
    elif fmt == "pdf":
        build_pdf(spec, out_path)
    elif fmt == "xlsx":
        build_xlsx(spec, out_path)
    else:
        raise ValueError(f"unknown format: {fmt}")

    print(f"generated: {out_path}")


if __name__ == "__main__":
    main()
