"""FR-PRE-003 DOCX 파싱: 표 및 머리말(헤더/푸터)을 포함하여 파싱한다.

python-docx는 document.paragraphs / document.tables를 따로 제공하지만, 그렇게
읽으면 "문단1 - 표1 - 문단2"처럼 실제 문서에 섞여 있는 순서 정보가 사라진다. 탐지
위치 인덱싱(FR-PRE-005)에서 "원문 내 위치"를 말하려면 원래 순서가 보존돼야 하므로,
document.element.body를 직접 순회해 문단/표를 등장 순서 그대로 Block 리스트로
만든다.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from .errors import CorruptedDocumentError, EncryptedDocumentError

# 레거시 OLE 복합 문서 시그니처. 암호가 설정된 .docx는 OOXML(zip)이 아니라
# CryptoAPI RC4/AES로 감싼 OLE 컨테이너로 저장되므로 zip으로 열리지 않는다.
_OLE_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def _is_encrypted_ole(path: Path) -> bool:
    try:
        with open(path, "rb") as f:
            head = f.read(8)
        return head == _OLE_SIGNATURE
    except OSError:
        return False


def _iter_block_items(parent):
    """document.element.body를 순회하며 Paragraph/Table을 등장 순서대로 yield."""
    parent_elm = parent.element.body
    for child in parent_elm.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def _table_to_rows(table: Table) -> list[list[str]]:
    rows = []
    for row in table.rows:
        rows.append([cell.text for cell in row.cells])
    return rows


def _heading_level(paragraph: Paragraph) -> int | None:
    style_name = paragraph.style.name if paragraph.style else ""
    if style_name and style_name.startswith("Heading"):
        try:
            return int(style_name.rsplit(" ", 1)[-1])
        except ValueError:
            return 0
    if style_name in ("Title",):
        return 0
    return None


def parse_docx(path: str | Path) -> list[dict]:
    """DOCX를 순서 보존 Block dict 리스트로 반환한다.

    반환되는 각 dict는 pipeline.Block으로 감싸지기 전 원시 형태다:
      {"type": "heading"|"paragraph"|"table"|"header"|"footer",
       "level": int|None, "text": str|None, "rows": list[list[str]]|None,
       "section": int|None}
    """
    path = Path(path)

    if _is_encrypted_ole(path):
        raise EncryptedDocumentError(f"암호로 보호된 DOCX 파일입니다: {path.name}")

    try:
        document = Document(str(path))
    except PackageNotFoundError as e:
        raise CorruptedDocumentError(f"손상되었거나 유효한 DOCX가 아닙니다: {path.name}") from e
    except zipfile.BadZipFile as e:
        raise CorruptedDocumentError(f"손상된 ZIP 컨테이너입니다: {path.name}") from e
    except KeyError as e:
        # 필수 OOXML part(document.xml 등)가 없는 잘림/손상 파일
        raise CorruptedDocumentError(f"필수 구성요소가 누락된 DOCX입니다: {path.name}") from e

    blocks: list[dict] = []

    # 머리말/꼬리말: 섹션마다 별도로 존재 (docgen.py는 sections[0]에만 기록하지만
    # 문서가 섹션을 나눠 사용할 경우를 대비해 전체 섹션을 순회한다)
    for section_idx, section in enumerate(document.sections):
        header_text = "\n".join(
            p.text for p in section.header.paragraphs if p.text.strip()
        )
        if header_text:
            blocks.append(
                {"type": "header", "level": None, "text": header_text, "rows": None, "section": section_idx}
            )
        footer_text = "\n".join(
            p.text for p in section.footer.paragraphs if p.text.strip()
        )
        if footer_text:
            blocks.append(
                {"type": "footer", "level": None, "text": footer_text, "rows": None, "section": section_idx}
            )

    # 본문: 문단/표를 원문 순서 그대로
    for item in _iter_block_items(document):
        if isinstance(item, Paragraph):
            if not item.text.strip():
                continue
            level = _heading_level(item)
            blocks.append(
                {
                    "type": "heading" if level is not None else "paragraph",
                    "level": level,
                    "text": item.text,
                    "rows": None,
                    "section": None,
                }
            )
        elif isinstance(item, Table):
            rows = _table_to_rows(item)
            if any(any(cell.strip() for cell in row) for row in rows):
                blocks.append(
                    {"type": "table", "level": None, "text": None, "rows": rows, "section": None}
                )

    return blocks
