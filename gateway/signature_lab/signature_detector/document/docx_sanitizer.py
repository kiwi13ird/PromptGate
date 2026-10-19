from copy import deepcopy
from pathlib import Path

from docx import Document

from signature_detector.masker import mask_text


SEPARATOR = " | "


def _build_row_text_and_spans(cells, headers=None):
    """
    표 한 행을 하나의 문자열로 합치고,
    각 셀이 row_text에서 차지하는 위치를 기록한다.

    예:
        ["생년월일", "1999.03.14"]

    row_text:
        "생년월일 | 1999.03.14"
    """

    parts = []
    spans = []

    current_offset = 0

    for index, cell in enumerate(cells):

        text = cell.text

        if index > 0:
            parts.append(SEPARATOR)
            current_offset += len(SEPARATOR)

        if headers and index < len(headers) and headers[index].strip():
            prefix = f"{headers[index].strip()}: "
            parts.append(prefix)
            current_offset += len(prefix)

        start = current_offset

        parts.append(text)
        current_offset += len(text)

        end = current_offset

        spans.append(
            {
                "cell_index": index,
                "start": start,
                "end": end,
            }
        )

    return "".join(parts), spans


def _find_cell_for_detection(result, cell_spans):
    """
    row 전체에서 탐지된 span이
    어느 cell 안에 포함되는지 찾는다.
    """

    for span in cell_spans:

        if (
            result.start >= span["start"]
            and result.end <= span["end"]
        ):
            return span

    return None


def _mask_paragraph(
    paragraph,
    analyzer,
    target_entities,
):
    """
    일반 paragraph 하나를 탐지하고 마스킹한다.
    """

    text = paragraph.text

    if not text.strip():
        return

    results = analyzer.analyze(
        text=text,
        language="en",
        entities=target_entities,
    )

    if not results:
        return

    masked_text = mask_text(
        text,
        results,
    )

    paragraph.text = masked_text


def _mask_table_row(
    row,
    analyzer,
    target_entities,
    headers=None,
):
    """
    Table Row 전체를 합쳐 탐지한 뒤,
    탐지 결과를 원래 Cell 위치로 역매핑하여 마스킹한다.
    """

    cells = row.cells

    row_text, cell_spans = _build_row_text_and_spans(
        cells,
        headers=headers,
    )

    if not row_text.strip():
        return

    results = analyzer.analyze(
        text=row_text,
        language="en",
        entities=target_entities,
    )

    if not results:
        return

    detections_by_cell = {}

    for result in results:

        cell_span = _find_cell_for_detection(
            result,
            cell_spans,
        )

        # 현재 MVP에서는 두 셀에 걸쳐 있는 탐지는 처리하지 않는다.
        if cell_span is None:
            continue

        cell_index = cell_span["cell_index"]

        # row 기준 offset → cell 내부 offset
        local_start = (
            result.start
            - cell_span["start"]
        )

        local_end = (
            result.end
            - cell_span["start"]
        )

        local_result = deepcopy(result)

        local_result.start = local_start
        local_result.end = local_end

        detections_by_cell.setdefault(
            cell_index,
            []
        ).append(local_result)

    # 실제 Cell 수정
    for cell_index, cell_results in detections_by_cell.items():

        cell = cells[cell_index]

        original_text = cell.text

        masked_text = mask_text(
            original_text,
            cell_results,
        )

        cell.text = masked_text


def _mask_table(
    table,
    analyzer,
    target_entities,
):
    """
    표 전체를 처리한다.
    """

    headers = [cell.text for cell in table.rows[0].cells] if table.rows else []

    for row_index, row in enumerate(table.rows):

        _mask_table_row(
            row=row,
            analyzer=analyzer,
            target_entities=target_entities,
            headers=headers if row_index > 0 else None,
        )

        # Cell 내부에 중첩된 Table이 있을 수도 있으므로 재귀 처리
        for cell in row.cells:

            for nested_table in cell.tables:

                _mask_table(
                    table=nested_table,
                    analyzer=analyzer,
                    target_entities=target_entities,
                )


def _sanitize_part(
    paragraphs,
    tables,
    analyzer,
    target_entities,
):
    """
    DOCX의 특정 영역을 공통적으로 처리한다.

    사용 대상:
    - 본문
    - Header
    - Footer
    """

    for paragraph in paragraphs:

        _mask_paragraph(
            paragraph=paragraph,
            analyzer=analyzer,
            target_entities=target_entities,
        )

    for table in tables:

        _mask_table(
            table=table,
            analyzer=analyzer,
            target_entities=target_entities,
        )


def sanitize_docx(
    input_path,
    output_path,
    analyzer,
    target_entities,
):
    """
    DOCX의 다음 영역을 검사하여 PII를 마스킹한다.

    지원:
    - 일반 본문 Paragraph
    - 일반 Table
    - Header Paragraph
    - Header Table
    - Footer Paragraph
    - Footer Table
    - 중첩 Table

    결과는 새로운 DOCX 파일로 저장한다.
    """

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(input_path)

    # --------------------------------------------------
    # 1. 본문
    # --------------------------------------------------

    _sanitize_part(
        paragraphs=doc.paragraphs,
        tables=doc.tables,
        analyzer=analyzer,
        target_entities=target_entities,
    )

    # --------------------------------------------------
    # 2. Header / Footer
    # --------------------------------------------------

    for section in doc.sections:

        header = section.header

        _sanitize_part(
            paragraphs=header.paragraphs,
            tables=header.tables,
            analyzer=analyzer,
            target_entities=target_entities,
        )

        footer = section.footer

        _sanitize_part(
            paragraphs=footer.paragraphs,
            tables=footer.tables,
            analyzer=analyzer,
            target_entities=target_entities,
        )

    # --------------------------------------------------
    # 3. 저장
    # --------------------------------------------------

    doc.save(output_path)
