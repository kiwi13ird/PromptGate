"""DOCX/XLSX를 원본 위치에 되돌릴 수 있는 segment로 다룬다."""

from __future__ import annotations

from copy import deepcopy
from dataclasses import dataclass
from difflib import SequenceMatcher
from io import BytesIO
from time import perf_counter

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.text.run import Run
from openpyxl import load_workbook

from signature_detector.contracts import CheckableSegment, SegmentFinding, SegmentResult, validate_segment_result
from signature_detector.policy import PolicyAction
from signature_detector.contextual_pii import canonical_label


@dataclass(frozen=True)
class SanitizedArtifacts:
    sanitized_bytes: bytes
    review_bytes: bytes
    segment_count: int
    elapsed_ms: float
    verify: object | None = None   # 2026-10-11 HWP·HWPX: 재검사 글자·바꾼 구간(hwp_segments.HwpVerify), 그 밖 형식은 None


def _docx_paragraphs(document):
    """지원 영역의 paragraph를 안정적인 순번과 함께 반환한다."""

    seen = set()

    def emit_paragraphs(paragraphs, prefix):
        for index, paragraph in enumerate(paragraphs):
            key = paragraph._p
            if key not in seen:
                seen.add(key)
                yield f"{prefix}:p{index}", paragraph

    def emit_tables(tables, prefix):
        for table_index, table in enumerate(tables):
            for row_index, row in enumerate(table.rows):
                for cell_index, cell in enumerate(row.cells):
                    cell_prefix = f"{prefix}:t{table_index}:r{row_index}:c{cell_index}"
                    yield from emit_paragraphs(cell.paragraphs, cell_prefix)
                    yield from emit_tables(cell.tables, cell_prefix)

    yield from emit_paragraphs(document.paragraphs, "body")
    yield from emit_tables(document.tables, "body")
    for section_index, section in enumerate(document.sections):
        yield from emit_paragraphs(section.header.paragraphs, f"header{section_index}")
        yield from emit_tables(section.header.tables, f"header{section_index}")
        yield from emit_paragraphs(section.footer.paragraphs, f"footer{section_index}")
        yield from emit_tables(section.footer.tables, f"footer{section_index}")


def extract_docx_segments(file_bytes: bytes) -> list[CheckableSegment]:
    document = Document(BytesIO(file_bytes))
    contexts = {}
    def scan(tables):
        for table in tables:
            headers = {}
            for row_index, row in enumerate(table.rows):
                if not any(cell.text.strip() for cell in row.cells):
                    headers.clear()
                for col, cell in enumerate(row.cells):
                    label = None
                    if col and row.cells[col - 1]._tc is not cell._tc:
                        candidate = row.cells[col - 1].text.strip()
                        if canonical_label(candidate):
                            label = candidate
                    if not label and row_index:
                        label = headers.get(col)
                    if canonical_label(cell.text.strip()):
                        headers[col] = cell.text.strip()
                        label = None
                    if label:
                        for paragraph in cell.paragraphs:
                            contexts[paragraph._p] = label
                    scan(cell.tables)
    scan(document.tables)
    for section in document.sections:
        scan(section.header.tables)
        scan(section.footer.tables)
    return [
        CheckableSegment(f"f:0:docx:{locator}", "file", paragraph.text, contexts.get(paragraph._p))
        for locator, paragraph in _docx_paragraphs(document)
        if paragraph.text.strip()
    ]


def _style_for_offset(runs, offset):
    cursor = 0
    for run in runs:
        end = cursor + len(run.text)
        if offset < end or (offset == end and end == len("".join(r.text for r in runs))):
            return deepcopy(run._r.rPr) if run._r.rPr is not None else None
        cursor = end
    return None


def _rewrite_docx_paragraph(paragraph, result: SegmentResult, review: bool) -> None:
    original = paragraph.text
    if review:
        validate_segment_result(CheckableSegment(result.segment_id, 'file', original), result)
        spans = [(f.start, f.end) for f in result.findings]
        if not spans:
            raise ValueError('Original review requires finding spans')
        old_runs = list(paragraph.runs)
        if ''.join(r.text for r in old_runs) != original:
            raise ValueError('Unsupported DOCX run mapping for original review')
        cursor = 0
        for old_run in old_runs:
            text = old_run.text
            boundaries = sorted({0, len(text)} | {max(0, min(len(text), p - cursor))
                for span in spans for p in span})
            for a, b in zip(boundaries, boundaries[1:]):
                node = deepcopy(old_run._r)
                old_run._r.addprevious(node)
                new_run = Run(node, paragraph)
                new_run.text = text[a:b]
                if any(start < cursor + b and end > cursor + a for start, end in spans):
                    new_run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            paragraph._p.remove(old_run._r)
            cursor += len(text)
        return
    if result.masked_text is None:
        raise ValueError("DOCX mask result requires masked_text")
    pieces = []
    for tag, original_start, original_end, masked_start, masked_end in SequenceMatcher(
        None, original, result.masked_text, autojunk=False
    ).get_opcodes():
        if tag == "equal":
            pieces.append((original[original_start:original_end], original_start, False))
        else:
            # Detection의 offset은 정규화 문자열 기준일 수 있으므로 원문에 적용하지 않는다.
            # 완성된 masked_text와 원문 diff만 사용한다.
            pieces.append((result.masked_text[masked_start:masked_end], original_start, True))

    old_runs = list(paragraph.runs)
    for run in old_runs:
        paragraph._p.remove(run._r)
    for text, style_offset, is_mask in pieces:
        if not text:
            continue
        run = paragraph.add_run(text)
        rpr = _style_for_offset(old_runs, style_offset)
        if rpr is not None:
            run._r.insert(0, rpr)
        if review and is_mask:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _docx_result_map(segments, results):
    segment_by_id = {segment.segment_id: segment for segment in segments}
    result_by_id = {result.segment_id: result for result in results}
    if set(result_by_id) != set(segment_by_id) or len(result_by_id) != len(results):
        raise ValueError("DOCX Detection results are missing, duplicated, or unexpected")
    for segment_id, result in result_by_id.items():
        # 파일 재구성에는 정규화 offset 대신 완성된 masked_text를 사용한다.
        if result.segment_id != segment_id:
            raise ValueError("Detection returned a different DOCX segment_id")
        if result.action in {PolicyAction.MASK_PARTIAL, PolicyAction.MASK_FULL}:
            if result.masked_text is None or result.masked_text == segment_by_id[segment_id].text:
                raise ValueError("DOCX mask action requires changed masked_text")
        if result.action is PolicyAction.BLOCK_NOTIFY:
            raise PermissionError("DOCX blocked by Signature policy")
    return result_by_id


def sanitize_docx_bytes(file_bytes: bytes, results: list[SegmentResult]) -> SanitizedArtifacts:
    started = perf_counter()
    segments = extract_docx_segments(file_bytes)
    result_by_id = _docx_result_map(segments, results)
    if not any(r.action in {PolicyAction.MASK_PARTIAL, PolicyAction.MASK_FULL} for r in results):
        return SanitizedArtifacts(file_bytes, file_bytes, len(segments), (perf_counter() - started) * 1000)
    outputs = []
    for review in (False, True):
        document = Document(BytesIO(file_bytes))
        for locator, paragraph in _docx_paragraphs(document):
            segment_id = f"f:0:docx:{locator}"
            result = result_by_id.get(segment_id)
            if result and result.action in {PolicyAction.MASK_PARTIAL, PolicyAction.MASK_FULL}:
                _rewrite_docx_paragraph(paragraph, result, review)
        stream = BytesIO()
        document.save(stream)
        outputs.append(stream.getvalue())
    return SanitizedArtifacts(outputs[0], outputs[1], len(segments), (perf_counter() - started) * 1000)


def extract_xlsx_segments(file_bytes: bytes) -> list[CheckableSegment]:
    workbook = load_workbook(BytesIO(file_bytes), data_only=False)
    segments = []
    for sheet_index, worksheet in enumerate(workbook.worksheets):
        headers = {}
        def merged_value(row_index, column):
            value = worksheet.cell(row_index, column).value
            if value is not None:
                return value
            for area in worksheet.merged_cells.ranges:
                if area.min_row <= row_index <= area.max_row and area.min_col <= column <= area.max_col:
                    return worksheet.cell(area.min_row, area.min_col).value
            return None
        for row in worksheet.iter_rows():
            if not any(cell.value is not None for cell in row):
                headers.clear()
            for cell in row:
                context = None
                if cell.column > 1:
                    label = str(merged_value(cell.row, cell.column - 1) or '')
                    if canonical_label(label):
                        context = label
                if not context and cell.row > 1:
                    context = headers.get(cell.column)
                value = cell.value
                header = merged_value(cell.row, cell.column)
                if canonical_label(str(header or '')):
                    headers[cell.column] = str(header)
                    context = None
                if canonical_label(context) == 'KR_DATE_OF_BIRTH':
                    from datetime import date, datetime
                    if isinstance(value, (date, datetime)):
                        value = value.strftime('%Y-%m-%d')
                if context and type(value) is int and value >= 0:
                    value = str(value)
                    if canonical_label(context) == 'KR_POSTAL_CODE' and cell.number_format == '00000':
                        value = value.zfill(5)
                if isinstance(value, str) and value.strip():
                    segments.append(
                        CheckableSegment(
                            f"f:0:xlsx:s{sheet_index}:r{cell.row - 1}:c{cell.column - 1}",
                            "file",
                            value,
                            context,
                        )
                    )
    return segments


def sanitize_xlsx_bytes(file_bytes: bytes, results: list[SegmentResult]) -> SanitizedArtifacts:
    started = perf_counter()
    segments = extract_xlsx_segments(file_bytes)
    segment_by_id = {segment.segment_id: segment for segment in segments}
    result_by_id = {result.segment_id: result for result in results}
    if set(result_by_id) != set(segment_by_id) or len(result_by_id) != len(results):
        raise ValueError("XLSX Detection results are missing, duplicated, or unexpected")
    for segment_id, result in result_by_id.items():
        # cell 전체 문자열을 교체하므로 정규화 offset을 원본 cell에 적용하지 않는다.
        if result.segment_id != segment_id:
            raise ValueError("Detection returned a different XLSX segment_id")
        if result.action in {PolicyAction.MASK_PARTIAL, PolicyAction.MASK_FULL}:
            if result.masked_text is None or result.masked_text == segment_by_id[segment_id].text:
                raise ValueError("XLSX mask action requires changed masked_text")
        if result.action is PolicyAction.BLOCK_NOTIFY:
            raise PermissionError("XLSX blocked by Signature policy")

    if not any(r.action in {PolicyAction.MASK_PARTIAL, PolicyAction.MASK_FULL} for r in results):
        return SanitizedArtifacts(file_bytes, file_bytes, len(segments), (perf_counter() - started) * 1000)
    outputs = []
    for review in (False, True):
        workbook = load_workbook(BytesIO(file_bytes), data_only=False)
        for segment_id, result in result_by_id.items():
            if result.action not in {PolicyAction.MASK_PARTIAL, PolicyAction.MASK_FULL}:
                continue
            _, _, _, sheet_token, row_token, column_token = segment_id.split(":")
            worksheet = workbook.worksheets[int(sheet_token[1:])]
            cell = worksheet.cell(row=int(row_token[1:]) + 1, column=int(column_token[1:]) + 1)
            if not review:
                cell.value = result.masked_text
        stream = BytesIO()
        workbook.save(stream)
        outputs.append(stream.getvalue())
    return SanitizedArtifacts(outputs[0], outputs[1], len(segments), (perf_counter() - started) * 1000)


def result_from_masking(segment: CheckableSegment, masking_result) -> SegmentResult:
    return SegmentResult(
        segment.segment_id,
        masking_result.action,
        masking_result.masked_text,
        tuple(
            SegmentFinding(
                finding.entity_type,
                finding.start,
                finding.end,
                finding.replacement,
                finding.action,
                finding.score,
            )
            for finding in masking_result.findings
        ),
    )
