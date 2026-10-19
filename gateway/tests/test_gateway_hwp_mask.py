"""HWP·HWPX 가림 3단계 Gateway 쪽 (2026-10-11, masking-design/SPEC.md §4). 가짜 값만 쓴다.

1. 모듈: HWPX 표시 치환(바뀐 문단만 linesegarray 삭제, zip 순서·mimetype, 미리보기), HWP 같은 길이(완전 가림 '[유형]' + 공백,
   표시가 더 길면 '*', 2026-10-12. 파일 크기·레코드 구조 같음, 안 쓰는 바이트 0, 크기 못 맞춤·제어 글자 → 거절), 부가 글자 가림
   요청 → 거절, 병합 칸 이름표(격자), 미리보기 그림
2. 팀 HWP·한컴 견본 왕복(통계만, 파일이 없으면 건너뜀): 전 ASCII 숫자 가림('*' / HWP는 표시 '[N]' + 공백도) → hwpdoc 다시
   읽기, 남은 값 0, 구조·크기 같음
3. 실제 mitmproxy flow(request() 훅 전체, 작업 스레드 포함) × ChatGPT 원본 PUT·multipart·Claude multipart:
   127.0.0.1 가짜 Detection(/detect) + 실제 detection_client.check. 가림 성공·재검사(바꾼 구간 안 적중 무시, 넘으면 403)·
   부가 글자·파일 전체에서만 잡힌 값·칸 무응답 표(§4.5 네 줄)·정책 조합·글자 상한
실행: python -m unittest tests.test_gateway_hwp_mask -v (Gateway 사본 폴더에서)
"""

from __future__ import annotations

import io
import json
import os
import re
import struct
import sys
import tempfile
import threading
import unittest
import zipfile
import zlib
from email.parser import BytesParser
from email.policy import default as email_default
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

_TMP = tempfile.mkdtemp(prefix="gw-hwp-mask-")
os.environ.setdefault("SIGNATURE_REVIEW_DIR", str(Path(_TMP) / "review"))
os.environ.setdefault("DASHBOARD_OUTBOX_PATH", str(Path(_TMP) / "outbox.sqlite3"))
os.environ.pop("DASHBOARD_EVENT_URL", None)
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import olefile  # noqa: E402
from docx import Document  # noqa: E402
from mitmproxy.test import tflow, tutils  # noqa: E402

import detection_client  # noqa: E402
import gateway  # noqa: E402
import gateway_file_masking  # noqa: E402
from gateway_file_classes import inspect_upload  # noqa: E402
from signature_detector.contracts import MaskRefused, SegmentFinding, SegmentResult  # noqa: E402
from signature_detector.document import hwp_mask, hwp_segments, hwpdoc  # noqa: E402
from signature_detector.policy import PolicyAction  # noqa: E402
from tests import hwp_synth as S  # noqa: E402
from tests.gw_support import FixedSettings, png_bytes, gif_bytes, run_hook  # noqa: E402

RRN = "900101-1234567"                     # 가상 주민등록번호
RRN_RE = re.compile(r"\d{6}-\d{7}")
EMAIL = "hong.gildong@example.test"        # 정책 ALLOW(그대로 둠)
EMAIL_RE = re.compile(r"[\w.]+@[\w.]+\.test")
ACCOUNT = "110-123-456789"                 # 칸 이름표(계좌번호)가 있어야 잡히는 값
ACCOUNT_RE = re.compile(r"\d{3}-\d{3}-\d{6}")
PHONE = "010-1234-5678"                    # 가상 전화번호(13자 < '[KR_PHONE_NUMBER]' 17자 → HWP는 '*')
HWP_RRN = "[KR_RRN]" + " " * 6             # HWP 완전 가림: 표시 + 원래 길이(14)까지 공백
ROSTER = [f"{i:04d} 가짜 직원 메모 {i * 7919 % 100000:05d} 비고 {'가나다라마바사'[i % 7]}" for i in range(400)]   # 압축 여유
WHOLE_ONLY = "WHOLE-ONLY-7731"             # 파일 전체 검사에서만 잡히는 값(이름표가 옆 칸이 아닌 경우 등)
DOWN = "DOWN-MARKER-0001"                  # 칸 검사에서 Detection 500
FILE_ID = "00000000-aaaa-8206-bbbb-0000000000a7"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
OLD_503_BODY = b'{"error":{"type":"file_masking_failed"}}'
LIMIT = 20 * 1024 * 1024
TEAM_DIR = Path(os.environ.get("GW_TEST_TEAM_HWP_DIR", "C:/Users/user/Documents/카카오톡 받은 파일"))
HANCOM_DIR = Path(os.environ.get("GW_TEST_HANCOM_DIR", "C:/Program Files (x86)/HNC"))


# ------------------------------------------------------------------ 가짜 문서

def zlib6(data: bytes) -> bytes:
    """한컴 저장본처럼 기본 수준으로 압축(같은 크기 재압축 여유를 흉내 냄 - 팀 HWP 19/19가 맞춤)."""
    c = zlib.compressobj(6, zlib.DEFLATED, -15)
    return c.compress(data) + c.flush()


def zlib1(data: bytes) -> bytes:
    """느슨한 압축(수준 1). 10/12: 수준 6 원본의 작은 구역은 표시 + 공백을 같은 크기로 다시 압축할 여유가 zlib 판마다
    달랐다(운영 zlib 1.3은 '*'로 되돌아가고 로컬 zlib-ng는 맞춤). 표시 시험이 판에 상관없이 같은 결과가 나오게 쓴다.
    팀 HWP 19개(한컴 저장본)는 운영과 같은 zlib 1.3.1에서도 표시가 모두 들어감(되돌아감 0)."""
    c = zlib.compressobj(1, zlib.DEFLATED, -15)
    return c.compress(data) + c.flush()


def hwp_doc(*, body=(), rows=(), header=None, summary=None, prvtext="", extra=None, compress_raw=zlib1):
    """본문 문단·2열 표(이름표, 값)·머리말 문단으로 HWP. 미리보기는 기본으로 빈 글자(분류의 미리보기 대조를 피함)."""
    first = body[0] if body else "가짜 문서"
    children = []
    if header is not None:
        children = [S.ext_ctrl(16, "head"), S.header_ctrl(1, [S.p(2, header)])]
    sec = S.para(0, S.text_units(first) + (children[0] if children else b""), *children[1:])
    if rows:
        cells = []
        for r, (label, value) in enumerate(rows):
            cells += [(r, 0, 1, 1, [S.p(2, label)]), (r, 1, 1, 1, [S.p(2, value)])]
        sec += S.para(0, S.ext_ctrl(11, "tbl "), S.table(1, cells, len(rows), 2))
    sec += b"".join(S.p(0, text) for text in body[1:])
    return S.build_hwp([sec], summary=summary, prvtext=prvtext, extra=extra, compress_raw=compress_raw)


def hwpx_doc(*, body=(), rows=(), header=None, meta=None, prvtext="", extra=None):
    paras = [S.hp_p(S.hp_t(text)) for text in body] or [S.hp_p(S.hp_t("가짜 문서"))]
    if header is not None:
        paras[0] = S.hp_p(S.hp_t(body[0] if body else "가짜 문서"), S.hp_run(
            '<hp:ctrl><hp:header id="1" applyPageType="BOTH"><hp:subList id="" textDirection="HORIZONTAL">'
            + S.hp_p(S.hp_t(header)) + "</hp:subList></hp:header></hp:ctrl>"))
    if rows:
        trs = [[S.hp_tc(r, 0, S.hp_p(S.hp_t(label))), S.hp_tc(r, 1, S.hp_p(S.hp_t(value)))]
               for r, (label, value) in enumerate(rows)]
        paras.insert(1, S.hp_p(S.hp_tbl(trs)))
    return S.build_hwpx(["".join(paras)], meta=meta, prvtext=prvtext, extra=extra)


PII_ROWS = [("성명", "김가짜"), ("주민등록번호", RRN)]
PII_BODY = ["가짜 인사 기록 카드", f"연락처 메모 {RRN} 끝"]


def pii_hwp(**kwargs):
    return hwp_doc(body=PII_BODY, rows=PII_ROWS, header=f"머리말 {RRN}", **kwargs)


def pii_hwpx(**kwargs):
    return hwpx_doc(body=PII_BODY, rows=PII_ROWS, header=f"머리말 {RRN}", **kwargs)


def reread(data):
    return hwpdoc.read(data, max_decompressed=LIMIT)


def finding(start, end, replacement, action="mask_full", entity="KR_RRN"):
    return SegmentFinding(entity, start, end, replacement, PolicyAction(action), 1.0)


def results_for(segments, rule):
    """rule(segment) → [(start, end, replacement, action[, entity])] - Detection 칸 판정(masked_text는 replacement로 만든 것)."""
    out = []
    for segment in segments:
        found = rule(segment)
        findings = tuple(finding(*item) for item in found)
        masked = segment.text
        for s, e, r, *_ in sorted(found, reverse=True):
            if r is not None:
                masked = masked[:s] + r + masked[e:]
        mask = [item[3] for item in found if item[3].startswith("mask")]
        action = PolicyAction("mask_full" if "mask_full" in mask else "mask_partial" if mask else "allow")
        out.append(SegmentResult(segment.segment_id, action, masked if masked != segment.text else None, findings))
    return out


def rrn_rule(replacement="[KR_RRN]"):
    return lambda seg: [(m.start(), m.end(), replacement, "mask_full") for m in RRN_RE.finditer(seg.text)]


def sanitize(kind, data, rule):
    doc = hwp_segments.read_document(data)
    segments = hwp_segments.extract_segments(doc, kind)
    return hwp_segments.sanitize_document(kind, data, results_for(segments, rule), doc=doc, segments=segments)


def section_records(data):
    ole = hwpdoc._ole_open(data)
    try:
        compressed = ole.openstream("FileHeader").read()[36] & 1
        out = {}
        for entry in ole.listdir(streams=True, storages=False):
            if entry[0] == "BodyText":
                raw = ole.openstream(entry).read()
                plain = hwpdoc._inflate(raw, hwpdoc._Budget(LIMIT)) if compressed else raw
                out["/".join(entry)] = [(t, l, s) for t, l, _, s in hwpdoc._records(plain)]
        return out
    finally:
        ole.close()


def nonzero_free_sectors(data):
    ole = olefile.OleFileIO(data)
    try:
        ss = ole.sectorsize
        count = (len(data) - ss) // ss
        return sum(1 for s, v in enumerate(ole.fat[:count])
                   if v == hwp_mask.FREESECT and any(data[(s + 1) * ss:(s + 2) * ss]))
    finally:
        ole.close()


# ------------------------------------------------------------------ 1. 모듈

class HwpxMaskTests(unittest.TestCase):
    def test_tokens_replace_values_and_only_changed_paragraphs_lose_linesegarray(self):
        original = pii_hwpx(prvtext=f"가짜 인사 기록 카드\r\n주민등록번호 {RRN}\r\n")
        art = sanitize("hwpx", original, rrn_rule())
        doc = reread(art.sanitized_bytes)
        texts = {p.locator: p.text for p in doc.paragraphs}
        self.assertEqual(texts["body:p2"], "연락처 메모 [KR_RRN] 끝")
        self.assertEqual(texts["body:t0:r1:c1:p0"], "[KR_RRN]")
        self.assertEqual(texts["header0:p0"], "머리말 [KR_RRN]")
        self.assertNotIn(RRN, "\n".join(doc.lines()))
        self.assertEqual(doc.preview_text().strip(), "")                      # 미리보기 글자 지움(줄바꿈은 남음)
        self.assertIn("\r\n", doc.preview_text())
        source, target = zipfile.ZipFile(io.BytesIO(original)), zipfile.ZipFile(io.BytesIO(art.sanitized_bytes))
        self.assertEqual([i.filename for i in target.infolist()], [i.filename for i in source.infolist()])
        self.assertEqual((target.infolist()[0].filename, target.infolist()[0].compress_type), ("mimetype", 0))
        self.assertEqual([i.date_time for i in target.infolist()], [i.date_time for i in source.infolist()])
        raw_xml = target.read("Contents/section0.xml")
        self.assertTrue(raw_xml.startswith(b'<?xml version="1.0" encoding="UTF-8" standalone="yes" ?>'))
        from lxml import etree
        hp = "{%s}" % S.HP
        has_lineseg = {}
        for p in etree.fromstring(raw_xml).iter(hp + "p"):
            own = "".join(t.text or "" for t in p.iter(hp + "t")
                          if next(t.iterancestors(hp + "p")) is p)
            has_lineseg[own] = p.find(hp + "linesegarray") is not None
        self.assertEqual({k: v for k, v in has_lineseg.items() if "[KR_RRN]" in k},
                         {"[KR_RRN]": False, "연락처 메모 [KR_RRN] 끝": False, "머리말 [KR_RRN]": False})   # 바뀐 문단: 삭제
        self.assertTrue(all(v for k, v in has_lineseg.items() if "[KR_RRN]" not in k))      # 안 바뀐 문단: 그대로
        self.assertEqual(target.comment, b"")
        # 재검사 입력: 다시 뽑은 글자 + 바꾼 구간(새 글자 좌표)
        self.assertEqual([art.verify.text[a:b] for a, b in art.verify.spans], ["[KR_RRN]"] * 3)
        self.assertEqual(art.verify.text, "\n".join(doc.lines()))
        # Gateway 분류도 그대로 검사 가능(바이트 배치 규칙 통과)
        fc = inspect_upload(art.sanitized_bytes, max_bytes=LIMIT, max_chars=80645, deadline=1e18)
        self.assertEqual((fc.kind, fc.group), ("hwpx", None))

    def test_value_split_across_runs_is_written_into_the_first_piece(self):
        data = S.build_hwpx([S.hp_p(S.hp_t("앞 900101"), S.hp_t("-1234567 뒤"))])
        art = sanitize("hwpx", data, rrn_rule())
        self.assertEqual([p.text for p in reread(art.sanitized_bytes).paragraphs], ["앞 [KR_RRN] 뒤"])

    def test_element_character_inside_the_span_is_refused(self):
        data = S.build_hwpx([S.hp_p(S.hp_t("번호 900101<hp:tab/>1234567"))])
        with self.assertRaises(MaskRefused) as caught:
            sanitize("hwpx", data, lambda seg: [(3, 17, "[KR_RRN]", "mask_full")] if seg.text.startswith("번호") else [])
        self.assertEqual(caught.exception.reason, "hwpx_masked_text_mismatch")

    def test_partial_mask_and_allowed_value(self):
        data = hwpx_doc(body=[f"{RRN} / {EMAIL}"])

        def rule(seg):
            out = [(m.start(), m.end(), m.group()[:8] + "*" * 6, "mask_partial") for m in RRN_RE.finditer(seg.text)]
            return out + [(m.start(), m.end(), None, "allow") for m in EMAIL_RE.finditer(seg.text)]
        art = sanitize("hwpx", data, rule)
        self.assertEqual(reread(art.sanitized_bytes).paragraphs[0].text, f"900101-1****** / {EMAIL}")

    def test_meta_text_needing_a_mask_is_refused(self):
        data = hwpx_doc(body=["평범한 메모"], meta={"creator": RRN})
        with self.assertRaises(MaskRefused) as caught:
            sanitize("hwpx", data, rrn_rule())
        self.assertEqual(caught.exception.reason, "hwpx_side_text_mask")


class HwpMaskTests(unittest.TestCase):
    def test_full_mask_token_with_spaces_keeps_size_records_and_zeroes_unused_bytes(self):
        big_picture = png_bytes(2, 2) + bytes((i * 7919) % 251 for i in range(6000))   # 미니 스트림 밖(4096바이트 이상)
        original = pii_hwp(prvtext=f"가짜 인사 기록 카드\r\n주민등록번호 {RRN}\r\n",
                           extra={"BinData/BIN0001.png": big_picture, "PrvImage": png_bytes(30, 30)})
        dirty = self.dirty_copy(original)
        art = sanitize("hwp", dirty, rrn_rule())
        masked = art.sanitized_bytes
        self.assertEqual(len(masked), len(dirty))
        self.assertEqual(section_records(masked), section_records(dirty))
        doc = reread(masked)
        self.assertNotIn(RRN, "\n".join(doc.lines()))
        self.assertIn(f"연락처 메모 {HWP_RRN} 끝", doc.lines())
        self.assertEqual([p.text for p in doc.paragraphs if p.locator.startswith("header0:")], [f"머리말 {HWP_RRN}"])
        self.assertEqual(doc.preview_text().replace("\x00", "").strip(), "")
        self.assertEqual(nonzero_free_sectors(dirty), 1)
        self.assertEqual(nonzero_free_sectors(masked), 0)
        self.assertNotIn(b"JUNK", masked)                                      # 빈 섹터·스트림 꼬리·미니 꼬리의 옛 바이트
        before, after = hwp_mask._all_streams(dirty), hwp_mask._all_streams(masked)
        changed = {name for name in before if before[name] != after[name]}
        self.assertEqual(changed, {"BodyText/Section0", "PrvText", "PrvImage"})
        image = after["PrvImage"]
        self.assertEqual((len(image), image[:8]), (len(before["PrvImage"]), b"\x89PNG\r\n\x1a\n"))
        self.assertEqual([art.verify.text[a:b] for a, b in art.verify.spans], [HWP_RRN] * 3)   # 표시 + 공백 모두 바꾼 구간

    @staticmethod
    def dirty_copy(data: bytes) -> bytes:
        """빈 섹터 하나(파일 끝에 덧붙임)·큰 스트림 마지막 섹터 꼬리·미니 스트림 꼬리에 옛 글자를 흉내 낸 바이트."""
        buf = bytearray(data) + b"JUNK" * 128
        ole = olefile.OleFileIO(bytes(data))
        ss, mss = ole.sectorsize, ole.minisectorsize
        for entry in ole.direntries:
            if entry is None or entry.entry_type != olefile.STGTY_STREAM or not entry.size:
                continue
            if entry.size >= ole.minisectorcutoff:
                sector, chain = entry.isectStart, []
                while sector not in (olefile.ENDOFCHAIN, olefile.FREESECT):
                    chain.append(sector)
                    sector = ole.fat[sector]
                tail = entry.size - (len(chain) - 1) * ss
                if tail < ss:
                    at = (chain[-1] + 1) * ss + tail
                    buf[at:at + 4] = b"JUNK"
            elif entry.name == "PrvText" and entry.size % mss:
                ole.loadminifat()
                container, sector = [], ole.root.isectStart
                while sector not in (olefile.ENDOFCHAIN, olefile.FREESECT):
                    container.append(sector)
                    sector = ole.fat[sector]
                mini, chain = entry.isectStart, []
                while mini not in (olefile.ENDOFCHAIN, olefile.FREESECT):
                    chain.append(mini)
                    mini = ole.minifat[mini]
                pos = chain[-1] * mss + entry.size % mss
                at = (container[pos // ss] + 1) * ss + pos % ss
                buf[at:at + 4] = b"JUNK"
        ole.close()
        return bytes(buf)

    def test_bytes_no_sector_can_reach_are_zeroed(self):
        # 파일 끝의 섹터 크기 미만 조각(어느 FAT 칸도 가리키지 못하는 자리)도 0으로, 파일 크기는 그대로
        original = hwp_doc(body=[f"번호 {RRN}"]) + b"JUNK" * 25
        art = sanitize("hwp", original, rrn_rule())
        self.assertEqual(len(art.sanitized_bytes), len(original))
        self.assertNotIn(b"JUNK", art.sanitized_bytes)
        self.assertNotIn(RRN, "\n".join(reread(art.sanitized_bytes).lines()))

    def test_partial_replacement_of_the_same_length_is_written_as_is(self):
        data = hwp_doc(body=[f"번호 {RRN}"])
        art = sanitize("hwp", data, lambda seg: [(m.start(), m.end(), "900101-*******", "mask_partial")
                                                 for m in RRN_RE.finditer(seg.text)])
        self.assertEqual(reread(art.sanitized_bytes).paragraphs[0].text, "번호 900101-*******")

    def test_allowed_value_stays(self):
        data = hwp_doc(body=[f"{RRN} {EMAIL}"])

        def rule(seg):
            return ([(m.start(), m.end(), "[KR_RRN]", "mask_full") for m in RRN_RE.finditer(seg.text)]
                    + [(m.start(), m.end(), None, "allow") for m in EMAIL_RE.finditer(seg.text)])
        self.assertEqual(reread(sanitize("hwp", data, rule).sanitized_bytes).paragraphs[0].text,
                         f"{HWP_RRN} {EMAIL}")

    def assert_hwp_mask(self, body, rule, expected, filler=ROSTER):
        """가린 문단 글자 = expected, 파일 크기·레코드 같음. 재검사 입력을 돌려준다. filler = 뒤에 붙일 문단(구역 압축 여유)."""
        data = hwp_doc(body=[body, *filler])
        art = sanitize("hwp", data, rule)
        self.assertEqual((len(art.sanitized_bytes), section_records(art.sanitized_bytes)), (len(data), section_records(data)))
        self.assertEqual(reread(art.sanitized_bytes).paragraphs[0].text, expected)
        return art

    def test_token_that_fits_is_padded_and_one_that_does_not_fit_is_stars(self):
        body = f"번호 {RRN} 전화 {PHONE} 계좌 {ACCOUNT} 메일 {EMAIL}"
        email = "[EMAIL_ADDRESS]".ljust(len(EMAIL))

        def rule(seg):
            out = []
            for pattern, entity in ((RRN_RE, "KR_RRN"), (re.compile(re.escape(PHONE)), "KR_PHONE_NUMBER"),
                                    (ACCOUNT_RE, "KR_BANK_ACCOUNT"), (EMAIL_RE, "EMAIL_ADDRESS")):
                out += [(m.start(), m.end(), f"[{entity}]", "mask_full", entity) for m in pattern.finditer(seg.text)]
            return out
        art = self.assert_hwp_mask(body, rule, f"번호 {HWP_RRN} 전화 {'*' * 13} 계좌 {'*' * 14} 메일 {email}")
        # 재검사 무시 구간 = 바꾼 글자 전부(표시 + 채운 공백 / '*')
        self.assertEqual([art.verify.text[a:b] for a, b in art.verify.spans], [HWP_RRN, "*" * 13, "*" * 14, email])

    def test_token_of_exactly_the_original_length_has_no_padding(self):
        value = "pw-secret!"                                              # 10자 = '[PASSWORD]' 10자
        self.assert_hwp_mask(f"암호 {value} 끝", lambda seg: [(3, 13, "[PASSWORD]", "mask_full", "PASSWORD")]
                             if seg.text.startswith("암호") else [], "암호 [PASSWORD] 끝")

    def test_token_is_entity_type_when_the_full_mask_replacement_is_not_a_token(self):
        def rule(seg):
            return [(m.start(), m.end(), "###", "mask_full") for m in RRN_RE.finditer(seg.text)]
        self.assert_hwp_mask(f"번호 {RRN}", rule, f"번호 {HWP_RRN}")

    def test_partial_policy_with_a_token_replacement_writes_the_token(self):
        # 부분 가림 규칙이 없는 유형은 Detection이 부분 가림에도 '[유형]'을 준다 - 프롬프트·DOCX·HWPX와 같은 글자
        def rule(seg):
            return [(m.start(), m.end(), "[KR_RRN]", "mask_partial") for m in RRN_RE.finditer(seg.text)]
        self.assert_hwp_mask(f"번호 {RRN}", rule, f"번호 {HWP_RRN}")

    def test_partial_replacement_of_another_length_stays_stars(self):
        value = "a@b.test"                                                # 부분 가림 'a*@b.test'는 한 글자 김 → 예전처럼 '*'
        self.assert_hwp_mask(f"메일 {value}", lambda seg: [(3, 11, "a*@b.test", "mask_partial", "EMAIL_ADDRESS")]
                             if seg.text.startswith("메일") else [], f"메일 {'*' * 8}")

    def test_four_byte_characters_count_as_two_units(self):
        emoji = "\U0001F600"
        value = "pw" + emoji * 4                                         # 6글자, UTF-16 10단위 → '[PASSWORD]'(10)가 들어감
        self.assert_hwp_mask(f"암호 {value} 끝", lambda seg: [(3, 9, "[PASSWORD]", "mask_full", "PASSWORD")]
                             if seg.text.startswith("암호") else [], "암호 [PASSWORD] 끝")
        value = "p" + emoji * 2                                          # 3글자, 5단위 < 10 → 단위마다 '*'
        self.assert_hwp_mask(f"암호 {value} 끝", lambda seg: [(3, 6, "[PASSWORD]", "mask_full", "PASSWORD")]
                             if seg.text.startswith("암호") else [], "암호 ***** 끝")
        self.assert_hwp_mask(f"암호 {value} 끝", lambda seg: [(3, 6, "p**", "mask_partial", "PASSWORD")]   # 부분 가림: 예전과 같음
                             if seg.text.startswith("암호") else [], "암호 p**** 끝")

    def test_section_that_keeps_its_size_only_with_stars_falls_back_to_stars(self):
        # 표시 + 공백으로는 원래 압축 크기를 못 맞춘 구역 → 그 구역의 완전 가림만 '*'(예전 결과).
        # 맞추고 못 맞추는 것은 zlib 판에 따라 달라서(10/12 운영 zlib 1.3 · 로컬 zlib-ng) 표시판만 못 맞추게 고정한다
        value = "pw-secret!"
        rule = lambda seg: [(3, 13, "[PASSWORD]", "mask_full", "PASSWORD")] if seg.text.startswith("암호") else []  # noqa: E731
        real = hwp_mask.deflate_exact
        token_units = "[PASSWORD]".encode("utf-16-le")
        only_stars = lambda data, target: None if token_units in data else real(data, target)  # noqa: E731
        with mock.patch.object(hwp_mask, "deflate_exact", only_stars):
            art = self.assert_hwp_mask(f"암호 {value} 끝", rule, "암호 ********** 끝", filler=())
        self.assertEqual([art.verify.text[a:b] for a, b in art.verify.spans], ["*" * 10])
        with mock.patch.object(hwp_mask, "deflate_exact", lambda data, target: None):      # '*'로도 못 맞추면 예전처럼 거절
            with self.assertRaises(MaskRefused) as caught:
                sanitize("hwp", hwp_doc(body=[f"암호 {value} 끝"]), rule)
        self.assertEqual(caught.exception.reason, "hwp_section_size")

    def test_tokens_and_spaces_keep_the_size_of_a_large_compressed_section(self):
        rows = [f"{i:04d} 가짜 직원 {i:06d}-{(i * 7919) % 10000000:07d} 메모" for i in range(400)]
        data = hwp_doc(body=["가짜 명부"] + rows)
        art = sanitize("hwp", data, rrn_rule())
        self.assertEqual(len(art.sanitized_bytes), len(data))
        self.assertEqual(section_records(art.sanitized_bytes), section_records(data))
        lines = reread(art.sanitized_bytes).lines()
        self.assertFalse(any(RRN_RE.search(line) for line in lines))
        self.assertEqual(sum(line.count(HWP_RRN) for line in lines), 400)

    def test_section_that_cannot_keep_its_size_is_refused(self):
        text = "A" * 3000                                             # 압축하면 수십 바이트
        data = hwp_doc(body=[text], compress_raw=S.deflate)
        noisy = "".join(chr(0x41 + (i * 7919) % 26) for i in range(3000))   # 같은 길이지만 압축이 안 되는 글자
        with self.assertRaises(MaskRefused) as caught:
            sanitize("hwp", data, lambda seg: [(0, 3000, noisy, "mask_partial")] if seg.text == text else [])
        self.assertEqual(caught.exception.reason, "hwp_section_size")

    def test_control_character_inside_the_span_is_refused(self):
        sec = S.para(0, S.text_units("번호 900101") + S.struct.pack("<H", 10) + S.text_units("1234567"))
        data = S.build_hwp([sec], prvtext="", compress_raw=zlib6)
        with self.assertRaises(MaskRefused) as caught:
            sanitize("hwp", data, lambda seg: [(3, 17, "[KR_RRN]", "mask_full")] if seg.text.startswith("번호") else [])
        self.assertEqual(caught.exception.reason, "hwp_control_in_span")

    def test_summary_text_needing_a_mask_is_refused(self):
        data = hwp_doc(body=["평범한 메모"], summary={2: f"제목 {RRN}"})
        with self.assertRaises(MaskRefused) as caught:
            sanitize("hwp", data, rrn_rule())
        self.assertEqual(caught.exception.reason, "hwp_side_text_mask")

    def test_preview_image_blanks(self):
        for original in (png_bytes(20, 20), png_bytes(300, 200), gif_bytes(),
                         gif_bytes()[:-1] + b"\x21\xfe\x05hello\x00\x3b"):
            blank = hwp_mask.blank_image_same_size(original)
            self.assertEqual(len(blank), len(original))
            self.assertEqual(blank[:3], original[:3])
        self.assertIsNone(hwp_mask.blank_image_same_size(b"BM" + b"\x00" * 60))     # BMP: 거절
        data = hwp_doc(body=[f"번호 {RRN}"], extra={"PrvImage": b"BM" + b"\x00" * 60})
        with self.assertRaises(MaskRefused) as caught:
            sanitize("hwp", data, rrn_rule())
        self.assertEqual(caught.exception.reason, "hwp_preview_image")


class CellLabelTests(unittest.TestCase):
    """병합 칸 이름표(명세 §4.1 격자 방식): 행 2칸 병합 '계좌번호' 아래 칸의 값도 이름표를 받는다."""

    def test_merged_label_covers_the_second_row(self):
        cells = [(0, 0, 2, 1, [S.p(2, "계좌번호")]), (0, 1, 1, 1, [S.p(2, ACCOUNT)]), (1, 1, 1, 1, [S.p(2, "220-456-789012")]),
                 (2, 0, 1, 1, [S.p(2, "비고")]), (2, 1, 1, 1, [S.p(2, "없음")])]
        sec = S.p(0, "가짜 계좌 양식") + S.para(0, S.ext_ctrl(11, "tbl "), S.table(1, cells, 3, 2))
        hwp = S.build_hwp([sec], prvtext="")
        trs = [[S.hp_tc(0, 0, S.hp_p(S.hp_t("계좌번호")), rowspan=2), S.hp_tc(0, 1, S.hp_p(S.hp_t(ACCOUNT)))],
               [S.hp_tc(1, 1, S.hp_p(S.hp_t("220-456-789012")))],
               [S.hp_tc(2, 0, S.hp_p(S.hp_t("비고"))), S.hp_tc(2, 1, S.hp_p(S.hp_t("없음")))]]
        hwpx = S.build_hwpx([S.hp_p(S.hp_t("가짜 계좌 양식")) + S.hp_p(S.hp_tbl(trs))])
        for kind, data in (("hwp", hwp), ("hwpx", hwpx)):
            with self.subTest(kind):
                segments = hwp_segments.extract_segments(hwp_segments.read_document(data), kind)
                context = {seg.text: seg.context for seg in segments}
                self.assertEqual(context[ACCOUNT], "계좌번호")
                self.assertEqual(context["220-456-789012"], "계좌번호")      # 정확한 주소 방식이면 None
                self.assertIsNone(context["없음"])                            # '비고'는 이름표가 아님
                self.assertTrue(all(isinstance(seg.block_order, int) for seg in segments if ":side:" not in seg.segment_id))


# ------------------------------------------------------------------ 2. 팀 HWP·한컴 견본 왕복(통계만)

class RealFileRoundTripTests(unittest.TestCase):
    """전 ASCII 숫자를 같은 길이 '*'(HWP)·'[NUM]'(HWPX)으로 → hwpdoc 다시 읽기. HWP는 완전 가림 표시('[N]' + 공백,
    세 자리 미만은 '*')로도 한 번 더. 글자는 출력하지 않는다(건수만)."""

    def run_folder(self, folder: Path, patterns, hwp_token=None):
        files = [f for pattern in patterns for f in sorted(folder.rglob(pattern))]
        if not files:
            self.skipTest(f"no files under {folder}")
        stats = {"files": 0, "masked": 0, "unsupported": 0, "refused": [], "token_spans": 0, "star_spans": 0}
        for number, path in enumerate(files):
            data = path.read_bytes()
            kind = "hwp" if data[:8] == hwpdoc.CFB_MAGIC else "hwpx"
            try:
                doc = hwp_segments.read_document(data, 64 << 20)
            except hwpdoc.HwpUnsupported:
                stats["unsupported"] += 1                       # 닫힌 목록 밖(1단계에서 이미 other)
                continue
            stats["files"] += 1
            edits = {}
            for pi, para in enumerate(doc.paragraphs):
                spans = [(m.start(), m.end(), "*" * len(m.group()) if kind == "hwp" else "[NUM]")
                         + ((hwp_token,) if kind == "hwp" and hwp_token else ())
                         for m in re.finditer(r"[0-9]+", para.text)]
                if spans:
                    edits[pi] = spans
            if not edits:
                continue
            try:
                mask = hwp_mask.mask_hwp if kind == "hwp" else __import__(
                    "signature_detector.document.hwpx_mask", fromlist=["mask_hwpx"]).mask_hwpx
                masked, new_texts, marks = mask(data, doc, edits, max_decompressed=64 << 20)
                verify = hwp_segments._recheck(kind, masked, doc, new_texts, marks, 64 << 20)
            except MaskRefused as exc:
                stats["refused"].append(exc.reason)
                continue
            stats["masked"] += 1
            again = hwp_segments.read_document(masked, 64 << 20)
            self.assertFalse(any(re.search(r"[0-9]", p.text) for p in again.paragraphs), f"file #{number}")
            self.assertTrue(all(set(verify.text[a:b]) <= set("*[NUM] ") for a, b in verify.spans))
            if kind == "hwp" and hwp_token:
                pieces = [verify.text[a:b] for a, b in verify.spans]
                self.assertTrue(all(p in ("*" * len(p), hwp_token.ljust(len(p))) for p in pieces), f"file #{number}")
                stats["token_spans"] += sum(p.startswith(hwp_token) for p in pieces)
                stats["star_spans"] += sum(not p.startswith(hwp_token) for p in pieces)
            if kind == "hwp":
                self.assertEqual(len(masked), len(data))
                self.assertEqual(section_records(masked), section_records(data))
                self.assertEqual(nonzero_free_sectors(masked), 0)
            else:
                source, target = zipfile.ZipFile(io.BytesIO(data)), zipfile.ZipFile(io.BytesIO(masked))
                self.assertIsNone(target.testzip())
                self.assertEqual(sorted(i.filename for i in target.infolist()),
                                 sorted(i.filename for i in source.infolist()))
                if "mimetype" in source.namelist():
                    self.assertEqual((target.infolist()[0].filename, target.infolist()[0].compress_type),
                                     ("mimetype", zipfile.ZIP_STORED))
        print(f"[ROUND-TRIP] {folder.name} hwp_token={hwp_token}: {stats['files']} files, masked {stats['masked']}, "
              f"closed-list other {stats['unsupported']}, refused {stats['refused']}"
              + (f", hwp spans token {stats['token_spans']} / '*' {stats['star_spans']}" if hwp_token else ""))
        return stats

    def test_team_hwp_files(self):
        for token in (None, "[N]"):
            with self.subTest(token=token):
                stats = self.run_folder(TEAM_DIR, ("*.hwp",), token)
                self.assertEqual(stats["refused"], [])
                self.assertEqual(stats["masked"], stats["files"])
                if token:
                    self.assertTrue(stats["token_spans"] and stats["star_spans"])

    def test_hancom_sample_files(self):
        for token in (None, "[N]"):
            with self.subTest(token=token):
                stats = self.run_folder(HANCOM_DIR, ("*.hwp", "*.hwtx"), token)
                self.assertEqual(stats["refused"], [])


# ------------------------------------------------------------------ 3. 실제 flow

class FakeDetection:
    """127.0.0.1 가짜 /detect. 가짜 주민번호 → rrn_action(mask_full '[KR_RRN]' 또는 mask_partial), 이메일 → allow,
    계좌 이름표 문맥 → mask_full, WHOLE_ONLY → 파일 전체에서만, DOWN → 칸 검사 500. verify_extra(text) → 재검사에 더할 finding."""

    def __init__(self):
        self.rrn_action = "mask_full"
        self.rrn_token = "[KR_RRN]"
        self.verify_extra = None
        self.detects = []
        outer = self

        class Handler(BaseHTTPRequestHandler):
            protocol_version = "HTTP/1.1"

            def log_message(self, *args):
                pass

            def reply(self, status, body):
                self.close_connection = True
                self.send_response(status)
                self.send_header("Connection", "close")
                self.send_header("Content-Type", "application/json")
                self.send_header("Content-Length", str(len(body)))
                self.end_headers()
                self.wfile.write(body)

            def do_POST(self):
                body = self.rfile.read(int(self.headers.get("Content-Length") or 0))
                status, reply = outer.detect(json.loads(body))
                self.reply(status, json.dumps(reply).encode())

        self.server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        self.server.daemon_threads = True
        self.url = f"http://127.0.0.1:{self.server.server_address[1]}"
        threading.Thread(target=self.server.serve_forever, kwargs={"poll_interval": 0.05}, daemon=True).start()

    def close(self):
        self.server.shutdown()
        self.server.server_close()

    def detect(self, request):
        rid, text, source = request["request_id"], request["normalized_text"], request.get("source") or {}
        self.detects.append(request)
        segment, verify = "#f:0:" in rid, rid.endswith("#verify")
        if segment and DOWN in text:
            return 500, {"detail": "boom"}
        found = []
        for m in RRN_RE.finditer(text):
            rep = self.rrn_token if self.rrn_action == "mask_full" else m.group()[:7] + "*" * 7
            found.append((m.start(), m.end(), "KR_RRN", rep, self.rrn_action))
        found += [(m.start(), m.end(), "EMAIL_ADDRESS", None, "allow") for m in EMAIL_RE.finditer(text)]
        if segment and source.get("context_label") == "KR_BANK_ACCOUNT":
            found += [(m.start(), m.end(), "KR_BANK_ACCOUNT", "[KR_BANK_ACCOUNT]", "mask_full")
                      for m in ACCOUNT_RE.finditer(text)]
        if not segment and not verify:
            found += [(m.start(), m.end(), "KR_DATE_OF_BIRTH", "[KR_DATE_OF_BIRTH]", "mask_full")
                      for m in re.finditer(re.escape(WHOLE_ONLY), text)]
        if verify and self.verify_extra:
            found += self.verify_extra(text)
        found.sort()
        masked = text
        for start, end, _, rep, _ in reversed(found):
            if rep is not None:
                masked = masked[:start] + rep + masked[end:]
        actions = {a for *_, a in found}
        action = "mask_full" if "mask_full" in actions else "mask_partial" if "mask_partial" in actions else "allow"
        counts = {}
        for _, _, entity, _, _ in found:
            counts[entity] = counts.get(entity, 0) + 1
        return 200, {
            "blocked": False, "reason": "signature" if found else "", "request_id": rid,
            "decided_by": "signature" if found else "none",
            "pii_hits": ",".join(f"{k}:{v}" for k, v in sorted(counts.items())) or None,
            "hash_status": "ok", "signature_status": "ok",
            "findings": [{"entity_type": e, "start": s, "end": t, "replacement": r, "policy_action": a, "score": 1.0}
                         for s, t, e, r, a in found],
            "masked_text": masked if masked != text else None,
            "decision": {"action": action, "reason_code": "pii" if action != "allow" else "none"}}


def multipart_body(boundary, parts):
    body = b""
    for filename, ctype, content in parts:
        body += (f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"{filename}\"\r\n"
                 f"Content-Type: {ctype}\r\n\r\n").encode() + content + b"\r\n"
    return body + f"--{boundary}--\r\n".encode()


def multipart_files(body, boundary):
    message = BytesParser(policy=email_default).parsebytes(
        f"Content-Type: multipart/form-data; boundary={boundary}\r\n\r\n".encode() + body)
    return [part.get_payload(decode=True) for part in message.iter_parts() if part.get_filename()]


class FlowBase(unittest.TestCase):
    target_llm = "ChatGPT"
    verify_suffix = ""

    def setUp(self):
        self.detection = FakeDetection()
        self.addCleanup(self.detection.close)
        self.reports = []

        def capture(flow, result, dashboard_context, *, enforcement_applied, http_status):
            self.reports.append(SimpleNamespace(result=result, context=dashboard_context,
                                                applied=enforcement_applied, http_status=http_status))
        for name, value in (("DETECTION_BASE_URL", self.detection.url), ("check_detection", detection_client.check),
                            ("report_dashboard_event", capture), ("send_to_next_module", lambda event: None)):
            patcher = mock.patch.object(gateway, name, value)
            patcher.start()
            self.addCleanup(patcher.stop)
        self.use_settings()
        gateway.CHATGPT_FILE_META.clear()
        gateway.CHATGPT_SANITIZED_FILE_SIZE.clear()
        getattr(gateway, "_dedupe_seen", {}).clear()

    def use_settings(self, **kwargs):
        patcher = mock.patch.object(gateway, "GATEWAY_SETTINGS", FixedSettings(**kwargs))
        patcher.start()
        self.addCleanup(patcher.stop)

    def whole_rid(self):
        return next(d["request_id"] for d in self.detection.detects if "#" not in d["request_id"])

    def verify_ids(self):
        return [d["request_id"] for d in self.detection.detects if d["request_id"].endswith("#verify")]

    def assert_refused(self, flow, status, reason_code):
        self.assertIsNotNone(flow.response)          # 응답을 Gateway가 정함 = 원본·가린 파일 모두 upstream으로 가지 않음
        self.assertEqual(flow.response.status_code, status)
        self.assertEqual(len(self.reports), 1, [r.result.action for r in self.reports])
        report = self.reports[0]
        self.assertEqual((report.result.action, report.applied, report.http_status), ("block", True, status))
        self.assertEqual(report.result.reason_code, reason_code)
        self.assertEqual(report.context["targetLlm"], self.target_llm)
        if status == 503:
            self.assertEqual(flow.response.raw_content, OLD_503_BODY)
        return report


class HwpMaskFlowMatrix:
    """세 경로 공통. 하위 클래스가 send(data, filename)·forwarded(flow)를 정한다."""

    def test_pii_is_masked_rechecked_and_forwarded(self):
        review_dir = gateway_file_masking.REVIEW_DIR
        before = len(list(review_dir.glob("*"))) if review_dir.exists() else 0
        for kind, data, token in (("hwp", pii_hwp(), HWP_RRN), ("hwpx", pii_hwpx(), "[KR_RRN]")):
            with self.subTest(kind):
                self.reports.clear()
                self.detection.detects.clear()
                flow = self.send(data, f"인사.{kind}")
                self.assertIsNone(flow.response)
                out = self.forwarded(flow)
                self.assertNotEqual(out, data)
                doc = reread(out)
                self.assertNotIn(RRN, "\n".join(doc.lines()))
                self.assertEqual(sum(line.count(token) for line in doc.lines()), 3)       # 본문·표·머리말
                if kind == "hwp":
                    self.assertEqual(len(out), len(data))
                self.assertNotIn(RRN.encode("utf-16-le"), flow.request.raw_content)
                report, = self.reports
                self.assertEqual((report.result.action, report.applied), ("mask_full", True))
                self.assertNotIn("#", report.result.request_id)                        # 기록 = 파일 전체 판정
                prefix = self.whole_rid() + self.verify_suffix
                self.assertEqual(self.verify_ids(), [f"{prefix}#verify", f"{prefix}#nows#verify"])
                verify = [d for d in self.detection.detects if d["request_id"].endswith("#verify")]
                self.assertEqual(verify[0]["normalized_text"], "\n".join(doc.lines()))
                self.assertEqual(verify[1]["normalized_text"].count("[KR_RRN]"), 3)                # 공백 뺀 재검사
                self.assertEqual(verify[0]["source"]["format"], kind)
                cells = [d for d in self.detection.detects if f"#f:0:{kind}:body:" in d["request_id"]]
                self.assertTrue(cells and all(isinstance(d["source"].get("block_order"), int) for d in cells))
        after = len(list(review_dir.glob("*"))) if review_dir.exists() else 0
        self.assertEqual(after, before)                                                 # 검토용 사본 안 만듦

    def test_partial_mask_policy(self):
        self.detection.rrn_action = "mask_partial"
        for kind, data in (("hwp", pii_hwp()), ("hwpx", pii_hwpx())):
            with self.subTest(kind):
                self.reports.clear()
                flow = self.send(data, f"인사.{kind}")
                self.assertIsNone(flow.response)
                lines = reread(self.forwarded(flow)).lines()
                self.assertIn("연락처 메모 900101-******* 끝", lines)
                self.assertEqual(self.reports[0].result.action, "mask_partial")

    def test_recheck_hit_inside_the_replaced_span_is_ignored_and_one_crossing_it_is_403(self):
        marked = re.compile(r"\[KR_RRN\] {0,6}")    # HWP 표시 + 채운 공백(공백 뺀 재검사에서는 표시만)

        def inside(text):       # 표시 일부 + 채운 공백: 바꾼 구간 안에 완전히 든 적중(표시 꼴이 아니라 구간 규칙으로 무시)
            return [(m.start() + 1, m.end(), "PASSWORD", "[PASSWORD]", "mask_full") for m in marked.finditer(text)]

        def crossing(text):     # 바꾼 구간을 넘는 적중(앞 두 글자 포함)
            return [(max(0, m.start() - 2), m.end(), "PASSWORD", "[PASSWORD]", "mask_full")
                    for m in marked.finditer(text)][:1]
        self.detection.verify_extra = inside
        flow = self.send(pii_hwp(), "a.hwp")
        self.assertIsNone(flow.response)
        self.assertEqual(sum(line.count(HWP_RRN) for line in reread(self.forwarded(flow)).lines()), 3)
        self.reports.clear()
        self.detection.verify_extra = crossing
        flow = self.send(pii_hwp(), "a.hwp")
        self.assert_refused(flow, 403, "mask_application_failed")

    def test_section_size_mismatch_is_403(self):
        # 구역 압축 크기를 못 맞춤(사용자 결정: 차단) → 403 + 이벤트 1건, 재검사 호출 없음, 원본·가린 파일 모두 안 나감
        with mock.patch.object(hwp_mask, "deflate_exact", lambda data, target: None):
            flow = self.send(pii_hwp(), "a.hwp")
        self.assert_refused(flow, 403, "mask_application_failed")
        self.assertEqual(self.verify_ids(), [])

    def test_recheck_residual_is_403(self):
        self.detection.verify_extra = lambda text: [(0, 2, "KR_NAME", "[KR_NAME]", "mask_full")]
        for kind, data in (("hwp", pii_hwp()), ("hwpx", pii_hwpx())):
            with self.subTest(kind):
                self.reports.clear()
                self.assert_refused(self.send(data, f"a.{kind}"), 403, "mask_application_failed")

    def test_side_text_needing_a_mask_is_403(self):
        for kind, data in (("hwp", hwp_doc(body=["평범한 메모"], summary={2: f"제목 {RRN}"})),
                           ("hwpx", hwpx_doc(body=["평범한 메모"], meta={"creator": RRN}))):
            with self.subTest(kind):
                self.reports.clear()
                report = self.assert_refused(self.send(data, f"a.{kind}"), 403, "mask_application_failed")
                self.assertEqual(report.result.pii_hits, "KR_RRN:1")
                self.assertEqual(self.verify_ids(), [])

    def test_value_only_the_whole_file_sees_is_403(self):
        for kind, data in (("hwp", hwp_doc(body=["생년월일 아래 칸", WHOLE_ONLY])),
                           ("hwpx", hwpx_doc(body=["생년월일 아래 칸", WHOLE_ONLY]))):
            with self.subTest(kind):
                self.reports.clear()
                self.assert_refused(self.send(data, f"a.{kind}"), 403, "mask_application_failed")

    def test_cell_label_value_is_masked_when_the_whole_file_passes(self):
        data = hwp_doc(body=["가짜 급여 양식"], rows=[("계좌번호", ACCOUNT)])
        flow = self.send(data, "a.hwp")
        self.assertIsNone(flow.response)
        lines = reread(self.forwarded(flow)).lines()
        self.assertNotIn(ACCOUNT, "\n".join(lines))
        self.assertIn("계좌번호 | " + "*" * 14, lines)       # '[KR_BANK_ACCOUNT]'(17자)는 14자에 안 들어감 → '*'
        report, = self.reports
        self.assertEqual((report.result.action, report.applied, report.result.pii_hits),
                         ("mask_full", True, "KR_BANK_ACCOUNT:1"))
        self.assertEqual(len(self.verify_ids()), 2)

    def test_cell_no_response_table(self):
        """명세 §4.5 네 줄."""
        whole_mask = hwp_doc(body=[f"번호 {RRN}", DOWN])
        allow_down = hwp_doc(body=["평범한 메모", DOWN])
        allow_down_changed = hwp_doc(body=["가짜 급여 양식", DOWN], rows=[("계좌번호", ACCOUNT)])
        cases = [("block", whole_mask, 503), ("forward", whole_mask, 503),          # ① 파일 전체 가림 + 칸 무응답
                 ("block", allow_down, 503),                                          # ② 통과 + 무응답 + 차단 설정
                 ("forward", allow_down, None),                                       # ③ 통과 + 무응답 + 원본 전송, 바꾼 칸 없음
                 ("forward", allow_down_changed, 503)]                                # ④ 통과 + 무응답 + 원본 전송, 바꾼 칸 있음
        for action, data, status in cases:
            with self.subTest(action=action, status=status, n=len(data)):
                self.reports.clear()
                self.use_settings(file=(action, 2.0))
                flow = self.send(data, "a.hwp")
                if status is None:
                    self.assertIsNone(flow.response)
                    self.assertEqual(self.forwarded(flow), data)
                    report, = self.reports
                    self.assertTrue(report.result.degraded)
                    self.assertEqual(report.result.failure_policy, "FAIL_OPEN")
                else:
                    self.assertEqual(flow.response.status_code, status)
                    self.assertEqual(self.reports[0].http_status, status)

    def test_recheck_over_the_char_limit_is_403(self):
        self.detection.rrn_token = "[" + "X" * 60 + "]"           # 표시가 값보다 긴 경우(HWPX는 글자가 늘어남)
        data = pii_hwpx()
        whole = len(inspect_upload(data, max_bytes=LIMIT, max_chars=80645, deadline=1e18).text)
        with mock.patch.object(gateway, "file_max_chars", lambda seconds: whole + 10):
            flow = self.send(data, "a.hwpx")
        self.assert_refused(flow, 403, "mask_application_failed")
        self.assertEqual(self.verify_ids(), [])


class ChatGptPutFlowTests(HwpMaskFlowMatrix, FlowBase):
    def send(self, content, filename):
        gateway.CHATGPT_FILE_META[FILE_ID] = filename
        req = tutils.treq(host="sdmntprkoreacentral.oaiusercontent.com", port=443, scheme=b"https", method=b"PUT",
                          path=f"/files/{FILE_ID}/raw?se=2026-10-11&sig=fake".encode(), content=content,
                          headers=((b"content-type", b"application/octet-stream"),))
        flow = tflow.tflow(req=req)
        run_hook(gateway.request(flow))
        return flow

    def forwarded(self, flow):
        return flow.request.raw_content

    def test_size_table_follows_the_masked_bytes(self):
        for kind, data in (("hwp", pii_hwp()), ("hwpx", pii_hwpx())):
            with self.subTest(kind):
                gateway.CHATGPT_SANITIZED_FILE_SIZE.clear()
                flow = self.send(data, f"a.{kind}")
                self.assertEqual(gateway.CHATGPT_SANITIZED_FILE_SIZE[FILE_ID], len(flow.request.raw_content))
                self.assertEqual(flow.request.headers["content-length"], str(len(flow.request.raw_content)))
                if kind == "hwp":
                    self.assertEqual(len(flow.request.raw_content), len(data))      # HWP는 파일 크기 그대로


class ChatGptMultipartFlowTests(HwpMaskFlowMatrix, FlowBase):
    boundary = "----gwhwpmaskboundary"

    def send(self, content, filename):
        req = tutils.treq(host="chatgpt.com", port=443, scheme=b"https", method=b"POST",
                          path=b"/backend-api/files/upload", content=multipart_body(
                              self.boundary, [(filename, "application/octet-stream", content)]),
                          headers=((b"content-type", f"multipart/form-data; boundary={self.boundary}".encode()),))
        flow = tflow.tflow(req=req)
        run_hook(gateway.request(flow))
        return flow

    def forwarded(self, flow):
        payload, = multipart_files(flow.request.raw_content, self.boundary)
        return payload


class ClaudeFlowTests(HwpMaskFlowMatrix, FlowBase):
    target_llm = "Claude"
    verify_suffix = "#file:0"
    boundary = "----claudehwpmaskboundary"

    def upload(self, *parts):
        req = tutils.treq(host="claude.ai", port=443, scheme=b"https", method=b"POST",
                          path=b"/api/00000000-0000-4000-8000-00000000c1a0/upload",
                          content=multipart_body(self.boundary, parts),
                          headers=((b"content-type", f"multipart/form-data; boundary={self.boundary}".encode()),))
        flow = tflow.tflow(req=req)
        run_hook(gateway.request(flow))
        return flow

    def send(self, content, filename):
        return self.upload((filename, "application/x-hwp", content))

    def forwarded(self, flow):
        payload, = multipart_files(flow.request.raw_content, self.boundary)
        return payload

    def test_hwp_hwpx_and_docx_in_one_upload_are_all_masked(self):
        document = Document()
        document.add_paragraph(f"주민등록번호 {RRN}")
        stream = io.BytesIO()
        document.save(stream)
        flow = self.upload(("a.hwp", "application/x-hwp", pii_hwp()), ("b.hwpx", "application/x-hwp", pii_hwpx()),
                           ("c.docx", DOCX, stream.getvalue()))
        self.assertIsNone(flow.response)
        hwp_out, hwpx_out, docx_out = multipart_files(flow.request.raw_content, self.boundary)
        for data in (hwp_out, hwpx_out):
            self.assertNotIn(RRN, "\n".join(reread(data).lines()))
        self.assertNotIn(RRN, "\n".join(p.text for p in Document(io.BytesIO(docx_out)).paragraphs))
        rid = self.whole_rid()
        self.assertEqual(self.verify_ids(), [f"{rid}#file:0#verify", f"{rid}#file:0#nows#verify",
                                             f"{rid}#file:1#verify", f"{rid}#file:1#nows#verify"])
        self.assertEqual([(r.result.action, r.applied) for r in self.reports], [("mask_full", True)])


if __name__ == "__main__":
    unittest.main()
