"""Claude 파일 경로 - ChatGPT와 같은 집행·대시보드 기록 (2026-10-10).

1. PDF·TXT처럼 칸별 가림을 못 하는 파일이 가림 판정 -> 403 mask_application_failed + 이벤트 (이전: 503, 이벤트 없음)
2. 처리 실패(칸 검사 무응답·가림 재현 실패·multipart 교체 실패·예상 못한 오류) -> 503 그대로 + 이벤트 (이전: 이벤트 없음)
3. 옛 /completion 본문의 attachments[].extracted_content·file_name도 칸마다 검사 (이전: prompt만 검사)
검토 반영(같은 날): 차단 요청에는 차단 이벤트 하나만, 요청 지표는 가장 강한 판정, 질문↔첨부 연결 번호,
Detection 무응답이면 남은 칸 재호출 없음, 실패 응답을 기록보다 먼저, 첨부 파일 이름 검사.
가린 파일 다시 검사(#verify)는 사용자 결정(10/10)으로 뺐다: Detection이 '[PASSWORD]' 같은 가림 표기를 다시 잡고,
옆 칸 이름표(생년월일 열 등)가 파일 전체에서는 잡히는데 칸에서는 안 잡혀 정상 가림 업로드도 503이 났다.
Gateway 처리 설정(같은 날): Detection 호출은 _inspect(질문·파일 규칙)를 지난다. 가짜 Detection도 넘겨받은
failure_policy를 지킨다(실제 detection_client.check와 같게). 무응답 넘기기 기대값은 새 규칙(받는 쪽 규칙)에 맞췄다.

가짜 Detection 응답과 실제 mitmproxy flow로 request() 훅 전체를 지난다. 실제 Detection·대시보드에 연결하지 않는다.
실행: ~/llm-gateway/.venv/bin/python -m unittest tests.test_claude_file_parity -v (Gateway 사본 폴더에서)
"""

from __future__ import annotations

import io
import json
import os
import re
import socket
import sys
import tempfile
import threading
import time
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

_TMP = tempfile.mkdtemp(prefix="gw-claude-parity-")
os.environ["SIGNATURE_REVIEW_DIR"] = str(Path(_TMP) / "review")
os.environ["DASHBOARD_OUTBOX_PATH"] = str(Path(_TMP) / "outbox.sqlite3")
os.environ.pop("DASHBOARD_EVENT_URL", None)
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document  # noqa: E402
from mitmproxy.test import tflow, tutils  # noqa: E402
from openpyxl import Workbook  # noqa: E402

import detection_client  # noqa: E402
import gateway  # noqa: E402
from detection_client import DetectionResult  # noqa: E402
from detection_client import RedactResult  # noqa: E402
try:
    from detection_client import failure_result  # noqa: E402
except ImportError:   # 운영본(이 묶음 전)으로 새 시험을 돌려 실패를 확인할 때. detection_client.check 안의 실패 판정과 같다
    def failure_result(reason, failure_policy, request_id):
        closed = failure_policy == "FAIL_CLOSED"
        return DetectionResult(
            blocked=closed, reason=f"inspection failed: {reason}; {failure_policy}", degraded=True,
            degraded_reason=reason, failure_policy=failure_policy, action="block" if closed else "allow_degraded",
            reason_code={"TIMEOUT": "detection_timeout", "PARSE_ERROR": "detection_parse_error",
                         "OTHER": "detection_error", "UNREACHABLE": "detection_unreachable"}[reason],
            request_id=request_id, http_status=503 if closed else None)
from tests.gw_support import FixedSettings, run_hook  # noqa: E402

RRN = "900101-1234567"              # 가상 주민등록번호
SECRET = "PROJECT-ORION-ROADMAP"    # 기밀 문서와 일치하는 표시(파일 전체·첨부 글자 검사에서 차단, 칸 검사에서는 안 잡힘)
CELL_BLOCK = "SEGMENT-ONLY-BLOCK"   # 칸 하나에서만 차단되는 표시
DOWN = "DETECTION-DOWN-MARKER"      # 칸 검사에서 Detection 무응답(잠정 통과 판정)을 흉내
CRASH = "CLIENT-CRASH-MARKER"       # 파일 전체 검사에서 Detection 호출이 예외를 올리는 경우
ORG = "00000000-0000-4000-8000-00000000c1a0"
CONV = "00000000-0000-4000-8000-00000000c0de"
UPLOAD_PATH = f"/api/{ORG}/upload"
COMPLETION_PATH = f"/api/organizations/{ORG}/chat_conversations/{CONV}/completion"
OLD_503_BODY = b'{"error":{"type":"file_masking_failed"}}'


def fake_detection(calls):
    def check(base_url, *, request_id, normalized_text, source=None, **kwargs):
        calls.append({"request_id": request_id, "text": normalized_text, "source": source or {}})
        is_segment = "#file:" in request_id   # DOCX/XLSX 칸 검사 번호: <요청>#file:<n>#<칸>
        if CRASH in normalized_text:
            raise ValueError("Invalid inspection timeout")
        if DOWN in normalized_text and is_segment:
            # 실제 detection_client.check처럼 넘겨받은 실패 정책을 따른다(파일 기본 = 차단 → FAIL_CLOSED)
            return unanswered(request_id, "TIMEOUT", kwargs.get("failure_policy", "FAIL_OPEN"))
        if SECRET in normalized_text and not is_segment:
            return DetectionResult(blocked=True, reason="doc match", request_id=request_id, decided_by="windowhash",
                                   doc_id="doc-fake-001", doc_span="0-21", input_span="0-21", match_rate=1.0,
                                   action="block", reason_code="doc_excerpt", http_status=403,
                                   user_message="보호 문서와 일치해 전송을 차단했습니다.")
        if CELL_BLOCK in normalized_text and is_segment:
            return DetectionResult(blocked=True, reason="signature", request_id=request_id, decided_by="signature",
                                   pii_hits="KR_RRN:1", action="block_notify", reason_code="pii", http_status=403)
        findings = [{"entity_type": "KR_RRN", "start": m.start(), "end": m.end(),
                     "replacement": "[KR_RRN]", "policy_action": "mask_full", "score": 1.0}
                    for m in re.finditer(re.escape(RRN), normalized_text)]
        if findings:
            return DetectionResult(blocked=False, reason="signature", request_id=request_id, decided_by="signature",
                                   findings=findings, masked_text=normalized_text.replace(RRN, "[KR_RRN]"),
                                   pii_hits=f"KR_RRN:{len(findings)}", action="mask_full", reason_code="pii")
        return DetectionResult(blocked=False, reason="", request_id=request_id, decided_by="none",
                               action="allow", reason_code="none")
    return check


def docx_bytes(*paragraphs):
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def pdf_bytes(text):
    """글자 한 줄짜리 PDF(표준 글꼴). pypdf가 글자를 뽑을 수 있는 최소 구조."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FILE_REF = r"^[A-Za-z0-9_-]{1,100}$"   # 대시보드 validation.mjs의 파일 번호 모양


def xlsx_bytes(title=None, cells=()):
    workbook = Workbook()
    sheet = workbook.active
    if title:
        sheet.title = title
    for ref, value in cells:
        sheet[ref] = value
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def unanswered(request_id, reason="TIMEOUT", policy="FAIL_OPEN"):
    """detection_client.check가 돌려주는 무응답 판정(실제 함수로 만든다)."""
    return failure_result(reason, policy, request_id)


def completion_body(prompt, *attachments):
    """attachments: (file_name, extracted_content)."""
    return {"prompt": prompt, "attachments": [
        {"file_name": name, "file_type": "txt", "file_size": len(text), "extracted_content": text}
        for name, text in attachments], "files": []}


class Base(unittest.TestCase):
    def setUp(self):
        self.calls, self.reports = [], []
        self._orig = (gateway.check_detection, gateway.report_dashboard_event, gateway.send_to_next_module)
        gateway.check_detection = fake_detection(self.calls)
        gateway.send_to_next_module = lambda event: None   # 옛 3000 포트 전송은 이 시험과 무관

        def capture(flow, result, dashboard_context, *, enforcement_applied, http_status):
            self.reports.append(SimpleNamespace(result=result, context=dashboard_context,
                                                applied=enforcement_applied, http_status=http_status))
        gateway.report_dashboard_event = capture

        # PDF 가림 API(2026-10-11): 이 묶음은 '지원하지 않음'을 돌려주는 가짜(→ 403 mask_application_failed, 재검사 없음).
        # 가림 성공·잔존·시간 초과 등은 tests/test_gateway_pdf_mask.py(가짜 Detection HTTP 서버)
        self.redact_calls = []
        self._orig_redact = gateway.redact_pdf_detection

        def redact_unsupported(base_url, prefix, data, source=None, timeout=2.0):
            self.redact_calls.append({"prefix": prefix, "bytes": len(data), "timeout": timeout})
            return RedactResult(status="unsupported", reason="fake_unsupported")
        gateway.redact_pdf_detection = redact_unsupported

    def tearDown(self):
        gateway.check_detection, gateway.report_dashboard_event, gateway.send_to_next_module = self._orig
        gateway.redact_pdf_detection = self._orig_redact

    def use_settings(self, **kwargs):
        """이 시험 동안 Gateway 처리 설정(gw_support.settings_doc 인자). 기본 = 질문 원본 전송 0.5초 · 파일 차단 2초."""
        patcher = mock.patch.object(gateway, "GATEWAY_SETTINGS", FixedSettings(**kwargs))
        patcher.start()
        self.addCleanup(patcher.stop)

    def detection_with(self, rule):
        """rule(request_id, text)가 판정을 돌려주면 그 판정, None이면 기본 가짜 판정. 호출은 모두 self.calls에 남는다.
        무응답 판정은 넘겨받은 failure_policy로 맞춘다(실제 detection_client.check와 같게)."""
        default = fake_detection(self.calls)

        def check(base_url, *, request_id, normalized_text, source=None, **kwargs):
            special = rule(request_id, normalized_text)
            if special is None:
                return default(base_url, request_id=request_id, normalized_text=normalized_text, source=source, **kwargs)
            self.calls.append({"request_id": request_id, "text": normalized_text, "source": source or {},
                               "timeout": kwargs.get("timeout")})
            policy = kwargs.get("failure_policy")
            if getattr(special, "degraded", False) and policy and special.failure_policy != policy:
                special = failure_result(special.degraded_reason, policy, special.request_id)
            return special
        gateway.check_detection = check

    def upload(self, *parts):
        """parts: (파일 이름, Content-Type, bytes). claude.ai multipart 업로드 하나."""
        boundary = "----claudeparityboundary"
        body = b""
        for filename, ctype, content in parts:
            body += (f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"{filename}\"\r\n"
                     f"Content-Type: {ctype}\r\n\r\n").encode() + content + b"\r\n"
        body += f"--{boundary}--\r\n".encode()
        req = tutils.treq(host="claude.ai", port=443, scheme=b"https", method=b"POST", path=UPLOAD_PATH.encode(),
                          content=body, headers=((b"content-type", f"multipart/form-data; boundary={boundary}".encode()),))
        flow = tflow.tflow(req=req)
        self.sent = body
        run_hook(gateway.request(flow))
        return flow

    def completion(self, body):
        content = json.dumps(body, ensure_ascii=False).encode()
        req = tutils.treq(host="claude.ai", port=443, scheme=b"https", method=b"POST", path=COMPLETION_PATH.encode(),
                          content=content, headers=((b"content-type", b"application/json"),))
        flow = tflow.tflow(req=req)
        self.sent = content
        run_hook(gateway.request(flow))
        return flow

    def assert_reports(self, *expected):
        """expected: (action, contentType, applied) 순서대로."""
        got = [(r.result.action, r.context["contentType"], r.applied) for r in self.reports]
        self.assertEqual(got, list(expected))
        for report in self.reports:
            self.assertEqual(report.context["targetLlm"], "Claude")
        return self.reports

    def assert_not_forwarded(self, flow, status):
        self.assertIsNotNone(flow.response)   # 응답을 Gateway가 만들었으므로 원본은 upstream으로 가지 않는다
        self.assertEqual(flow.response.status_code, status)


class UploadMaskTests(Base):
    """가림 판정이 났는데 가릴 수 없는 파일 -> 차단 + 이벤트. 나머지는 그대로."""

    def test_pdf_with_pii_is_blocked_with_an_event_when_the_redaction_is_unsupported(self):
        # 2026-10-11: 가림 판정 PDF는 Detection /redact/pdf로 간다. '지원하지 않음'이면 403 + 이벤트
        flow = self.upload(("a.pdf", "application/pdf", pdf_bytes(f"RRN {RRN}")))
        self.assert_not_forwarded(flow, 403)
        self.assertEqual(json.loads(flow.response.get_text())["error"]["reason_code"], "mask_application_failed")
        report = self.assert_reports(("block", "file", True))[0]
        self.assertEqual(report.result.reason_code, "mask_application_failed")
        self.assertEqual(report.http_status, 403)
        self.assertEqual(len(self.calls), 1)   # 파일 전체 검사 한 번(재검사 없음)
        self.assertEqual([c["prefix"] for c in self.redact_calls], [f"{self.calls[0]['request_id']}#file:0"])

    def test_txt_with_pii_is_blocked_with_an_event_instead_of_503(self):
        flow = self.upload(("a.txt", "text/plain", f"주민등록번호 {RRN}".encode()))
        self.assert_not_forwarded(flow, 403)
        self.assertEqual(self.assert_reports(("block", "file", True))[0].result.reason_code, "mask_application_failed")

    def test_clean_pdf_passes_unchanged_with_an_allow_event(self):
        flow = self.upload(("c.pdf", "application/pdf", pdf_bytes("meeting notes")))
        self.assertIsNone(flow.response)
        self.assertEqual(flow.request.raw_content, self.sent)
        self.assert_reports(("allow", "file", False))

    def test_pii_docx_is_masked_and_reported(self):
        flow = self.upload(("d.docx", DOCX, docx_bytes("안내", f"주민등록번호 {RRN}")))
        self.assertIsNone(flow.response)
        self.assertNotIn(RRN.encode(), flow.request.raw_content)
        self.assertEqual(flow.request.headers["content-length"], str(len(flow.request.raw_content)))
        report = self.assert_reports(("mask_full", "file", True))[0]
        self.assertNotIn("#", report.result.request_id)   # 이벤트는 파일 전체 판정(요청 번호 그대로)

    def test_confidential_document_is_blocked_before_segments(self):
        flow = self.upload(("s.docx", DOCX, docx_bytes(f"{SECRET} 3분기 계획", f"주민등록번호 {RRN}")))
        self.assert_not_forwarded(flow, 403)
        self.assertEqual(len(self.calls), 1)
        self.assert_reports(("block", "file", True))

    def test_a_blocked_cell_blocks_the_upload(self):
        flow = self.upload(("c.docx", DOCX, docx_bytes("보통 문단", f"{CELL_BLOCK} 내부 값")))
        self.assert_not_forwarded(flow, 403)
        self.assert_reports(("block_notify", "file", True))

    def test_docx_and_pdf_needing_a_mask_block_the_whole_request(self):
        # PDF 가림이 안 되면(지원하지 않음) DOCX만 가려 보내지 않는다. PDF를 먼저 처리하므로 칸 검사도 하지 않는다
        flow = self.upload(("d.docx", DOCX, docx_bytes(f"주민등록번호 {RRN}")),
                           ("a.pdf", "application/pdf", pdf_bytes(f"RRN {RRN}")))
        self.assert_not_forwarded(flow, 403)
        self.assertEqual(self.assert_reports(("block", "file", True))[0].result.reason_code, "mask_application_failed")
        self.assertEqual(len(self.calls), 1)
        self.assertEqual(len(self.redact_calls), 1)

    def test_clean_docx_and_clean_pdf_pass_together(self):
        flow = self.upload(("d.docx", DOCX, docx_bytes("회의록")), ("c.pdf", "application/pdf", pdf_bytes("agenda")))
        self.assertIsNone(flow.response)
        self.assertEqual(flow.request.raw_content, self.sent)
        self.assert_reports(("allow", "file", False))


class UploadFailureTests(Base):
    """처리 실패 -> 503 그대로(원본 안 보냄) + BLOCK 이벤트."""

    def assert_failure(self, flow, reason_code):
        self.assert_not_forwarded(flow, 503)
        self.assertEqual(flow.response.raw_content, OLD_503_BODY)   # 응답은 기존과 같다
        report = self.assert_reports(("block", "file", True))[0]
        self.assertEqual(report.http_status, 503)
        self.assertEqual(report.result.reason_code, reason_code)
        self.assertNotIn("#", report.result.request_id)
        return report

    def test_detection_failure_during_cell_check(self):
        flow = self.upload(("d.docx", DOCX, docx_bytes("안내", f"{DOWN} 문단")))
        report = self.assert_failure(flow, "detection_timeout")
        self.assertTrue(report.result.degraded)
        self.assertEqual(report.result.failure_policy, "FAIL_CLOSED")
        self.assertEqual(report.result.request_id, self.calls[0]["request_id"])   # 파일 전체 판정과 같은 번호

    def test_multipart_replacement_failure(self):
        def broken(*args, **kwargs):
            raise ValueError("file payload mismatch")
        with mock.patch.object(gateway, "replace_claude_files", broken):
            flow = self.upload(("d.docx", DOCX, docx_bytes(f"주민등록번호 {RRN}")))
        report = self.assert_failure(flow, "file_masking_failed")
        self.assertEqual((report.result.decided_by, report.result.pii_hits), ("signature", "KR_RRN:1"))

    def test_mask_not_reproduced_by_cells(self):
        def unchanged(filename, file_bytes, request_id, detect):
            return SimpleNamespace(sanitized_bytes=file_bytes)
        with mock.patch.object(gateway, "sanitize_uploaded_file", unchanged):
            flow = self.upload(("d.docx", DOCX, docx_bytes(f"주민등록번호 {RRN}")))
        self.assert_failure(flow, "file_masking_failed")

    def test_unexpected_exception_before_detection_result(self):
        flow = self.upload(("c.txt", "text/plain", f"{CRASH} memo".encode()))
        report = self.assert_failure(flow, "file_masking_failed")
        self.assertEqual(report.result.decided_by, "none")


class CompletionAttachmentTests(Base):
    """옛 /completion 본문의 attachments[].extracted_content - 항목마다 따로 검사."""

    def body(self, prompt, *attachments):
        return {"prompt": prompt, "attachments": [
            {"file_name": name, "file_type": "txt", "file_size": len(text), "extracted_content": text}
            for name, text in attachments], "files": []}

    def forwarded(self, flow):
        self.assertIsNone(flow.response)
        return json.loads(flow.request.get_text())

    def test_attachment_with_pii_is_masked_and_prompt_untouched(self):
        flow = self.completion(self.body("요약해줘", ("paste.txt", f"주민등록번호 {RRN} 연락처 확인")))
        sent = self.forwarded(flow)
        self.assertEqual(sent["prompt"], "요약해줘")
        self.assertNotIn(RRN, flow.request.get_text())
        self.assertEqual(sent["attachments"][0]["extracted_content"], "주민등록번호 [KR_RRN] 연락처 확인")
        self.assertEqual(sent["attachments"][0]["file_name"], "paste.txt")
        self.assertEqual(flow.request.headers["content-length"], str(len(flow.request.raw_content)))
        self.assert_reports(("allow", "prompt", False), ("mask_full", "file", True))
        prompt_call, attachment_call, name_call = self.calls
        self.assertEqual(prompt_call["text"], "요약해줘")             # 합치지 않는다
        self.assertEqual(attachment_call["request_id"], f"{prompt_call['request_id']}#attachment:0")
        self.assertEqual(attachment_call["source"], {"service": "Claude", "endpoint": COMPLETION_PATH,
                                                     "content_type": "file", "filename": "paste.txt"})
        self.assertEqual((name_call["text"], name_call["request_id"]),
                         ("paste.txt", f"{prompt_call['request_id']}#attachment:0:name"))

    def test_attachment_needing_a_block_blocks_the_whole_request(self):
        flow = self.completion(self.body("검토해줘", ("plan.txt", f"{SECRET} 일정표")))
        self.assert_not_forwarded(flow, 403)
        self.assert_reports(("block", "file", True))   # 집행되지 않은 질문 통과는 기록하지 않는다

    def test_empty_prompt_with_attachment_is_still_inspected(self):
        flow = self.completion(self.body("", ("", f"주민등록번호 {RRN}")))
        sent = self.forwarded(flow)
        self.assertNotIn(RRN, flow.request.get_text())
        self.assertEqual(sent["attachments"][0]["extracted_content"], "주민등록번호 [KR_RRN]")
        self.assertEqual(len(self.calls), 1)
        self.assertIsNone(self.calls[0]["source"]["filename"])
        self.assertRegex(self.calls[0]["request_id"], r"#attachment:0$")
        self.assert_reports(("mask_full", "file", True))

    def test_each_attachment_is_inspected_and_masked_separately(self):
        flow = self.completion(self.body("두 파일 비교", ("a.txt", "회의 메모"), ("b.txt", f"명단 {RRN}"),
                                         ("c.txt", "   ")))
        sent = self.forwarded(flow)
        self.assertEqual([c["text"] for c in self.calls],
                         ["두 파일 비교", "회의 메모", "a.txt", f"명단 {RRN}", "b.txt", "c.txt"])
        self.assertEqual(sent["attachments"][0]["extracted_content"], "회의 메모")
        self.assertEqual(sent["attachments"][1]["extracted_content"], "명단 [KR_RRN]")
        self.assertEqual(sent["attachments"][2]["extracted_content"], "   ")   # 빈 칸은 검사·변경하지 않음
        # 첨부 항목마다 이벤트 하나(글자·이름 칸 중 강한 판정). 세 번째는 이름만 있는 항목
        self.assert_reports(("allow", "prompt", False), ("allow", "file", False), ("mask_full", "file", True),
                            ("allow", "file", False))

    def test_prompt_and_attachment_are_both_masked(self):
        flow = self.completion(self.body(f"내 번호 {RRN}", ("p.txt", f"주민 {RRN}")))
        sent = self.forwarded(flow)
        self.assertNotIn(RRN, flow.request.get_text())
        self.assertEqual(sent["prompt"], "내 번호 [KR_RRN]")
        self.assertEqual(sent["attachments"][0]["extracted_content"], "주민 [KR_RRN]")
        self.assert_reports(("mask_full", "prompt", True), ("mask_full", "file", True))

    def test_attachment_mask_that_cannot_be_applied_blocks(self):
        def no_masked_text(*args, **kwargs):
            result = fake_detection([])(*args, **kwargs)
            if result.action == "mask_full" and "#attachment:" in kwargs["request_id"]:
                result.masked_text, result.mask_spans = None, []
            self.calls.append({"request_id": kwargs["request_id"], "text": kwargs["normalized_text"]})
            return result
        gateway.check_detection = no_masked_text
        flow = self.completion(self.body("요약", ("a.txt", f"번호 {RRN}")))
        self.assert_not_forwarded(flow, 403)
        self.assertEqual(self.assert_reports(("block", "file", True))[0].result.reason_code, "mask_application_failed")


class CompletionUnchangedTests(Base):
    """첨부 글자가 없는 요청은 이전과 같다 - Detection 1회(prompt), 이벤트 1건."""

    def test_plain_prompt_allow(self):
        body = {"prompt": "안녕", "files": []}
        flow = self.completion(body)
        self.assertIsNone(flow.response)
        self.assertEqual(flow.request.raw_content, self.sent)
        self.assertEqual(len(self.calls), 1)
        self.assertNotIn("#", self.calls[0]["request_id"])
        self.assertEqual(self.calls[0]["source"]["content_type"], "prompt")
        self.assert_reports(("allow", "prompt", False))

    def test_plain_prompt_mask(self):
        flow = self.completion({"prompt": f"번호 {RRN}", "attachments": []})
        self.assertIsNone(flow.response)
        self.assertEqual(json.loads(flow.request.get_text()), {"prompt": "번호 [KR_RRN]", "attachments": []})
        self.assertEqual(len(self.calls), 1)
        self.assert_reports(("mask_full", "prompt", True))

    def test_plain_prompt_block(self):
        flow = self.completion({"prompt": f"{SECRET} 요약", "attachments": [{"file_name": "x", "extracted_content": ""}]})
        self.assert_not_forwarded(flow, 403)
        self.assertEqual(len(self.calls), 1)
        self.assert_reports(("block", "prompt", True))

    def test_no_prompt_and_no_attachment_text_or_name_is_not_inspected(self):
        flow = self.completion({"prompt": "", "attachments": [{"file_name": "  "}, {"extracted_content": " "}, "x"],
                                "files": ["x"]})
        self.assertIsNone(flow.response)
        self.assertEqual(flow.request.raw_content, self.sent)
        self.assertEqual(self.calls, [])
        self.assertEqual(self.reports, [])


class BlockedRequestRecordTests(Base):
    """검토 1: 뒤 칸이 요청을 막으면 차단 이벤트 하나만 남는다. 앞 칸의 가림(가린 뒤 보냄)은 실제로 일어나지 않았다."""

    def test_later_attachment_block_drops_the_prompt_mask_and_earlier_attachment_mask(self):
        flow = self.completion(completion_body(f"my id {RRN}", ("a.txt", f"first {RRN}"), ("b.txt", f"{SECRET} plan")))
        self.assert_not_forwarded(flow, 403)
        report, = self.assert_reports(("block", "file", True))
        base = self.calls[0]["request_id"]
        self.assertEqual(report.result.request_id, f"{base}#attachment:1")
        self.assertEqual(report.context["fileId"], f"{base}-attachment-1")
        self.assertNotIn("attachedFileIds", report.context)
        self.assertEqual([c["text"] for c in self.calls], [f"my id {RRN}", f"first {RRN}", "a.txt", f"{SECRET} plan"])

    def test_clean_prompt_masked_attachment_then_block_records_only_the_block(self):
        flow = self.completion(completion_body("hi", ("a.txt", f"x {RRN}"), ("b.txt", f"{SECRET} plan")))
        self.assert_not_forwarded(flow, 403)
        self.assert_reports(("block", "file", True))

    def test_mask_that_cannot_be_applied_after_a_prompt_mask_records_only_the_block(self):
        def rule(request_id, text):
            if request_id.endswith("#attachment:0"):   # 가리라고 했지만 가린 글자·위치가 없다
                return DetectionResult(blocked=False, reason="signature", request_id=request_id, decided_by="signature",
                                       pii_hits="KR_RRN:1", action="mask_full", reason_code="pii")
            return None
        self.detection_with(rule)
        flow = self.completion(completion_body(f"번호 {RRN}", ("a.txt", f"명단 {RRN}")))
        self.assert_not_forwarded(flow, 403)
        report, = self.assert_reports(("block", "file", True))
        self.assertEqual(report.result.reason_code, "mask_application_failed")

    def test_attachment_text_equal_to_the_prompt_is_masked_in_both_places(self):
        text = f"번호 {RRN}"
        flow = self.completion(completion_body(text, ("a.txt", text)))
        self.assertIsNone(flow.response)
        sent = json.loads(flow.request.get_text())
        self.assertNotIn(RRN, flow.request.get_text())
        self.assertEqual((sent["prompt"], sent["attachments"][0]["extracted_content"]), ("번호 [KR_RRN]", "번호 [KR_RRN]"))
        # 첨부 칸은 원본 본문에서 읽어 검사한다(질문을 먼저 가린 본문에서 읽으면 이미 가려져 ALLOW로 남았다)
        self.assert_reports(("mask_full", "prompt", True), ("mask_full", "file", True))


class RequestMetricTests(Base):
    """검토 2: 요청 지표(action_counts·detector_counts·degraded_requests)는 칸 중 가장 강한 판정, degraded는 하나라도."""

    def deltas(self, body):
        metrics = gateway._gateway_metrics
        before = (dict(metrics["action_counts"]), metrics["degraded_requests"], dict(metrics["detector_counts"]))
        with mock.patch.object(gateway, "_emit_gateway_metrics", lambda: None):
            flow = self.completion(body)
            if flow.response is None:
                flow.response = tutils.tresp(status_code=200)
            gateway.response(flow)

        def changed(now, then):
            return {key: value - then.get(key, 0) for key, value in now.items() if value - then.get(key, 0)}
        return (changed(dict(metrics["action_counts"]), before[0]), metrics["degraded_requests"] - before[1],
                changed(dict(metrics["detector_counts"]), before[2]))

    def test_masked_prompt_with_a_clean_attachment_counts_as_mask(self):
        self.assertEqual(self.deltas(completion_body(f"rrn {RRN}", ("n.txt", "plain notes"))),
                         ({"MASK": 1}, 0, {"signature": 1}))

    def test_unanswered_prompt_check_with_a_clean_attachment_counts_as_degraded(self):
        self.detection_with(lambda request_id, text: unanswered(request_id) if "#" not in request_id else None)
        self.assertEqual(self.deltas(completion_body("hi", ("n.txt", "plain notes"))),
                         ({"ALLOW_DEGRADED": 1}, 1, {"none": 1}))

    def test_unanswered_attachment_check_after_a_masked_prompt_counts_both(self):
        self.use_settings(file=("forward", 2.0))   # 첨부 무응답을 원본 전송으로 처리하는 설정
        self.detection_with(lambda request_id, text: unanswered(request_id, "PARSE_ERROR")
                            if request_id.endswith("#attachment:0") else None)
        self.assertEqual(self.deltas(completion_body(f"rrn {RRN}", ("", "plain notes"))),
                         ({"MASK": 1}, 1, {"signature": 1}))

    def test_clean_prompt_with_a_masked_attachment_counts_as_mask(self):
        self.assertEqual(self.deltas(completion_body("hello", ("", f"id {RRN}"))), ({"MASK": 1}, 0, {"signature": 1}))

    def test_prompt_only_request_is_unchanged(self):
        self.assertEqual(self.deltas({"prompt": f"rrn {RRN}", "attachments": []}), ({"MASK": 1}, 0, {"signature": 1}))


class AttachmentLinkTests(Base):
    """검토 3: 같은 본문의 질문과 첨부를 번호로 잇는다. 질문 attachedFileIds = 파일 fileId들."""

    def test_prompt_and_attachment_events_carry_the_same_link_ids(self):
        flow = self.completion({"prompt": f"번호 {RRN}", "attachments": [
            {"extracted_content": "  "}, {"file_name": "a.txt", "extracted_content": "회의 메모"},
            {"file_name": "b.txt", "extracted_content": f"명단 {RRN}"}]})
        self.assertIsNone(flow.response)
        prompt, first, second = self.assert_reports(("mask_full", "prompt", True), ("allow", "file", False),
                                                    ("mask_full", "file", True))
        base = prompt.result.request_id
        self.assertEqual(prompt.context["attachedFileIds"], [f"{base}-attachment-1", f"{base}-attachment-2"])
        self.assertEqual([first.context["fileId"], second.context["fileId"]], prompt.context["attachedFileIds"])
        for ref in prompt.context["attachedFileIds"]:
            self.assertRegex(ref, FILE_REF)
            self.assertNotIn("#", ref)

    def test_attachment_without_prompt_still_has_a_file_id(self):
        self.completion(completion_body("", ("", f"번호 {RRN}")))
        report, = self.assert_reports(("mask_full", "file", True))
        self.assertRegex(report.context["fileId"], FILE_REF)
        self.assertTrue(report.context["fileId"].endswith("-attachment-0"))

    def test_prompt_only_request_has_no_link_ids(self):
        self.completion({"prompt": f"번호 {RRN}", "attachments": []})
        report, = self.assert_reports(("mask_full", "prompt", True))
        self.assertNotIn("attachedFileIds", report.context)

    def test_link_ids_are_left_out_when_the_request_id_is_not_a_file_reference(self):
        self.detection_with(lambda request_id, text: DetectionResult(
            blocked=False, reason="", request_id="odd id/1", decided_by="none", action="allow", reason_code="none")
            if "#" not in request_id else None)
        self.completion(completion_body("요약", ("", f"번호 {RRN}")))
        prompt, file = self.assert_reports(("allow", "prompt", False), ("mask_full", "file", True))
        self.assertNotIn("attachedFileIds", prompt.context)
        self.assertNotIn("fileId", file.context)


class DetectionDownTests(Base):
    """검토 4 + Gateway 처리 설정(10/10): 같은 요청 안의 Detection 무응답 넘기기(_inspect).
    연결 안 됨(UNREACHABLE)은 다시 부르지 않는다. 시간 초과(TIMEOUT)는 다음 호출의 대기 시간이 같거나 짧을 때만
    넘긴다. 넘길 때는 받는 쪽 규칙(첨부 = 파일 규칙)으로 판정을 새로 만든다. 이벤트는 칸마다."""

    def real_check_against_a_silent_socket(self, body):
        server = socket.socket()
        server.bind(("127.0.0.1", 0))
        server.listen(16)
        held = []

        def accept():
            while True:
                try:
                    held.append(server.accept()[0])   # 연결만 받고 답하지 않는다
                except OSError:
                    return
        threading.Thread(target=accept, daemon=True).start()
        calls = []

        def counted(*args, **kwargs):
            calls.append((kwargs["request_id"], kwargs["timeout"], kwargs["failure_policy"]))
            return detection_client.check(*args, **kwargs)
        gateway.check_detection = counted
        started = time.monotonic()
        try:
            with mock.patch.object(gateway, "DETECTION_BASE_URL", f"http://127.0.0.1:{server.getsockname()[1]}"):
                flow = self.completion(body)
        finally:
            server.close()
            for conn in held:
                conn.close()
        return flow, calls, time.monotonic() - started

    def test_prompt_timeout_is_retried_for_attachments_with_the_longer_file_wait(self):
        # 기본: 질문 원본 전송 0.5초 · 파일 차단 2초. 질문 0.5초 시간 초과는 2초 대기인 첨부에 넘기지 않는다
        self.detection_with(lambda request_id, text: unanswered(request_id) if "#" not in request_id else None)
        flow = self.completion(completion_body("요약", ("a.txt", "메모 1"), ("b.txt", f"명단 {RRN}")))
        self.assertIsNone(flow.response)
        self.assertEqual(len(self.calls), 5)   # 질문 + 첨부 2개 × (글자·이름)
        self.assertNotIn(RRN, flow.request.get_text())
        self.assert_reports(("allow_degraded", "prompt", False), ("allow", "file", False), ("mask_full", "file", True))

    def test_prompt_timeout_is_not_repeated_when_the_file_wait_is_not_longer(self):
        self.use_settings(file=("forward", 0.5))
        self.detection_with(lambda request_id, text: unanswered(request_id) if "#" not in request_id else None)
        flow = self.completion(completion_body("요약", ("a.txt", "메모 1"), ("b.txt", f"명단 {RRN}"), ("c.txt", "메모 3")))
        self.assertIsNone(flow.response)
        self.assertEqual(flow.request.raw_content, self.sent)   # 파일 처리 설정(원본 전송) 그대로 통과
        self.assertEqual(len(self.calls), 1)
        reports = self.assert_reports(*[("allow_degraded", "prompt", False)] + [("allow_degraded", "file", False)] * 3)
        base = self.calls[0]["request_id"]
        self.assertEqual([r.result.request_id for r in reports],
                         [base] + [f"{base}#attachment:{i}" for i in range(3)])
        self.assertTrue(all(r.result.degraded_reason == "TIMEOUT" for r in reports))

    def test_unreachable_prompt_check_is_carried_to_attachments_under_the_file_rule(self):
        # 질문은 원본 전송이지만 첨부는 파일 규칙(차단): 호출 없이 FAIL_CLOSED → 503 + 첨부 차단 이벤트 1건
        self.detection_with(lambda request_id, text: unanswered(request_id, "UNREACHABLE") if "#" not in request_id else None)
        flow = self.completion(completion_body("요약", ("a.txt", "메모 1"), ("b.txt", "메모 2")))
        self.assertEqual(len(self.calls), 1)
        self.assert_not_forwarded(flow, 503)
        report, = self.assert_reports(("block", "file", True))
        self.assertEqual((report.result.failure_policy, report.result.reason_code, report.http_status),
                         ("FAIL_CLOSED", "detection_unreachable", 503))
        self.assertEqual(report.result.request_id, f"{self.calls[0]['request_id']}#attachment:0")

    def test_unanswered_attachment_check_skips_the_remaining_parts(self):
        self.use_settings(file=("forward", 2.0))
        self.detection_with(lambda request_id, text: unanswered(request_id, "UNREACHABLE")
                            if request_id.endswith("#attachment:0") else None)
        flow = self.completion(completion_body("요약", ("a.txt", "메모 1"), ("b.txt", f"명단 {RRN}")))
        self.assertEqual(len(self.calls), 2)   # 질문 + 첫 첨부 글자
        self.assertEqual(flow.request.raw_content, self.sent)
        self.assert_reports(("allow", "prompt", False), ("allow_degraded", "file", False),
                            ("allow_degraded", "file", False))

    def test_unanswered_attachment_check_under_the_block_rule_blocks(self):
        self.detection_with(lambda request_id, text: unanswered(request_id, "UNREACHABLE")
                            if request_id.endswith("#attachment:0") else None)
        flow = self.completion(completion_body("요약", ("a.txt", "메모 1"), ("b.txt", f"명단 {RRN}")))
        self.assertEqual(len(self.calls), 2)
        self.assert_not_forwarded(flow, 503)
        report, = self.assert_reports(("block", "file", True))
        self.assertEqual(report.result.failure_policy, "FAIL_CLOSED")

    def test_other_inspection_failures_are_not_skipped(self):
        self.use_settings(file=("forward", 2.0))
        self.detection_with(lambda request_id, text: unanswered(request_id, "PARSE_ERROR")
                            if request_id.endswith("#attachment:0") else None)
        flow = self.completion(completion_body("요약", ("a.txt", "메모 1"), ("b.txt", f"명단 {RRN}")))
        self.assertEqual(len(self.calls), 5)
        self.assertNotIn(RRN, flow.request.get_text())

    def test_real_client_waits_for_one_timeout_only(self):
        self.use_settings(file=("forward", 0.5))
        flow, calls, _ = self.real_check_against_a_silent_socket(
            completion_body("요약", ("a.txt", "메모 1"), ("b.txt", "메모 2"), ("c.txt", "메모 3")))
        self.assertEqual(len(calls), 1)
        self.assertIsNone(flow.response)
        reports = self.assert_reports(*[("allow_degraded", "prompt", False)] + [("allow_degraded", "file", False)] * 3)
        self.assertTrue(all(r.result.degraded_reason == "TIMEOUT" for r in reports))

    def test_real_client_retries_the_attachment_with_the_file_wait_and_blocks(self):
        # 명세 §11.1 무응답 넘기기: 질문 forward 0.5초 시간 초과 → 첨부는 2초로 다시 부르고 시간 초과 → 파일 차단 503
        # (이전 work 방식은 질문 판정을 복사해 첨부가 검사 없이 나갔다)
        flow, calls, elapsed = self.real_check_against_a_silent_socket(
            completion_body("요약", ("a.txt", "메모 1"), ("b.txt", "메모 2")))
        self.assertEqual([(timeout, policy) for _, timeout, policy in calls],
                         [(0.5, "FAIL_OPEN"), (2.0, "FAIL_CLOSED")])
        self.assertGreater(elapsed, 2.4)
        self.assertLess(elapsed, 4.5)        # 남은 칸은 다시 기다리지 않는다
        self.assert_not_forwarded(flow, 503)
        report, = self.assert_reports(("block", "file", True))
        self.assertEqual((report.result.degraded_reason, report.result.failure_policy), ("TIMEOUT", "FAIL_CLOSED"))


class FailClosedRecordTests(Base):
    """검토 5: 실패 응답(503)을 먼저 정한다. 기록 단계에서 예외가 나도 원본은 나가지 않는다."""

    def test_unreadable_entity_count_still_blocks_the_upload(self):
        def rule(request_id, text):
            if "#" not in request_id:   # 파일 전체: 가림. 건수 칸이 isdigit()은 통과하지만 int()는 못 읽는 글자
                return DetectionResult(blocked=False, reason="signature", request_id=request_id, decided_by="signature",
                                       pii_hits="KR_RRN:²", action="mask_full", reason_code="pii")
            return DetectionResult(blocked=False, reason="", request_id=request_id, decided_by="none",
                                   action="allow", reason_code="none")   # 칸은 깨끗 -> 가림 재현 실패
        self.detection_with(rule)
        flow = self.upload(("d.docx", DOCX, docx_bytes("안내", f"주민등록번호 {RRN}")))
        self.assert_not_forwarded(flow, 503)
        self.assertEqual(flow.response.raw_content, OLD_503_BODY)
        report, = self.assert_reports(("block", "file", True))
        self.assertEqual((report.result.reason_code, report.result.pii_hits), ("file_masking_failed", None))

    def test_failure_while_recording_never_forwards_the_upload(self):
        with mock.patch.object(gateway, "_file_failure_result", mock.Mock(side_effect=RuntimeError("x"))):
            flow = self.upload(("c.txt", "text/plain", f"{CRASH} memo".encode()))
        self.assert_not_forwarded(flow, 503)
        self.assertEqual(flow.response.raw_content, OLD_503_BODY)
        self.assertEqual(self.reports, [])
        self.assertEqual(flow.metadata["gw_action"], "block")

    def test_merged_counts_read_ascii_digits_only(self):
        self.assertEqual(gateway._merged_pii_hits([SimpleNamespace(pii_hits="KR_RRN:²,EMAIL_ADDRESS:2")]),
                         "EMAIL_ADDRESS:2")


class MaskedUploadForwardTests(Base):
    """사용자 결정(10/10): 가린 파일 다시 검사(#verify)를 뺐다. 칸 가림이 재현되면 이전처럼 바로 전달 + MASK 기록."""

    def test_masked_docx_is_forwarded_without_a_whole_file_recheck(self):
        flow = self.upload(("d.docx", DOCX, docx_bytes("안내", f"주민등록번호 {RRN}")))
        self.assertIsNone(flow.response)
        self.assertNotIn(RRN.encode(), flow.request.raw_content)
        self.assertFalse([c for c in self.calls if "#verify" in c["request_id"]])
        self.assertTrue(all("#file:" in c["request_id"] for c in self.calls[1:]))   # 파일 전체 1회 + 칸 검사만
        self.assert_reports(("mask_full", "file", True))


class AttachmentFileNameTests(Base):
    """검토 7: attachments[].file_name도 Claude에 그대로 가므로 따로 검사한다(요청 번호 끝 ':name'). 이벤트는 항목마다 하나."""

    def test_pii_in_a_file_name_is_masked(self):
        flow = self.completion(completion_body("요약해줘", (f"hong_{RRN}.txt", "plain memo")))
        self.assertIsNone(flow.response)
        sent = json.loads(flow.request.get_text())
        self.assertNotIn(RRN, flow.request.get_text())
        self.assertEqual(sent["attachments"][0]["file_name"], "hong_[KR_RRN].txt")
        self.assertEqual(sent["attachments"][0]["extracted_content"], "plain memo")
        _, file = self.assert_reports(("allow", "prompt", False), ("mask_full", "file", True))
        self.assertTrue(file.result.request_id.endswith("#attachment:0:name"))

    def test_confidential_marker_in_a_file_name_blocks(self):
        flow = self.completion(completion_body("요약해줘", (f"{SECRET}.txt", "plain memo")))
        self.assert_not_forwarded(flow, 403)
        self.assert_reports(("block", "file", True))

    def test_attachment_with_only_a_name_is_inspected(self):
        flow = self.completion({"prompt": "", "attachments": [{"file_name": f"{RRN}.txt"}]})
        self.assertIsNone(flow.response)
        self.assertEqual(json.loads(flow.request.get_text())["attachments"][0]["file_name"], "[KR_RRN].txt")
        self.assert_reports(("mask_full", "file", True))

    def test_text_and_name_both_masked_give_one_event_with_summed_counts(self):
        flow = self.completion(completion_body("요약", (f"{RRN}.txt", f"번호 {RRN}")))
        self.assertNotIn(RRN, flow.request.get_text())
        _, file = self.assert_reports(("allow", "prompt", False), ("mask_full", "file", True))
        self.assertEqual(file.result.pii_hits, "KR_RRN:2")
        self.assertTrue(file.result.request_id.endswith("#attachment:0"))


# 대시보드 Security Event v1 검사(local-dashboard/server/validation.mjs)의 허용 필드·조건 일부를 옮긴 것
_EVENT_FIELDS = {
    "type", "eventId", "requestId", "timestamp", "clientId", "userId", "tenantId", "targetLlm", "contentType", "host",
    "endpoint", "detectionOutcome", "recommendedAction", "finalAction", "severity", "enforcement", "maskEvidence",
    "processingMs", "degraded", "degradedReason", "failurePolicy", "detectorVersion", "rulesetVersion", "modelVersion",
    "reasonCode", "policyVersion", "detections", "fileId", "attachedFileIds", "viewerKeyHash",
}


class DashboardEventShapeTests(Base):
    """실제 report_dashboard_event·build_dashboard_event로 만든 이벤트가 대시보드 검사를 통과하는 모양인지."""

    def setUp(self):
        super().setUp()
        gateway.report_dashboard_event = self._orig[1]
        os.environ["DASHBOARD_EVENT_URL"] = "http://127.0.0.1:9/ingest"
        os.environ["DASHBOARD_INGEST_TOKEN"] = "test-token-not-used"
        os.environ["DASHBOARD_REPORT_ALLOW_EVENTS"] = "true"

    def tearDown(self):
        for key in ("DASHBOARD_EVENT_URL", "DASHBOARD_INGEST_TOKEN", "DASHBOARD_REPORT_ALLOW_EVENTS"):
            os.environ.pop(key, None)
        super().tearDown()

    def events(self, flow):
        events = flow.metadata.get("gw_dashboard_events", [])
        dump = os.environ.get("CLAUDE_PARITY_EVENT_DUMP")   # node로 실제 validation.mjs 검사를 돌릴 때만
        if dump:
            with open(dump, "a", encoding="utf-8") as handle:
                for event in events:
                    handle.write(json.dumps({"case": self.id(), "event": event}, ensure_ascii=False) + "\n")
        for event in events:
            self.assertLessEqual(set(event), _EVENT_FIELDS)
            self.assertEqual(event["targetLlm"], "Claude")
            self.assertIn(event["contentType"], {"prompt", "file"})
            self.assertIn(event["finalAction"], {"ALLOW", "BLOCK", "MASK"})
            self.assertIn(event["detectionOutcome"], {"MATCH", "NO_MATCH", "UNAVAILABLE"})
            if event["detectionOutcome"] == "MATCH":
                self.assertTrue(event["detections"])
            if "reasonCode" in event:
                self.assertTrue(0 < len(event["reasonCode"]) <= 100)
            if "maskEvidence" in event:
                self.assertEqual((event["finalAction"], event["contentType"]), ("MASK", "prompt"))
            if "degradedReason" in event or "failurePolicy" in event:
                self.assertIs(event["degraded"], True)
            status = event.get("enforcement", {}).get("httpStatus")
            if status is not None:
                self.assertTrue(isinstance(status, int) and 100 <= status <= 599)
        return events

    def test_pdf_mask_block_event(self):
        event, = self.events(self.upload(("a.pdf", "application/pdf", pdf_bytes(f"RRN {RRN}"))))
        self.assertEqual((event["finalAction"], event["reasonCode"], event["enforcement"]),
                         ("BLOCK", "mask_application_failed", {"applied": True, "httpStatus": 403}))
        self.assertEqual(event["detections"][0]["entityCounts"], [{"entityType": "KR_RRN", "count": 1}])

    def test_cell_detection_failure_event(self):
        event, = self.events(self.upload(("d.docx", DOCX, docx_bytes(f"{DOWN} 문단"))))
        self.assertEqual((event["finalAction"], event["reasonCode"], event["enforcement"]["httpStatus"]),
                         ("BLOCK", "detection_timeout", 503))
        self.assertEqual((event["detectionOutcome"], event["degradedReason"], event["failurePolicy"]),
                         ("UNAVAILABLE", "TIMEOUT", "FAIL_CLOSED"))

    def test_replacement_failure_event_keeps_the_pii_counts(self):
        with mock.patch.object(gateway, "replace_claude_files", mock.Mock(side_effect=ValueError("x"))):
            event, = self.events(self.upload(("d.docx", DOCX, docx_bytes(f"주민등록번호 {RRN}"))))
        self.assertEqual((event["finalAction"], event["reasonCode"], event["enforcement"]["httpStatus"]),
                         ("BLOCK", "file_masking_failed", 503))
        self.assertEqual(event["detectionOutcome"], "MATCH")
        self.assertEqual(event["detections"][0]["entityCounts"], [{"entityType": "KR_RRN", "count": 1}])

    def test_unexpected_exception_event(self):
        event, = self.events(self.upload(("c.txt", "text/plain", f"{CRASH}".encode())))
        self.assertEqual((event["finalAction"], event["reasonCode"], event["detectionOutcome"], event["degraded"]),
                         ("BLOCK", "file_masking_failed", "NO_MATCH", False))

    def test_attachment_mask_event_has_no_mask_evidence(self):
        prompt_event, file_event = self.events(self.completion(
            {"prompt": f"번호 {RRN}", "attachments": [{"file_name": "a.txt", "extracted_content": f"명단 {RRN}"}]}))
        self.assertEqual((prompt_event["contentType"], prompt_event["finalAction"]), ("prompt", "MASK"))
        self.assertIn("maskEvidence", prompt_event)                 # 질문 가림 근거는 이전과 같이 남는다
        self.assertEqual((file_event["contentType"], file_event["finalAction"]), ("file", "MASK"))
        self.assertNotIn("maskEvidence", file_event)
        self.assertTrue(file_event["requestId"].startswith(prompt_event["requestId"] + "#attachment:"))

    def test_attachment_block_event(self):
        file_event, = self.events(self.completion(
            {"prompt": "요약", "attachments": [{"file_name": "a.txt", "extracted_content": f"{SECRET} 일정"}]}))
        self.assertEqual((file_event["finalAction"], file_event["detectionOutcome"], file_event["enforcement"]),
                         ("BLOCK", "MATCH", {"applied": True, "httpStatus": 403}))
        self.assertRegex(file_event["fileId"], r"^[A-Za-z0-9_-]{1,100}-attachment-0$")

    def test_later_block_leaves_only_the_block_event(self):
        event, = self.events(self.completion(completion_body(
            f"my id {RRN}", ("a.txt", f"first {RRN}"), ("b.txt", f"{SECRET} plan"))))
        self.assertEqual((event["contentType"], event["finalAction"], event["enforcement"]),
                         ("file", "BLOCK", {"applied": True, "httpStatus": 403}))
        self.assertNotIn("maskEvidence", event)

    def test_prompt_and_attachment_events_are_linked(self):
        prompt_event, file_event = self.events(self.completion(completion_body(f"번호 {RRN}", ("a.txt", f"명단 {RRN}"))))
        self.assertEqual(prompt_event["attachedFileIds"], [file_event["fileId"]])
        self.assertRegex(file_event["fileId"], FILE_REF)
        self.assertEqual(file_event["fileId"], f"{prompt_event['requestId']}-attachment-0")
        self.assertIn("maskEvidence", prompt_event)
        self.assertNotIn("attachedFileIds", file_event)

    def test_unanswered_prompt_check_gives_one_unavailable_event_per_part(self):
        self.use_settings(file=("forward", 0.5))   # 첨부도 같은 대기 시간·원본 전송 → 다시 부르지 않음
        self.detection_with(lambda request_id, text: unanswered(request_id) if "#" not in request_id else None)
        prompt_event, file_event = self.events(self.completion(completion_body("요약", ("a.txt", "메모"))))
        for event in (prompt_event, file_event):
            self.assertEqual((event["finalAction"], event["detectionOutcome"], event["degradedReason"],
                              event["failurePolicy"]), ("ALLOW", "UNAVAILABLE", "TIMEOUT", "FAIL_OPEN"))
        self.assertEqual(file_event["requestId"], f"{prompt_event['requestId']}#attachment:0")

    def test_file_name_mask_event(self):
        _, file_event = self.events(self.completion(completion_body("요약", (f"hong_{RRN}.txt", "plain memo"))))
        self.assertEqual((file_event["contentType"], file_event["finalAction"]), ("file", "MASK"))
        self.assertEqual(file_event["detections"][0]["entityCounts"], [{"entityType": "KR_RRN", "count": 1}])
        self.assertNotIn("maskEvidence", file_event)


if __name__ == "__main__":
    unittest.main()
