"""PDF 가림 2단계 Gateway 쪽 (2026-10-11, masking-design/SPEC.md §3.5·§3.6).

가림 판정 PDF → Detection POST /redact/pdf → 돌려받은 PDF를 같은 추출기로 다시 뽑아 재검사(글자 그대로 + 공백 뺀 글자)
→ 본문 교체(ChatGPT 원본 PUT·multipart, Claude multipart) + MASK 기록. 실패는 403 mask_application_failed 또는 503,
원본은 절대 보내지 않는다.

가짜 Detection은 127.0.0.1 임시 포트 HTTP 서버로 /detect·/redact/pdf 계약(명세 §3.2)을 흉내 낸다 - 실제
detection_client.check·redact_pdf가 그 서버를 부른다. 요청은 실제 mitmproxy flow로 request() 훅 전체를 지난다(작업 스레드 포함).
가짜 값만 쓴다. 실제 Detection·대시보드에 연결하지 않는다.
실행: python -m unittest tests.test_gateway_pdf_mask -v (Gateway 사본 폴더에서)
"""

from __future__ import annotations

import io
import json
import os
import re
import sys
import tempfile
import threading
import time
import unittest
from email.parser import BytesParser
from email.policy import default as email_default
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from types import SimpleNamespace
from unittest import mock
from urllib.parse import unquote

_TMP = tempfile.mkdtemp(prefix="gw-pdf-mask-")
os.environ["SIGNATURE_REVIEW_DIR"] = str(Path(_TMP) / "review")
os.environ["DASHBOARD_OUTBOX_PATH"] = str(Path(_TMP) / "outbox.sqlite3")
os.environ.pop("DASHBOARD_EVENT_URL", None)
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document  # noqa: E402
from mitmproxy.test import tflow, tutils  # noqa: E402

import detection_client  # noqa: E402
import gateway  # noqa: E402
from tests.gw_support import FixedSettings, pdf_from, run_hook, scanned_pdf, text_pdf  # noqa: E402

RRN = "900101-1234567"                   # 가상 주민등록번호
RRN_RE = re.compile(r"\d{6}-\d{7}")
SECRET = "PROJECT-ORION-ROADMAP"         # 파일 전체 검사에서 보호 문서 일치(가짜)
FILE_ID = "00000000-aaaa-8206-bbbb-000000000077"
CHATGPT_BOUNDARY = "----gwpdfmaskboundary"
CLAUDE_BOUNDARY = "----claudepdfmaskboundary"
CLAUDE_UPLOAD = "/api/00000000-0000-4000-8000-00000000c1a0/upload"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PREFIX_RE = re.compile(r"^[A-Za-z0-9#:_-]{1,150}$")
OLD_503_BODY = b'{"error":{"type":"file_masking_failed"}}'


def pii_pdf():
    return text_pdf(f"RRN {RRN} end")


def masked_pdf():
    return text_pdf("RRN removed end")   # Detection이 값을 지운 결과(검은 상자 자리, 글자 없음)


def docx_bytes(*paragraphs):
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


class FakeDetection:
    """/detect: 가짜 주민번호 → mask_full finding, SECRET → 차단(파일 전체만). /redact/pdf: mode에 따라 계약대로 답한다."""

    def __init__(self):
        self.mode = "masked"
        self.masked = masked_pdf()
        self.delay = 0.0
        self.verify_status = 200
        self.detects, self.redacts, self.contract_errors = [], [], []
        outer = self

        class Handler(BaseHTTPRequestHandler):
            protocol_version = "HTTP/1.1"

            def log_message(self, *args):
                pass

            def reply(self, status, body, content_type="application/json", headers=None):
                self.close_connection = True   # 요청 하나에 연결 하나(Gateway가 먼저 끊어도 다음 줄을 읽지 않게)
                try:
                    self.send_response(status)
                    self.send_header("Connection", "close")
                    self.send_header("Content-Type", content_type)
                    self.send_header("Content-Length", str(len(body)))
                    for name, value in (headers or {}).items():
                        self.send_header(name, value)
                    self.end_headers()
                    self.wfile.write(body)
                except OSError:
                    pass   # Gateway가 시한으로 먼저 끊음

            def do_POST(self):
                body = self.rfile.read(int(self.headers.get("Content-Length") or 0))
                if self.path == "/detect":
                    outer.detect(self, body)
                elif self.path == "/redact/pdf":
                    outer.redact(self, body)
                else:
                    self.reply(404, b'{"detail":"Not Found"}')

        self.server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        self.server.daemon_threads = True
        self.url = f"http://127.0.0.1:{self.server.server_address[1]}"
        self.thread = threading.Thread(target=self.server.serve_forever, kwargs={"poll_interval": 0.05}, daemon=True)
        self.thread.start()

    def close(self):
        self.server.shutdown()
        self.server.server_close()

    def detect(self, handler, body):
        request = json.loads(body)
        rid, text = request["request_id"], request["normalized_text"]
        self.detects.append(request)
        if "#verify" in rid and self.verify_status != 200:
            handler.reply(self.verify_status, b'{"detail":"boom"}')
            return
        if SECRET in text and "#" not in rid:
            reply = {"blocked": True, "reason": "doc match", "request_id": rid, "decided_by": "windowhash",
                     "hash_status": "ok", "signature_status": "ok", "findings": [],
                     "decision": {"action": "block", "reason_code": "doc_excerpt", "http_status": 403,
                                  "user_message": "보호 문서와 일치해 전송을 차단했습니다."}}
        else:
            findings = [{"entity_type": "KR_RRN", "start": m.start(), "end": m.end(), "replacement": "[KR_RRN]",
                         "policy_action": "mask_full", "score": 1.0} for m in RRN_RE.finditer(text)]
            masked = RRN_RE.sub("[KR_RRN]", text) if findings else None
            reply = {"blocked": False, "reason": "signature" if findings else "", "request_id": rid,
                     "decided_by": "signature" if findings else "none",
                     "pii_hits": f"KR_RRN:{len(findings)}" if findings else None,
                     "hash_status": "ok", "signature_status": "ok", "findings": findings, "masked_text": masked,
                     "decision": {"action": "mask_full" if findings else "allow",
                                  "reason_code": "pii" if findings else "none"}}
        handler.reply(200, json.dumps(reply).encode())

    def redact(self, handler, body):
        headers = {k.lower(): v for k, v in handler.headers.items()}
        self.redacts.append({"headers": headers, "body": body})
        # 계약(명세 §3.2) 확인: 어긋나면 실제 Detection처럼 4xx
        deadline = headers.get("x-deadline-ms", "")
        if not PREFIX_RE.fullmatch(headers.get("x-request-prefix", "")):
            self.contract_errors.append("prefix")
            return handler.reply(400, b'{"detail":"bad_request_prefix"}')
        if not deadline.isdigit() or not 100 <= int(deadline) <= 30000:
            self.contract_errors.append("deadline")
            return handler.reply(400, b'{"detail":"bad_deadline"}')
        if headers.get("content-type") != "application/pdf" or not body.startswith(b"%PDF-"):
            self.contract_errors.append("content")
            return handler.reply(415, b'{"detail":"content_type_not_pdf"}')
        if self.delay:
            time.sleep(self.delay)
        prefix = headers["x-request-prefix"]
        if self.mode == "masked":
            return handler.reply(200, self.masked, "application/pdf", {
                "X-Redact-Status": "masked", "X-Redact-Pages": "1", "X-Redact-Text-Pages": "1",
                "X-Redact-Masked-Pages": "1", "X-Redact-Pii-Hits": "KR_RRN:1", "X-Redact-Elapsed-Ms": "3.0"})
        if self.mode == "http404":
            return handler.reply(404, b'{"detail":"Not Found"}')
        decision = None
        if self.mode == "blocked":
            decision = {"blocked": True, "reason": "doc match", "request_id": f"{prefix}#f:0:pdf:body:p0",
                        "decided_by": "windowhash", "doc_id": "doc-fake-001", "match_rate": 1.0, "findings": [],
                        "decision": {"action": "block", "reason_code": "doc_excerpt", "http_status": 403,
                                     "user_message": "보호 문서와 일치해 전송을 차단했습니다."}}
        reply = {"status": self.mode, "pages": 1, "text_pages": 1, "masked_pages": 0, "pii_hits": "KR_RRN:1",
                 "residual": "KR_RRN:1" if self.mode == "residual" else None, "decision": decision,
                 "detail": f"fake_{self.mode}", "elapsed_ms": 1.0}
        handler.reply(200, json.dumps(reply).encode(), "application/json", {"X-Redact-Status": self.mode})


class Base(unittest.TestCase):
    """세 경로 공통 준비: 가짜 Detection 서버, 기록 가로채기, 설정 고정."""

    def setUp(self):
        self.detection = FakeDetection()
        self.addCleanup(self.detection.close)
        self.reports = []

        def capture(flow, result, dashboard_context, *, enforcement_applied, http_status):
            self.reports.append(SimpleNamespace(result=result, context=dashboard_context,
                                                applied=enforcement_applied, http_status=http_status))
        self._patch(gateway, "DETECTION_BASE_URL", self.detection.url)
        self._patch(gateway, "check_detection", detection_client.check)
        self._patch(gateway, "redact_pdf_detection", detection_client.redact_pdf)
        self._patch(gateway, "report_dashboard_event", capture)
        self._patch(gateway, "send_to_next_module", lambda event: None)
        self.use_settings()
        gateway.CHATGPT_FILE_META.clear()
        gateway.CHATGPT_SANITIZED_FILE_SIZE.clear()

    def _patch(self, target, name, value):
        patcher = mock.patch.object(target, name, value)
        patcher.start()
        self.addCleanup(patcher.stop)

    def use_settings(self, **kwargs):
        self._patch(gateway, "GATEWAY_SETTINGS", FixedSettings(**kwargs))

    def whole_rid(self):
        return self.detection.detects[0]["request_id"]

    def verify_ids(self):
        return [d["request_id"] for d in self.detection.detects if "#verify" in d["request_id"]]

    def assert_blocked(self, flow, status, reason_code):
        self.assertIsNotNone(flow.response)   # Gateway가 응답을 정했다 = 원본은 upstream으로 가지 않는다
        self.assertEqual(flow.response.status_code, status)
        self.assertEqual(len(self.reports), 1, [r.result.action for r in self.reports])
        report = self.reports[0]
        self.assertEqual((report.result.action, report.applied, report.http_status), ("block", True, status))
        self.assertEqual(report.result.reason_code, reason_code)
        if status == 503:
            self.assertEqual(flow.response.raw_content, OLD_503_BODY)
        else:
            self.assertEqual(json.loads(flow.response.get_text())["error"]["reason_code"], reason_code)
        return report


class PdfMaskMatrix:
    """세 경로(ChatGPT 원본 PUT·ChatGPT multipart·Claude multipart)에 같은 계약 시험. 하위 클래스가 send·forwarded·
    prefix_suffix를 정한다."""

    target_llm = "ChatGPT"
    prefix_suffix = ""

    def send(self, content, filename="pii.pdf"):
        raise NotImplementedError

    def forwarded(self, flow) -> bytes:
        raise NotImplementedError

    # ---------------------------------------------------------------- 성공

    def test_masked_pdf_is_verified_and_forwarded(self):
        original = pii_pdf()
        flow = self.send(original)
        self.assertIsNone(flow.response)
        self.assertEqual(self.forwarded(flow), self.detection.masked)
        self.assertNotIn(RRN.encode(), flow.request.raw_content)
        self.assertEqual(flow.request.headers["content-length"], str(len(flow.request.raw_content)))
        report, = self.reports
        self.assertEqual((report.result.action, report.applied, report.context["targetLlm"]),
                         ("mask_full", True, self.target_llm))
        self.assertNotIn("#", report.result.request_id)   # 기록은 파일 전체 판정
        self.assertEqual(report.result.pii_hits, "KR_RRN:1")
        # 가림 호출 계약: 원본 바이트 그대로, 기록 번호 앞부분, Detection 시한 = 파일 대기 시간(2.0초)의 80%
        call, = self.detection.redacts
        prefix = self.whole_rid() + self.prefix_suffix
        self.assertEqual(call["body"], original)
        self.assertEqual(call["headers"]["x-request-prefix"], prefix)
        self.assertEqual(call["headers"]["x-deadline-ms"], "1600")
        self.assertEqual(call["headers"]["x-source-service"], self.target_llm)
        self.assertEqual(unquote(call["headers"]["x-source-filename"]), "pii.pdf")
        self.assertNotIn("?", call["headers"]["x-source-endpoint"])
        self.assertEqual(self.detection.contract_errors, [])
        # 재검사 두 번(글자 그대로, 공백 뺀 글자) - 돌려받은 PDF의 글자로, 보고서가 빼는 '#verify'로 끝나는 번호
        self.assertEqual(self.verify_ids(), [f"{prefix}#verify", f"{prefix}#nows#verify"])
        verify = [d for d in self.detection.detects if "#verify" in d["request_id"]]
        self.assertEqual(verify[0]["normalized_text"], "RRN removed end")
        self.assertEqual(verify[1]["normalized_text"], "RRNremovedend")
        self.assertTrue(all(v["source"]["signature_original_offsets"] for v in verify))

    # ---------------------------------------------------------------- 403 mask_application_failed

    def test_refusals_are_403_mask_application_failed(self):
        cases = {
            "unsupported": ("unsupported", None),
            "residual": ("residual", None),
            "detection_left_the_value": ("masked", text_pdf(f"RRN {RRN} end")),            # 재검사 ① 글자 그대로
            "value_split_by_space": ("masked", text_pdf("RRN 900101- 1234567 end")),       # 재검사 ② 공백 뺀 글자
            "masked_reply_has_no_text": ("masked", scanned_pdf()),                          # 재분류: 스캔 PDF
            "clean_but_original_still_has_the_value": ("clean", None),                       # 원본 재검사에 남음
        }
        for name, (mode, reply) in cases.items():
            with self.subTest(name):
                self.reports.clear()
                self.detection.mode = mode
                self.detection.masked = reply or masked_pdf()
                flow = self.send(pii_pdf())
                self.assert_blocked(flow, 403, "mask_application_failed")
                self.assertEqual(self.reports[0].context["targetLlm"], self.target_llm)

    def test_masked_pdf_over_the_gateway_byte_limit_is_refused(self):
        original = pii_pdf()
        big = text_pdf("RRN removed end")
        self.detection.masked = big[:9] + b"%" + b"x" * 9000 + b"\n" + big[9:]
        self._patch(gateway, "GATEWAY_FILE_MAX_BYTES", 8192)
        self.assertLess(len(original), 8192)
        flow = self.send(original)
        self.assert_blocked(flow, 403, "mask_application_failed")

    def test_page_block_is_403_with_the_page_decision(self):
        self.detection.mode = "blocked"
        flow = self.send(pii_pdf())
        report = self.assert_blocked(flow, 403, "doc_excerpt")
        self.assertEqual(report.result.request_id, f"{self.whole_rid()}{self.prefix_suffix}#f:0:pdf:body:p0")
        self.assertEqual(self.verify_ids(), [])

    # ---------------------------------------------------------------- 503

    def test_timeout_is_503_even_when_the_file_setting_forwards(self):
        self.detection.mode = "masked"
        self.detection.delay = 1.5
        for action in ("block", "forward"):
            with self.subTest(action):
                self.reports.clear()
                self.use_settings(file=(action, 0.5))
                started = time.monotonic()
                flow = self.send(pii_pdf())
                elapsed = time.monotonic() - started
                report = self.assert_blocked(flow, 503, "detection_timeout")
                self.assertEqual(report.result.degraded_reason, "TIMEOUT")
                self.assertLess(elapsed, 1.3)   # 총시한(0.5초) 안에 끊는다
                self.assertEqual(self.detection.redacts[-1]["headers"]["x-deadline-ms"], "400")

    def test_oversize_reply_is_503(self):
        self._patch(detection_client, "REDACT_MAX_PDF_BYTES", 8192)
        self.detection.masked = b"%PDF-1.4\n%" + b"x" * 20000 + b"\n"
        flow = self.send(pii_pdf())
        self.assert_blocked(flow, 503, "file_masking_failed")
        self.assertEqual(self.verify_ids(), [])

    def test_unavailable_answers_are_503(self):
        for mode in ("busy", "timeout", "error", "http404", "verify_500"):
            with self.subTest(mode):
                self.reports.clear()
                self.detection.mode = "masked" if mode == "verify_500" else mode
                self.detection.verify_status = 500 if mode == "verify_500" else 200
                flow = self.send(pii_pdf())
                report = self.assert_blocked(flow, 503, "detection_error" if mode == "verify_500"
                                             else "file_masking_failed")
                if mode != "verify_500":
                    self.assertEqual(report.result.pii_hits, "KR_RRN:1")   # 파일 전체 판정의 건수로 기록

    def test_scanned_pdf_is_still_blocked_without_a_redaction_call(self):
        flow = self.send(scanned_pdf())
        self.assertEqual(flow.response.status_code, 403)
        self.assertEqual(self.detection.redacts, [])
        self.assertEqual(self.reports[0].result.reason_code, "file_uninspectable_scanned_pdf")

    def test_clean_pdf_does_not_call_the_redaction(self):
        original = text_pdf("Quarterly meeting notes")
        flow = self.send(original)
        self.assertIsNone(flow.response)
        self.assertEqual(self.forwarded(flow), original)
        self.assertEqual(self.detection.redacts, [])
        self.assertEqual([r.result.action for r in self.reports], ["allow"])


class ChatGptRawPutTests(PdfMaskMatrix, Base):
    def send(self, content, filename="pii.pdf"):
        gateway.CHATGPT_FILE_META[FILE_ID] = filename
        req = tutils.treq(host="sdmntprkoreacentral.oaiusercontent.com", port=443, scheme=b"https", method=b"PUT",
                          path=f"/files/{FILE_ID}/raw?se=2026-10-11&sig=fake".encode(), content=content,
                          headers=((b"content-type", b"application/octet-stream"),))
        req.headers["content-length"] = str(len(content))
        flow = tflow.tflow(req=req)
        run_hook(gateway.request(flow))
        return flow

    def forwarded(self, flow):
        return flow.request.raw_content

    def test_completion_file_size_is_rewritten_to_the_masked_size(self):
        original = pii_pdf()
        flow = self.send(original)
        self.assertIsNone(flow.response)
        self.assertEqual(gateway.CHATGPT_SANITIZED_FILE_SIZE[FILE_ID], len(self.detection.masked))
        self.assertNotEqual(len(original), len(self.detection.masked))
        claim = tflow.tflow(req=tutils.treq(
            host="chatgpt.com", port=443, scheme=b"https", method=b"POST", path=b"/backend-api/files/process_upload_stream",
            content=json.dumps({"file_id": FILE_ID, "file_size": len(original)}).encode(),
            headers=((b"content-type", b"application/json"),)))
        run_hook(gateway.request(claim))
        self.assertIsNone(claim.response)
        self.assertEqual(json.loads(claim.request.raw_content)["file_size"], len(self.detection.masked))

    def test_refused_upload_leaves_no_size_entry(self):
        self.detection.mode = "residual"
        self.send(pii_pdf())
        self.assertNotIn(FILE_ID, gateway.CHATGPT_SANITIZED_FILE_SIZE)


def multipart_body(boundary, parts):
    body = b""
    for filename, ctype, content in parts:
        body += (f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"{filename}\"\r\n"
                 f"Content-Type: {ctype}\r\n\r\n").encode() + content + b"\r\n"
    return body + f"--{boundary}--\r\n".encode()


def multipart_files(body, boundary):
    message = BytesParser(policy=email_default).parsebytes(
        f"Content-Type: multipart/form-data; boundary={boundary}\r\n\r\n".encode() + body)
    return [part.get_payload(decode=True) for part in message.iter_parts() if part.get_filename()]


class ChatGptMultipartTests(PdfMaskMatrix, Base):
    def send(self, content, filename="pii.pdf"):
        body = multipart_body(CHATGPT_BOUNDARY, [(filename, "application/pdf", content)])
        req = tutils.treq(host="chatgpt.com", port=443, scheme=b"https", method=b"POST",
                          path=b"/backend-api/files/upload?trace=fake", content=body,
                          headers=((b"content-type", f"multipart/form-data; boundary={CHATGPT_BOUNDARY}".encode()),))
        flow = tflow.tflow(req=req)
        run_hook(gateway.request(flow))
        return flow

    def forwarded(self, flow):
        payload, = multipart_files(flow.request.raw_content, CHATGPT_BOUNDARY)
        return payload

    def test_boundary_inside_the_masked_pdf_is_503(self):
        self.detection.masked = pdf_from(b"BT /F1 12 Tf 72 720 Td (RRN removed end) Tj ET",
                                         extra_objects=[f"(--{CHATGPT_BOUNDARY})".encode()])
        flow = self.send(pii_pdf())
        self.assert_blocked(flow, 503, "file_masking_failed")


class ClaudeMultipartTests(PdfMaskMatrix, Base):
    target_llm = "Claude"
    prefix_suffix = "#file:0"

    def upload(self, *parts):
        body = multipart_body(CLAUDE_BOUNDARY, parts)
        req = tutils.treq(host="claude.ai", port=443, scheme=b"https", method=b"POST", path=CLAUDE_UPLOAD.encode(),
                          content=body,
                          headers=((b"content-type", f"multipart/form-data; boundary={CLAUDE_BOUNDARY}".encode()),))
        flow = tflow.tflow(req=req)
        run_hook(gateway.request(flow))
        return flow

    def send(self, content, filename="pii.pdf"):
        return self.upload((filename, "application/pdf", content))

    def forwarded(self, flow):
        payload, = multipart_files(flow.request.raw_content, CLAUDE_BOUNDARY)
        return payload

    def test_docx_and_pdf_are_both_masked(self):
        flow = self.upload(("d.docx", DOCX, docx_bytes("안내", f"주민등록번호 {RRN}")),
                           ("a.pdf", "application/pdf", pii_pdf()))
        self.assertIsNone(flow.response)
        docx_out, pdf_out = multipart_files(flow.request.raw_content, CLAUDE_BOUNDARY)
        self.assertEqual(pdf_out, self.detection.masked)
        self.assertNotIn(RRN.encode(), flow.request.raw_content)
        self.assertNotIn(RRN, "\n".join(p.text for p in Document(io.BytesIO(docx_out)).paragraphs))
        self.assertEqual(self.detection.redacts[0]["headers"]["x-request-prefix"], f"{self.whole_rid()}#file:1")
        self.assertEqual([(r.result.action, r.applied) for r in self.reports], [("mask_full", True)])

    def test_clean_pdf_next_to_a_masked_docx_is_verified_and_kept(self):
        clean = text_pdf("agenda notes")
        self.detection.mode = "clean"
        flow = self.upload(("d.docx", DOCX, docx_bytes(f"주민등록번호 {RRN}")), ("c.pdf", "application/pdf", clean))
        self.assertIsNone(flow.response)
        _, pdf_out = multipart_files(flow.request.raw_content, CLAUDE_BOUNDARY)
        self.assertEqual(pdf_out, clean)
        self.assertEqual(self.verify_ids(), [f"{self.whole_rid()}#file:1#verify", f"{self.whole_rid()}#file:1#nows#verify"])
        self.assertEqual([(r.result.action, r.applied) for r in self.reports], [("mask_full", True)])

    def test_boundary_inside_the_masked_pdf_is_503(self):
        self.detection.masked = pdf_from(b"BT /F1 12 Tf 72 720 Td (RRN removed end) Tj ET",
                                         extra_objects=[f"(--{CLAUDE_BOUNDARY})".encode()])
        flow = self.send(pii_pdf())
        self.assert_blocked(flow, 503, "file_masking_failed")


if __name__ == "__main__":
    unittest.main()
