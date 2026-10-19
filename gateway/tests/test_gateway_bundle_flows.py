"""Gateway 처리 설정 묶음(2026-10-10) - 실제 mitmproxy flow로 request()·response() 훅을 지난다.

검사할 수 없는 파일 집행(ChatGPT 원본 PUT·multipart, Claude 한 파일·여러 파일), 대기 시간·실패 처리(질문·파일),
칸 무응답, 글자·바이트 상한, 이미지가 붙은 ChatGPT 질문, 처리 시간(processingMs = Gateway 자체 처리 시간),
Claude rpc 구조 기록(값 없음), 이벤트 모양·확장 프로그램 알림 표시, 호출 지점별 설정, 배포 순서 근거.
가짜 값만 쓴다. Detection·대시보드는 가짜(또는 로컬 시험 서버)다.
실행: python -m unittest tests.test_gateway_bundle_flows -v (Gateway 사본 폴더에서)
"""

from __future__ import annotations

import ast
import asyncio
import contextlib
import io
import json
import os
import re
import shutil
import socket
import subprocess
import sys
import tempfile
import threading
import time
import unittest
import urllib.request
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

_TMP = tempfile.mkdtemp(prefix="gw-bundle-flows-")
os.environ["SIGNATURE_REVIEW_DIR"] = str(Path(_TMP) / "review")
os.environ["DASHBOARD_OUTBOX_PATH"] = str(Path(_TMP) / "outbox.sqlite3")
os.environ.pop("DASHBOARD_EVENT_URL", None)
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from docx import Document  # noqa: E402
from mitmproxy.test import tflow, tutils  # noqa: E402
from openpyxl import Workbook  # noqa: E402

import detection_client  # noqa: E402
import gateway  # noqa: E402
from detection_client import DetectionResult  # noqa: E402
from extension_event_view import extension_event_view  # noqa: E402
try:
    from detection_client import failure_result  # noqa: E402
except ImportError:   # 운영본(이 묶음 전)으로 새 시험을 돌려 실패를 확인할 때. detection_client.check 안의 실패 판정과 같다
    def failure_result(reason, failure_policy, request_id):
        closed = failure_policy == "FAIL_CLOSED"
        return DetectionResult(
            blocked=closed, reason=f"inspection failed: {reason}; {failure_policy}", degraded=True,
            degraded_reason=reason, failure_policy=failure_policy, action="block" if closed else "allow_degraded",
            reason_code={"TIMEOUT": "detection_timeout", "PARSE_ERROR": "detection_parse_error",
                         "OTHER": "detection_error", "UNREACHABLE": "detection_unreachable"}[reason],
            request_id=request_id, http_status=503 if closed else None)
try:
    from gateway_settings import SettingsHolder, env_defaults  # noqa: E402
except ImportError:
    SettingsHolder = env_defaults = None
from tests.gw_support import (  # noqa: E402
    REVISION_2, FixedSettings, mpo_bytes, nested_text_boxes, png_bytes, raw_docx, scanned_pdf, settings_doc, text_pdf,
)
from tests.gw_support import run_hook  # noqa: E402

RRN = "900101-1234567"                 # 가상 주민등록번호
SECRET = "PROJECT-ORION-ROADMAP"       # 파일 전체·첨부 검사에서 기밀 문서 일치(차단)
CELL_ONLY = "CELL-ONLY-ACCOUNT-0042"   # 칸 검사에서만 잡히는 값(가림)
DOWN = "DETECTION-DOWN-MARKER"         # 칸 검사에서 Detection 무응답
EMAIL = "hong.gildong@example.test"
PHONE = "010-1234-5678"
FILE_ID = "00000000-aaaa-8206-bbbb-000000000001"
ORG = "00000000-0000-4000-8000-00000000c1a0"
CONV = "00000000-0000-4000-8000-00000000c0de"
COMPLETION_PATH = f"/api/organizations/{ORG}/chat_conversations/{CONV}/completion"
CLAUDE_UPLOAD_PATH = f"/api/{ORG}/upload"
CFB = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 504   # HWP 5 등(검사할 수 없는 형식)
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SAMPLE_KO_PDF = Path(os.environ.get("GW_TEST_SAMPLE_KO_PDF", ROOT / "tests" / "fixtures" / "sample_ko.pdf"))
DASHBOARD_SERVER_DIR = Path(os.environ.get(
    "GW_TEST_DASHBOARD_SERVER_DIR",
    "C:/Users/user/AppData/Local/Temp/claude/C--Users-user-Desktop-kisia-project/"
    "32d4e37a-7639-4ae3-a018-7182e5d25155/scratchpad/dash-work/server"))
_SEGMENT_RE = re.compile(r"#(?:file:\d+#)?f:\d+:(?:docx|xlsx):")


def fake_detection(calls, delay=0.0):
    """가짜 Detection. 넘겨받은 대기 시간·실패 정책을 기록하고, 무응답이면 그 정책으로 판정을 만든다."""
    def check(base_url, *, request_id, normalized_text, source=None, timeout=None, failure_policy=None, **_):
        calls.append({"request_id": request_id, "text": normalized_text, "source": source or {},
                      "timeout": timeout, "failure_policy": failure_policy})
        if delay:
            time.sleep(delay)
        segment = bool(_SEGMENT_RE.search(request_id))
        if DOWN in normalized_text and segment:
            return failure_result("TIMEOUT", failure_policy or "FAIL_OPEN", request_id)
        if SECRET in normalized_text and not segment:
            return DetectionResult(blocked=True, reason="doc match", request_id=request_id, decided_by="windowhash",
                                   doc_id="doc-fake-001", doc_span="0-21", input_span="0-21", match_rate=1.0,
                                   action="block", reason_code="doc_excerpt", http_status=403,
                                   user_message="보호 문서와 일치해 전송을 차단했습니다.")
        found = [(m.start(), m.end(), "KR_RRN") for m in re.finditer(re.escape(RRN), normalized_text)]
        if segment:
            found += [(m.start(), m.end(), "KR_BANK_ACCOUNT") for m in re.finditer(re.escape(CELL_ONLY), normalized_text)]
        if found:
            found.sort()
            masked = normalized_text
            for start, end, entity in reversed(found):
                masked = masked[:start] + f"[{entity}]" + masked[end:]
            counts = {}
            for _, _, entity in found:
                counts[entity] = counts.get(entity, 0) + 1
            findings = [{"entity_type": entity, "start": start, "end": end, "replacement": f"[{entity}]",
                         "policy_action": "mask_full", "score": 1.0} for start, end, entity in found]
            return DetectionResult(blocked=False, reason="signature", request_id=request_id, decided_by="signature",
                                   findings=findings, masked_text=masked, action="mask_full", reason_code="pii",
                                   pii_hits=",".join(f"{k}:{v}" for k, v in sorted(counts.items())))
        return DetectionResult(blocked=False, reason="", request_id=request_id, decided_by="none", action="allow",
                               reason_code="none")
    return check


def docx_bytes(*paragraphs):
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def xlsx_bytes(cells):
    workbook = Workbook()
    sheet = workbook.active
    for ref, value in cells:
        sheet[ref] = value
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


FILE_MAX_CHARS_2S = 80645   # gateway_settings.file_max_chars(2.0): 기본 파일 대기 2초에 검사할 수 있는 글자


def too_large_text():
    return b"x" * (FILE_MAX_CHARS_2S + 1)


GROUP_FILES = {
    "image": png_bytes,
    "scanned_pdf": scanned_pdf,
    "other": lambda: CFB,
    "too_large": too_large_text,
}
GROUP_REASON = {"image": "file_uninspectable_image", "scanned_pdf": "file_uninspectable_scanned_pdf",
                "other": "file_uninspectable_other", "too_large": "file_too_large"}


def group_settings(group, action):
    """그룹별 차단/허용 설정. too_large는 파일 처리 설정(차단/원본 전송)."""
    if group == "too_large":
        return {"file": ("block" if action == "block" else "forward", 2.0)}
    return {group: action}


class Base(unittest.TestCase):
    def setUp(self):
        self.calls, self.reports = [], []
        self._orig = (gateway.check_detection, gateway.report_dashboard_event, gateway.send_to_next_module)
        gateway.check_detection = fake_detection(self.calls)
        gateway.send_to_next_module = lambda event: None

        def capture(flow, result, dashboard_context, *, enforcement_applied, http_status):
            self.reports.append(SimpleNamespace(result=result, context=dashboard_context,
                                                applied=enforcement_applied, http_status=http_status))
        gateway.report_dashboard_event = capture
        gateway.CHATGPT_FILE_META.clear()
        gateway.CHATGPT_SANITIZED_FILE_SIZE.clear()
        getattr(gateway, "_dedupe_seen", {}).clear()   # 10/11: 같은 업로드 재시도 기록 한 번만 - 시험마다 따로
        # PDF 가림 API(2026-10-11): '지원하지 않음'을 돌려주는 가짜(→ 403 mask_application_failed, 이전 PDF 403과 같은 결과).
        # 가림 성공·잔존·시간 초과 등은 tests/test_gateway_pdf_mask.py
        self._orig_redact = gateway.redact_pdf_detection
        gateway.redact_pdf_detection = lambda *args, **kwargs: detection_client.RedactResult(
            status="unsupported", reason="fake_unsupported")
        # 시스템 상태(10/12): 지표 전송기 값을 만드는 시험이 실제 중계(127.0.0.1:8766)를 부르지 않게 한다. 중계 시험은 tests/test_gateway_relay_status.py
        relay = mock.patch.object(gateway, "fetch_relay_status", return_value={"ok": True, "connected": 2, "registered": 1})
        relay.start()
        self.addCleanup(relay.stop)

    def tearDown(self):
        gateway.check_detection, gateway.report_dashboard_event, gateway.send_to_next_module = self._orig
        gateway.redact_pdf_detection = self._orig_redact

    def use_settings(self, **kwargs):
        patcher = mock.patch.object(gateway, "GATEWAY_SETTINGS", FixedSettings(**kwargs))
        patcher.start()
        self.addCleanup(patcher.stop)

    # -- 요청 만들기 (gateway.request까지 지난 flow를 돌려준다)
    def run_flow(self, req):
        flow = tflow.tflow(req=req)
        self.sent = req.raw_content
        run_hook(gateway.request(flow))
        return flow

    def chatgpt_put(self, content, filename=None):
        if filename:
            gateway.CHATGPT_FILE_META[FILE_ID] = filename
        req = tutils.treq(host="sdmntprkoreacentral.oaiusercontent.com", port=443, scheme=b"https", method=b"PUT",
                          path=f"/files/{FILE_ID}/raw?se=2026-10-10".encode(), content=content,
                          headers=((b"content-type", b"application/octet-stream"),))
        return self.run_flow(req)

    def chatgpt_multipart(self, content, filename="upload.bin", field="file"):
        boundary = "----gwbundleboundary"
        body = (f"--{boundary}\r\nContent-Disposition: form-data; name=\"{field}\"; filename=\"{filename}\"\r\n"
                f"Content-Type: application/octet-stream\r\n\r\n").encode() + content + f"\r\n--{boundary}--\r\n".encode()
        req = tutils.treq(host="chatgpt.com", port=443, scheme=b"https", method=b"POST", path=b"/backend-api/files/upload",
                          content=body, headers=((b"content-type", f"multipart/form-data; boundary={boundary}".encode()),))
        return self.run_flow(req)

    def claude_upload(self, *parts, raw_body=None, boundary="----gwbundleclaude"):
        body = b""
        for filename, ctype, content in parts:
            body += (f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"{filename}\"\r\n"
                     f"Content-Type: {ctype}\r\n\r\n").encode() + content + b"\r\n"
        body += f"--{boundary}--\r\n".encode()
        req = tutils.treq(host="claude.ai", port=443, scheme=b"https", method=b"POST", path=CLAUDE_UPLOAD_PATH.encode(),
                          content=raw_body if raw_body is not None else body,
                          headers=((b"content-type", f"multipart/form-data; boundary={boundary}".encode()),))
        return self.run_flow(req)

    def claude_completion(self, body):
        req = tutils.treq(host="claude.ai", port=443, scheme=b"https", method=b"POST", path=COMPLETION_PATH.encode(),
                          content=json.dumps(body, ensure_ascii=False).encode(),
                          headers=((b"content-type", b"application/json"),))
        return self.run_flow(req)

    def chatgpt_conversation(self, message):
        req = tutils.treq(host="chatgpt.com", port=443, scheme=b"https", method=b"POST", path=b"/backend-api/f/conversation",
                          content=json.dumps({"action": "next", "messages": [message]}, ensure_ascii=False).encode(),
                          headers=((b"content-type", b"application/json"),))
        return self.run_flow(req)

    def chatgpt_prompt(self, text):
        return self.chatgpt_conversation({"author": {"role": "user"},
                                          "content": {"content_type": "text", "parts": [text]}})

    def finish(self, flow, *, upstream_seconds=0.0):
        """upstream(LLM)을 흉내: 기다린 뒤 응답을 받고 response() 훅."""
        if flow.response is None:
            time.sleep(upstream_seconds)
            flow.response = tutils.tresp(status_code=200)
        with mock.patch.object(gateway, "_emit_gateway_metrics", lambda: None):
            gateway.response(flow)
        return flow

    def got(self):
        return [(r.result.action, r.context["contentType"], r.applied, getattr(r.result, "reason_code", None))
                for r in self.reports]

    def assert_blocked(self, flow, status=403):
        self.assertIsNotNone(flow.response)
        self.assertEqual(flow.response.status_code, status)

    def forwarded_docx_text(self, flow):
        """전달되는 본문 속 DOCX의 글자(압축돼 있어 바이트로는 볼 수 없다). 원본 PUT·Claude multipart."""
        if flow.request.host == "claude.ai":
            data = gateway.extract_claude_file_upload_info(flow)[0]["file_bytes"]
        else:
            data = flow.request.raw_content
        return chr(10).join(paragraph.text for paragraph in Document(io.BytesIO(data)).paragraphs)


class UninspectableEnforcementTests(Base):
    """명세 §4.6·§11.1 검사 불가 집행: 경로 4 × 그룹 4 × 차단/허용."""

    def run_path(self, path, data):
        if path == "chatgpt_put":
            return self.chatgpt_put(data, "upload.bin")
        if path == "chatgpt_multipart":
            return self.chatgpt_multipart(data, "upload.bin")
        if path == "claude_single":
            return self.claude_upload(("upload.bin", "application/octet-stream", data))
        return self.claude_upload(("memo.txt", "text/plain", "plain meeting memo".encode()),
                                  ("upload.bin", "application/octet-stream", data))

    def test_every_path_group_and_action(self):
        for path in ("chatgpt_put", "chatgpt_multipart", "claude_single", "claude_multi"):
            for group, make in GROUP_FILES.items():
                for action in ("block", "allow"):
                    with self.subTest(path=path, group=group, action=action):
                        self.calls.clear()
                        self.reports.clear()
                        with mock.patch.object(gateway, "GATEWAY_SETTINGS",
                                               FixedSettings(**group_settings(group, action))):
                            flow = self.run_path(path, make())
                        llm = "ChatGPT" if path.startswith("chatgpt") else "Claude"
                        reason = GROUP_REASON[group]
                        if action == "block":
                            self.assert_blocked(flow, 403)
                            body = json.loads(flow.response.get_text())
                            self.assertEqual((body["error"]["type"], body["error"]["reason_code"]),
                                             ("policy_violation", reason))
                            self.assertEqual(self.calls, [])          # Detection 전에 끝난다(여러 파일이어도)
                            self.assertEqual(self.got(), [("block", "file", True, reason)])
                            self.assertEqual(self.reports[0].http_status, 403)
                            self.assertEqual(flow.metadata["gw_action"], "block")
                        else:
                            self.assertIsNone(flow.response)
                            self.assertEqual(flow.request.raw_content, self.sent)   # 원본 그대로 전달
                            if path == "claude_multi":
                                self.assertEqual(len(self.calls), 1)               # 검사 가능한 파일만
                                self.assertEqual(self.got(), [("allow", "file", False, "none"),
                                                              ("allow", "file", False, reason)])
                            else:
                                self.assertEqual(self.calls, [])
                                self.assertEqual(self.got(), [("allow", "file", False, reason)])
                            self.assertEqual(flow.metadata["gw_action"], "allow")
                        self.assertTrue(all(r.context["targetLlm"] == llm for r in self.reports))
                        if path == "chatgpt_put":
                            self.assertEqual(self.reports[-1].context["fileId"], gateway.canonical_file_ref(FILE_ID))

    def test_image_allowed_with_a_confidential_pdf_leaves_one_block_event(self):
        flow = self.claude_upload(("p.png", "image/png", png_bytes()),
                                  ("s.pdf", "application/pdf", text_pdf(f"{SECRET} plan")))
        self.assert_blocked(flow, 403)
        self.assertEqual(self.got(), [("block", "file", True, "doc_excerpt")])   # 보류한 이미지 허용 기록은 버림

    def test_image_allowed_with_a_clean_pdf_leaves_two_events(self):
        flow = self.claude_upload(("p.png", "image/png", png_bytes()), ("c.pdf", "application/pdf", text_pdf("agenda")))
        self.assertIsNone(flow.response)
        self.assertEqual(self.got(), [("allow", "file", False, "none"),
                                      ("allow", "file", False, "file_uninspectable_image")])
        image_report = self.reports[1]
        self.assertRegex(image_report.result.request_id, r"#upload:0$")

    def test_empty_text_document_is_other_not_pass(self):
        flow = self.claude_upload(("e.docx", DOCX, docx_bytes()))
        self.assert_blocked(flow, 403)
        self.assertEqual(self.got(), [("block", "file", True, "file_uninspectable_other")])
        flow = self.chatgpt_put(docx_bytes(), "e.docx")
        self.assert_blocked(flow, 403)

    def test_zero_byte_file_passes_without_an_event(self):
        flow = self.chatgpt_put(b"", "e.txt")
        self.assertIsNone(flow.response)
        self.assertEqual((self.calls, self.reports), ([], []))

    def test_text_named_like_an_image_is_inspected(self):
        flow = self.chatgpt_put(f"주민번호 {RRN}".encode(), "photo.png")
        self.assertEqual(len(self.calls), 1)
        self.assertIn(RRN, self.calls[0]["text"])
        self.assert_blocked(flow, 403)   # 글자 파일은 칸 가림을 못 하므로 가림 판정 → 차단(mask_application_failed)
        self.assertEqual(self.got(), [("block", "file", True, "mask_application_failed")])

    def test_renamed_docx_is_masked_as_docx(self):
        original = docx_bytes("안내", f"주민번호 {RRN}")
        flow = self.claude_upload(("notes.txt", "text/plain", original))
        self.assertIsNone(flow.response)
        self.assertNotIn(RRN, self.forwarded_docx_text(flow))
        self.assertIn("[KR_RRN]", self.forwarded_docx_text(flow))
        self.assertEqual(self.got(), [("mask_full", "file", True, "pii")])


class ParseFailureTests(Base):
    def test_unparsed_claude_multipart_with_a_file_header_is_other(self):
        raw = (b"--WRONG\r\nContent-Disposition: form-data; name=\"file\"; filename=\"a.txt\"\r\n\r\n"
               + f"RRN {RRN}".encode() + b"\r\n--WRONG--\r\n")
        flow = self.claude_upload(raw_body=raw, boundary="expected-boundary")
        self.assert_blocked(flow, 403)
        self.assertEqual(self.got(), [("block", "file", True, "file_uninspectable_other")])
        self.assertEqual(self.calls, [])

    def test_multipart_without_file_parts_still_passes(self):
        raw = b"--b1\r\nContent-Disposition: form-data; name=\"note\"\r\n\r\nhello\r\n--b1--\r\n"
        flow = self.claude_upload(raw_body=raw, boundary="b1")
        self.assertIsNone(flow.response)
        self.assertEqual((self.calls, self.reports), ([], []))

    def test_part_with_an_empty_filename_passes_like_before(self):
        # 10/11 검토: filename=""(파일 고르지 않은 칸)은 해석 실패가 아니다. 운영본처럼 통과·기록 없음
        raw = (b"--b2\r\nContent-Disposition: form-data; name=\"file\"; filename=\"\"\r\n"
               b"Content-Type: application/octet-stream\r\n\r\n\r\n--b2--\r\n")
        flow = self.claude_upload(raw_body=raw, boundary="b2")
        self.assertIsNone(flow.response)
        self.assertEqual((self.calls, self.reports), ([], []))

    def test_unparsed_multipart_filename_forms(self):
        for header in (b'filename="a.txt"', b"filename=a.txt", b"filename*=UTF-8''%EB%A9%94%EB%AA%A8.txt",
                       b'filename=""; filename*=UTF-8\'\'a.txt'):
            with self.subTest(header=header):
                self.reports.clear()
                raw = (b"--WRONG\r\nContent-Disposition: form-data; name=\"file\"; " + header + b"\r\n\r\n"
                       + f"RRN {RRN}".encode() + b"\r\n--WRONG--\r\n")
                flow = self.claude_upload(raw_body=raw, boundary="expected-boundary")
                self.assert_blocked(flow, 403)
                self.assertEqual(self.got(), [("block", "file", True, "file_uninspectable_other")])

    def test_chatgpt_multipart_without_a_file_part_is_other(self):
        flow = self.chatgpt_multipart(b"plain", field="attachment")
        self.assert_blocked(flow, 403)
        self.assertEqual(self.got(), [("block", "file", True, "file_uninspectable_other")])

    def test_chatgpt_processing_errors_fail_closed_with_an_event(self):
        for target in ("_classify_upload", "create_chatgpt_raw_file_event"):
            with self.subTest(target=target):
                self.reports.clear()
                with mock.patch.object(gateway, target, side_effect=RuntimeError("boom")):
                    flow = self.chatgpt_put(f"memo {RRN}".encode(), "m.txt")
                self.assert_blocked(flow, 503)
                self.assertEqual(flow.response.raw_content, b'{"error":{"type":"file_masking_failed"}}')
                self.assertEqual(self.got(), [("block", "file", True, "file_masking_failed")])
                self.assertEqual(self.reports[0].context["targetLlm"], "ChatGPT")
        with mock.patch.object(gateway, "create_file_event", side_effect=RuntimeError("boom")):
            flow = self.chatgpt_multipart(f"memo {RRN}".encode(), "m.txt")
        self.assert_blocked(flow, 503)


class KoreanPdfFlowTests(Base):
    def setUp(self):
        super().setUp()
        if not SAMPLE_KO_PDF.exists():
            self.skipTest("sample_ko.pdf not available")
        self.pdf = SAMPLE_KO_PDF.read_bytes()

    def test_chatgpt_and_claude_send_hangul_to_detection(self):
        self.chatgpt_put(self.pdf, "ko.pdf")
        self.claude_upload(("ko.pdf", "application/pdf", self.pdf))
        self.assertEqual(len(self.calls), 2)
        for call in self.calls:
            self.assertGreaterEqual(len(re.findall("[가-힣]", call["text"])), 81)


class ReviewFixFlowTests(Base):
    """10/11 검토 반영(파일 분류)을 실제 flow로: ChatGPT 원본 PUT·Claude multipart."""

    def setUp(self):
        super().setUp()
        self.use_settings()      # 기본값: 이미지 허용, 파일 차단 2초

    def test_repeated_json_key_value_reaches_detection(self):
        data = ('{"note": "%s", "note": "ok"}' % RRN).encode()
        self.chatgpt_put(data, "a.json")
        self.claude_upload(("a.json", "application/json", data))
        self.assertEqual(len(self.calls), 2)
        for call in self.calls:
            self.assertIn(RRN, call["text"])

    def test_multi_picture_jpeg_is_allowed_as_an_image(self):
        for path in ("chatgpt_put", "claude"):
            with self.subTest(path=path):
                self.reports.clear()
                if path == "chatgpt_put":
                    flow = self.chatgpt_put(mpo_bytes(), "photo.jpg")
                else:
                    flow = self.claude_upload(("photo.jpg", "image/jpeg", mpo_bytes()))
                self.assertIsNone(flow.response)
                self.assertEqual(flow.request.raw_content, self.sent)
                self.assertEqual(self.got(), [("allow", "file", False, "file_uninspectable_image")])
        self.assertEqual(self.calls, [])

    def test_terminal_colored_log_is_inspected(self):
        log = f"\x1b[32mINFO\x1b[0m user {RRN}\n\x1b[31mERROR\x1b[0m done\n".encode()
        self.chatgpt_put(log, "build.log")
        self.claude_upload(("build.log", "text/plain", log))
        self.assertEqual(len(self.calls), 2)
        for call in self.calls:
            self.assertIn(RRN, call["text"])

    def test_nested_text_box_docx_ends_within_the_file_wait(self):
        # 수 KB DOCX(글상자 84겹, w:t 20만 개): 이전 work는 약 18초 동안 프록시 전체를 멈췄다
        data = raw_docx(nested_text_boxes(84, 200_000))
        started = time.perf_counter()
        flow = self.claude_upload(("deep.docx", DOCX, data))
        elapsed = time.perf_counter() - started
        self.assertLess(elapsed, 2.0 + 0.5)
        self.assert_blocked(flow, 403)
        # 로컬 libxml2는 시간 상한으로, 운영 서버 libxml2 2.14+는 256겹 넘는 XML 거부(검사 불가)로 막는다. 둘 다 Detection 호출 없이 차단
        self.assertIn(self.got(), ([("block", "file", True, "file_too_large")],
                                   [("block", "file", True, "file_uninspectable_other")]))
        self.assertEqual(self.calls, [])


class SlowDetection:
    """대기 시간보다 늦게 답하는 로컬 Detection 흉내(실제 detection_client.check가 붙는다)."""

    def __init__(self, delay):
        self.delay = delay
        self.requests = []
        outer = self

        class Handler(BaseHTTPRequestHandler):
            def do_POST(self):
                body = json.loads(self.rfile.read(int(self.headers.get("content-length") or 0)))
                outer.requests.append(body["request_id"])
                time.sleep(outer.delay)
                payload = json.dumps({"blocked": False, "reason": "", "request_id": body["request_id"],
                                      "decided_by": "none", "decision": {"action": "allow"}}).encode()
                try:
                    self.send_response(200)
                    self.send_header("content-type", "application/json")
                    self.send_header("content-length", str(len(payload)))
                    self.end_headers()
                    self.wfile.write(payload)
                except OSError:
                    pass

            def log_message(self, *args):
                pass
        self.server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        self.server.daemon_threads = True
        threading.Thread(target=self.server.serve_forever, daemon=True).start()
        self.url = f"http://127.0.0.1:{self.server.server_address[1]}"

    def close(self):
        self.server.shutdown()
        self.server.server_close()


class TimeoutPolicyTests(Base):
    """명세 §1.1·§4.4: 질문·파일 각각의 대기 시간과 처리(원본 전송/차단). 실제 클라이언트 + 느린 로컬 서버."""

    def setUp(self):
        super().setUp()
        gateway.check_detection = self._orig[0]           # 실제 detection_client.check
        self.server = SlowDetection(delay=1.0)
        self.addCleanup(self.server.close)
        patcher = mock.patch.object(gateway, "DETECTION_BASE_URL", self.server.url)
        patcher.start()
        self.addCleanup(patcher.stop)

    def timed(self, action):
        started = time.monotonic()
        flow = action()
        return flow, time.monotonic() - started

    def test_prompt_forward_and_block(self):
        self.use_settings(prompt=("forward", 0.2))
        flow, elapsed = self.timed(lambda: self.chatgpt_prompt("회의 일정 정리해줘"))
        self.assertIsNone(flow.response)
        self.assertLess(elapsed, 0.8)
        report, = self.reports
        self.assertEqual((report.result.action, report.result.degraded_reason, report.result.failure_policy,
                          report.result.reason_code), ("allow_degraded", "TIMEOUT", "FAIL_OPEN", "detection_timeout"))
        self.reports.clear()
        with mock.patch.object(gateway, "GATEWAY_SETTINGS", FixedSettings(prompt=("block", 0.2))):
            flow, elapsed = self.timed(lambda: self.chatgpt_prompt("회의 일정 정리해줘"))
        self.assert_blocked(flow, 503)
        self.assertEqual(json.loads(flow.response.get_text())["error"]["type"], "inspection_unavailable")
        self.assertLess(elapsed, 0.8)
        self.assertEqual((self.reports[0].result.action, self.reports[0].result.failure_policy), ("block", "FAIL_CLOSED"))

    def test_file_forward_and_block(self):
        with mock.patch.object(gateway, "GATEWAY_SETTINGS", FixedSettings(file=("forward", 0.2))):
            flow, elapsed = self.timed(lambda: self.chatgpt_put(b"plain memo text", "m.txt"))
        self.assertIsNone(flow.response)
        self.assertEqual(flow.request.raw_content, self.sent)
        self.assertLess(elapsed, 0.8)
        self.assertEqual(self.got(), [("allow_degraded", "file", False, "detection_timeout")])
        self.reports.clear()
        with mock.patch.object(gateway, "GATEWAY_SETTINGS", FixedSettings(file=("block", 0.2))):
            flow, elapsed = self.timed(lambda: self.chatgpt_put(b"plain memo text", "m.txt"))
        self.assert_blocked(flow, 503)
        self.assertLess(elapsed, 0.8)
        self.assertEqual(self.got(), [("block", "file", True, "detection_timeout")])

    def test_longer_wait_lets_the_slow_server_answer(self):
        self.server.delay = 0.3
        with mock.patch.object(gateway, "GATEWAY_SETTINGS", FixedSettings(prompt=("block", 1.0))):
            flow = self.chatgpt_prompt("안녕")
        self.assertIsNone(flow.response)
        self.assertEqual(self.reports[0].result.action, "allow")


class LimitFlowTests(Base):
    def test_text_over_the_wait_limit_never_reaches_detection(self):
        flow = self.chatgpt_put(too_large_text(), "big.txt")
        self.assert_blocked(flow, 403)
        self.assertEqual(json.loads(flow.response.get_text())["error"]["message"], gateway.FILE_TOO_LARGE_MESSAGE)
        self.assertEqual(self.calls, [])

    def test_claude_attachment_text_over_the_limit(self):
        flow = self.claude_completion({"prompt": "요약", "attachments": [
            {"file_name": "big.txt", "extracted_content": "가" * (FILE_MAX_CHARS_2S + 1)}]})
        self.assert_blocked(flow, 403)
        self.assertEqual([c["text"] for c in self.calls], ["요약"])   # 첨부 글자는 Detection에 보내지 않음
        self.assertEqual(self.got(), [("block", "file", True, "file_too_large")])

    def test_sent_length_over_the_limit_with_forward_skips_cells(self):
        # 공백이 많아 분류 단계(공백 뺀 글자)는 통과하지만 보낼 문자열 길이는 상한을 넘는 DOCX
        self.use_settings(file=("forward", 2.0))
        data = docx_bytes(*["x " * 1000] * 45)
        flow = self.chatgpt_put(data, "spaces.docx")
        self.assertIsNone(flow.response)
        self.assertEqual(flow.request.raw_content, self.sent)
        self.assertEqual(self.calls, [])                      # 파일 전체·칸 모두 호출 없음
        self.assertEqual(self.got(), [("allow", "file", False, "file_too_large")])

    def test_bytes_over_the_limit_are_not_extracted(self):
        import gateway_file_classes
        with mock.patch.object(gateway, "GATEWAY_FILE_MAX_BYTES", 4096), \
                mock.patch.object(gateway_file_classes, "extract_pdf", side_effect=AssertionError("extracted")):
            flow = self.chatgpt_put(text_pdf("hello") + b"%" * 5000, "big.pdf")
        self.assert_blocked(flow, 403)
        self.assertEqual(self.got(), [("block", "file", True, "file_too_large")])


class SilentDetection:
    """연결은 받고 답하지 않는 Detection."""

    def __init__(self):
        self.server = socket.socket()
        self.server.bind(("127.0.0.1", 0))
        self.server.listen(32)
        self.held = []
        threading.Thread(target=self._accept, daemon=True).start()
        self.url = f"http://127.0.0.1:{self.server.getsockname()[1]}"

    def _accept(self):
        while True:
            try:
                self.held.append(self.server.accept()[0])
            except OSError:
                return

    def close(self):
        self.server.close()
        for conn in self.held:
            conn.close()


class CellNoResponseTests(Base):
    """명세 §4.4 칸 검사 무응답 규칙(ChatGPT·Claude 같게)."""

    def real_silent(self):
        calls = []
        silent = SilentDetection()
        self.addCleanup(silent.close)

        def counted(*args, **kwargs):
            calls.append(kwargs["request_id"])
            return detection_client.check(*args, **kwargs)
        gateway.check_detection = counted
        patcher = mock.patch.object(gateway, "DETECTION_BASE_URL", silent.url)
        patcher.start()
        self.addCleanup(patcher.stop)
        return calls

    def xlsx(self):
        return xlsx_bytes([(f"A{i}", f"이름 {i}") for i in range(1, 6)] + [(f"B{i}", f"메모 {i}") for i in range(1, 6)])

    def test_silent_detection_waits_once_for_many_cells(self):
        for name, send in (("chatgpt", lambda: self.chatgpt_put(self.xlsx(), "s.xlsx")),
                           ("claude", lambda: self.claude_upload(("s.xlsx", XLSX, self.xlsx())))):
            for action in ("block", "forward"):
                with self.subTest(path=name, action=action):
                    self.reports.clear()
                    calls = self.real_silent()
                    with mock.patch.object(gateway, "GATEWAY_SETTINGS", FixedSettings(file=(action, 0.3))):
                        started = time.monotonic()
                        flow = send()
                        elapsed = time.monotonic() - started
                    self.assertEqual(len(calls), 1)            # 파일 전체 1회, 칸 10개는 호출 없음
                    self.assertLess(elapsed, 0.9)              # 대기 1회(0.3초)
                    if action == "block":
                        self.assert_blocked(flow, 503)
                        self.assertEqual([(r.result.action, r.result.failure_policy) for r in self.reports],
                                         [("block", "FAIL_CLOSED")])
                    else:
                        self.assertIsNone(flow.response)
                        self.assertEqual(flow.request.raw_content, self.sent)
                        self.assertEqual([(r.result.action, r.result.failure_policy, r.result.degraded_reason)
                                          for r in self.reports], [("allow_degraded", "FAIL_OPEN", "TIMEOUT")])

    def docx_with_down_cell(self):
        return docx_bytes(f"계좌 {CELL_ONLY}", f"{DOWN} 문단", f"끝 계좌 {CELL_ONLY}")

    def test_forward_applies_earlier_masks_and_records_degraded_mask(self):
        self.use_settings(file=("forward", 2.0))
        for name, send in (("chatgpt", lambda: self.chatgpt_put(self.docx_with_down_cell(), "a.docx")),
                           ("claude", lambda: self.claude_upload(("a.docx", DOCX, self.docx_with_down_cell())))):
            with self.subTest(path=name):
                self.calls.clear()
                self.reports.clear()
                flow = send()
                self.assertIsNone(flow.response)
                text = self.forwarded_docx_text(flow)
                self.assertEqual(text.count(CELL_ONLY), 1)     # 앞 칸만 가림, 뒤 칸은 검사하지 못해 그대로
                self.assertIn("[KR_BANK_ACCOUNT]", text.splitlines()[0])
                self.assertIn(CELL_ONLY, text.splitlines()[2])
                self.assertEqual(len(self.calls), 3)          # 파일 전체 + 칸 0(가림) + 칸 1(무응답). 칸 2는 호출 없음
                report, = self.reports
                self.assertEqual((report.result.action, report.result.degraded, report.result.failure_policy,
                                  report.result.reason_code, report.applied),
                                 ("mask_full", True, "FAIL_OPEN", "detection_timeout", True))
                self.assertEqual(flow.metadata["gw_action"], "mask")
                self.assertTrue(flow.metadata["gw_degraded"])

    def test_block_rule_blocks_on_the_unanswered_cell(self):
        for name, send in (("chatgpt", lambda: self.chatgpt_put(self.docx_with_down_cell(), "a.docx")),
                           ("claude", lambda: self.claude_upload(("a.docx", DOCX, self.docx_with_down_cell())))):
            with self.subTest(path=name):
                self.calls.clear()
                self.reports.clear()
                flow = send()
                self.assert_blocked(flow, 503)
                self.assertEqual(len(self.calls), 3)
                report, = self.reports
                self.assertEqual((report.result.action, report.result.failure_policy, report.result.reason_code),
                                 ("block", "FAIL_CLOSED", "detection_timeout"))

    def test_whole_file_mask_with_an_unanswered_cell_is_503_even_with_forward(self):
        self.use_settings(file=("forward", 2.0))
        data = docx_bytes(f"주민번호 {RRN}", f"{DOWN} 문단")
        for name, send in (("chatgpt", lambda: self.chatgpt_put(data, "a.docx")),
                           ("claude", lambda: self.claude_upload(("a.docx", DOCX, data)))):
            with self.subTest(path=name):
                self.reports.clear()
                flow = send()
                self.assert_blocked(flow, 503)
                report, = self.reports
                self.assertEqual((report.result.action, report.http_status, report.result.reason_code),
                                 ("block", 503, "detection_timeout"))


class MultimodalPromptTests(Base):
    """명세 §4.9: 이미지가 붙은 ChatGPT 질문(multimodal_text)의 글자도 검사·가림."""

    def message(self, text):
        return {"id": "m1", "author": {"role": "user"}, "content": {"content_type": "multimodal_text", "parts": [
            {"content_type": "image_asset_pointer", "asset_pointer": "sediment://file_00000000aaaa8206bbbb000000000001",
             "size_bytes": 1234, "width": 10, "height": 10},
            text]}}

    def test_question_text_is_inspected_and_masked(self):
        flow = self.chatgpt_conversation(self.message(f"이 사진 속 주민번호 {RRN} 확인"))
        self.assertEqual([c["text"] for c in self.calls], [f"이 사진 속 주민번호 {RRN} 확인"])
        self.assertEqual(self.calls[0]["source"]["content_type"], "prompt")
        self.assertIsNone(flow.response)
        sent = json.loads(flow.request.get_text())
        parts = sent["messages"][0]["content"]["parts"]
        self.assertEqual(parts[1], "이 사진 속 주민번호 [KR_RRN] 확인")
        self.assertEqual(parts[0]["asset_pointer"], "sediment://file_00000000aaaa8206bbbb000000000001")
        report, = self.reports
        self.assertEqual((report.result.action, report.context["contentType"]), ("mask_full", "prompt"))
        self.assertEqual(report.context["attachedFileIds"], [gateway.canonical_file_ref(FILE_ID)])

    def test_confidential_text_with_an_image_is_blocked(self):
        flow = self.chatgpt_conversation(self.message(f"{SECRET} 요약"))
        self.assert_blocked(flow, 403)


class LatencyTests(Base):
    """사용자 결정(10/10): 처리 시간 = Gateway 자체 처리(요청 시작 → request() 끝). LLM 응답 시간은 빠진다."""

    def setUp(self):
        super().setUp()
        gateway.report_dashboard_event = self._orig[1]     # 실제 이벤트 만들기
        self.env = mock.patch.dict(os.environ, {"DASHBOARD_EVENT_URL": "http://127.0.0.1:9/api/events/ingest",
                                                "DASHBOARD_INGEST_TOKEN": "test-token-not-used",
                                                "DASHBOARD_REPORT_ALLOW_EVENTS": "true"})
        self.env.start()
        self.addCleanup(self.env.stop)
        self.queued, self.extension = [], []
        for target, sink in (("enqueue", self.queued),):
            patcher = mock.patch.object(gateway._gateway_event_outbox, target,
                                        side_effect=lambda event, sink=sink: sink.append(event) or SimpleNamespace(status="queued"))
            patcher.start()
            self.addCleanup(patcher.stop)
        patcher = mock.patch.object(gateway, "send_extension_event", side_effect=self.extension.append)
        patcher.start()
        self.addCleanup(patcher.stop)
        latencies = mock.patch.object(gateway, "_gateway_latencies_ms", gateway.deque(maxlen=50))
        latencies.start()
        self.addCleanup(latencies.stop)

    def test_slow_upstream_is_not_in_processing_time(self):
        flow = self.finish(self.chatgpt_prompt("안녕하세요"), upstream_seconds=0.6)
        event, = self.queued
        self.assertLess(event["processingMs"], 300)
        self.assertGreaterEqual(flow.metadata["gw_round_trip_ms"], 600)
        self.assertEqual(flow.metadata["gw_latency_ms"], event["processingMs"])
        self.assertEqual(list(gateway._gateway_latencies_ms), [event["processingMs"]])   # 평균·p95의 표본
        snapshot = gateway._gateway_metrics_snapshot()
        self.assertEqual((snapshot["avg_latency_ms"], snapshot["p95_latency_ms"]),
                         (round(event["processingMs"], 2), round(float(event["processingMs"]), 2)))

    def test_gateway_own_waiting_is_in_processing_time(self):
        gateway.check_detection = fake_detection(self.calls, delay=0.25)
        flow = self.finish(self.chatgpt_prompt("안녕하세요"), upstream_seconds=0.3)
        event, = self.queued
        self.assertGreaterEqual(event["processingMs"], 250)
        self.assertLess(event["processingMs"], flow.metadata["gw_round_trip_ms"] - 250)

    def test_blocked_request_measures_until_the_gateway_answer(self):
        flow = self.chatgpt_prompt(f"{SECRET} 요약")
        time.sleep(0.3)                                    # 응답 훅이 늦게 와도
        self.finish(flow)
        event, = self.queued
        self.assertEqual(event["finalAction"], "BLOCK")
        self.assertLess(event["processingMs"], 250)
        self.assertGreaterEqual(flow.metadata["gw_round_trip_ms"], 300)

    def test_pass_through_requests_are_not_latency_samples(self):
        # 10/11 검토: 정적 파일·폴링처럼 Gateway가 아무것도 하지 않은 요청(0.0x ms)이 표본에 섞이면
        # 가장 가까운 순위 p95가 약 0이 되어 검사 비용이 보이지 않는다. 집계(요청 수·완료 수)는 모든 요청
        gateway.check_detection = fake_detection(self.calls, delay=0.25)
        before = gateway._gateway_metrics_snapshot()
        for index in range(60):
            req = tutils.treq(host="cdn.oaistatic.com", port=443, scheme=b"https", method=b"GET",
                              path=f"/assets/a{index}.js".encode(), content=b"")
            self.finish(self.run_flow(req))
        poll = tutils.treq(host="chatgpt.com", port=443, scheme=b"https", method=b"GET",
                           path=b"/backend-api/conversations?offset=0", content=b"")
        self.finish(self.run_flow(poll))
        prompt = self.finish(self.chatgpt_prompt("회의 요약"))
        blocked = self.finish(self.chatgpt_prompt(f"{SECRET} 요약"))
        after = gateway._gateway_metrics_snapshot()
        self.assertEqual(after["total_requests"] - before["total_requests"], 63)
        self.assertEqual(after["completed_requests"] - before["completed_requests"], 63)
        self.assertEqual(list(gateway._gateway_latencies_ms),
                         [prompt.metadata["gw_latency_ms"], blocked.metadata["gw_latency_ms"]])
        self.assertGreaterEqual(after["p95_latency_ms"], 250)
        self.assertGreaterEqual(after["avg_latency_ms"], 250)


class EndToEndLatencyTests(unittest.TestCase):
    """실제 mitmproxy(DumpMaster)에 gateway 애드온을 올리고 느린 upstream(1초)을 둔다. 지표 지연은 Gateway 처리만."""

    def test_proxy_with_a_slow_upstream(self):
        from mitmproxy import options
        from mitmproxy.tools.dump import DumpMaster

        class SlowUpstream(BaseHTTPRequestHandler):
            def do_POST(self):
                self.rfile.read(int(self.headers.get("content-length") or 0))
                time.sleep(1.0)
                body = b'{"ok":true}'
                self.send_response(200)
                self.send_header("content-type", "application/json")
                self.send_header("content-length", str(len(body)))
                self.end_headers()
                self.wfile.write(body)

            def log_message(self, *args):
                pass
        upstream = ThreadingHTTPServer(("127.0.0.1", 0), SlowUpstream)
        threading.Thread(target=upstream.serve_forever, daemon=True).start()
        self.addCleanup(upstream.server_close)
        self.addCleanup(upstream.shutdown)

        class Redirect:
            """gateway 다음 순서: 검사가 끝난 요청을 로컬 upstream으로 보낸다. 응답 머리에서 원래 주소로 되돌린다."""

            def request(self, flow):
                flow.metadata["original_host"] = flow.request.host
                flow.request.host, flow.request.port, flow.request.scheme = "127.0.0.1", upstream.server_address[1], "http"

            def responseheaders(self, flow):
                flow.request.host, flow.request.port = flow.metadata["original_host"], 80

        calls = []
        port_probe = socket.socket()
        port_probe.bind(("127.0.0.1", 0))
        port = port_probe.getsockname()[1]
        port_probe.close()
        state = {}

        async def run_master():
            opts = options.Options(listen_host="127.0.0.1", listen_port=port)
            master = DumpMaster(opts, with_termlog=False, with_dumper=False)
            master.addons.add(gateway, Redirect())
            state["master"] = master
            await master.run()

        loop = asyncio.new_event_loop()
        thread = threading.Thread(target=lambda: loop.run_until_complete(run_master()), daemon=True)
        latencies = gateway.deque(maxlen=50)
        with mock.patch.object(gateway, "check_detection", fake_detection(calls)), \
                mock.patch.object(gateway, "send_to_next_module", lambda event: None), \
                mock.patch.object(gateway, "_gateway_latencies_ms", latencies):
            thread.start()
            deadline = time.monotonic() + 15
            while time.monotonic() < deadline:
                with contextlib.suppress(OSError), socket.create_connection(("127.0.0.1", port), timeout=0.2):
                    break
                time.sleep(0.1)
            opener = urllib.request.build_opener(urllib.request.ProxyHandler({"http": f"http://127.0.0.1:{port}"}))
            request = urllib.request.Request(f"http://claude.ai{COMPLETION_PATH}", method="POST",
                                             data=json.dumps({"prompt": "회의 요약"}).encode(),
                                             headers={"content-type": "application/json"})
            started = time.monotonic()
            with opener.open(request, timeout=10) as response:
                self.assertEqual(response.status, 200)
                response.read()
            client_seconds = time.monotonic() - started
            time.sleep(0.2)
            loop.call_soon_threadsafe(state["master"].shutdown)
            thread.join(10)
        self.assertEqual(len(calls), 1)                    # 질문 검사는 실제로 했다
        self.assertGreaterEqual(client_seconds, 1.0)       # 사용자는 LLM 응답까지 1초 넘게 기다렸고
        self.assertEqual(len(latencies), 1)
        self.assertLess(latencies[0], 500)                 # 지표 지연은 Gateway 처리만(이전: 1,000ms 넘음)


class ClaudeRpcShapeTests(Base):
    """사용자 결정(10/10): 새 claude.ai rpc 요청은 구조만 기록(값 없음), 전달에 영향 없음."""

    PATH = f"/api/organizations/{ORG}/claudeai-rpc/anthropic.claudeai.v1.ConversationService/PerformAction"

    def setUp(self):
        super().setUp()
        gateway._claude_rpc_shape_counts.clear()
        gateway._claude_rpc_shape_window[0] = 0.0

    def rpc(self, body: bytes, content_type=b"application/json", path=None):
        req = tutils.treq(host="claude.ai", port=443, scheme=b"https", method=b"POST", path=(path or self.PATH).encode(),
                          content=body, headers=((b"content-type", content_type),))
        out = io.StringIO()
        with contextlib.redirect_stdout(out):
            flow = self.run_flow(req)
        return flow, [line for line in out.getvalue().splitlines() if line.startswith("[CLAUDE-RPC]")]

    def body(self):
        return json.dumps({
            "conversationId": CONV,
            "action": {"type": "send_message", "prompt": f"내 주민번호는 {RRN}이고 메일은 {EMAIL}, 전화 {PHONE}",
                       "attachments": [{"name": f"{RRN}.txt", "text": "기밀 메모 내용"}], "count": 3, "flag": True,
                       "extra": None},
            "metadata": {RRN: 1, "홍길동": "x"},
        }, ensure_ascii=False).encode()

    def test_logs_structure_without_values(self):
        flow, lines = self.rpc(self.body())
        self.assertIsNone(flow.response)
        self.assertEqual(flow.request.raw_content, self.sent)   # 전달 본문 그대로
        self.assertEqual((self.calls, self.reports), ([], []))
        line, = lines
        for value in (RRN, EMAIL, PHONE, "주민번호", "기밀 메모", "홍길동", CONV, ORG, "send_message"):
            self.assertNotIn(value, line)
        record = json.loads(line[len("[CLAUDE-RPC] "):])
        self.assertEqual(record["method"], "anthropic.claudeai.v1.ConversationService/PerformAction")
        self.assertEqual((record["content_type"], record["bytes"], record["format"]),
                         ("application/json", len(self.body()), "json"))
        shape = record["shape"]
        self.assertEqual(shape["conversationId"], f"str:{len(CONV)}")
        self.assertEqual(shape["action"]["prompt"], f"str:{len(json.loads(self.body())['action']['prompt'])}")
        self.assertEqual(shape["action"]["attachments"][0], "<list:1>")
        self.assertEqual((shape["action"]["count"], shape["action"]["flag"], shape["action"]["extra"]),
                         ("num", "bool", "null"))
        self.assertEqual(shape["metadata"], {"<map:2>": "num"})    # 값 같은 키가 있는 객체는 키 수만(10/11)

    def test_data_like_keys_are_not_logged(self):
        # 10/11 검토: 파일 이름·점·붙임표 들어간 이름·메시지 번호 같은 키는 값이다. 필드 이름 모양만 남긴다
        data_keys = ["salary_2026_hong.xlsx", "kim.cheolsu", "hong-gildong", "msg_01AbCdEf", "x9Kq2LmN"]
        body = json.dumps({
            "conversationUuid": CONV,
            "parent_message_uuid": CONV,
            "sha256Hash": "0" * 64,
            "files": {key: {"size": 1} for key in data_keys},
            "people": {"kim.cheolsu": 1, "hong_gildong": 2},    # 값 같은 키가 하나라도 있으면 다른 키도 숨김
            "byId": {"msg_01AbCdEf": "x"},
            "many": {f"field{i}": i for i in range(gateway._CLAUDE_RPC_SHAPE_MAX_KEYS + 1)},
        }).encode()
        _, lines = self.rpc(body)
        line, = lines
        for key in data_keys + ["hong_gildong", "field0"]:
            self.assertNotIn(key, line)
        shape = json.loads(line[len("[CLAUDE-RPC] "):])["shape"]
        self.assertEqual(set(shape), {"conversationUuid", "parent_message_uuid", "sha256Hash", "files", "people",
                                      "byId", "many"})
        self.assertEqual(shape["files"], {f"<map:{len(data_keys)}>": {"size": "num"}})
        self.assertEqual(shape["people"], {"<map:2>": "num"})
        self.assertEqual(shape["byId"], {"<map:1>": "str:1"})
        self.assertEqual(shape["many"], {f"<map:{gateway._CLAUDE_RPC_SHAPE_MAX_KEYS + 1}>": "num"})

    def test_rate_limit_per_method_and_switch(self):
        for _ in range(gateway.CLAUDE_RPC_SHAPE_PER_HOUR):
            self.assertEqual(len(self.rpc(b"{}")[1]), 1)
        self.assertEqual(self.rpc(b"{}")[1], [])
        other = self.PATH.replace("PerformAction", "StreamTimeline")
        self.assertEqual(len(self.rpc(b"{}", path=other)[1]), 1)     # 메서드마다 따로 센다
        gateway._claude_rpc_shape_counts.clear()
        with mock.patch.dict(os.environ, {"GATEWAY_LOG_CLAUDE_RPC_SHAPE": "0"}):
            self.assertEqual(self.rpc(b"{}")[1], [])

    def test_binary_and_envelope_bodies(self):
        # 10/11: application/proto는 JSON 모양(non_json) 대신 protobuf 잎 구조(값 없음). 자세한 시험은 test_gateway_dedupe_proto
        _, lines = self.rpc(b"\x0a\x05hello\x12\x0e" + RRN.encode(), content_type=b"application/proto")
        record = json.loads(lines[0][len("[CLAUDE-RPC] "):])
        self.assertEqual(record["format"], "proto")
        self.assertNotIn("shape", record)
        self.assertEqual([(leaf["path"], leaf["kind"], leaf["len"]) for leaf in record["leaves"]],
                         [("1", "text", 5), ("2", "text", len(RRN))])
        self.assertNotIn(RRN, lines[0])
        self.assertNotIn("hello", lines[0])
        frame = json.dumps({"prompt": RRN}).encode()
        _, lines = self.rpc(b"\x00" + len(frame).to_bytes(4, "big") + frame, content_type=b"application/connect+json")
        record = json.loads(lines[0][len("[CLAUDE-RPC] "):])
        self.assertTrue(record["format"].startswith("envelope:"))
        self.assertEqual(record["shape"], {"prompt": f"str:{len(RRN)}"})

    def test_deep_and_wide_bodies_are_bounded(self):
        deep = {"a": 1}
        for _ in range(50):
            deep = {"n": deep}
        wide = {f"k{i}": i for i in range(500)}
        _, lines = self.rpc(json.dumps({"deep": deep, "wide": wide, "list": list(range(10000))}).encode())
        self.assertLess(len(lines[0]), 20000)
        record = json.loads(lines[0][len("[CLAUDE-RPC] "):])
        self.assertEqual(record["shape"]["wide"], {"<map:500>": "num"})   # 키가 많으면 지도로 보고 키 수만

    def test_logger_failure_never_changes_forwarding(self):
        with mock.patch.object(gateway, "_claude_rpc_body_shape", side_effect=RuntimeError("boom")):
            flow, lines = self.rpc(self.body())
        self.assertIsNone(flow.response)
        self.assertEqual(flow.request.raw_content, self.sent)


class EventShapeTests(Base):
    """명세 §4.7 이벤트 모양(대시보드 Security Event v1 그대로). 실제 report_dashboard_event·mapper로 만든다."""

    def setUp(self):
        super().setUp()
        gateway.report_dashboard_event = self._orig[1]
        self.env = mock.patch.dict(os.environ, {"DASHBOARD_EVENT_URL": "http://127.0.0.1:9/api/events/ingest",
                                                "DASHBOARD_INGEST_TOKEN": "test-token-not-used"})
        self.env.start()
        self.addCleanup(self.env.stop)
        os.environ.pop("DASHBOARD_REPORT_ALLOW_EVENTS", None)   # 기본: 보통 허용 이벤트는 보내지 않음

    def events(self, flow):
        return flow.metadata.get("gw_dashboard_events", [])

    def collect(self):
        rows = []
        self.use_settings(file=("block", 2.0), image="block")
        rows.append(("uninspectable block", self.events(self.chatgpt_put(png_bytes(), "p.png"))))
        with mock.patch.object(gateway, "GATEWAY_SETTINGS", FixedSettings(image="allow")):
            rows.append(("uninspectable allow", self.events(self.chatgpt_put(png_bytes(), "p.png"))))
            rows.append(("too large block", self.events(self.chatgpt_put(too_large_text(), "b.txt"))))
        with mock.patch.object(gateway, "GATEWAY_SETTINGS", FixedSettings(file=("forward", 2.0))):
            rows.append(("too large allow", self.events(self.chatgpt_put(too_large_text(), "b.txt"))))
            rows.append(("partial mask then unanswered", self.events(self.chatgpt_put(
                docx_bytes(f"계좌 {CELL_ONLY}", f"{DOWN} 문단"), "a.docx"))))
        gateway.check_detection = lambda *a, **k: failure_result("TIMEOUT", k["failure_policy"], k["request_id"])
        rows.append(("delay block", self.events(self.chatgpt_put(b"memo", "m.txt"))))
        with mock.patch.object(gateway, "GATEWAY_SETTINGS", FixedSettings(file=("forward", 2.0))):
            rows.append(("delay forward", self.events(self.chatgpt_put(b"memo", "m.txt"))))
        return rows

    def test_shapes(self):
        expected = {
            "uninspectable block": ("BLOCK", "UNAVAILABLE", False, "file_uninspectable_image", {"applied": True, "httpStatus": 403}),
            "uninspectable allow": ("ALLOW", "UNAVAILABLE", False, "file_uninspectable_image", None),
            "too large block": ("BLOCK", "UNAVAILABLE", False, "file_too_large", {"applied": True, "httpStatus": 403}),
            "too large allow": ("ALLOW", "UNAVAILABLE", False, "file_too_large", None),
            "delay block": ("BLOCK", "UNAVAILABLE", True, "detection_timeout", {"applied": True, "httpStatus": 503}),
            "delay forward": ("ALLOW", "UNAVAILABLE", True, "detection_timeout", None),
            "partial mask then unanswered": ("MASK", "UNAVAILABLE", True, "detection_timeout", {"applied": True}),
        }
        rows = self.collect()
        dump = os.environ.get("GW_BUNDLE_EVENT_DUMP")
        for name, events in rows:
            with self.subTest(case=name):
                event, = events
                self.assertEqual((event["finalAction"], event["detectionOutcome"], event["degraded"],
                                  event["reasonCode"], event.get("enforcement")), expected[name])
                self.assertEqual(event["detections"], [])
                self.assertNotIn("notify", event)
                if not event["degraded"]:
                    self.assertNotIn("degradedReason", event)
                    self.assertNotIn("failurePolicy", event)
                if dump:
                    with open(dump, "a", encoding="utf-8") as handle:
                        handle.write(json.dumps({"case": f"bundle.{name.replace(' ', '_')}", "event": event},
                                                ensure_ascii=False) + "\n")

    @unittest.skipUnless(shutil.which("node") and (DASHBOARD_SERVER_DIR / "validation.mjs").exists(),
                         "node or dashboard server copy not available")
    def test_dashboard_validation_accepts_every_shape(self):
        rows = self.collect()
        events = [event for _, items in rows for event in items]
        module = (DASHBOARD_SERVER_DIR / "validation.mjs").as_uri()
        script = (f"import {{ validateCanonicalEvent }} from '{module}';\n"
                  "import { readFileSync } from 'node:fs';\n"
                  "const events = JSON.parse(readFileSync(0, 'utf8'));\n"
                  "const out = events.map((event) => { try { const v = validateCanonicalEvent(event);"
                  " return { ok: true, outcome: v.detectionOutcome, action: v.finalAction }; }"
                  " catch (error) { return { ok: false, error: String(error.message), path: error.path }; } });\n"
                  "console.log(JSON.stringify(out));\n")
        folder = Path(tempfile.mkdtemp(prefix="gw-node-"))
        self.addCleanup(shutil.rmtree, folder, True)
        (folder / "validate.mjs").write_text(script, encoding="utf-8")
        result = subprocess.run(["node", str(folder / "validate.mjs")], input=json.dumps(events), capture_output=True,
                                text=True, encoding="utf-8", timeout=60)
        self.assertEqual(result.returncode, 0, result.stderr)
        verdicts = json.loads(result.stdout)
        self.assertEqual(len(verdicts), 7)
        self.assertTrue(all(v["ok"] for v in verdicts), verdicts)
        self.assertTrue(all(v["outcome"] == "UNAVAILABLE" for v in verdicts))


class ExtensionNotifyTests(Base):
    """명세 §4.8: 확장용 사본에만 notify. 대시보드 사본에는 없음. 웹소켓 보기에 사유·종류·표시."""

    def setUp(self):
        super().setUp()
        gateway.report_dashboard_event = self._orig[1]
        self.env = mock.patch.dict(os.environ, {"DASHBOARD_EVENT_URL": "http://127.0.0.1:9/api/events/ingest",
                                                "DASHBOARD_INGEST_TOKEN": "test-token-not-used",
                                                "DASHBOARD_REPORT_ALLOW_EVENTS": "true"})
        self.env.start()
        self.addCleanup(self.env.stop)
        self.queued, self.extension = [], []
        patcher = mock.patch.object(gateway._gateway_event_outbox, "enqueue",
                                    side_effect=lambda event: self.queued.append(event) or SimpleNamespace(status="queued"))
        patcher.start()
        self.addCleanup(patcher.stop)
        patcher = mock.patch.object(gateway, "send_extension_event", side_effect=self.extension.append)
        patcher.start()
        self.addCleanup(patcher.stop)

    def notify_for(self, send, **settings):
        self.queued.clear()
        self.extension.clear()
        getattr(gateway, "_dedupe_seen", {}).clear()   # 10/11: 같은 파일 번호를 쓰는 서로 다른 상황(재시도 아님)
        with mock.patch.object(gateway, "GATEWAY_SETTINGS", FixedSettings(**settings)):
            self.finish(send())
        self.assertTrue(self.extension)
        self.assertTrue(all("notify" not in event for event in self.queued))
        return [event["notify"] for event in self.extension]

    def test_notify_flags(self):
        self.assertEqual(self.notify_for(lambda: self.chatgpt_prompt(f"{SECRET} 요약")), [True])            # BLOCK
        self.assertEqual(self.notify_for(lambda: self.chatgpt_prompt(f"번호 {RRN}")), [True])               # MASK
        self.assertEqual(self.notify_for(lambda: self.chatgpt_prompt("안녕"), notify=True), [False])         # 보통 ALLOW
        self.assertEqual(self.notify_for(lambda: self.chatgpt_put(png_bytes(), "p.png")), [False])
        self.assertEqual(self.notify_for(lambda: self.chatgpt_put(png_bytes(), "p.png"), notify=True), [True])
        self.assertEqual(self.notify_for(lambda: self.chatgpt_put(too_large_text(), "b.txt"),
                                         file=("forward", 2.0), notify=True), [True])
        gateway.check_detection = lambda *a, **k: failure_result("TIMEOUT", k["failure_policy"], k["request_id"])
        self.assertEqual(self.notify_for(lambda: self.chatgpt_prompt("안녕"), notify=True), [True])          # 지연 ALLOW
        self.assertEqual(self.notify_for(lambda: self.chatgpt_prompt("안녕"), notify=False), [False])

    def test_extension_view(self):
        base = {"eventId": "evt-1", "clientId": "c1", "targetLlm": "ChatGPT", "contentType": "file"}
        view = extension_event_view({**base, "finalAction": "ALLOW", "detectionOutcome": "UNAVAILABLE",
                                     "degraded": False, "reasonCode": "file_uninspectable_image", "notify": True})
        self.assertEqual((view["reasonCode"], view["contentType"], view["notify"], view["action"]),
                         ("file_uninspectable_image", "file", True, "ALLOW"))
        self.assertNotIn("inspectionMessage", view)           # 검사 불가는 장애가 아니다
        view = extension_event_view({**base, "finalAction": "BLOCK", "detectionOutcome": "UNAVAILABLE",
                                     "degraded": True, "reasonCode": "detection_timeout", "notify": True})
        self.assertIn("inspectionMessage", view)
        view = extension_event_view({**base, "finalAction": "BLOCK", "reasonCode": "Bad Code!", "contentType": "x",
                                     "notify": "yes"})
        for key in ("reasonCode", "contentType", "notify"):
            self.assertNotIn(key, view)


class CallSiteSettingsTests(Base):
    """명세 §4.4: 모든 Detection 호출이 요청 시작 때 고정한 설정의 질문/파일 규칙을 쓴다."""

    def setUp(self):
        super().setUp()
        self.use_settings(prompt=("block", 1.25), file=("forward", 3.5))

    def rules(self):
        return [(c["timeout"], c["failure_policy"]) for c in self.calls]

    PROMPT = (1.25, "FAIL_CLOSED")
    FILE = (3.5, "FAIL_OPEN")

    def test_each_call_site(self):
        self.chatgpt_prompt("안녕")
        self.assertEqual(self.rules(), [self.PROMPT])
        self.calls.clear()
        self.claude_completion({"prompt": "요약", "attachments": [{"file_name": "a.txt", "extracted_content": "메모"}]})
        self.assertEqual(self.rules(), [self.PROMPT, self.FILE, self.FILE])
        self.calls.clear()
        self.chatgpt_put(docx_bytes("안내", f"번호 {RRN}"), "a.docx")
        self.assertEqual(set(self.rules()), {self.FILE})
        self.assertGreater(len(self.calls), 2)                      # 파일 전체 + 칸
        self.calls.clear()
        self.claude_upload(("a.docx", DOCX, docx_bytes("안내", f"번호 {RRN}")))
        self.assertEqual(set(self.rules()), {self.FILE})
        self.assertGreater(len(self.calls), 2)
        self.calls.clear()
        req = tutils.treq(host="api.openai.com", port=443, scheme=b"https", method=b"POST", path=b"/v1/responses",
                          content=json.dumps({"model": "m", "input": "안녕"}).encode(),
                          headers=((b"content-type", b"application/json"),))
        self.run_flow(req)
        self.assertEqual(self.rules(), [self.PROMPT])
        self.calls.clear()
        flow = tflow.tflow()
        flow.metadata["gw_settings"] = gateway.GATEWAY_SETTINGS.current()
        gateway._inspect(flow, "image", request_id="r1", normalized_text="x", source={})
        self.assertEqual(self.rules(), [self.FILE])                 # 모르는 종류 = 파일 규칙

    def test_settings_are_fixed_at_request_start(self):
        first = FixedSettings(prompt=("forward", 0.5), file=("block", 2.0))
        later = FixedSettings(prompt=("forward", 9.0), file=("forward", 9.0))
        calls = []
        default = fake_detection(calls)

        def switching(*args, **kwargs):
            gateway.GATEWAY_SETTINGS = later                        # 지표 응답으로 도중에 바뀜
            return default(*args, **kwargs)
        with mock.patch.object(gateway, "GATEWAY_SETTINGS", first):
            gateway.check_detection = switching
            self.claude_completion({"prompt": "요약", "attachments": [{"extracted_content": "메모"}]})
        self.assertEqual([(c["timeout"], c["failure_policy"]) for c in calls],
                         [(0.5, "FAIL_OPEN"), (2.0, "FAIL_CLOSED")])

    def test_settings_applied_from_a_metrics_response(self):
        holder = SettingsHolder(env_defaults({}), cache_path=None, metrics_url=None, logger=lambda m: None)
        with mock.patch.object(gateway, "GATEWAY_SETTINGS", holder):
            self.chatgpt_prompt("하나")
            response = SimpleNamespace(status="sent", http_status=202, payload={
                "accepted": True, "gateway_settings": settings_doc(prompt=("block", 1.5), version=5)})
            sent = []
            publisher = gateway.MetricsPublisher(gateway._publisher_snapshot, sender=lambda s: sent.append(s) or response,
                                                 heartbeat=30, min_interval=0.05, logger=lambda m: None,
                                                 on_payload=holder.offer, on_status=holder.on_metrics_status)
            holder.set_on_change(publisher.notify)
            publisher.start()
            deadline = time.monotonic() + 5
            while len(sent) < 2 and time.monotonic() < deadline:
                time.sleep(0.02)
            publisher.stop()
            self.chatgpt_prompt("둘")
            bad = settings_doc(prompt=("block", 99), version=6, revision=REVISION_2)
            holder.offer({"accepted": True, "gateway_settings": bad})
            self.chatgpt_prompt("셋")
        self.assertEqual(self.rules(), [(0.5, "FAIL_OPEN"), (1.5, "FAIL_CLOSED"), (1.5, "FAIL_CLOSED")])
        self.assertEqual(sent[1]["settings_applied"]["version"], 5)   # 적용 보고(다음 지표)
        self.assertEqual(holder.report()["rejected"], {"version": 6, "revision": REVISION_2,
                                                       "field": "timeout.prompt.seconds"})

    def test_only_the_inspect_helper_and_the_gemini_upload_call_detection(self):
        tree = ast.parse((ROOT / "gateway.py").read_text(encoding="utf-8"))
        owners = []
        for function in ast.walk(tree):
            if isinstance(function, (ast.FunctionDef, ast.AsyncFunctionDef)):
                for node in ast.walk(function):
                    if isinstance(node, ast.Call) and getattr(node.func, "id", None) == "check_detection":
                        owners.append(function.name)
        # _verify_masked(2026-10-11, 가린 파일 재검사)는 일부러 _inspect를 거치지 않는다: 글자 상한 지름길·무응답 넘기기·
        # FAIL_OPEN이 '검사 없는 통과'를 돌려주므로 항상 FAIL_CLOSED·파일 대기 시간으로 직접 부른다(명세 §3.5.1)
        self.assertEqual(sorted(owners), ["_inspect", "_verify_masked", "handle_gemini_resumable_upload"])


class DeployOrderTests(unittest.TestCase):
    """명세 §12 근거: 새 gateway.py는 옛 보조 모듈과 함께 불러지지 않는다(그래서 멈춘 동안 모든 파일을 바꾼다)."""

    OLD_METRICS_DELIVERY = (
        "class MetricsPublisher:\n"
        "    def __init__(self, snapshot_factory, *, sender=None, heartbeat=10.0, min_interval=1.0, logger=print):\n"
        "        pass\n"
    )

    def staged(self):
        folder = Path(tempfile.mkdtemp(prefix="gw-deploy-"))
        self.addCleanup(shutil.rmtree, folder, True)
        for path in ROOT.glob("*.py"):
            shutil.copy2(path, folder / path.name)
        shutil.copytree(ROOT / "signature_detector", folder / "signature_detector",
                        ignore=shutil.ignore_patterns("__pycache__"))
        return folder

    def import_gateway(self, folder, extra_env=None, code="import gateway; print('import ok')"):
        env = {**os.environ, "DASHBOARD_OUTBOX_PATH": str(folder / "outbox.sqlite3"),
               "SIGNATURE_REVIEW_DIR": str(folder / "review"), "PYTHONIOENCODING": "utf-8"}
        for key in ("DASHBOARD_EVENT_URL", "DASHBOARD_INGEST_TOKEN", "GATEWAY_SETTINGS_CACHE_PATH"):
            env.pop(key, None)
        env.update(extra_env or {})
        return subprocess.run([sys.executable, "-c", code], cwd=folder, env=env, capture_output=True, text=True,
                              encoding="utf-8", errors="replace", timeout=120)

    def test_new_gateway_with_the_old_metrics_module_fails_to_import(self):
        folder = self.staged()
        (folder / "gateway_metrics_delivery.py").write_text(self.OLD_METRICS_DELIVERY, encoding="utf-8")
        result = self.import_gateway(folder)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("TypeError", result.stderr)

    def test_all_new_files_import(self):
        result = self.import_gateway(self.staged())
        self.assertEqual(result.returncode, 0, result.stderr[-2000:])
        self.assertIn("import ok", result.stdout)
        self.assertIn("[SETTINGS] source=ENV_DEFAULT version=0", result.stdout)

    def test_environment_fallback_and_disk_cache_at_start(self):
        folder = self.staged()
        code = ("import gateway; s = gateway.GATEWAY_SETTINGS.current(); "
                "print('RESULT', s.source, s.version, s.prompt.action, s.file.action, s.file.seconds, s.image)")
        result = self.import_gateway(folder, {"GATEWAY_FILE_TIMEOUT_SECONDS": "3.5", "GATEWAY_UNINSPECTABLE_IMAGE": "block"},
                                     code)
        self.assertIn("RESULT ENV_DEFAULT 0 forward block 3.5 block", result.stdout, result.stderr[-2000:])
        dashboard = {"DASHBOARD_EVENT_URL": "http://127.0.0.1:19100/api/events/ingest",
                     "DASHBOARD_INGEST_TOKEN": "test-token-not-used",
                     "GATEWAY_SETTINGS_CACHE_PATH": str(folder / "cache" / "settings.json")}
        holder = SettingsHolder(env_defaults({}), cache_path=folder / "cache" / "settings.json",
                                metrics_url="http://127.0.0.1:19100/api/traffic/metrics", logger=lambda m: None)
        holder.offer({"gateway_settings": settings_doc(file=("forward", 4.5), version=9)})
        result = self.import_gateway(folder, dashboard, code)
        self.assertIn("RESULT CACHE 9 forward forward 4.5 allow", result.stdout, result.stderr[-2000:])
        result = self.import_gateway(folder, {**dashboard, "DASHBOARD_EVENT_URL": "http://127.0.0.1:9/api/events/ingest"},
                                     code)
        self.assertIn("RESULT ENV_DEFAULT 0", result.stdout)   # 다른 대시보드 주소의 사본은 쓰지 않음


if __name__ == "__main__":
    unittest.main()
