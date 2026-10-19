"""업로드 파일 분류·추출(gateway_file_classes.inspect_upload, 2026-10-10). 가짜 값만 쓴다.

명세 §4.5·§5·§11.1(이미지 분류·형식 분류·숨은 글자·PDF 분류·상한·손상 입력).
실행: python -m unittest tests.test_gateway_file_classes -v (Gateway 사본 폴더에서)
"""

from __future__ import annotations

import io
import json
import os
import random
import re
import subprocess
import sys
import time
import unittest
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from docx import Document  # noqa: E402
from openpyxl import Workbook  # noqa: E402

import gateway_file_classes  # noqa: E402
from gateway_file_classes import inspect_upload  # noqa: E402
from gateway_settings import file_max_chars  # noqa: E402
from tests.gw_support import (  # noqa: E402
    annotation_pdf, attachment_pdf, empty_pdf, form_field_pdf, gif_bytes, jpeg_bytes, mpo_bytes, nested_text_boxes,
    png_bytes, raw_docx, rewrite_zip, scanned_pdf, text_pdf, vector_only_pdf, webp_bytes, wrapper_chain, xfa_pdf,
    zip_bomb_declared, zip_bytes,
)

RRN = "900101-1234567"
SAMPLE_KO_PDF = Path(os.environ.get("GW_TEST_SAMPLE_KO_PDF", ROOT / "tests" / "fixtures" / "sample_ko.pdf"))
HANGUL = re.compile("[가-힣]")
MAX_BYTES = 20 * 1024 * 1024
MAX_CHARS = file_max_chars(2.0)
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def classify(data, *, max_bytes=MAX_BYTES, max_chars=MAX_CHARS, seconds=2.0, clock=time.monotonic, **kwargs):
    return inspect_upload(data, max_bytes=max_bytes, max_chars=max_chars, deadline=clock() + seconds, clock=clock,
                          **kwargs)


def docx_bytes(*paragraphs, header=None, footer=None, picture=False):
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if header:
        document.sections[0].header.paragraphs[0].text = header
    if footer:
        document.sections[0].footer.paragraphs[0].text = footer
    if picture:
        document.add_picture(io.BytesIO(png_bytes(4, 4)))
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def docx_with_body_xml(inner_xml: str) -> bytes:
    """본문 XML을 직접 넣은 DOCX(글상자 등 python-docx가 만들지 못하는 구조)."""
    base = docx_bytes()
    document = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:document xmlns:w="{W}" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        f'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
        f'xmlns:v="urn:schemas-microsoft-com:vml">'
        f'<w:body>{inner_xml}<w:sectPr/></w:body></w:document>'
    )
    return rewrite_zip(base, {"word/document.xml": document.encode("utf-8")})


def text_box_paragraph(text: str) -> str:
    box = f"<w:txbxContent><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:txbxContent>"
    return ("<w:p><w:r><mc:AlternateContent>"
            f"<mc:Choice Requires=\"wps\"><w:drawing><wps:txbx>{box}</wps:txbx></w:drawing></mc:Choice>"
            f"<mc:Fallback><w:pict><v:textbox>{box}</v:textbox></w:pict></mc:Fallback>"
            "</mc:AlternateContent></w:r></w:p>")


def xlsx_bytes(cells=(), title=None, media=False):
    workbook = Workbook()
    sheet = workbook.active
    if title:
        sheet.title = title
    for ref, value in cells:
        sheet[ref] = value
    stream = io.BytesIO()
    workbook.save(stream)
    data = stream.getvalue()
    return rewrite_zip(data, {"xl/media/image1.png": png_bytes()}) if media else data


class ImageTests(unittest.TestCase):
    def test_real_image_structures(self):
        for name, data in (("png", png_bytes()), ("jpeg", jpeg_bytes()), ("jpeg", jpeg_bytes(progressive=True)),
                           ("gif", gif_bytes()), ("webp", webp_bytes()), ("png", png_bytes(trailer=b"\x00" * 16))):
            result = classify(data, filename="photo.bin", mime="application/octet-stream")
            self.assertEqual((result.kind, result.group), (name, "image"), name)
            self.assertIsNone(result.text)

    def test_pil_images_when_available(self):
        try:
            from PIL import Image
        except ImportError:
            self.skipTest("PIL not installed")
        for fmt, kind in (("PNG", "png"), ("JPEG", "jpeg"), ("GIF", "gif"), ("WEBP", "webp")):
            stream = io.BytesIO()
            Image.new("RGB", (40, 30), (20, 120, 220)).save(stream, fmt)
            self.assertEqual(classify(stream.getvalue()).group, "image", fmt)
        stream = io.BytesIO()
        Image.new("RGB", (40, 30), (20, 120, 220)).save(stream, "JPEG", progressive=True)
        self.assertEqual(classify(stream.getvalue()).group, "image")
        for fmt in ("BMP", "TIFF"):
            stream = io.BytesIO()
            Image.new("RGB", (40, 30)).save(stream, fmt)
            self.assertEqual(classify(stream.getvalue()).group, "other", fmt)

    def test_large_jpeg_structure_check_is_fast(self):
        try:
            from PIL import Image
        except ImportError:
            self.skipTest("PIL not installed")
        noise = Image.frombytes("RGB", (4000, 3000), random.Random(7).randbytes(4000 * 3000 * 3))
        stream = io.BytesIO()
        noise.save(stream, "JPEG", quality=95)
        data = stream.getvalue()
        self.assertGreater(len(data), 10 * 1024 * 1024)
        started = time.perf_counter()
        result = classify(data)
        elapsed_ms = (time.perf_counter() - started) * 1000
        self.assertEqual(result.group, "image")
        self.assertLess(elapsed_ms, 50)

    def test_bypasses_are_inspected_or_rejected(self):
        bm_csv = f"BM,이름,주민번호\n홍길동,{RRN}\n".encode("utf-8")
        result = classify(bm_csv, filename="x.bmp", mime="image/bmp")
        self.assertEqual((result.kind, result.group), ("text", None))
        self.assertIn(RRN, result.text)
        named_png = f"주민번호 {RRN}".encode("utf-8")
        result = classify(named_png, filename="x.png", mime="image/png")
        self.assertEqual((result.kind, result.group), ("text", None))
        self.assertIn(RRN, result.text)
        svg = f'<svg xmlns="http://www.w3.org/2000/svg"><text x="1" y="9">RRN {RRN}</text></svg>'.encode()
        result = classify(svg, filename="x.svg", mime="image/svg+xml")
        self.assertEqual((result.group, result.detail), (None, "markup"))
        self.assertIn(RRN, result.text)
        cases = {
            "png+zip": (png_bytes(trailer=zip_bytes({"secret.txt": f"RRN {RRN}"})), "png_trailer"),
            "jpeg+pdf": (jpeg_bytes() + text_pdf(f"RRN {RRN}"), "jpeg_trailer"),
            "png magic + text": (b"\x89PNG\r\n\x1a\n" + f"주민번호 {RRN}".encode(), "broken_image"),
            "gif + text": (gif_bytes() + b"tail", "gif_trailer"),
            "webp size beyond data": (webp_bytes()[:-2], "broken_image"),
        }
        for name, (data, detail) in cases.items():
            result = classify(data, filename="x.png", mime="image/png")
            self.assertEqual((result.group, result.detail), ("other", detail), name)

    def test_multi_picture_jpeg_is_an_image(self):
        # 10/11 검토: Ultra HDR·이득 지도 사진(APP2 MPF + 끝 표지 뒤 둘째 JPEG)은 그림 여러 장이라 image
        for name, data in (("mpo", mpo_bytes()), ("mpo + zero padding", mpo_bytes(b"\x00" * 32)),
                           ("three jpegs", jpeg_bytes() + jpeg_bytes() + jpeg_bytes(progressive=True))):
            result = classify(data)
            self.assertEqual((result.kind, result.group, result.detail), ("jpeg", "image", "jpeg"), name)
        mp4 = b"\x00\x00\x00\x18ftypmp42\x00\x00\x00\x00mp42isom" + b"\x00" * 16
        cases = {
            "mpo + text": mpo_bytes(f"RRN {RRN}".encode()),
            "mpo + zip": mpo_bytes(zip_bytes({"secret.txt": f"RRN {RRN}"})),
            "jpeg + broken jpeg": jpeg_bytes() + b"\xff\xd8\xff" + f"RRN {RRN}".encode(),
            "jpeg + jpeg + pdf": jpeg_bytes() + jpeg_bytes() + text_pdf(f"RRN {RRN}"),
            "jpeg + zeros + jpeg": jpeg_bytes() + b"\x00" * 4 + jpeg_bytes(),
            "motion photo (jpeg + mp4, 명세 §14-15 기본값)": jpeg_bytes() + mp4,
        }
        for name, data in cases.items():
            result = classify(data)
            self.assertEqual((result.group, result.detail), ("other", "jpeg_trailer"), name)

    def test_pil_multi_picture_jpeg_when_available(self):
        try:
            from PIL import Image
        except ImportError:
            self.skipTest("PIL not installed")
        stream = io.BytesIO()
        Image.new("RGB", (64, 48), (200, 30, 30)).save(stream, "MPO", save_all=True,
                                                      append_images=[Image.new("RGB", (32, 24), (10, 10, 10))])
        data = stream.getvalue()
        self.assertIn(b"MPF\x00", data[:4096])
        self.assertGreaterEqual(data.count(b"\xff\xd8\xff"), 2)
        self.assertEqual(classify(data).group, "image")

    def test_other_image_formats_are_other(self):
        bmp = b"BM" + (54).to_bytes(4, "little") + b"\x00" * 48
        tiff = b"II*\x00\x08\x00\x00\x00" + b"\x00" * 32
        heic = b"\x00\x00\x00\x18ftypheic\x00\x00\x00\x00mif1heic" + b"\x00" * 32
        for name, data in (("bmp", bmp), ("tiff", tiff), ("heic", heic)):
            self.assertEqual(classify(data).group, "other", name)


class FormatTests(unittest.TestCase):
    def test_containers_without_inspection_are_other(self):
        cases = {
            "hwp(cfb)": (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 504, "cfb"),
            # 2026-10-11 1단계: HWPX는 hwpdoc으로 읽는다. 본문 문단이 없는 HWPX는 empty(other)
            "hwpx": (zip_bytes({"mimetype": b"application/hwp+zip", "Contents/section0.xml": b"<x/>"}), "empty"),
            "pptx": (zip_bytes({"[Content_Types].xml": b"<x/>", "ppt/presentation.xml": b"<x/>"}), "pptx"),
            "zip": (zip_bytes({"a.txt": f"RRN {RRN}"}), "zip"),
            "rtf": (b"{\\rtf1\\ansi " + "\\'c1\\'d6".encode() + b"}", "rtf"),
            "broken zip": (b"PK\x03\x04" + b"\x00" * 64, None),
            "binary": (bytes(range(256)) * 4, "unknown"),
        }
        for name, (data, detail) in cases.items():
            result = classify(data)
            self.assertEqual(result.group, "other", name)
            if detail:
                self.assertEqual(result.detail, detail, name)
        self.assertTrue(classify(b"PK\x03\x04" + b"\x00" * 64).detail.startswith("error:"))

    def test_korean_text_in_cp949_and_utf8(self):
        for name, data in (("cp949 txt", f"주민등록번호 {RRN}\n".encode("cp949")),
                           ("cp949 csv", f"이름,주민번호\n홍길동,{RRN}\n".encode("cp949")),
                           ("md", f"# 회의록\n- 주민번호 {RRN}\n".encode("utf-8")),
                           ("log", f"2026-10-10 INFO 사용자 {RRN}\n".encode("utf-8")),
                           ("utf8 bom", "﻿주민번호 ".encode("utf-8") + RRN.encode())):
            result = classify(data)
            self.assertEqual((result.kind, result.group), ("text", None), name)
            self.assertIn(RRN, result.text)
            self.assertTrue(HANGUL.search(result.text), name)
            self.assertNotIn("﻿", result.text)

    def test_json_escapes_and_html_entities_are_decoded(self):
        data = json.dumps({"이름": "홍길동", "memo": f"주민번호 {RRN}", "n": [1, 2.5, True, None]}).encode()
        self.assertIn(b"\\u", data)   # ensure_ascii 기본: 한글이 \uXXXX
        result = classify(data)
        self.assertEqual((result.group, result.detail), (None, "json"))
        self.assertEqual(result.text.split("\n"), ["이름", "홍길동", "memo", f"주민번호 {RRN}", "n", "1", "2.5"])
        result = classify(b"<p>&#xC8FC;&#xBBFC;&#xBC88;&#xD638; 900101&#45;1234567</p>")
        self.assertEqual((result.group, result.detail), (None, "markup"))
        self.assertIn(f"주민번호 {RRN}", result.text)
        self.assertEqual(classify(b"{not json").detail, "text")

    def test_repeated_json_keys_keep_every_value(self):
        # 10/11 검토: json.loads는 같은 키의 마지막 값만 남기는데 원본(앞 값 포함)은 그대로 나간다 → 모든 값을 검사
        cases = {
            "top level": '{"note": "%s", "note": "ok"}',
            "nested": '{"a": {"note": "%s", "note": "ok"}}',
            "array of objects": '[{"id": 1}, {"memo": "%s", "memo": ""}]',
            "hangul keys": '{"주민번호": "%s", "주민번호": "없음"}',
            "repeated object values": '{"p": {"v": "%s"}, "p": {"v": "x"}}',
        }
        for name, template in cases.items():
            result = classify((template % RRN).encode("utf-8"))
            self.assertEqual((result.group, result.detail), (None, "json"), name)
            self.assertIn(RRN, result.text, name)
        self.assertEqual(classify(b'{"k": "first", "k": "second"}').text.split("\n"), ["k", "first", "k", "second"])

    def test_terminal_colored_log_is_text(self):
        # 10/11 검토(결정): 색을 넣은 빌드·CI 로그(ESC·BEL·BS·VT)는 글자 파일로 검사(명세 §11.1 '.log → 검사')
        log = (f"\x1b[32mINFO\x1b[0m 2026-10-10 user {RRN}\x07\n"
               "\x1b[1;31mERROR\x1b[0m retry\x08\x0b\n").encode("utf-8")
        result = classify(log)
        self.assertEqual((result.kind, result.group, result.detail), ("text", None, "text"))
        self.assertIn(RRN, result.text)
        self.assertEqual(classify(f"\x1b[33m경고\x1b[0m 주민번호 {RRN}".encode("cp949")).group, None)
        # 그 밖 제어 문자는 그대로 other(NUL·UTF-16 글자·바이너리 머리)
        for name, data in (("nul", f"user {RRN}\x00".encode()), ("utf-16", f"주민번호 {RRN}".encode("utf-16-le")),
                           ("soh", b"\x01\x02 header " + RRN.encode()), ("shift out", b"a\x0eb"),
                           ("file separator", b"a\x1cb")):
            self.assertEqual(classify(data).group, "other", name)

    def test_renamed_files_use_their_real_format(self):
        result = classify(docx_bytes(f"주민번호 {RRN}"), filename="x.txt", mime="text/plain")
        self.assertEqual((result.kind, result.group), ("docx", None))
        self.assertIn(RRN, result.text)
        result = classify(text_pdf(f"RRN {RRN}"), filename="x.docx", mime="application/msword")
        self.assertEqual((result.kind, result.group), ("pdf", None))
        self.assertIn(RRN, result.text)

    def test_zero_bytes_pass_without_text(self):
        result = classify(b"")
        self.assertEqual((result.group, result.text), (None, ""))

    def test_whitespace_only_text_is_other(self):
        self.assertEqual((classify(b"  \n\t ").group, classify(b"  \n\t ").detail), ("other", "empty"))


class HiddenTextTests(unittest.TestCase):
    def test_header_and_footer_only(self):
        result = classify(docx_bytes(header=f"머리글 {RRN}", footer="바닥글 비밀"))
        self.assertEqual(result.group, None)
        self.assertIn(RRN, result.text)
        self.assertIn("바닥글 비밀", result.text)

    def test_text_box_only_without_the_vml_duplicate(self):
        result = classify(docx_with_body_xml(text_box_paragraph(f"글상자 {RRN}")))
        self.assertEqual(result.group, None)
        self.assertEqual(result.text.count(RRN), 1)   # mc:Fallback(VML) 사본은 한 번만

    def test_footnote_only(self):
        footnotes = (f'<w:footnotes xmlns:w="{W}"><w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/>'
                     f'</w:r></w:p></w:footnote><w:footnote w:id="1"><w:p><w:r><w:t>각주 {RRN}</w:t></w:r></w:p>'
                     f'</w:footnote></w:footnotes>')
        result = classify(rewrite_zip(docx_bytes(), {"word/footnotes.xml": footnotes.encode()}))
        self.assertEqual(result.group, None)
        self.assertIn(f"각주 {RRN}", result.text)

    def test_body_order_tables_nested_tables_and_content_controls(self):
        document = Document()
        document.add_paragraph("첫 문단")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "항목"
        inner = table.cell(0, 1).add_table(rows=1, cols=1)
        inner.cell(0, 0).text = f"표 안 표 {RRN}"
        document.add_paragraph("끝 문단")
        stream = io.BytesIO()
        document.save(stream)
        result = classify(stream.getvalue())
        lines = result.text.split("\n")
        self.assertEqual(lines[0], "첫 문단")
        self.assertIn(f"표 안 표 {RRN}", result.text)
        self.assertEqual(lines[-1], "끝 문단")
        sdt = (f"<w:sdt><w:sdtContent><w:p><w:r><w:t>내용 제어 {RRN}</w:t></w:r></w:p></w:sdtContent></w:sdt>")
        self.assertIn(f"내용 제어 {RRN}", classify(docx_with_body_xml(sdt)).text)

    def test_picture_only_docx_is_an_image_and_empty_docx_is_other(self):
        self.assertEqual(classify(docx_bytes(picture=True)).group, "image")
        self.assertEqual((classify(docx_bytes()).group, classify(docx_bytes()).detail), ("other", "empty"))

    def test_xlsx_emptiness_uses_cell_values_only(self):
        self.assertEqual(classify(xlsx_bytes(media=True)).group, "image")
        self.assertEqual(classify(xlsx_bytes(title="명단")).group, "other")   # 시트 이름 줄만 있음
        result = classify(xlsx_bytes(cells=[("A1", "이름"), ("B1", RRN)], title="명단"))
        self.assertEqual(result.group, None)
        self.assertIn("[Sheet: 명단]", result.text)
        self.assertIn(RRN, result.text)


class PdfTests(unittest.TestCase):
    def test_korean_cmap_pdf_yields_hangul(self):
        if not SAMPLE_KO_PDF.exists():
            self.skipTest("sample_ko.pdf not available")
        data = SAMPLE_KO_PDF.read_bytes()
        result = classify(data)
        self.assertEqual((result.kind, result.group), ("pdf", None))
        self.assertGreaterEqual(len(HANGUL.findall(result.text)), 81)
        from pypdf import PdfReader
        legacy = "".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
        self.assertEqual(len(HANGUL.findall(legacy)), 0)   # 이전 추출기(pypdf)는 한글 0자

    def test_plain_pdf_matches_the_previous_extractor(self):
        from pypdf import PdfReader
        data = text_pdf(f"RRN {RRN} meeting notes")
        legacy = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(data)).pages)
        self.assertEqual(classify(data).text.split(), legacy.split())

    def test_pdf_groups(self):
        from pypdf import PdfReader, PdfWriter
        writer = PdfWriter()
        writer.append(PdfReader(io.BytesIO(text_pdf(f"RRN {RRN}"))))
        writer.encrypt(user_password="pw-test", owner_password="owner-test", algorithm="RC4-128")
        encrypted = io.BytesIO()
        writer.write(encrypted)
        cases = {
            "encrypted (user password)": (encrypted.getvalue(), "other"),
            "scanned": (scanned_pdf(), "scanned_pdf"),
            "vector only": (vector_only_pdf(), "scanned_pdf"),
            "empty": (empty_pdf(), "other"),
            "attachment": (attachment_pdf("hello"), "other"),
            "xfa": (xfa_pdf("hello"), "other"),
            "broken": (b"%PDF-1.4\n" + b"\x00garbage" * 20, "other"),
        }
        for name, (data, group) in cases.items():
            self.assertEqual(classify(data).group, group, name)
        self.assertEqual(classify(attachment_pdf("hello")).detail, "pdf_attachment")
        self.assertEqual(classify(scanned_pdf()).image_only_pages, 1)

    def test_annotation_form_and_offpage_text_is_read(self):
        self.assertIn(f"ANNOT {RRN}", classify(annotation_pdf(f"ANNOT {RRN}")).text)
        self.assertIn(f"FORM {RRN}", classify(form_field_pdf(f"FORM {RRN}")).text)
        offpage = classify(text_pdf(f"OFFPAGE {RRN}", x=700, y=100))   # MediaBox(612) 밖
        self.assertIn(RRN, offpage.text)


class LimitTests(unittest.TestCase):
    def test_character_limit(self):
        self.assertEqual(classify(b"x" * MAX_CHARS).group, None)
        result = classify(b"x" * (MAX_CHARS + 1))
        self.assertEqual((result.group, result.detail), ("too_large", "chars"))
        self.assertEqual(classify(b"x " * MAX_CHARS).group, None)   # 공백은 세지 않는다
        self.assertEqual(classify(xlsx_bytes(cells=[("A1", "가" * 600)]), max_chars=500).group, "too_large")
        self.assertEqual(classify(docx_bytes("가" * 600), max_chars=500).group, "too_large")

    def test_byte_limit_skips_extraction(self):
        data = text_pdf("hello") + b"%" * 5000
        with mock.patch.object(gateway_file_classes, "extract_pdf", side_effect=AssertionError("extracted")):
            result = classify(data, max_bytes=4096)
        self.assertEqual((result.group, result.detail), ("too_large", "bytes"))
        self.assertEqual(classify(png_bytes(), max_bytes=10).group, "image")   # 이미지 판정은 바이트 상한보다 먼저

    def test_declared_uncompressed_size(self):
        result = classify(zip_bomb_declared(MAX_BYTES + 1))
        self.assertEqual((result.group, result.detail), ("too_large", "zip_bytes"))

    def test_extraction_time_limit(self):
        def clock_after(calls):
            values = iter([0.0] * calls)
            return lambda: next(values, 100.0)
        result = inspect_upload(text_pdf("hello"), max_bytes=MAX_BYTES, max_chars=MAX_CHARS, deadline=1.0,
                                clock=clock_after(1))
        self.assertEqual((result.group, result.detail), ("too_large", "pdf_time"))
        result = inspect_upload(b"plain text", max_bytes=MAX_BYTES, max_chars=MAX_CHARS, deadline=1.0,
                                clock=clock_after(1))
        self.assertEqual((result.group, result.detail), ("too_large", "time"))
        self.assertEqual(inspect_upload(b"plain text", max_bytes=MAX_BYTES, max_chars=MAX_CHARS, deadline=1.0,
                                        clock=clock_after(5)).group, None)

    def test_pdf_character_limit(self):
        result = classify(text_pdf("x" * 300), max_chars=100)
        self.assertEqual((result.group, result.detail), ("too_large", "pdf_chars"))

    def test_docx_extraction_ends_within_the_wait(self):
        def depth_refused(result):
            return (result.kind, result.group, result.detail) == ("unknown", "other", "error:XMLSyntaxError")

        # 10/11 검토: 수 KB DOCX의 겹친 글상자(84겹)·깊은 감싸기(240겹)가 추출에 7~18초(요청 처리 함수가 동기라
        # 그동안 모든 사용자 요청이 멈춘다). 대기 시간 + 여유 안에서 끝나야 한다
        for name, body in (("nested text boxes", nested_text_boxes(84, 200_000)),
                           ("wrapper chain", wrapper_chain(240, 200_000))):
            data = raw_docx(body)
            self.assertLess(len(data), 20_000, name)
            for seconds, max_chars in ((0.5, 10_000_000), (2.0, MAX_CHARS)):
                started = time.perf_counter()
                result = classify(data, seconds=seconds, max_chars=max_chars)
                elapsed = time.perf_counter() - started
                self.assertLess(elapsed, seconds + 0.5, (name, seconds))
                if depth_refused(result):         # libxml2 2.14+(운영 서버)는 256겹 넘는 XML을 읽기 전에 거부 → 검사 불가로 차단
                    continue
                self.assertEqual(result.kind, "docx", (name, seconds))
                if result.group is None:          # 시간 안에 다 읽었으면 w:t 200,000개를 한 번씩
                    self.assertEqual(result.text.count("a"), 200_000, (name, seconds))
                else:
                    self.assertEqual(result.group, "too_large", (name, seconds))
            if not depth_refused(result):
                self.assertEqual((result.group, result.detail), ("too_large", "chars"), name)   # 2초: 글자 상한

    def test_docx_deadline_is_checked_during_extraction(self):
        calls = []

        def clock():
            calls.append(1)
            return 0.0 if len(calls) <= 10 else 100.0
        data = raw_docx("<w:p><w:r><w:t>x</w:t></w:r></w:p>" * 50)
        result = inspect_upload(data, max_bytes=MAX_BYTES, max_chars=MAX_CHARS, deadline=1.0, clock=clock)
        self.assertEqual((result.kind, result.group, result.detail), ("docx", "too_large", "time"))

    def test_xlsx_deadline_is_checked_during_extraction(self):
        calls = []

        def clock():
            calls.append(1)
            return 0.0 if len(calls) <= 3 else 100.0
        data = xlsx_bytes(cells=[(f"A{row}", f"v{row}") for row in range(1, 51)])
        result = inspect_upload(data, max_bytes=MAX_BYTES, max_chars=MAX_CHARS, deadline=1.0, clock=clock)
        self.assertEqual((result.kind, result.group, result.detail), ("xlsx", "too_large", "time"))


def _reference_paragraph(paragraph, budget=None):
    """10/11 이전 방식(글자·글상자마다 가장 가까운 문단·mc:Fallback 조상을 찾음). 새 한 번 훑기와 결과가 같아야 한다."""
    del budget
    fc = gateway_file_classes

    def owner(node):
        parent = node.getparent()
        while parent is not None:
            if parent.tag in (fc._MC_FALLBACK, fc._W_P):
                return parent
            parent = parent.getparent()
        return None
    text = "".join(node.text or "" for node in paragraph.iter(fc._W_T) if owner(node) is paragraph)
    boxes = [box for box in paragraph.iter(fc._W_TXBX) if owner(box) is paragraph]
    return text, boxes


class DocxWalkEquivalenceTests(unittest.TestCase):
    """문단 글자 모으기를 한 번 훑기로 바꿈(10/11). 이전 방식과 같은 글자·순서인지 여러 구조로 대조."""

    def corpus(self):
        box = lambda text: (f"<w:r><mc:AlternateContent><mc:Choice Requires=\"wps\"><w:drawing><wps:txbx>"
                            f"<w:txbxContent><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:txbxContent></wps:txbx>"
                            f"</w:drawing></mc:Choice><mc:Fallback><w:pict><v:textbox><w:txbxContent><w:p><w:r>"
                            f"<w:t>{text}</w:t></w:r></w:p></w:txbxContent></v:textbox></w:pict></mc:Fallback>"
                            f"</mc:AlternateContent></w:r>")
        table_in_box = ("<w:p><w:r><w:t>앞</w:t></w:r><w:r><w:txbxContent><w:tbl><w:tr><w:tc><w:p><w:r><w:t>칸</w:t>"
                        "</w:r>" + box("칸 안 글상자") + "</w:p></w:tc></w:tr></w:tbl></w:txbxContent></w:r>"
                        "<w:r><w:t>뒤</w:t></w:r></w:p>")
        bodies = {
            "text box with vml duplicate": text_box_paragraph(f"글상자 {RRN}"),
            "text around a box": f"<w:p><w:r><w:t>앞 </w:t></w:r>{box('가운데')}<w:r><w:t> 뒤</w:t></w:r></w:p>",
            "table inside a text box": table_in_box,
            "box in a table cell": ("<w:tbl><w:tr><w:tc><w:p><w:r><w:t>칸1</w:t></w:r>" + box("칸 글상자")
                                    + "</w:p></w:tc><w:tc><w:p><w:r><w:t>칸2</w:t></w:r></w:p></w:tc></w:tr></w:tbl>"),
            "loose run directly in a box": ("<w:p><w:r><w:t>밖</w:t><w:txbxContent><w:r><w:t>문단 없는 글</w:t></w:r>"
                                            "<w:p><w:r><w:t>안</w:t></w:r></w:p></w:txbxContent></w:r></w:p>"),
            "box in a box without a paragraph": ("<w:p><w:r><w:txbxContent><w:txbxContent><w:p><w:r><w:t>겹</w:t>"
                                                 "</w:r></w:p></w:txbxContent></w:txbxContent></w:r></w:p>"),
            "nested boxes": nested_text_boxes(5, 3, "n"),
            "wrapper chain": wrapper_chain(10, 3, "c"),
            "two choices and a fallback": ("<w:p><w:r><mc:AlternateContent><mc:Choice Requires=\"a\"><w:t>c1</w:t>"
                                           "</mc:Choice><mc:Choice Requires=\"b\"><w:t>c2</w:t></mc:Choice>"
                                           "<mc:Fallback><w:t>fb</w:t></mc:Fallback></mc:AlternateContent></w:r></w:p>"),
            "box inside a fallback paragraph": ("<w:p><w:r><mc:AlternateContent><mc:Choice Requires=\"a\"><w:t>c</w:t>"
                                                "</mc:Choice><mc:Fallback><w:p><w:r><w:t>fp</w:t></w:r>" + box("fb 글상자")
                                                + "</w:p></mc:Fallback></mc:AlternateContent></w:r></w:p>"),
            "run level content control and link": ("<w:p><w:sdt><w:sdtContent><w:r><w:t>내용 제어</w:t></w:r>"
                                                   "</w:sdtContent></w:sdt><w:hyperlink><w:r><w:t> 연결</w:t></w:r>"
                                                   "</w:hyperlink><w:smartTag><w:r><w:t> 태그</w:t></w:r></w:smartTag>"
                                                   "</w:p>"),
            "text inside text": "<w:p><w:r><w:t>바깥<w:t>안쪽</w:t></w:t></w:r></w:p>",
        }
        cases = {name: docx_with_body_xml(body) for name, body in bodies.items()}
        document = Document()
        document.add_paragraph("첫 문단")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "항목"
        table.cell(0, 1).add_table(rows=1, cols=1).cell(0, 0).text = f"표 안 표 {RRN}"
        stream = io.BytesIO()
        document.save(stream)
        cases["python-docx table"] = stream.getvalue()
        cases["header and footer"] = docx_bytes("본문", header="머리", footer="바닥")
        return cases

    def test_same_text_as_the_previous_walk(self):
        for name, data in self.corpus().items():
            with self.subTest(name=name):
                new = classify(data)
                with mock.patch.object(gateway_file_classes, "_paragraph", _reference_paragraph):
                    old = classify(data)
                self.assertEqual((new.group, new.detail), (old.group, old.detail))
                self.assertEqual(new.text, old.text)
                self.assertTrue(new.text)


def _mutations(seed, samples, count):
    rng = random.Random(seed)
    for _ in range(count):
        data = bytearray(rng.choice(samples))
        operation = rng.randrange(4)
        position = rng.randrange(len(data))
        if operation == 0:
            for _ in range(rng.randint(1, 16)):
                data[rng.randrange(len(data))] = rng.randrange(256)
        elif operation == 1:
            data = data[:position]
        elif operation == 2:
            data[position:position] = rng.randbytes(rng.randint(1, 64))
        else:
            data[position:position + rng.randint(1, 64)] = b"\x00" * rng.randint(1, 64)
        yield bytes(data)


class CorruptInputTests(unittest.TestCase):
    def samples(self):
        samples = [text_pdf(f"RRN {RRN}"), scanned_pdf(), annotation_pdf("x"), attachment_pdf("y"),
                   docx_bytes(f"주민번호 {RRN}", header="h"), xlsx_bytes(cells=[("A1", RRN)]), png_bytes(),
                   jpeg_bytes()]
        if SAMPLE_KO_PDF.exists():
            samples.append(SAMPLE_KO_PDF.read_bytes())
        return samples

    def test_mutated_files_never_raise(self):
        groups = {None, "image", "scanned_pdf", "other", "too_large"}
        for data in _mutations(11, self.samples(), 300):
            result = classify(data)
            self.assertIn(result.group, groups)

    def test_mutated_pdfs_do_not_kill_the_process(self):
        # 손상 PDF를 자식 프로세스에서(PDFium C 코드가 프로세스를 죽이는지). 운영 판(Linux·3.12)에서도 돌린다
        script = (
            "import sys, time, random\n"
            f"sys.path.insert(0, {str(ROOT)!r})\n"
            "from tests.test_gateway_file_classes import _mutations, CorruptInputTests\n"
            "from gateway_file_classes import inspect_upload\n"
            "samples = [s for s in CorruptInputTests().samples() if s.startswith(b'%PDF')]\n"
            "count = 0\n"
            "for data in _mutations(23, samples, int(sys.argv[1])):\n"
            "    inspect_upload(data, max_bytes=20 << 20, max_chars=80645, deadline=time.monotonic() + 2)\n"
            "    count += 1\n"
            "print('done', count)\n"
        )
        count = int(os.environ.get("GW_TEST_PDF_FUZZ_COUNT", "300"))
        result = subprocess.run([sys.executable, "-c", script, str(count)], cwd=ROOT, capture_output=True, text=True,
                                timeout=600)
        self.assertEqual(result.returncode, 0, result.stderr[-2000:])
        self.assertIn(f"done {count}", result.stdout)

    def test_unexpected_exception_is_other(self):
        with mock.patch.object(gateway_file_classes, "_text_class", side_effect=RuntimeError("boom")):
            result = classify(b"plain text")
        self.assertEqual((result.group, result.detail), ("other", "error:RuntimeError"))
        with mock.patch.object(gateway_file_classes, "_docx_text", side_effect=MemoryError()):
            result = classify(docx_bytes("x"))
        self.assertEqual((result.group, result.detail), ("other", "error:MemoryError"))


if __name__ == "__main__":
    unittest.main()
