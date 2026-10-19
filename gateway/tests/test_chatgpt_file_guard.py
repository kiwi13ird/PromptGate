"""ChatGPT 파일 업로드 - 파일 전체 검사 → 칸별 가림 → 집행·대시보드 기록 (2026-10-08).

가짜 Detection 응답과 실제 mitmproxy flow로 request() 훅 전체를 지난다. 실제 Detection·대시보드에 연결하지 않는다.
실행: ~/llm-gateway/.venv/bin/python -m unittest tests.test_chatgpt_file_guard -v (Gateway 사본 폴더에서)
"""

from __future__ import annotations

import io
import json
import os
import re
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

_TMP = tempfile.mkdtemp(prefix="gw-file-guard-")
os.environ["SIGNATURE_REVIEW_DIR"] = str(Path(_TMP) / "review")
os.environ["DASHBOARD_OUTBOX_PATH"] = str(Path(_TMP) / "outbox.sqlite3")
os.environ.pop("DASHBOARD_EVENT_URL", None)
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document  # noqa: E402
from mitmproxy.test import tflow, tutils  # noqa: E402

import gateway  # noqa: E402
from detection_client import RedactResult  # noqa: E402
from tests.gw_support import run_hook  # noqa: E402

RRN = "900101-1234567"          # 가상 주민등록번호
SECRET = "PROJECT-ORION-ROADMAP"  # 파일 전체로 보면 기밀 문서와 일치하는 표시(가짜 문서 일치)
CELL_BLOCK = "SEGMENT-ONLY-BLOCK"  # 칸 하나에서만 차단되는 표시
ACCOUNT = "110-245-678901"         # 가상 계좌번호
FILE_ID = "00000000-aaaa-8206-bbbb-000000000001"


def fake_detection(calls):
    def check(base_url, *, request_id, normalized_text, source=None, **_):
        calls.append({"request_id": request_id, "text": normalized_text, "source": source or {}})
        is_whole = "#" not in request_id
        def result(**values):
            base = dict(request_id=request_id, degraded=False, http_status=None, user_message=None,
                        reason_code=None, total_ms=0.1, findings=[], masked_text=None, pii_hits=None)
            base.update(values)
            return SimpleNamespace(**base)

        if SECRET in normalized_text and is_whole:
            return result(action="block", blocked=True, decided_by="windowhash",
                          http_status=403, reason_code="confidential_document")
        if CELL_BLOCK in normalized_text and not is_whole:
            return result(action="block_notify", blocked=True, decided_by="signature",
                          http_status=403, reason_code="pii")
        findings = [{"entity_type": "KR_RRN", "start": m.start(), "end": m.end(),
                     "replacement": "[KR_RRN]", "policy_action": "mask_full", "score": 1.0}
                    for m in re.finditer(re.escape(RRN), normalized_text)]
        if ACCOUNT in normalized_text and not is_whole:   # 칸 이름표를 봐야 잡히는 값 - 파일 전체 검사에서는 안 잡힘
            start = normalized_text.index(ACCOUNT)
            findings.append({"entity_type": "KR_BANK_ACCOUNT", "start": start, "end": start + len(ACCOUNT),
                             "replacement": "[KR_BANK_ACCOUNT]", "policy_action": "mask_full", "score": 1.0})
            findings.sort(key=lambda f: f["start"])
        if findings:
            masked = normalized_text
            for f in reversed(findings):
                masked = masked[:f["start"]] + f["replacement"] + masked[f["end"]:]
            counts = {}
            for f in findings:
                counts[f["entity_type"]] = counts.get(f["entity_type"], 0) + 1
            return result(action="mask_full", blocked=False, decided_by="signature", findings=findings, masked_text=masked,
                          pii_hits=",".join(f"{k}:{v}" for k, v in sorted(counts.items())), reason_code="pii")
        return result(action="allow", blocked=False, decided_by="none")
    return check


def docx_bytes(*paragraphs, table=None):
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def pdf_bytes(text):
    """글자 한 줄짜리 PDF(표준 글꼴). pypdf가 글자를 뽑을 수 있는 최소 구조."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


class Base(unittest.TestCase):
    def setUp(self):
        self.calls, self.reports = [], []
        self._orig = (gateway.check_detection, gateway.report_dashboard_event, gateway.send_to_next_module)
        gateway.check_detection = fake_detection(self.calls)
        gateway.send_to_next_module = lambda event: None   # 옛 3000 포트 전송은 이 시험과 무관

        def capture(flow, result, dashboard_context, *, enforcement_applied, http_status):
            self.reports.append(SimpleNamespace(result=result, context=dashboard_context,
                                                applied=enforcement_applied, http_status=http_status))
        gateway.report_dashboard_event = capture
        gateway.CHATGPT_FILE_META.clear()
        gateway.CHATGPT_SANITIZED_FILE_SIZE.clear()

        # PDF 가림 API(2026-10-11): 이 묶음은 '지원하지 않음'을 돌려주는 가짜(→ 403 mask_application_failed, 재검사 없음).
        # 가림 성공·잔존·시간 초과 등은 tests/test_gateway_pdf_mask.py(가짜 Detection HTTP 서버)
        self.redact_calls = []
        self._orig_redact = gateway.redact_pdf_detection

        def redact_unsupported(base_url, prefix, data, source=None, timeout=2.0):
            self.redact_calls.append({"prefix": prefix, "bytes": len(data), "timeout": timeout})
            return RedactResult(status="unsupported", reason="fake_unsupported")
        gateway.redact_pdf_detection = redact_unsupported

    def tearDown(self):
        gateway.check_detection, gateway.report_dashboard_event, gateway.send_to_next_module = self._orig
        gateway.redact_pdf_detection = self._orig_redact

    def raw_put(self, content, filename):
        gateway.CHATGPT_FILE_META[FILE_ID] = filename
        req = tutils.treq(host="sdmntprkoreacentral.oaiusercontent.com", port=443, scheme=b"https", method=b"PUT",
                          path=f"/files/{FILE_ID}/raw?se=2026-10-08".encode(), content=content,
                          headers=((b"content-type", b"application/octet-stream"),))
        req.headers["content-length"] = str(len(content))
        flow = tflow.tflow(req=req)
        run_hook(gateway.request(flow))
        return flow

    def multipart(self, content, filename):
        boundary = "----gwtestboundary"
        body = (f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"{filename}\"\r\n"
                f"Content-Type: application/octet-stream\r\n\r\n").encode() + content + f"\r\n--{boundary}--\r\n".encode()
        req = tutils.treq(host="chatgpt.com", port=443, scheme=b"https", method=b"POST",
                          path=b"/backend-api/files/upload", content=body,
                          headers=((b"content-type", f"multipart/form-data; boundary={boundary}".encode()),))
        flow = tflow.tflow(req=req)
        run_hook(gateway.request(flow))
        return flow

    def assert_one_report(self, action, applied, content_type="file"):
        self.assertEqual(len(self.reports), 1, [r.result.action for r in self.reports])
        report = self.reports[0]
        self.assertEqual(report.result.action, action)
        self.assertEqual(report.applied, applied)
        self.assertEqual(report.context["targetLlm"], "ChatGPT")
        self.assertEqual(report.context["contentType"], content_type)
        return report


class RawUploadTests(Base):
    def test_whole_file_is_checked_first_with_the_base_request_id(self):
        self.raw_put(docx_bytes("안내문", f"주민등록번호 {RRN}"), "a.docx")
        whole = self.calls[0]
        self.assertNotIn("#", whole["request_id"])
        self.assertIn(RRN, whole["text"])
        self.assertEqual(whole["source"]["content_type"], "file")
        self.assertTrue(all("#" in c["request_id"] for c in self.calls[1:]))

    def test_pii_docx_is_masked_and_reported(self):
        original = docx_bytes("안내문", f"주민등록번호 {RRN}")
        flow = self.raw_put(original, "a.docx")
        self.assertIsNone(flow.response)
        self.assertNotEqual(flow.request.raw_content, original)
        self.assertNotIn(RRN.encode(), flow.request.raw_content)
        self.assertEqual(gateway.CHATGPT_SANITIZED_FILE_SIZE[FILE_ID], len(flow.request.raw_content))
        report = self.assert_one_report("mask_full", applied=True)
        self.assertNotIn("#", report.result.request_id)   # 대시보드 이벤트는 파일 전체 판정(요청 번호 그대로)

    def test_pii_inside_a_table_cell_is_masked(self):
        flow = self.raw_put(docx_bytes("명단", table=[["항목", "값"], ["주민등록번호", RRN]]), "t.docx")
        self.assertIsNone(flow.response)
        self.assertNotIn(RRN.encode(), flow.request.raw_content)
        self.assert_one_report("mask_full", applied=True)

    def test_confidential_document_is_blocked_before_segments(self):
        flow = self.raw_put(docx_bytes(f"{SECRET} 3분기 계획", f"주민등록번호 {RRN}"), "s.docx")
        self.assertEqual(flow.response.status_code, 403)
        self.assertEqual(len(self.calls), 1)               # 칸별 검사 전에 끝남
        self.assert_one_report("block", applied=True)

    def test_a_blocked_cell_blocks_the_upload_with_an_event(self):
        flow = self.raw_put(docx_bytes("보통 문단", f"{CELL_BLOCK} 내부 값"), "c.docx")
        self.assertEqual(flow.response.status_code, 403)
        report = self.assert_one_report("block_notify", applied=True)
        self.assertIn("#", report.result.request_id)

    def test_clean_docx_passes_unchanged_and_is_reported_as_allow(self):
        original = docx_bytes("회의록", "다음 주 일정 공유")
        flow = self.raw_put(original, "clean.docx")
        self.assertIsNone(flow.response)
        self.assertEqual(flow.request.raw_content, original)
        self.assertNotIn(FILE_ID, gateway.CHATGPT_SANITIZED_FILE_SIZE)
        self.assert_one_report("allow", applied=False)

    def test_clean_pdf_now_passes_instead_of_503(self):
        original = pdf_bytes("Quarterly meeting notes")
        flow = self.raw_put(original, "notes.pdf")
        self.assertIsNone(flow.response)
        self.assertEqual(flow.request.raw_content, original)
        self.assert_one_report("allow", applied=False)

    def test_pdf_with_pii_is_blocked_when_the_redaction_is_unsupported(self):
        # 2026-10-11: 가림 판정 PDF는 Detection /redact/pdf로 간다. '지원하지 않음'이면 403(원본 안 보냄)
        original = pdf_bytes(f"RRN {RRN}")
        flow = self.raw_put(original, "pii.pdf")
        self.assertEqual(flow.response.status_code, 403)
        report = self.assert_one_report("block", applied=True)
        self.assertEqual(report.result.reason_code, "mask_application_failed")
        self.assertEqual([c["prefix"] for c in self.redact_calls], [self.calls[0]["request_id"]])
        self.assertEqual(len(self.calls), 1)   # 재검사 없음
        self.assertNotIn(FILE_ID, gateway.CHATGPT_SANITIZED_FILE_SIZE)

    def test_clean_txt_passes(self):
        original = "plain memo without secrets".encode()
        flow = self.raw_put(original, "memo.txt")
        self.assertIsNone(flow.response)
        self.assertEqual(flow.request.raw_content, original)
        self.assert_one_report("allow", applied=False)

    def test_masking_failure_still_fails_closed(self):
        def broken(*args, **kwargs):
            raise ValueError("sanitizer broke")
        original_sanitize = gateway.sanitize_for_upload
        gateway.sanitize_for_upload = broken
        try:
            flow = self.raw_put(docx_bytes(f"주민등록번호 {RRN}"), "a.docx")
        finally:
            gateway.sanitize_for_upload = original_sanitize
        self.assertEqual(flow.response.status_code, 503)
        self.assertEqual(self.reports, [])


class MultipartUploadTests(Base):
    def test_pii_docx_is_masked_in_the_multipart_body_and_reported(self):
        flow = self.multipart(docx_bytes(f"주민등록번호 {RRN}"), "m.docx")
        self.assertIsNone(flow.response)
        self.assertNotIn(RRN.encode(), flow.request.raw_content)
        self.assert_one_report("mask_full", applied=True)

    def test_clean_docx_is_reported_as_allow(self):
        flow = self.multipart(docx_bytes("회의록"), "m.docx")
        self.assertIsNone(flow.response)
        self.assert_one_report("allow", applied=False)

    def test_confidential_document_is_blocked(self):
        flow = self.multipart(docx_bytes(f"{SECRET} 설계"), "m.docx")
        self.assertEqual(flow.response.status_code, 403)
        self.assert_one_report("block", applied=True)


class LinkTests(Base):
    """질문과 첨부 파일 연결 번호 · 칸별 본문 순번 · 건수 합계 (2026-10-09)."""

    def test_canonical_file_ref(self):
        ref = "00000000aaaa8206bbbb000000000001"
        for value in (FILE_ID, f"file_{ref}", f"file-{ref}", f"sediment://file_{ref}", ref.upper()):
            self.assertEqual(gateway.canonical_file_ref(value), ref, value)
        self.assertEqual(gateway.canonical_file_ref("file-AbC9xyz"), "AbC9xyz")
        for bad in (None, "", "../etc", "a b", "x" * 101, 12):
            self.assertIsNone(gateway.canonical_file_ref(bad), bad)

    def test_file_event_carries_the_file_number(self):
        self.raw_put(docx_bytes(f"주민등록번호 {RRN}"), "a.docx")
        report = self.assert_one_report("mask_full", applied=True)
        self.assertEqual(report.context["fileId"], gateway.canonical_file_ref(FILE_ID))

    def test_segments_carry_body_order_of_paragraphs_and_tables(self):
        document = Document()
        document.add_paragraph("제목")                       # 본문 0
        t = document.add_table(rows=1, cols=2)                # 본문 1
        t.cell(0, 0).text, t.cell(0, 1).text = "주민등록번호", RRN
        document.add_paragraph("중간 문단")                   # 본문 2
        t2 = document.add_table(rows=1, cols=1)               # 본문 3
        t2.cell(0, 0).text = "끝 표"
        stream = io.BytesIO()
        document.save(stream)
        self.raw_put(stream.getvalue(), "order.docx")
        orders = {c["request_id"].split("#", 1)[1]: c["source"].get("block_order") for c in self.calls if "#" in c["request_id"]}
        self.assertEqual(orders["f:0:docx:body:p0"], 0)
        self.assertEqual(orders["f:0:docx:body:t0:r0:c1:p0"], 1)
        self.assertEqual(orders["f:0:docx:body:p1"], 2)
        self.assertEqual(orders["f:0:docx:body:t1:r0:c0:p0"], 3)

    def test_xlsx_segments_have_no_body_order(self):
        from openpyxl import Workbook
        workbook = Workbook()
        sheet = workbook.active
        sheet["A1"], sheet["B1"] = "이름", "주민등록번호"
        sheet["A2"], sheet["B2"] = "홍길동", RRN
        stream = io.BytesIO()
        workbook.save(stream)
        self.raw_put(stream.getvalue(), "s.xlsx")
        self.assertTrue(all("block_order" not in c["source"] for c in self.calls))
        self.assert_one_report("mask_full", applied=True)

    def test_masked_counts_are_the_sum_of_masked_cells(self):
        flow = self.raw_put(docx_bytes(f"주민등록번호 {RRN}", table=[["급여 계좌", ACCOUNT]]), "acct.docx")
        self.assertNotIn(ACCOUNT.encode(), flow.request.raw_content)
        report = self.assert_one_report("mask_full", applied=True)
        self.assertEqual(report.result.pii_hits, "KR_BANK_ACCOUNT:1,KR_RRN:1")   # 파일 전체 검사만으로는 KR_RRN:1
        self.assertNotIn("#", report.result.request_id)

    def conversation(self, message):
        req = tutils.treq(host="chatgpt.com", port=443, scheme=b"https", method=b"POST", path=b"/backend-api/f/conversation",
                          content=json.dumps({"action": "next", "messages": [message]}).encode(),
                          headers=((b"content-type", b"application/json"),))
        run_hook(gateway.request(tflow.tflow(req=req)))

    def test_prompt_event_carries_attached_file_numbers(self):
        self.conversation({
            "id": "m1", "author": {"role": "user"},
            "content": {"content_type": "text", "parts": ["이 파일 요약해줘"]},
            "metadata": {"attachments": [{"id": "file_00000000aaaa8206bbbb000000000001", "name": "a.docx"},
                                          {"id": "file_00000000aaaa8206bbbb000000000001", "name": "a.docx"}]},
        })
        report = self.assert_one_report("allow", applied=False, content_type="prompt")
        self.assertEqual(report.context["attachedFileIds"], [gateway.canonical_file_ref(FILE_ID)])

    def test_prompt_without_attachments_has_no_numbers(self):
        self.conversation({"author": {"role": "user"}, "content": {"content_type": "text", "parts": ["안녕"]}})
        self.assertEqual(self.assert_one_report("allow", applied=False, content_type="prompt").context["attachedFileIds"], [])

    def test_real_event_builder_adds_the_numbers(self):
        from detection_client import DetectionResult
        real_report = self._orig[1]
        os.environ["DASHBOARD_EVENT_URL"] = "http://127.0.0.1:9/ingest"
        os.environ["DASHBOARD_INGEST_TOKEN"] = "test-token-not-used"
        os.environ["DASHBOARD_REPORT_ALLOW_EVENTS"] = "true"
        try:
            flow = tflow.tflow(req=tutils.treq(host="sdmntprkoreacentral.oaiusercontent.com", path=b"/files/x/raw"))
            masked = DetectionResult(blocked=False, reason="", request_id="req-1", decided_by="signature",
                                     pii_hits="KR_RRN:1", action="mask_full", reason_code="pii")
            real_report(flow, masked, {"identity": {"client_id": "chrome-a"}, "targetLlm": "ChatGPT", "contentType": "file",
                                       "fileId": "00000000aaaa8206bbbb000000000001"}, enforcement_applied=True, http_status=None)
            allowed = DetectionResult(blocked=False, reason="", request_id="req-2", action="allow")
            real_report(flow, allowed, {"identity": {"client_id": "chrome-a"}, "targetLlm": "ChatGPT", "contentType": "prompt",
                                        "attachedFileIds": ["00000000aaaa8206bbbb000000000001"]}, enforcement_applied=False, http_status=None)
        finally:
            for key in ("DASHBOARD_EVENT_URL", "DASHBOARD_INGEST_TOKEN", "DASHBOARD_REPORT_ALLOW_EVENTS"):
                os.environ.pop(key, None)
        events = flow.metadata.get("gw_dashboard_events", [])
        file_event = next(e for e in events if e.get("requestId") == "req-1")
        self.assertEqual(file_event["fileId"], "00000000aaaa8206bbbb000000000001")
        self.assertNotIn("attachedFileIds", file_event)
        prompt_event = next(e for e in events if e.get("requestId") == "req-2")
        self.assertEqual(prompt_event["attachedFileIds"], ["00000000aaaa8206bbbb000000000001"])
        self.assertNotIn("fileId", prompt_event)


if __name__ == "__main__":
    unittest.main()
