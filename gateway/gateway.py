from mitmproxy import http

import asyncio
import contextlib
import hashlib
import json
import os
import re
import struct
import threading
import time
import uuid
import requests
import zipfile
from collections import Counter, deque
from concurrent.futures import ThreadPoolExecutor
from urllib.parse import unquote_plus

from datetime import datetime, timezone
from io import BytesIO
from collections import OrderedDict
import copy
from email.parser import BytesParser
from email.policy import default

from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook

from detection_client import check as check_detection
from detection_client import DetectionResult, failure_result
from detection_client import redact_pdf as redact_pdf_detection
from event_adapter import checkable_parts, part_request_id
from client_identity import read_and_strip_identity
from dashboard_event_mapper import UNINSPECTED_REASONS as _UNINSPECTED_REASONS, build_dashboard_event
from gateway_evidence import capture_mask
from dashboard_event_outbox import EventOutbox
from gateway_review_link import request_review_url
from gateway_dashboard_metadata import attach_metadata, finalize_metadata
from dashboard_event_sender import (
    dashboard_integration_enabled,
    metrics_url as dashboard_metrics_url,
    report_allow_events_enabled,
    send_dashboard_event,
    send_dashboard_metrics,
)
from extension_event_sender import send_extension_event
from gateway_relay_status import fetch_relay_status
from normalizer import normalize_text
from gateway_file_masking import FileMaskOutcome, load_review_artifact, replace_multipart_file, sanitize_for_upload
from gateway_settings import (
    MAX_SECONDS as SETTINGS_MAX_SECONDS,
    SettingsHolder,
    cache_path_from_env,
    env_defaults,
    file_max_bytes_from_env,
    file_max_chars,
)
from gateway_file_classes import FileClass, inspect_upload
from gateway_client_tracker import ClientTracker, normalize_client_id, normalize_ip
from gateway_metrics_ext import METRICS_EXT_VERSION, MetricsExtNegotiator
from gateway_access import (
    CLIENT_ID_HEADER as ACCESS_CLIENT_ID_HEADER,
    EXEMPT_HOSTS as ACCESS_EXEMPT_HOSTS,
    REASON_BLOCKED as ACCESS_REASON_BLOCKED,
    STATUS_PATH as ACCESS_STATUS_PATH,
    AccessBlockResult,
    AccessHolder,
    ConnectionClientIds,
    access_status,
    block_response as access_block_response,
    cache_path_from_env as access_cache_path_from_env,
    status_response as access_status_response,
    cors_preflight,
    decide as access_decide,
    exempt_peer as access_exempt_peer,
    top_level_navigation,
)
import gateway_proto
import gateway_my_requests_relay as my_requests_relay
from signature_detector.file_protocol import completion_summary, request_size_summary, matching_file_keys
from signature_detector.upload_identity import IdentityDiagnostic, identifiers as upload_identifiers
from signature_detector.contextual_pii import canonical_label
from signature_detector.claude_upload import replace_files as replace_claude_files
from signature_detector.gateway_runtime import sanitize_uploaded_file, BlockedBySignaturePolicy
from signature_detector.contracts import MaskRefused
_upload_identity_diagnostic = IdentityDiagnostic()
DETECTION_BASE_URL = os.environ.get("DETECTION_BASE_URL", "http://10.0.3.168:8000")

# Gateway 처리 설정(2026-10-10): 대시보드 지표 응답으로 받는다. 받기 전에는 디스크 사본 → 환경변수 기본값.
# 환경변수 값이 틀리면 여기서 ValueError(시작 실패, detection_client.py와 같은 원칙)
GATEWAY_FILE_MAX_BYTES = file_max_bytes_from_env()
GATEWAY_SETTINGS = SettingsHolder(
    env_defaults(),
    cache_path=cache_path_from_env(os.path.dirname(os.path.abspath(__file__))),
    metrics_url=dashboard_metrics_url(),
)
# 접속 차단(2026-10-12): 확장 프로그램 ID 차단 목록·'확장 ID 없는 접속 차단'. 같은 지표 응답의 gateway_access 칸으로 받는다.
# 받기 전에는 디스크 사본(var/gateway-access.json), 그것도 없거나 못 읽으면 차단 없음(fail-open, 경고 로그)
GATEWAY_ACCESS = AccessHolder(
    cache_path=access_cache_path_from_env(os.path.dirname(os.path.abspath(__file__))),
    metrics_url=dashboard_metrics_url(),
)


# ============================================================
# Gateway 설정
# ============================================================

OPENAI_HOST = "api.openai.com"

OPENAI_RESPONSES_ENDPOINT = "/v1/responses"
OPENAI_FILES_ENDPOINT = "/v1/files"

# ============================================================
# LLM Web Service 식별
# ============================================================

LLM_HOSTS = {
    "chatgpt.com": "ChatGPT",
    "openai.com": "ChatGPT/OpenAI",
    "oaistatic.com": "ChatGPT/OpenAI",
    "oaiusercontent.com": "ChatGPT/OpenAI",
    "gemini.google.com": "Gemini",
    "claude.ai": "Claude",
    "anthropic.com": "Claude/Anthropic",
}


def service_for(host: str | None) -> str | None:
    host = (host or "").lower()

    for domain, service in LLM_HOSTS.items():
        if host == domain or host.endswith("." + domain):
            return service

    return None


# 다음 모듈
NEXT_MODULE_URL = (
    "http://127.0.0.1:3000/internal/gateway/events"
)

CHATGPT_FILE_META = {}
CHATGPT_SANITIZED_FILE_SIZE = {}
GEMINI_FILE_META = {}

# 공유 상태 잠금(2026-10-11, 파일 업로드 작업 스레드): 업로드는 작업 스레드(request 훅 아래 "파일 업로드 작업 스레드"),
# 질문·완료 요청·응답 훅은 이벤트 루프에서 돈다. 둘이 함께 쓰고 훑거나 여러 단계로 바꾸는 상태는 이 잠금 안에서만 다룬다:
#   CHATGPT_SANITIZED_FILE_SIZE(작업 스레드가 넣고 루프가 훑음), CHATGPT_UPLOADED_REFS(넣기·순서 바꾸기·지우기를 한 번에),
#   _upload_identity_diagnostic(remember·compare·resolve_reservation이 안쪽 표를 훑고 바꿈),
#   CHATGPT_FILE_META 쓰기(2026-10-12: 넣고 오래된 것을 지우기를 한 번에, _remember_chatgpt_file_name)
# 잠금 안에서는 메모리 일만 한다(본문 해석·네트워크·파일 처리 없음).
# CHATGPT_FILE_META 읽기·GEMINI_FILE_META는 열쇠 하나 읽기·쓰기뿐이라(훑지 않음, GIL 아래 원자적) 잠그지 않는다.
# 그 밖 공유 상태는 원래 잠금이 있다: _dedupe_lock, _gateway_metrics_lock, _claude_rpc_shape_lock, SettingsHolder, EventOutbox
_upload_state_lock = threading.RLock()


def _sanitized_sizes() -> dict:
    """CHATGPT_SANITIZED_FILE_SIZE 사본(훑는 쪽은 사본을 쓴다. 작업 스레드가 도중에 넣어도 RuntimeError가 나지 않게)."""
    with _upload_state_lock:
        return dict(CHATGPT_SANITIZED_FILE_SIZE)


# ============================================================
# Client IP
# ============================================================

def get_client_ip(flow: http.HTTPFlow) -> str | None:

    if flow.client_conn.peername:
        return flow.client_conn.peername[0]

    return None


# ============================================================
# 파일 텍스트 추출
# TXT / PDF / DOCX / XLSX
# ============================================================

def _json_body(flow: http.HTTPFlow) -> dict | None:
    try:
        data = json.loads(flow.request.get_text() or "")
        return data if isinstance(data, dict) else None
    except (json.JSONDecodeError, TypeError):
        return None


def _safe_file_protocol_metadata(data):
    """Return only non-content fields needed to diagnose file protocol failures."""
    collected = {}
    interesting = ("status", "success", "error", "size", "length", "bytes", "hash", "checksum")
    safe_states = {"success", "failed", "error", "ready", "uploaded", "processed", "complete", "completed"}

    def walk(value, path=""):
        if isinstance(value, dict):
            for key, child in value.items():
                child_path = f"{path}.{key}" if path else str(key)
                key_l = str(key).lower()
                if any(token in key_l for token in interesting):
                    if isinstance(child, bool) or isinstance(child, (int, float)):
                        collected[child_path] = child
                    elif isinstance(child, str):
                        normalized = child.strip().lower()
                        collected[child_path] = normalized if normalized in safe_states else f"<str:{len(child)}>"
                    elif child is None:
                        collected[child_path] = None
                walk(child, child_path)
        elif isinstance(value, list):
            for index, child in enumerate(value):
                walk(child, f"{path}[{index}]")

    walk(data)
    return collected


def _remember_chatgpt_file_metadata(flow: http.HTTPFlow):
    data = _json_body(flow)
    if not data:
        return

    found_ids = []
    found_names = []

    def walk(obj):
        if isinstance(obj, dict):
            for key, value in obj.items():
                key_l = str(key).lower()
                if key_l in {"file_id", "fileid"} and isinstance(value, str):
                    found_ids.append(value)
                elif key_l == "id" and isinstance(value, str) and len(value) >= 16:
                    found_ids.append(value)
                elif key_l in {"filename", "file_name", "name"} and isinstance(value, str) and "." in value:
                    found_names.append(value)
                walk(value)
        elif isinstance(obj, list):
            for item in obj:
                walk(item)

    walk(data)

    if not found_ids or not found_names:
        return

    filename = _clean_file_name(found_names[0])
    if filename and _remember_chatgpt_file_name(found_ids, filename):
        print("→ Remembered ChatGPT file metadata (same request body)")   # 번호·이름은 로그에 남기지 않는다


# ============================================================
# ChatGPT 파일 이름 (2026-10-12)
#
# 운영 실측(10/11 02:42~02:44·13:41~14:34 UTC, 경로·키 이름만): 웹 업로드는
#   1) POST /backend-api/files/upload_reservations (옛 모양 POST /backend-api/files)
#      요청 JSON에 파일 이름, 응답에 서버가 준 번호(reservation_id 'file_<32자>' / 옛 모양 file_id)와
#      upload_url(…oaiusercontent.com/files/<UUID>/raw?…)
#   2) PUT …oaiusercontent.com/files/<UUID>/raw  파일 바이트(Gateway가 검사·가림)
#   3) POST …/upload_reservations/file_<32자>/claim_and_finish (옛 모양 …/process_upload_stream)  완료
# 이전: 이름을 요청 본문 하나에서 번호와 함께 찾았다(_remember_chatgpt_file_metadata). 번호는 응답에만 있어 한 번도
#   기억하지 못했고(옛 모양 경로는 아예 보지 않음), 원본 PUT 기록 이름이 늘 '<UUID>.<형식>'이었다
#   (10/11 대시보드 ChatGPT 파일 기록 46건 모두 번호, Claude는 이름).
# 지금: 요청 훅이 예약 요청 JSON에서 이름 하나를 읽어 flow에 두고(_note_chatgpt_reservation), 응답 훅이 응답의
#   upload_url 번호와 이어 기억한다(_link_chatgpt_reservation_name). 응답이 브라우저에 가기 전에 기억하므로 그 주소로
#   보내는 PUT보다 늘 먼저다(작업 스레드 순서·동시 업로드와 무관). 열쇠는 정규화 번호(canonical_file_ref)라
#   UUID·file_<32자> 어느 모양으로도 찾는다. 이름이 둘 이상이거나 upload_url이 둘 이상이면(한 요청에 여러 파일)
#   잇지 않는다(맞는 짝을 알 수 없음 → 지금처럼 번호로 기록). 로그에는 이름·번호를 남기지 않는다(키 이름·개수만).
# ============================================================
CHATGPT_FILE_META_MAX = 2000           # 기억하는 번호 수 상한(넘으면 오래된 것부터 지움)
CHATGPT_FILE_SHAPE_PER_HOUR = 30       # 예약 요청·응답 구조 줄 시간당 상한
_CHATGPT_FILE_NAME_MAX = 255
_chatgpt_file_shape_lock = threading.Lock()
_chatgpt_file_shape_window = [0.0, 0]  # [창 시작 시각, 그 창에서 남긴 줄 수]


def _clean_file_name(value) -> str | None:
    """파일 이름 하나(경로가 붙어 오면 마지막 조각, 제어 문자 뺌, 255자까지). 쓸 수 없으면 None."""
    if not isinstance(value, str):
        return None
    name = re.split(r"[\\/]", value.strip())[-1]
    name = "".join(ch for ch in name if ch >= " " and ch != "\x7f").strip()
    if not name or name in {".", ".."}:
        return None
    return name[:_CHATGPT_FILE_NAME_MAX]


def _remember_chatgpt_file_name(refs, filename: str) -> int:
    """번호들(정규화) → 이름. 기억한 열쇠 수. 쓰기는 잠금 안에서 넣기 + 상한 넘은 것 지우기."""
    keys = {key for key in (canonical_file_ref(ref) for ref in refs or ()) if key}
    if not keys or not filename:
        return 0
    with _upload_state_lock:
        for key in keys:
            CHATGPT_FILE_META.pop(key, None)       # 다시 넣어 가장 새것으로
            CHATGPT_FILE_META[key] = filename
        while len(CHATGPT_FILE_META) > CHATGPT_FILE_META_MAX:
            del CHATGPT_FILE_META[next(iter(CHATGPT_FILE_META))]
    return len(keys)


def _chatgpt_file_name(file_id) -> str | None:
    """원본 PUT 주소·경로의 번호로 기억한 이름. 그대로 → 정규화 번호 순서로 찾는다."""
    if not isinstance(file_id, str) or not file_id:
        return None
    return CHATGPT_FILE_META.get(file_id) or CHATGPT_FILE_META.get(canonical_file_ref(file_id) or "")


def _chatgpt_reservation_name(data) -> tuple[str | None, str | None, int]:
    """예약 요청 JSON에서 파일 이름 하나. (이름, 찾은 키 이름, 서로 다른 이름 수).
    키는 대소문자·'_'·'-'를 무시해 filename(file_name·fileName·filename)을 먼저 보고, 없으면 점이 든 name을 본다
    (name은 다른 뜻일 수 있어 파일 이름 모양만). 이름이 0개·2개 이상이면 (None, None, 개수)."""
    strong: dict[str, str] = {}
    weak: dict[str, str] = {}

    def walk(value, depth=0):
        if depth > 8:
            return
        if isinstance(value, dict):
            for key, child in list(value.items())[:64]:
                if isinstance(child, str):
                    normalized = str(key).lower().replace("_", "").replace("-", "")
                    if normalized in {"filename", "name"}:
                        name = _clean_file_name(child)
                        if name and normalized == "filename":
                            strong.setdefault(name, str(key)[:40])
                        elif name and "." in name.strip("."):
                            weak.setdefault(name, str(key)[:40])
                else:
                    walk(child, depth + 1)
        elif isinstance(value, list):
            for child in value[:64]:
                walk(child, depth + 1)

    walk(data)
    found = strong or weak
    if len(found) != 1:
        return None, None, len(found)
    (name, key), = found.items()
    return name, key, 1


def _is_chatgpt_file_reservation(flow: http.HTTPFlow) -> bool:
    """파일 예약(이름을 싣고 upload_url을 받는) 요청인지. 완료 요청(claim_and_finish·process_upload_stream)은 아니다.
    경로는 넓게 본다(/backend-api/files 아래 JSON POST): 잇기는 응답에 upload_url이 하나일 때만 한다."""
    req = flow.request
    if (req.method or "").upper() != "POST" or service_for(req.host) != "ChatGPT":
        return False
    path = req.path.split("?", 1)[0]
    if path not in {"/backend-api/files", "/backend-api/files/"} and not path.startswith("/backend-api/files/"):
        return False
    if "claim_and_finish" in path or path.endswith(("/process_upload_stream", "/uploaded")):
        return False
    return (req.headers.get("content-type", "") or "").lower().startswith("application/json")


def _shape_path(path: str) -> str:
    """구조 줄용 경로: '?' 앞, 번호 모양 조각(숫자가 든 16자 이상)은 <id>."""
    return "/".join("<id>" if len(part) >= 16 and any(ch.isdigit() for ch in part) else part
                    for part in path.split("?", 1)[0].split("/"))[:200]


def _body_shape(raw: bytes | None, content_type: str) -> dict:
    """본문 구조만(키 이름·값 종류·길이). 값은 남기지 않는다."""
    raw = raw or b""
    record = {"content_type": (content_type or "").split(";", 1)[0].strip()[:100], "bytes": len(raw)}
    if len(raw) > 262144:
        record["format"] = "too_large"
        return record
    try:
        record["shape"] = _json_shape(json.loads(raw.decode("utf-8")))
        record["format"] = "json"
    except (ValueError, UnicodeDecodeError, RecursionError):
        record["format"] = "non_json"
    return record


def _chatgpt_file_shape_enabled() -> bool:
    return (os.getenv("GATEWAY_LOG_CHATGPT_FILE_SHAPE") or "1").strip().lower() not in {"0", "false", "off", "no"}


def _chatgpt_file_shape_rate_ok() -> bool:
    now = time.time()
    with _chatgpt_file_shape_lock:
        if now - _chatgpt_file_shape_window[0] >= 3600:
            _chatgpt_file_shape_window[0], _chatgpt_file_shape_window[1] = now, 0
        if _chatgpt_file_shape_window[1] >= CHATGPT_FILE_SHAPE_PER_HOUR:
            return False
        _chatgpt_file_shape_window[1] += 1
        return True


def _chatgpt_multipart_file_name(path: str, part_filename: str | None) -> str | None:
    """ChatGPT multipart 기록 이름: part의 filename(브라우저가 이름 없는 Blob에 붙이는 'blob'·빈 이름은 이름이 아님)
    → 경로의 파일 번호로 기억한 이름(예약에서 이은 것) → 지금처럼 part 값 그대로."""
    cleaned = _clean_file_name(part_filename)
    if cleaned and cleaned.lower() != "blob":
        return part_filename
    segments = path.split("?", 1)[0].split("/")[3:]   # '/backend-api/files/' 뒤
    for segment in segments:
        remembered = _chatgpt_file_name(segment) if len(segment) >= 16 else None
        if remembered:
            return remembered
    return part_filename


def _note_chatgpt_reservation(flow: http.HTTPFlow) -> None:
    """요청 훅: 예약 요청이면 이름(하나일 때만)과 찾은 키 이름을 flow에 둔다. 응답 훅이 잇는다."""
    try:
        if not _is_chatgpt_file_reservation(flow):
            return
        name, key, count = _chatgpt_reservation_name(_json_body(flow))
        flow.metadata["gw_chatgpt_reservation"] = {"name": name, "name_key": key, "name_count": count}
    except Exception as exc:
        print("[CHATGPT-FILE-NAME] reservation read failed error_type=", type(exc).__name__)


def _link_chatgpt_reservation_name(flow: http.HTTPFlow) -> None:
    """응답 훅: 예약 응답의 upload_url 번호(하나일 때만)와 서버 번호들에 요청의 이름을 잇는다.
    upload_url이 없으면 서버 번호가 하나일 때만 잇는다(파일 바이트를 다른 길로 보낼 때 찾게). 결과는 값 없이 한 줄."""
    info = flow.metadata.get("gw_chatgpt_reservation")
    if not isinstance(info, dict) or flow.response is None:
        return
    try:
        status = flow.response.status_code
        response_raw = flow.response.get_content(strict=False) or b""
        if _chatgpt_file_shape_enabled() and _chatgpt_file_shape_rate_ok():
            record = {"method": flow.request.method, "path": _shape_path(flow.request.path), "status": status,
                      "request": _body_shape(flow.request.get_content(strict=False),
                                             flow.request.headers.get("content-type", "")),
                      "response": _body_shape(response_raw, flow.response.headers.get("content-type", ""))}
            print("[CHATGPT-FILE-SHAPE] " + json.dumps(record, ensure_ascii=True, separators=(",", ":"))[:8000])
        data = None
        if 200 <= status < 300 and len(response_raw) <= 262144:
            try:
                data = json.loads(response_raw.decode("utf-8"))
            except (ValueError, UnicodeDecodeError, RecursionError):
                data = None
        ids, raw_ids = upload_identifiers(data) if data is not None else ({}, set())
        server_refs = {ref for values in ids.values() for ref in values}
        name, linked = info.get("name"), 0
        if name and len(raw_ids) == 1:
            linked = _remember_chatgpt_file_name(raw_ids | server_refs, name)
        elif name and not raw_ids and len({canonical_file_ref(ref) for ref in server_refs} - {None}) == 1:
            linked = _remember_chatgpt_file_name(server_refs, name)
        print("[CHATGPT-FILE-NAME] reservation", json.dumps({
            "status": status, "name_found": bool(name), "name_key": info.get("name_key"),
            "name_count": info.get("name_count"), "id_fields": sorted(ids), "raw_url_count": len(raw_ids),
            "linked": bool(linked)}, ensure_ascii=True, separators=(",", ":")))
    except Exception as exc:
        print("[CHATGPT-FILE-NAME] reservation link failed error_type=", type(exc).__name__)


def _extract_chatgpt_raw_file_id(path: str) -> str | None:
    clean_path = path.split("?", 1)[0]
    parts = clean_path.strip("/").split("/")
    if len(parts) >= 3 and parts[0] == "files" and parts[2] == "raw":
        return parts[1]
    return None


def _extract_chatgpt_claim_file_id(path: str) -> str | None:
    clean_path = path.split("?", 1)[0]
    match = re.search(r"/upload_reservations/file_([^/]+)/claim_and_finish$", clean_path)
    return match.group(1) if match else None


def _completion_file_ids(data) -> list[str]:
    found = []

    def walk(value):
        if isinstance(value, dict):
            for key, child in value.items():
                if str(key).lower() in {"file_id", "fileid", "id"} and isinstance(child, str):
                    found.append(child.removeprefix("file_"))
                walk(child)
        elif isinstance(value, list):
            for child in value:
                walk(child)

    walk(data)
    return found


def _rewrite_completion_file_size(flow: http.HTTPFlow, file_id: str | None = None) -> tuple[int, int] | None:
    data = _json_body(flow)
    if not data:
        return None

    candidate_ids = ([file_id] if file_id else []) + _completion_file_ids(data)
    sizes = _sanitized_sizes()
    matched_ids = matching_file_keys(candidate_ids, sizes)
    if not matched_ids:
        with _upload_state_lock:
            linked_raw_id = _upload_identity_diagnostic.resolve_reservation(
                get_client_ip(flow), file_id, sizes)
        if linked_raw_id is not None:
            matched_ids = [linked_raw_id]
    if len(matched_ids) != 1:
        return None
    matched_id = matched_ids[0]
    sanitized_size = sizes[matched_id]

    # Never apply one file's size to several fields in a batch request.
    def count_sizes(value):
        if isinstance(value, dict):
            return sum(1 if str(k).lower() == 'file_size' and type(v) is int
                       else count_sizes(v) for k, v in value.items())
        if isinstance(value, list):
            return sum(count_sizes(v) for v in value)
        return 0
    if count_sizes(data) != 1:
        return None

    previous_sizes = []

    def rewrite(value):
        if isinstance(value, dict):
            for key, child in value.items():
                if str(key).lower() == "file_size" and isinstance(child, (int, float)):
                    previous_sizes.append(int(child))
                    value[key] = sanitized_size
                else:
                    rewrite(child)
        elif isinstance(value, list):
            for child in value:
                rewrite(child)

    rewrite(data)
    if not previous_sizes:
        return None

    encoded = json.dumps(data, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    flow.request.raw_content = encoded
    flow.request.headers["content-length"] = str(len(encoded))
    # Keep the mapping available for retries, as completion HTTP 200 is not final success.
    return previous_sizes[0], sanitized_size


def _guess_extension_from_content_type(content_type: str) -> str | None:
    content_type = (content_type or "").lower().split(";", 1)[0].strip()
    return {
        "text/plain": "txt",
        "application/pdf": "pdf",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    }.get(content_type)



def _infer_chatgpt_file_extension(
    file_bytes: bytes,
    content_type: str = "",
    filename: str | None = None,
) -> str | None:
    """
    ChatGPT Web의 실제 PUT에는 원본 파일명이 없을 수 있으므로
    파일명 → MIME → 파일 시그니처 순서로 TXT/PDF/DOCX/XLSX를 판별한다.
    """

    # 1. 원본 filename이 확보된 경우
    if filename and "." in filename:
        ext = filename.lower().rsplit(".", 1)[-1]
        if ext in {"txt", "pdf", "docx", "xlsx"}:
            return ext

    # 2. Content-Type
    ext = _guess_extension_from_content_type(content_type)
    if ext:
        return ext

    # 3. PDF magic bytes
    if file_bytes.startswith(b"%PDF-"):
        return "pdf"

    # 4. DOCX/XLSX는 ZIP 컨테이너이므로 내부 구조로 구분
    if file_bytes.startswith(b"PK"):
        try:
            with zipfile.ZipFile(BytesIO(file_bytes)) as zf:
                names = zf.namelist()

                if any(name.startswith("word/") for name in names):
                    return "docx"

                if any(name.startswith("xl/") for name in names):
                    return "xlsx"

        except (zipfile.BadZipFile, OSError):
            pass

    # 5. TXT fallback
    try:
        decoded = file_bytes.decode("utf-8")
        if decoded:
            # 바이너리 파일을 TXT로 오인하지 않도록 NUL 문자는 제외
            if "\x00" not in decoded:
                return "txt"
    except UnicodeDecodeError:
        pass

    return None



def create_chatgpt_raw_file_event(
    flow: http.HTTPFlow,
    file_id: str,
    filename: str | None,
    *,
    file_class: FileClass | None = None,
    request_id: str | None = None,
) -> dict:
    """file_class가 있으면(2026-10-10, inspect_upload 결과) 그 글자를 쓴다. 이름은 기록용(실제 형식은 file_class.kind)."""
    file_bytes = flow.request.raw_content or b""

    if file_class is not None:
        filename = filename or f"{file_id}.{file_class.kind}"
        extracted_text = file_class.text
    else:
        content_type = flow.request.headers.get("content-type", "")

        extension = _infer_chatgpt_file_extension(
            file_bytes=file_bytes,
            content_type=content_type,
            filename=filename,
        )

        if not extension:
            filename = filename or f"{file_id}.unknown"
        else:
            # 실제 filename이 없거나 확장자가 지원되지 않는 경우
            # 추출 가능한 확장자를 붙인다.
            if not filename or filename.lower().rsplit(".", 1)[-1] not in {
                "txt", "pdf", "docx", "xlsx"
            }:
                filename = f"{file_id}.{extension}"

        extracted_text = extract_file_text(filename, file_bytes)

    file_content = {
        "type": "file",
        "filename": filename,
        "content": extracted_text if extracted_text is not None else "",
        "file_id": file_id,
        "body_size": len(file_bytes),
    }

    return {
        "request_id": request_id or str(uuid.uuid4()),
        "timestamp": datetime.now(timezone.utc).isoformat(),
        "client": {"ip": get_client_ip(flow)},
        "service": "ChatGPT",
        "request": {
            "method": flow.request.method,
            "endpoint": flow.request.path.split("?", 1)[0],
        },
        "contents": [file_content],
        "metadata": {
            "content_count": 1,
            "body_size": len(file_bytes),
            "file_id": file_id,
        },
    }


def extract_file_text(
    filename: str | None,
    file_bytes: bytes,
) -> str | None:

    if not filename:
        return None

    extension = filename.lower().rsplit(".", 1)[-1]

    # --------------------------------------------------------
    # TXT
    # --------------------------------------------------------

    if extension == "txt":

        try:
            return file_bytes.decode(
                "utf-8",
                errors="ignore",
            )

        except Exception as e:

            print(
                "❌ TXT text extraction failed:",
                e,
            )

            return None

    # --------------------------------------------------------
    # PDF
    # --------------------------------------------------------

    if extension == "pdf":

        try:

            reader = PdfReader(
                BytesIO(file_bytes)
            )

            texts = []

            for page in reader.pages:

                text = page.extract_text()

                if text:
                    texts.append(text)

            return "\n".join(texts)

        except Exception as e:

            print(
                "❌ PDF text extraction failed:",
                e,
            )

            return None

    # --------------------------------------------------------
    # DOCX
    # --------------------------------------------------------

    if extension == "docx":

        try:

            document = Document(
                BytesIO(file_bytes)
            )

            texts = []

            # 일반 문단
            for paragraph in document.paragraphs:

                if paragraph.text:
                    texts.append(
                        paragraph.text
                    )

            # 표 내부 텍스트
            for table in document.tables:

                for row in table.rows:

                    values = []

                    for cell in row.cells:

                        if cell.text:
                            values.append(
                                cell.text
                            )

                    if values:
                        texts.append(
                            " | ".join(values)
                        )

            return "\n".join(texts)

        except Exception as e:

            print(
                "❌ DOCX text extraction failed:",
                e,
            )

            return None

    # --------------------------------------------------------
    # XLSX
    # --------------------------------------------------------

    if extension == "xlsx":

        try:

            workbook = load_workbook(
                BytesIO(file_bytes),
                read_only=True,
                data_only=True,
            )

            texts = []

            for sheet in workbook.worksheets:

                # Sheet 이름도 정보로 포함
                texts.append(
                    f"[Sheet: {sheet.title}]"
                )

                for row in sheet.iter_rows(
                    values_only=True
                ):

                    values = [
                        str(value)
                        for value in row
                        if value is not None
                    ]

                    if values:

                        texts.append(
                            " | ".join(values)
                        )

            return "\n".join(texts)

        except Exception as e:

            print(
                "❌ XLSX text extraction failed:",
                e,
            )

            return None

    # --------------------------------------------------------
    # CSV / JSON / HTML / HTM / RTF / ODT(plain XML fallback)
    # --------------------------------------------------------

    if extension in {"csv", "json", "html", "htm", "rtf"}:

        try:
            return file_bytes.decode(
                "utf-8",
                errors="ignore",
            )
        except Exception as e:
            print(
                f"❌ {extension.upper()} text extraction failed:",
                e,
            )
            return None

    # ODT는 ZIP 기반 XML 문서이므로 기존 zipfile import를 이용해
    # content.xml의 텍스트를 간단히 추출한다.
    if extension == "odt":

        try:
            import re as _re
            with zipfile.ZipFile(BytesIO(file_bytes)) as archive:
                xml = archive.read("content.xml").decode(
                    "utf-8",
                    errors="ignore",
                )
            xml = _re.sub(r"<[^>]+>", " ", xml)
            return " ".join(xml.split())
        except Exception as e:
            print(
                "❌ ODT text extraction failed:",
                e,
            )
            return None

    # --------------------------------------------------------
    # 지원하지 않는 형식
    # --------------------------------------------------------

    print(
        f"⚠️ Unsupported file type: {filename}"
    )

    return None


# ============================================================
# Responses API 콘텐츠 추출
# ============================================================

def extract_content(
    body: dict,
) -> list[dict]:

    results = []

    input_data = body.get("input")

    # --------------------------------------------------------
    # input = 문자열
    # --------------------------------------------------------

    if isinstance(input_data, str):

        results.append({
            "type": "prompt",
            "content": input_data,
        })

    # --------------------------------------------------------
    # input = 배열
    # --------------------------------------------------------

    elif isinstance(input_data, list):

        for item in input_data:

            # 문자열
            if isinstance(item, str):

                results.append({
                    "type": "prompt",
                    "content": item,
                })

                continue

            # 객체가 아니면 무시
            if not isinstance(item, dict):
                continue

            content = item.get("content")

            # ------------------------------------------------
            # content = 문자열
            # ------------------------------------------------

            if isinstance(content, str):

                results.append({
                    "type": "prompt",
                    "content": content,
                })

            # ------------------------------------------------
            # content = 배열
            # ------------------------------------------------

            elif isinstance(content, list):

                for content_item in content:

                    if not isinstance(
                        content_item,
                        dict,
                    ):
                        continue

                    content_type = content_item.get(
                        "type"
                    )

                    # ----------------------------------------
                    # Prompt
                    # ----------------------------------------

                    if content_type == "input_text":

                        results.append({
                            "type": "prompt",
                            "content": content_item.get(
                                "text",
                                "",
                            ),
                        })

                    # ----------------------------------------
                    # File
                    # ----------------------------------------

                    elif content_type == "input_file":

                        results.append({
                            "type": "file",
                            "file_id": content_item.get(
                                "file_id"
                            ),
                            "filename": content_item.get(
                                "filename"
                            ),
                        })

                    # ----------------------------------------
                    # Image
                    # ----------------------------------------

                    elif content_type in (
                        "input_image",
                        "image_url",
                    ):

                        results.append({
                            "type": "image",
                            "content": content_item,
                        })

    return results


# ============================================================
# Responses API Gateway Event
# ============================================================

def create_responses_event(
    flow: http.HTTPFlow,
    body: dict,
) -> dict:

    contents = extract_content(body)

    event = {

        "request_id": str(
            uuid.uuid4()
        ),

        "timestamp": (
            datetime.now(
                timezone.utc
            ).isoformat()
        ),

        "client": {
            "ip": get_client_ip(flow),
        },

        "service": "OpenAI",

        "request": {
            "method": flow.request.method,
            "endpoint": _url_path(flow.request.path),
            "model": body.get("model"),
        },

        "contents": contents,

        "metadata": {
            "content_count": len(contents),
            "body_size": len(
                flow.request.raw_content or b""
            ),
        },
    }

    return event


# ============================================================
# Gemini Web Resumable File Upload
# ============================================================

def _remember_gemini_upload_filename(flow: http.HTTPFlow) -> str | None:
    """
    Gemini의 resumable upload 시작 요청에서 파일명을 기억한다.
    실제 파일 내용은 로그에 출력하지 않는다.
    """
    body = flow.request.raw_content or b""
    if not body:
        return None

    try:
        decoded = unquote_plus(body.decode("utf-8", errors="ignore")).strip()
    except Exception:
        return None

    prefix = "File name:"
    if not decoded.startswith(prefix):
        return None

    filename = decoded[len(prefix):].strip()
    if not filename:
        return None

    GEMINI_FILE_META[get_client_ip(flow) or "unknown"] = filename
    print("→ Remembered Gemini file metadata:", filename)
    return filename


def _get_gemini_upload_filename(flow: http.HTTPFlow) -> str | None:
    return GEMINI_FILE_META.get(get_client_ip(flow) or "unknown")


def _create_gemini_resumable_file_event(
    flow: http.HTTPFlow,
    filename: str | None,
) -> dict:
    file_bytes = flow.request.raw_content or b""

    # 파일명으로 확장자를 우선 판별하고,
    # 없으면 기존 시그니처 판별 로직을 재사용한다.
    extension = _infer_chatgpt_file_extension(
        file_bytes=file_bytes,
        content_type=flow.request.headers.get("content-type", ""),
        filename=filename,
    )

    if not extension:
        filename = filename or "gemini_upload.unknown"
    elif not filename or "." not in filename:
        filename = f"gemini_upload.{extension}"

    extracted_text = extract_file_text(filename, file_bytes)

    file_content = {
        "type": "file",
        "filename": filename,
        "content": extracted_text if extracted_text is not None else "",
        "body_size": len(file_bytes),
    }

    return {
        "request_id": str(uuid.uuid4()),
        "timestamp": datetime.now(timezone.utc).isoformat(),
        "client": {"ip": get_client_ip(flow)},
        "service": "Gemini",
        "request": {
            "method": flow.request.method,
            "endpoint": flow.request.path.split("?", 1)[0],
        },
        "contents": [file_content],
        "metadata": {
            "content_count": 1,
            "body_size": len(file_bytes),
        },
    }


def handle_gemini_resumable_upload(
    flow: http.HTTPFlow,
    identity: dict,
) -> bool:
    """
    Gemini Web의 push.clients6.google.com resumable upload를 처리한다.

    반환값:
      True  = Detection에 의해 요청을 종료함(BLOCK 등)
      False = 정상적으로 다음 단계로 진행
    """
    req = flow.request

    if req.host != "push.clients6.google.com" or req.method != "POST":
        return False

    raw_path = req.path.split("?", 1)[0]
    query = req.query

    # 업로드 시작 요청: 파일명만 전달
    if raw_path == "/upload/" and "upload_protocol" not in query:
        remembered = _remember_gemini_upload_filename(flow)
        if remembered:
            print("\n===== GEMINI FILE UPLOAD INIT =====")
            print("Filename:", remembered)
            print("→ Waiting for resumable file bytes")
        return False

    # 실제 resumable file bytes
    if (
        raw_path != "/upload/"
        or query.get("upload_protocol") != "resumable"
        or not query.get("upload_id")
    ):
        return False

    file_bytes = req.raw_content or b""
    if not file_bytes:
        return False

    filename = _get_gemini_upload_filename(flow)

    # DOCX/XLSX/기타 지원 파일을 시그니처로 보완
    if not filename:
        extension = _infer_chatgpt_file_extension(
            file_bytes=file_bytes,
            content_type=req.headers.get("content-type", ""),
        )
        filename = f"gemini_upload.{extension}" if extension else None

    print("\n===== GEMINI FILE UPLOAD DETECTED =====")
    print("Host:", req.host)
    print("Path:", raw_path)
    print("Filename:", filename or "(unknown)")
    print("Body size:", len(file_bytes))
    print("Upload protocol: resumable")

    if not filename:
        print("→ Filename could not be resolved; PASS THROUGH")
        return False

    event = _create_gemini_resumable_file_event(flow, filename)
    print_gateway_event(event)
    send_to_next_module(event)

    normalized_text = extract_file_checkable_text(event)

    if not normalized_text:
        print("→ No checkable text extracted; PASS THROUGH")
        return False

    print("→ Extracted file text length:", len(normalized_text))
    print("→ Sending extracted Gemini file text to Detection Module")

    result = check_detection(
        DETECTION_BASE_URL,
        request_id=event["request_id"],
        normalized_text=normalized_text,
        source={
            "service": "Gemini",
            "endpoint": raw_path,
            "content_type": "file",
            "filename": filename,
        },
    )

    print("\n===== GEMINI FILE DETECTION RESULT =====")
    print_detection_summary(result)

    if enforce_detection_result(
        flow,
        result,
        dashboard_context={
            "identity": identity,
            "targetLlm": "Gemini",
            "contentType": "file",
        },
    ):
        return True

    print("→ Gemini file upload: PASS")
    return False



# ============================================================
# Multipart File Upload 파싱
# ============================================================

def extract_file_upload_info(
    flow: http.HTTPFlow,
) -> dict:

    req = flow.request

    content_type = req.headers.get(
        "content-type",
        "",
    )

    # ========================================================
    # Multipart File Upload 파싱
    # ========================================================

    # OpenAI API는 아래의 기존 상세 파싱 로직으로 처리
    # (나머지 LLM 웹서비스는 현재 여기서 통과)
    body = req.raw_content or b""

    filename = None
    purpose = None
    file_bytes = None

    # --------------------------------------------------------
    # Content-Type 확인
    # --------------------------------------------------------

    if not content_type.startswith(
        "multipart/form-data"
    ):

        return {
            "filename": None,
            "purpose": None,
            "body_size": len(body),
            "file_bytes": None,
        }

    # --------------------------------------------------------
    # Multipart parser용 가상 MIME Header 생성
    # --------------------------------------------------------

    try:

        raw_message = (
            b"Content-Type: "
            + content_type.encode()
            + b"\r\n"
            + b"\r\n"
            + body
        )

        message = BytesParser(
            policy=default
        ).parsebytes(
            raw_message
        )

        # ----------------------------------------------------
        # Multipart 각 Part 처리
        # ----------------------------------------------------

        for part in message.iter_parts():

            disposition = part.get(
                "Content-Disposition",
                "",
            )

            part_name = part.get_param(
                "name",
                header="Content-Disposition",
            )

            # -----------------------------------------------
            # purpose
            # -----------------------------------------------

            if part_name == "purpose":

                payload = part.get_payload(
                    decode=True
                )

                if payload:

                    purpose = payload.decode(
                        "utf-8",
                        errors="ignore",
                    ).strip()

            # -----------------------------------------------
            # file
            # -----------------------------------------------

            elif part_name == "file":

                filename = part.get_filename()

                file_bytes = part.get_payload(
                    decode=True
                )

    except Exception as e:

        print(
            "❌ Multipart parsing failed:",
            e,
        )

    return {
        "filename": filename,
        "purpose": purpose,
        "body_size": len(body),
        "file_bytes": file_bytes,
    }


# ============================================================
# File Upload Gateway Event
# ============================================================

def create_file_event(
    flow: http.HTTPFlow,
    *,
    file_class: FileClass | None = None,
    request_id: str | None = None,
) -> dict:
    """file_class 없음(기본) = OpenAI /v1/files: 파일 이름 확장자로 추출(지금 그대로).
    file_class 있음 = ChatGPT 웹 multipart(2026-10-10): inspect_upload가 바이트로 정한 형식의 글자."""

    file_info = extract_file_upload_info(
        flow
    )

    filename = file_info["filename"]
    file_bytes = file_info["file_bytes"]

    # --------------------------------------------------------
    # 실제 파일 텍스트 추출
    # --------------------------------------------------------

    extracted_text = None

    if file_class is not None:

        extracted_text = file_class.text

    elif file_bytes:

        extracted_text = extract_file_text(
            filename,
            file_bytes,
        )

    # --------------------------------------------------------
    # Content 생성
    # --------------------------------------------------------

    file_content = {
        "type": "file",
        "filename": filename,
        "purpose": file_info["purpose"],
    }

    # 텍스트 추출 성공 시 content 추가
    if extracted_text is not None:

        file_content["content"] = extracted_text

    # --------------------------------------------------------
    # Event 생성
    # --------------------------------------------------------

    event = {

        "request_id": request_id or str(
            uuid.uuid4()
        ),

        "timestamp": (
            datetime.now(
                timezone.utc
            ).isoformat()
        ),

        "client": {
            "ip": get_client_ip(flow),
        },

        "service": "OpenAI",

        "request": {
            "method": flow.request.method,
            "endpoint": _url_path(flow.request.path),
        },

        "contents": [
            file_content
        ],

        "metadata": {
            "content_count": 1,
            "body_size": file_info["body_size"],
        },
    }

    return event


# ============================================================
# Detection Module 전달용 파일 텍스트 추출
# ============================================================

def extract_file_checkable_text(event: dict) -> str | None:
    """
    Gateway Event의 file content 중 실제 파일에서 추출한 텍스트만
    Detection Module 전달용 문자열로 합친다.
    """
    texts = []

    for item in event.get("contents", []):
        if not isinstance(item, dict):
            continue

        if item.get("type") != "file":
            continue

        content = item.get("content")

        if isinstance(content, str) and content.strip():
            texts.append(content.strip())

    if not texts:
        return None

    return "\n".join(texts)


# ============================================================
# 다음 모듈로 Event 전달
# ============================================================

def send_to_next_module(
    event: dict,
):

    try:

        session = requests.Session()

        # 시스템 Proxy 환경변수 무시
        # Gateway → localhost 통신이
        # 다시 Gateway를 거치지 않도록 함
        session.trust_env = False

        response = session.post(
            NEXT_MODULE_URL,
            json=event,
            timeout=3,
        )

        print(
            "\n===== NEXT MODULE RESPONSE ====="
        )

        print(
            "Status:",
            response.status_code,
        )

        print(
            "Response bytes:",
            len(response.content or b""),
        )

    except requests.RequestException as e:

        print(
            "\n❌ Failed to send Gateway Event"
        )

        print(
            "Error:",
            e,
        )


# ============================================================
# Event 출력
# ============================================================

def print_gateway_event(
    event: dict,
):
    """Log metadata only; never print prompt or extracted file content."""
    request = event.get("request") or {}
    metadata = event.get("metadata") or {}
    content_types = [
        item.get("type")
        for item in event.get("contents", [])
        if isinstance(item, dict) and item.get("type")
    ]
    print("\n===== GATEWAY EVENT METADATA =====")
    print(
        json.dumps(
            {
                "request_id": event.get("request_id"),
                "timestamp": event.get("timestamp"),
                "service": event.get("service"),
                "method": request.get("method"),
                "endpoint": request.get("endpoint"),
                "content_types": content_types,
                "content_count": metadata.get("content_count"),
                "body_size": metadata.get("body_size"),
            },
            ensure_ascii=False,
        )
    )


def print_detection_summary(result):
    """Log only non-sensitive decision metadata."""
    print("\n===== DETECTION RESULT METADATA =====")
    print("Request ID:", getattr(result, "request_id", None))
    print("Decided by:", getattr(result, "decided_by", None))
    print("Action:", getattr(result, "action", None))
    print("Blocked:", getattr(result, "blocked", None))
    print("Degraded:", getattr(result, "degraded", None))
    print("HTTP status:", getattr(result, "http_status", None))
    print("Total ms:", getattr(result, "total_ms", None))


def _masked_text(text: str, spans: list[dict]) -> str | None:
    """Apply metadata-only Detection spans without ever logging matched values."""
    validated = []
    for span in spans or []:
        if not isinstance(span, dict):
            return None
        start, end = span.get("start"), span.get("end")
        entity_type = str(span.get("entity_type") or "SENSITIVE")
        if not isinstance(start, int) or not isinstance(end, int) or start < 0 or end <= start or end > len(text):
            return None
        validated.append((start, end, entity_type))
    validated.sort()
    if not validated or any(left[1] > right[0] for left, right in zip(validated, validated[1:])):
        return None
    result = text
    for start, end, entity_type in reversed(validated):
        result = result[:start] + f"[MASKED:{entity_type}]" + result[end:]
    return result


def _replace_matching_json_strings(value, normalized_target: str, replacement: str) -> int:
    count = 0
    if isinstance(value, list):
        for index, item in enumerate(value):
            if isinstance(item, str) and normalize_text(item) == normalized_target:
                value[index] = replacement
                count += 1
            else:
                count += _replace_matching_json_strings(item, normalized_target, replacement)
    elif isinstance(value, dict):
        for key, item in list(value.items()):
            if isinstance(item, str) and normalize_text(item) == normalized_target:
                value[key] = replacement
                count += 1
            else:
                count += _replace_matching_json_strings(item, normalized_target, replacement)
    return count


def _apply_json_mask(flow: http.HTTPFlow, normalized_text: str, result, body: dict | None = None) -> bool:
    replacement = getattr(result, "masked_text", None)
    if replacement is None:
        replacement = _masked_text(normalized_text, getattr(result, "mask_spans", None) or [])
    if replacement is None:
        return False
    if body is None:
        body = _json_body(flow)
    if not isinstance(body, dict):
        return False
    if _replace_matching_json_strings(body, normalized_text, replacement) < 1:
        return False
    flow.request.set_text(json.dumps(body, ensure_ascii=False, separators=(",", ":")))
    capture_mask(flow, normalized_text, replacement, result, "JSON_TEXT")
    return True


def _apply_gemini_mask(flow: http.HTTPFlow, normalized_text: str, result) -> bool:
    replacement = getattr(result, "masked_text", None)
    if replacement is None:
        replacement = _masked_text(normalized_text, getattr(result, "mask_spans", None) or [])
    if replacement is None:
        return False
    try:
        form = flow.request.urlencoded_form
        outer = json.loads(form.get("f.req") or "")
        if not isinstance(outer, list) or len(outer) < 2 or not isinstance(outer[1], str):
            return False
        nested = json.loads(outer[1])
        if _replace_matching_json_strings(nested, normalized_text, replacement) < 1:
            return False
        outer[1] = json.dumps(nested, ensure_ascii=False, separators=(",", ":"))
        form["f.req"] = json.dumps(outer, ensure_ascii=False, separators=(",", ":"))
        capture_mask(flow, normalized_text, replacement, result, "GEMINI_FORM_TEXT")
        return True
    except (ValueError, TypeError):
        return False


# ============================================================
# ChatGPT Web Request 사용자 프롬프트 추출
# ============================================================

def extract_chatgpt_prompt(flow: http.HTTPFlow) -> str | None:
    """
    실제 ChatGPT 메시지 전송 요청인
    /backend-api/f/conversation 의 Request Body에서
    사용자(role=user)의 텍스트 프롬프트만 추출한다.

    대상 구조:
    messages[].author.role == "user"
    messages[].content.content_type == "text" | "multimodal_text"
    messages[].content.parts[]

    2026-10-10: 이미지가 붙은 질문(multimodal_text)도 검사한다. parts에는 질문 글자(문자열)와
    이미지 칸(asset_pointer를 담은 dict)이 섞여 있고, 문자열만 고른다(이미지 칸은 그대로 간다).
    """
    req = flow.request

    if req.method.upper() != "POST":
        return None

    if req.path.split("?", 1)[0] != "/backend-api/f/conversation":
        return None

    content_type = req.headers.get("content-type", "").lower()
    if "application/json" not in content_type:
        return None

    try:
        body_text = req.get_text(strict=False)
        if not body_text:
            return None

        body = json.loads(body_text)

    except (json.JSONDecodeError, TypeError, UnicodeDecodeError) as e:
        print("❌ ChatGPT conversation JSON parsing failed:", e)
        return None

    messages = body.get("messages")
    if not isinstance(messages, list):
        return None

    prompts = []

    for message in messages:
        if not isinstance(message, dict):
            continue

        author = message.get("author")
        if not isinstance(author, dict):
            continue

        if author.get("role") != "user":
            continue

        content = message.get("content")
        if not isinstance(content, dict):
            continue

        if content.get("content_type") not in {"text", "multimodal_text"}:
            continue

        parts = content.get("parts")
        if not isinstance(parts, list):
            continue

        for part in parts:
            if isinstance(part, str) and part.strip():
                prompts.append(part)

    if not prompts:
        return None

    # 실제 전송에서 마지막 user part가 이번 요청의 새 프롬프트다. 하나씩 검사해야
    # MASK 좌표를 요청 JSON의 정확한 문자열에 다시 적용할 수 있다.
    return prompts[-1]


# ============================================================
# Gemini Web Request 구조 확인용 DEBUG
# ============================================================

def debug_gemini_request_structure(flow: http.HTTPFlow):
    """
    Gemini 웹 요청의 구조만 확인하기 위한 DEBUG 함수.
    프롬프트/쿠키/Authorization/세션값/전체 body는 출력하지 않는다.
    """
    req = flow.request

    if req.method.upper() != "POST":
        return

    path = req.path.split("?", 1)[0]

    if not (
        req.host == "gemini.google.com"
        and (
            "StreamGenerate" in path
            or "batchexecute" in path
        )
    ):
        return

    raw_body = req.raw_content or b""
    content_type = req.headers.get("content-type", "")

    print("\n===== GEMINI REQUEST STRUCTURE DEBUG =====")
    print("Endpoint:", path)
    print("Content-Type:", content_type.split(";", 1)[0])
    print("Body size:", len(raw_body))

    if "application/json" in content_type.lower():
        try:
            body = json.loads(req.get_text(strict=False) or "")
            if isinstance(body, dict):
                print("Body format: JSON object")
                print("Top-level keys:", list(body.keys())[:30])
            elif isinstance(body, list):
                print("Body format: JSON array")
                print("Array length:", len(body))
            else:
                print("Body format: JSON", type(body).__name__)
        except Exception as e:
            print("JSON parse:", type(e).__name__)

    elif "application/x-www-form-urlencoded" in content_type.lower():
        try:
            from urllib.parse import parse_qs

            decoded = raw_body.decode("utf-8", errors="ignore")
            fields = parse_qs(decoded, keep_blank_values=True)

            print("Body format: x-www-form-urlencoded")
            print("Field names:", list(fields.keys())[:50])

            field_lengths = {
                key: [len(v) for v in values[:5]]
                for key, values in fields.items()
            }
            print("Field value lengths:", field_lengths)

        except Exception as e:
            print("Form parse:", type(e).__name__)

    else:
        print("Body format: other/unknown")

        printable = sum(
            1 for b in raw_body
            if b in (9, 10, 13) or 32 <= b <= 126
        )

        print(
            "Printable byte ratio:",
            round(printable / len(raw_body), 3) if raw_body else 0,
        )

    print("===== END GEMINI REQUEST STRUCTURE DEBUG =====")



# ============================================================
# Gemini StreamGenerate f.req 구조 확인용 DEBUG
# ============================================================

def debug_gemini_request_response(flow: http.HTTPFlow, phase: str):
    """
    Gemini request/response를 같은 flow.id 기준으로 추적한다.
    실제 body/query 값은 출력하지 않고 구조와 길이만 출력한다.
    """
    req = flow.request
    if req.host != "gemini.google.com" or req.method.upper() != "POST":
        return

    path = req.path.split("?", 1)[0]
    if "BardChatUi" not in path:
        return

    try:
        client_conn = getattr(flow, "client_conn", None)
        server_conn = getattr(flow, "server_conn", None)

        print("\n===== GEMINI REQUEST/RESPONSE FLOW DEBUG =====")
        print("Phase:", phase)
        print("Flow ID:", getattr(flow, "id", None))
        print("Client connection ID:", getattr(client_conn, "id", None))
        print("Server connection ID:", getattr(server_conn, "id", None))
        print("Method:", req.method)
        print("Host:", req.host)
        print("Path:", path)

        if path.endswith("/batchexecute"):
            print("Endpoint type: BATCHEXECUTE")
        elif "assistant" in path.lower():
            print("Endpoint type: ASSISTANT")
        else:
            print("Endpoint type: OTHER_BARDCHATUI")

        if phase == "REQUEST":
            body = req.raw_content or b""
            ctype = req.headers.get("content-type", "")
            print("Content-Type:", ctype.split(";", 1)[0].strip())
            print("Request body size:", len(body))

            if "application/x-www-form-urlencoded" in ctype.lower():
                from urllib.parse import parse_qs
                try:
                    fields = parse_qs(
                        body.decode("utf-8", errors="ignore"),
                        keep_blank_values=True,
                    )
                    print("Request form fields:", list(fields.keys())[:30])
                    for name, values in fields.items():
                        print(
                            f"Request field {name}: "
                            f"{len(values[0]) if values else 0} chars"
                        )
                except Exception as e:
                    print("Request form inspection failed:", type(e).__name__)

        elif phase == "RESPONSE":
            response = flow.response
            if response is None:
                print("Response: NONE")
            else:
                ctype = response.headers.get("content-type", "")
                body = response.raw_content or b""
                print("Response status:", response.status_code)
                print("Response Content-Type:", ctype.split(";", 1)[0].strip())
                print("Response body size:", len(body))

        print("===== END GEMINI REQUEST/RESPONSE FLOW DEBUG =====\n")
    except Exception as e:
        print("\n===== GEMINI REQUEST/RESPONSE FLOW DEBUG ERROR =====")
        print("Phase:", phase)
        print("Error type:", type(e).__name__)
        print("===== END GEMINI REQUEST/RESPONSE FLOW DEBUG ERROR =====\n")


def debug_gemini_flow_identity(flow: http.HTTPFlow):
    """
    Gemini request flow를 정확히 식별하기 위한 DEBUG.

    목적:
    - mitmproxy 한 줄 로그의 assistant... 요청과
      addon 내부에서 처리되는 flow를 매칭
    - batchexecute와 assistant 요청이 실제로 별도 flow인지 확인
    - request body 내용이나 query string은 출력하지 않음
    """

    req = flow.request

    if req.host != "gemini.google.com":
        return

    if req.method.upper() != "POST":
        return

    path = req.path.split("?", 1)[0]

    if "BardChatUi" not in path:
        return

    try:
        raw_body = req.raw_content or b""
        content_type = req.headers.get("content-type", "")
        content_type_base = content_type.split(";", 1)[0].strip()

        # flow 객체에서 안전하게 확인 가능한 식별 정보만 사용
        flow_id = getattr(flow, "id", None)

        client_conn = getattr(flow, "client_conn", None)
        server_conn = getattr(flow, "server_conn", None)

        client_id = getattr(client_conn, "id", None)
        server_id = getattr(server_conn, "id", None)

        client_addr = getattr(client_conn, "address", None)
        server_addr = getattr(server_conn, "address", None)

        print("\n===== GEMINI FLOW IDENTITY DEBUG =====")
        print("Flow ID:", flow_id)
        print("Client connection ID:", client_id)
        print("Server connection ID:", server_id)
        print("Client address:", client_addr)
        print("Server address:", server_addr)
        print("Method:", req.method)
        print("Host:", req.host)
        print("Path:", path)
        print("Path type:", type(path).__name__)
        print("Content-Type:", content_type_base)
        print("Body size:", len(raw_body))

        if path.endswith("/batchexecute"):
            print("Endpoint type: BATCHEXECUTE")
        elif "assistant" in path.lower():
            print("Endpoint type: ASSISTANT")
        else:
            print("Endpoint type: OTHER_BARDCHATUI")

        # body의 구조적 정보만 출력
        if "application/x-www-form-urlencoded" in content_type.lower():
            try:
                from urllib.parse import parse_qs

                decoded_body = raw_body.decode(
                    "utf-8",
                    errors="ignore",
                )

                fields = parse_qs(
                    decoded_body,
                    keep_blank_values=True,
                )

                print(
                    "Form field names:",
                    list(fields.keys())[:30],
                )

                for field_name, values in fields.items():
                    value_length = (
                        len(values[0])
                        if values
                        else 0
                    )

                    print(
                        f"Field {field_name}: "
                        f"{value_length} chars"
                    )

            except Exception as e:
                print(
                    "Form inspection failed:",
                    type(e).__name__,
                )

        print("===== END GEMINI FLOW IDENTITY DEBUG =====\n")

    except Exception as e:
        print("\n===== GEMINI FLOW IDENTITY DEBUG ERROR =====")
        print("Error type:", type(e).__name__)
        print("===== END GEMINI FLOW IDENTITY DEBUG ERROR =====\n")


def debug_gemini_assistant_path(flow: http.HTTPFlow):
    """
    Gemini 요청 중 assistant 관련 endpoint가 실제로 어떤 path로
    들어오는지 확인하기 위한 최우선 DEBUG.

    보안상 query string과 body 내용은 출력하지 않는다.
    """

    req = flow.request

    if req.host != "gemini.google.com":
        return

    if req.method.upper() != "POST":
        return

    path = req.path.split("?", 1)[0]

    # Gemini BardChatUi/data 계열 POST를 넓게 확인하고,
    # 특히 assistant 문자열이 포함된 요청을 별도로 표시한다.
    if "BardChatUi" not in path and "assistant" not in path.lower():
        return

    try:
        content_type = req.headers.get("content-type", "")
        content_type_base = content_type.split(";", 1)[0].strip()
        raw_body = req.raw_content or b""

        is_assistant = "assistant" in path.lower()

        print("\n===== GEMINI PATH DEBUG =====")

        if is_assistant:
            print("!!! ASSISTANT ENDPOINT DETECTED !!!")

        print("Method:", req.method)
        print("Host:", req.host)
        print("Path:", path)
        print("Path repr:", repr(path))
        print("Content-Type:", content_type_base)
        print("Body size:", len(raw_body))

        if "application/x-www-form-urlencoded" in content_type.lower():
            try:
                from urllib.parse import parse_qs

                decoded_body = raw_body.decode(
                    "utf-8",
                    errors="ignore",
                )

                fields = parse_qs(
                    decoded_body,
                    keep_blank_values=True,
                )

                print("Form field names:", list(fields.keys())[:30])

                for field_name, values in fields.items():
                    if field_name == "f.req":
                        print(
                            "Field f.req:",
                            f"{len(values[0]) if values else 0} chars"
                        )
                    elif field_name == "at":
                        print(
                            "Field at:",
                            f"{len(values[0]) if values else 0} chars"
                        )
                    else:
                        print(
                            f"Field {field_name}:",
                            f"{len(values[0]) if values else 0} chars"
                        )

            except Exception as e:
                print(
                    "Form inspection failed:",
                    type(e).__name__,
                )

        print("===== END GEMINI PATH DEBUG =====\n")

    except Exception as e:
        print("\n===== GEMINI PATH DEBUG ERROR =====")
        print("Error type:", type(e).__name__)
        print("===== END GEMINI PATH DEBUG ERROR =====\n")


def debug_gemini_f_req_structure(flow: http.HTTPFlow):
    """
    Gemini f.req 구조를 분석한다.

    핵심:
    - f.req 전체 값을 출력하지 않는다.
    - 문자열의 실제 내용도 출력하지 않는다.
    - 대신 문자열이 JSON으로 다시 파싱되는지 재귀적으로 확인한다.
    - Gemini 요청에서 흔히 보이는 "JSON 안에 JSON 문자열" 구조를 찾는다.
    """

    req = flow.request
    path = req.path.split("?", 1)[0]

    if not (
        req.host == "gemini.google.com"
        and req.method.upper() == "POST"
        and "BardChatUi" in path
    ):
        return

    try:
        from urllib.parse import parse_qs

        raw_body = req.raw_content or b""
        content_type = req.headers.get("content-type", "")

        if "application/x-www-form-urlencoded" not in content_type.lower():
            return

        decoded_body = raw_body.decode(
            "utf-8",
            errors="ignore",
        )

        fields = parse_qs(
            decoded_body,
            keep_blank_values=True,
        )

        f_req_values = fields.get("f.req")

        if not f_req_values:
            return

        f_req = f_req_values[0]

        print("\n===== GEMINI NESTED f.req DEBUG =====")
        print("Endpoint:", path)
        print("f.req length:", len(f_req))

        try:
            data = json.loads(f_req)
        except Exception as e:
            print("Outer JSON parse: FAILED")
            print("Error type:", type(e).__name__)
            print("===== END GEMINI NESTED f.req DEBUG =====\n")
            return

        print("Outer JSON parse: SUCCESS")
        print("Outer root type:", type(data).__name__)

        def inspect_value(
            value,
            path="$",
            depth=0,
            max_depth=12,
            max_items=30,
        ):
            indent = "  " * depth

            if depth > max_depth:
                print(
                    f"{indent}{path}: "
                    "<MAX_DEPTH_REACHED>"
                )
                return

            if isinstance(value, dict):
                print(
                    f"{indent}{path}: "
                    f"object ({len(value)} keys)"
                )

                for index, (key, child) in enumerate(
                    value.items()
                ):
                    if index >= max_items:
                        print(
                            f"{indent}  ... "
                            f"<remaining {len(value) - max_items} keys omitted>"
                        )
                        break

                    safe_key = str(key)
                    if len(safe_key) > 80:
                        safe_key = safe_key[:80] + "..."

                    inspect_value(
                        child,
                        f"{path}.{safe_key}",
                        depth + 1,
                        max_depth,
                        max_items,
                    )

            elif isinstance(value, list):
                print(
                    f"{indent}{path}: "
                    f"array ({len(value)} items)"
                )

                for index, child in enumerate(
                    value[:max_items]
                ):
                    inspect_value(
                        child,
                        f"{path}[{index}]",
                        depth + 1,
                        max_depth,
                        max_items,
                    )

                if len(value) > max_items:
                    print(
                        f"{indent}  ... "
                        f"<remaining {len(value) - max_items} items omitted>"
                    )

            elif isinstance(value, str):
                print(
                    f"{indent}{path}: "
                    f"string (length={len(value)})"
                )

                # 핵심: 문자열 자체가 JSON이면 재귀적으로 구조만 확인.
                stripped = value.strip()

                if (
                    len(stripped) >= 2
                    and stripped[0] in "[{"
                    and stripped[-1] in "]}"
                ):
                    try:
                        nested = json.loads(stripped)

                        print(
                            f"{indent}  "
                            "Nested JSON parse: SUCCESS"
                        )
                        print(
                            f"{indent}  "
                            "Nested JSON root type:",
                            type(nested).__name__,
                        )

                        inspect_value(
                            nested,
                            f"{path}::<parsed>",
                            depth + 2,
                            max_depth,
                            max_items,
                        )

                    except Exception:
                        print(
                            f"{indent}  "
                            "Nested JSON parse: FAILED"
                        )

            elif isinstance(value, bool):
                print(
                    f"{indent}{path}: boolean"
                )

            elif value is None:
                print(
                    f"{indent}{path}: null"
                )

            elif isinstance(value, (int, float)):
                print(
                    f"{indent}{path}: number"
                )

            else:
                print(
                    f"{indent}{path}: "
                    f"{type(value).__name__}"
                )

        inspect_value(data)

        print("===== END GEMINI NESTED f.req DEBUG =====\n")

    except Exception as e:
        print("\n===== GEMINI NESTED f.req DEBUG ERROR =====")
        print("Error type:", type(e).__name__)
        print("===== END GEMINI NESTED f.req DEBUG ERROR =====\n")


# ============================================================
# Detection 호출
# ============================================================

# 같은 요청 안에서 다시 불러도 같을 무응답(연결 안 됨·제한시간 초과). PARSE_ERROR·OTHER는 입력에 따라 다를 수 있어 넘기지 않는다
_DETECTION_DOWN_REASONS = {"TIMEOUT", "UNREACHABLE"}
FILE_TOO_LARGE_MESSAGE = "파일이 커서 제한 시간 안에 검사할 수 없어 보안 정책에 따라 전송을 차단했습니다."
UNINSPECTABLE_MESSAGES = {
    "image": "이미지 파일은 내용을 검사할 수 없어 보안 정책에 따라 전송을 차단했습니다.",
    "scanned_pdf": "스캔한 PDF는 글자를 읽을 수 없어 보안 정책에 따라 전송을 차단했습니다.",
    "other": "이 파일 형식은 내용을 검사할 수 없어 보안 정책에 따라 전송을 차단했습니다.",
}


def _flow_settings(flow: http.HTTPFlow):
    """요청 하나는 한 판만: request() 시작 때 고정한 설정. 없으면(직접 부른 시험 등) 지금 설정."""
    settings = getattr(flow, "metadata", {}).get("gw_settings")
    return settings if settings is not None else GATEWAY_SETTINGS.current()


def too_large_result(action: str, request_id: str) -> DetectionResult:
    """파일이 대기 시간 안에 검사할 수 있는 양(글자·바이트·추출 시간)을 넘음. Detection을 부르지 않는다.
    action: 파일 처리 설정 'block'(403) | 'forward'(원본 전달 + 기록)."""
    blocked = action == "block"
    return DetectionResult(
        blocked=blocked, reason="gateway: file too large to inspect in time", request_id=request_id,
        decided_by="none", action="block" if blocked else "allow", reason_code="file_too_large",
        http_status=403 if blocked else None, user_message=FILE_TOO_LARGE_MESSAGE if blocked else None,
    )


def uninspected_result(flow: http.HTTPFlow, group: str, request_id: str) -> DetectionResult:
    """검사할 수 없는 파일(image·scanned_pdf·other) 또는 상한 초과(too_large)의 판정. 설정에 따라 차단·허용."""
    settings = _flow_settings(flow)
    if group == "too_large":
        return too_large_result(settings.file.action, request_id)
    group = group if group in UNINSPECTABLE_MESSAGES else "other"
    blocked = settings.uninspectable_action(group) == "block"
    return DetectionResult(
        blocked=blocked, reason="gateway: file not inspectable", request_id=request_id, decided_by="none",
        action="block" if blocked else "allow", reason_code=f"file_uninspectable_{group}",
        user_message=UNINSPECTABLE_MESSAGES[group] if blocked else None, http_status=403 if blocked else None,
    )


def _inspect(flow: http.HTTPFlow, kind: str, *, request_id: str, normalized_text: str, source: dict):
    """모든 Detection 호출(Gemini 업로드 제외)이 지나는 곳. kind: 'prompt' | 'file'(그 밖은 더 엄격한 파일 규칙).

    - 대기 시간·실패 처리(FAIL_OPEN/FAIL_CLOSED)는 요청 시작 때 고정한 설정의 그 종류 규칙
    - 파일: 글자 수가 대기 시간 안에 검사할 수 있는 양을 넘으면 Detection을 부르지 않는다(file_too_large)
    - 같은 요청 안의 무응답 넘기기: UNREACHABLE은 그대로, TIMEOUT은 이번 대기 시간이 같거나 짧을 때만 넘긴다.
      넘길 때는 받는 쪽 종류의 규칙으로 판정을 새로 만든다(질문 무응답 뒤 첨부 = 파일 규칙)
    """
    settings = _flow_settings(flow)
    flow.metadata["gw_inspected"] = True     # 지연 표본(_gateway_did_work): Gateway가 검사한 요청
    kind = kind if kind in ("prompt", "file") else "file"
    rule = settings.timeout_for(kind)
    policy = rule.failure_policy
    if kind == "file" and len(normalized_text or "") > file_max_chars(rule.seconds):
        print("→ Detection call skipped; file text exceeds the limit chars=", len(normalized_text))
        return too_large_result(rule.action, request_id)
    down = flow.metadata.get("gw_detection_down")
    if down and (down[0] == "UNREACHABLE" or rule.seconds <= down[1]):
        print("→ Detection call skipped; earlier call failed:", down[0])
        return failure_result(down[0], policy, request_id)
    result = check_detection(
        DETECTION_BASE_URL,
        request_id=request_id,
        normalized_text=normalized_text,
        source=source,
        timeout=rule.seconds,
        failure_policy=policy,
    )
    if getattr(result, "degraded", False) and getattr(result, "degraded_reason", None) in _DETECTION_DOWN_REASONS:
        flow.metadata["gw_detection_down"] = (result.degraded_reason, rule.seconds)
    return result


def detect_web_prompt(
    flow: http.HTTPFlow,
    service: str,
    prompt: str,
):
    """
    웹 LLM에서 추출한 사용자 프롬프트만 Detection Module로 전달.
    """

    request_id = str(uuid.uuid4())

    result = _inspect(
        flow,
        "prompt",
        request_id=request_id,
        normalized_text=prompt,
        source={
            "service": service,
            "endpoint": _url_path(flow.request.path),
            "content_type": "prompt",
        },
    )

    print("\n===== WEB LLM PROMPT METADATA =====")
    print("Request ID:", request_id)
    print("Prompt length:", len(prompt))
    print_detection_summary(result)

    return result



# ============================================================
# Detection 결과에 따른 BLOCK / PASS 집행
# ============================================================

def enforce_detection_result(
    flow: http.HTTPFlow,
    result,
    dashboard_context: dict | None = None,
    mask_callback=None,
) -> bool:
    """
    Detection Module의 응답에 따라 요청을 차단하거나 통과시킨다.

    Detection 응답의 decision 값을 그대로 집행한다(2026-09-15).
      action        "block" -> 차단 / "allow" -> 통과 / "allow_degraded" -> Detection 무응답, 통과(잠정 정책)
      user_message  사용자에게 보여줄 문구. 문서명·구간은 들어 있지 않다
      http_status   차단 응답 상태 코드(403)
    문서명·구간·일치율(doc_id, doc_span, input_span, score)과 reason은 사용자 응답에 넣지 않고
    Gateway 로그에만 남긴다 - 사용자가 어떤 문서의 어느 부분이 걸렸는지 알 수 없게.

    Returns:
        True  -> 요청을 BLOCK함
        False -> 요청을 PASS함
    """
    action = getattr(result, "action", None)
    if not action:  # 옛 클라이언트 결과(action 없음) 호환
        action = "block" if getattr(result, "blocked", False) else "allow"

    request_id = getattr(result, "request_id", None)
    detector = _result_detector(result)
    degraded = bool(getattr(result, "degraded", False))
    flow.metadata["gw_request_id"] = request_id
    flow.metadata["gw_detection_result"] = result
    flow.metadata["gw_action"] = action
    flow.metadata["gw_detector"] = detector
    flow.metadata["gw_degraded"] = degraded

    if action in {"block", "block_notify"}:
        print("\n===== GATEWAY ACTION =====")
        print("→ BLOCK")
        print("Request ID:", request_id)
        print("Decided by:", getattr(result, "decided_by", None))
        status = getattr(result, "http_status", None) or 403

        review_url = request_review_url(request_id)

        flow.response = http.Response.make(
            status,
            json.dumps(
                {
                    "error": {
                        "type": "inspection_unavailable" if degraded else "policy_violation",
                        "message": getattr(result, "user_message", None)
                        or "보안 정책에 따라 외부 전송을 차단했습니다. 관리자는 차단 이유 보기에서 확인할 수 있습니다.",
                        "reason_code": getattr(result, "reason_code", None),
                        "request_id": request_id,
                        **({"review_url": review_url, "review_label": "차단 이유 보기 · 관리자 로그인 필요"} if review_url else {}),
                    }
                },
                ensure_ascii=False,
            ).encode("utf-8"),
            {
                "Content-Type": "application/json; charset=utf-8",
                **({"X-LLMSGW-Review-URL": review_url} if review_url else {}),
            },
        )

        report_dashboard_event(
            flow,
            result,
            dashboard_context,
            enforcement_applied=True,
            http_status=status,
        )

        return True

    if action in {"mask", "mask_partial", "mask_full"}:
        applied = False
        try:
            applied = bool(mask_callback and mask_callback(result))
        except Exception as exc:
            print("MASK application failed:", type(exc).__name__)
        if not applied:
            # A requested mask that cannot be proven applied must never leak the
            # original value. Convert only this request to a fail-closed block.
            result.action = "block"
            result.blocked = True
            result.reason_code = "mask_application_failed"
            result.user_message = "민감정보 마스킹에 실패하여 전송이 차단되었습니다. 문의 시 요청 번호를 알려주세요."
            result.http_status = 403
            return enforce_detection_result(flow, result, dashboard_context)

        flow.metadata["gw_action"] = "mask"
        print("\n===== GATEWAY ACTION =====")
        print("→ MASK")
        print("Request ID:", request_id)
        report_dashboard_event(
            flow,
            result,
            dashboard_context,
            enforcement_applied=True,
            http_status=None,
        )
        return False

    print("\n===== GATEWAY ACTION =====")
    if action == "allow_degraded":
        print("→ PASS (Detection 무응답 - 잠정 통과)")
        print("Request ID:", request_id)
    elif action == "allow_notify":
        print("→ PASS WITH NOTIFICATION")
    else:
        print("→ PASS")

    report_dashboard_event(
        flow,
        result,
        dashboard_context,
        enforcement_applied=False,
        http_status=None,
    )

    return False


def _url_path(path) -> str:
    """주소의 경로만('?'·'#' 뒤를 뗀다, 2026-10-16). 대시보드 이벤트 endpoint·Detection source·로그 'Path:' 줄에 쓴다.
    ChatGPT 직접 업로드 주소(files.oaiusercontent.com …/raw?…&sig=…)·파일 내려받기 주소(estuary/content?…&sig=…)의
    서명·토큰이 이벤트(RDS)·저널에 남지 않게 한다. 호스트는 따로(host 칸·'Host:' 줄) 남는다."""
    return str(path or "").split("?", 1)[0].split("#", 1)[0]


def report_dashboard_event(
    flow: http.HTTPFlow,
    result,
    dashboard_context: dict | None,
    *,
    enforcement_applied: bool,
    http_status: int | None,
):
    """Best-effort Dashboard delivery that cannot change BLOCK/PASS."""
    if not dashboard_context or not dashboard_integration_enabled():
        return

    request_id = getattr(result, "request_id", None)
    try:
        event, skip_reason = build_dashboard_event(
            result=result,
            request_id=request_id or "unknown-request",
            identity=dashboard_context.get("identity") or {},
            target_llm=dashboard_context["targetLlm"],
            host=flow.request.host,
            endpoint=_url_path(flow.request.path),   # 10/16: 쿼리(업로드 주소 서명·토큰)는 싣지 않는다
            content_type=dashboard_context["contentType"],
            enforcement_applied=enforcement_applied,
            http_status=http_status,
            report_allow_events=report_allow_events_enabled(),
        )
    except Exception as exc:
        print(
            "[DASHBOARD] event mapping failed "
            f"requestId={request_id or 'unknown'} error_type={type(exc).__name__}"
        )
        return

    if event is None:
        print(
            "[DASHBOARD] event skipped "
            f"requestId={request_id or 'unknown'} reason={skip_reason}"
        )
        return

    # 한 요청에 prompt/file 등 여러 검사 항목이 있을 수 있으므로 모두 보존한다.
    try:
        if os.getenv('DASHBOARD_EXTENDED_METADATA', '0') == '1':
            attach_metadata(event, result, dashboard_context, _gateway_metrics_snapshot())
    except Exception as exc:
        print(f"[DASHBOARD] extended metadata omitted error_type={type(exc).__name__}")
    proof = flow.metadata.get("gw_mask_evidence", {}).get(request_id)
    if proof and event["finalAction"] == "MASK" and enforcement_applied:
        event["maskEvidence"] = proof
    # 질문과 첨부 파일 연결(2026-10-09): 파일 이벤트는 파일 번호, 질문 이벤트는 붙은 파일 번호들. 번호만 보낸다
    if dashboard_context.get("fileId"):
        event["fileId"] = dashboard_context["fileId"]
    if dashboard_context.get("attachedFileIds"):
        event["attachedFileIds"] = list(dashboard_context["attachedFileIds"])[:20]
    # 내 요청 기록(2026-10-09): 이 요청을 보낸 브라우저의 열람 키 지문. 키 값은 보내지 않는다
    viewer_key_hash = (dashboard_context.get("identity") or {}).get("viewer_key_hash")
    if isinstance(viewer_key_hash, str) and re.fullmatch(r"[0-9a-f]{64}", viewer_key_hash):
        event["viewerKeyHash"] = viewer_key_hash
    # 응답까지 받은 후 실제 Gateway latency/HTTP status를 확정해서 전송한다.
    flow.metadata.setdefault("gw_dashboard_events", []).append(event)
    # 같은 업로드 재시도 기록 한 번만(2026-10-11): 파일 신원은 이벤트에 싣지 않고 옆에 둔다(대시보드는 모르는 칸이면 거부)
    if dashboard_context.get("fileIdentity"):
        flow.metadata.setdefault("gw_dedupe_identity", {})[id(event)] = dashboard_context["fileIdentity"]


def _extension_notify(event: dict, settings) -> bool:
    """확장 프로그램 알림 표시(2026-10-10). 차단·가림은 늘, 허용은 '검사 없이 전송됨'(검사 불가·상한 초과·지연)이고
    허용 알림 설정이 켜졌을 때만. 확장용 사본에만 싣는다(대시보드는 모르는 칸이면 이벤트 전체를 거부)."""
    action = event.get("finalAction")
    if action in {"BLOCK", "MASK"}:
        return True
    if action == "ALLOW":
        uninspected = event.get("reasonCode") in _UNINSPECTED_REASONS or event.get("degraded") is True
        return bool(getattr(settings, "notify_on_allow", False)) and uninspected
    return False


# ============================================================
# 같은 업로드 재시도: 시도마다 막고, 기록·알림은 한 번만 (2026-10-11)
# ChatGPT는 막힌 업로드(403·503)를 같은 파일 번호로 1초 안팎에 두 번 더 보낸다(10/11 실측: 개인정보 PDF·HWP·HWPX
# 파일마다 차단 이벤트·확장 알림 3번). Claude 옛 화면은 같은 파일을 wiggle/upload-file·convert_document 두 곳에 보냈다.
# 집행은 그대로(시도마다 차단·503, 원본은 한 번도 안 나감). 대시보드 이벤트와 확장 프로그램 알림만 처음 한 번 보낸다.
#
# 열쇠 = 서비스 + 사용자(clientId·IP) + 파일 신원 + 최종 처리 + 사유 코드
#   파일 신원: ChatGPT 원본 PUT = 주소의 파일 번호 / 그 밖 = 파일 바이트 sha256 + 이름(여러 파일이면 모두) /
#   파일 part를 못 읽은 multipart = 경계 문자열을 뺀 본문 sha256(재시도마다 경계가 바뀐다). 메모리 안에서만, 로그에는 열쇠 해시 앞 12자
# 같은 업로드 = 앞 시도의 기록 시각(응답 직전) → 이번 시도 첫 바이트 시각이 창 안. 억제한 시도도 시각을 늦춘다(재시도가 이어지면 한 묶음)
#   첫 바이트 = mitmproxy timestamp_start(요청 머리를 받은 때). request 훅은 본문을 다 받은 뒤 불리므로(stream_large_bodies 없음)
#   훅 시작 시각으로 재면 업로드 시간이 간격에 들어간다(10/11 검토: 1 MiB를 4.6초에 걸쳐 보내는 재시도가 3번 기록·알림)
#   → 간격 = 재시도 대기만. 아래 창 근거(응답 → 다음 응답 = 대기 + 다음 시도 처리)보다 늘 짧다
# 창 근거(journalctl 10/01~10/11, ChatGPT 원본 PUT에서 같은 파일 번호가 두 번 이상 온 24건(2번 6·3번 18), 연속 간격 42개, 모두 403·503 뒤):
#   로그 시각은 stdout 버퍼가 비는 때(mitmdump가 응답 줄을 쓸 때)라 시도마다 응답 시각에 가깝다 → 간격 = 재시도 대기 + 다음 시도 처리
#   첫 재시도 0.236~1.868 s, 둘째 0.461~2.174 s, 한 파일 최대 3번. 처리 시간이 ms인 것만 보면 최대 1.208 s.
#   최대 2.174 s(10/07, DOCX 2만 자 칸 검사 시간 포함)는 재시도 대기의 상한이다. 창 = 그 2배(표본 42개의 최댓값이라 여유).
#   Claude wiggle·convert_document 쌍 간격 0.191~0.645 s(10/01, 4쌍, 모두 403)도 창 안
# 메모리: 열쇠 최대 DEDUPE_MAX_KEYS개, 마지막 시각에서 DEDUPE_RETENTION_SECONDS 지나면 지움
#   (보관 = 창 + 요청 하나가 기다릴 수 있는 최대: 추출 시간 + Detection 대기, 각각 설정 상한)
#   본문 받기 + 처리가 보관 - 창(60초)을 넘는 재시도는 앞 기록이 지워졌을 수 있어 따로 기록된다(안전한 쪽)
# 기록한 첫 시도의 requestId를 열쇠 항목에 둔다: 억제한 시도의 403 본문·차단 이유 보기 링크 번호는 대시보드에 없으므로
#   [DEDUPE] 줄의 recordedRequestId로 대시보드 이벤트(첫 시도)와 잇는다
# ============================================================

DEDUPE_MEASURED_MAX_GAP_SECONDS = 2.174
DEDUPE_WINDOW_SECONDS = 2 * DEDUPE_MEASURED_MAX_GAP_SECONDS
DEDUPE_MAX_KEYS = 4096
DEDUPE_RETENTION_SECONDS = DEDUPE_WINDOW_SECONDS + 2 * SETTINGS_MAX_SECONDS
_dedupe_seen: "OrderedDict[str, list]" = OrderedDict()   # 열쇠 해시 → [마지막 시각(perf_counter), 억제 수, 기록한 requestId]
_dedupe_lock = threading.Lock()
_MULTIPART_BOUNDARY_RE = re.compile(r'(?i)\bboundary="?([^";\s]+)"?')


def upload_identity(*, file_ref=None, files=None, body=None, content_type=None) -> str | None:
    """같은 업로드인지 가르는 파일 신원(메모리 안에서만 쓴다). 만들 수 없으면 None(기록을 줄이지 않음).

    file_ref: ChatGPT 원본 PUT 주소의 파일 번호(정규화). files: [(이름, 바이트)] → 바이트 sha256 + 이름.
    body: 파일 part를 못 읽은 multipart 본문 → 경계 문자열을 뺀 본문 sha256."""
    try:
        if file_ref:
            return "ref:" + str(file_ref)
        digest = hashlib.sha256()
        if files:
            for filename, data in files:
                name = str(filename or "").encode("utf-8", "surrogatepass")
                digest.update(len(name).to_bytes(4, "big") + name + hashlib.sha256(data or b"").digest())
            return "sha256:" + digest.hexdigest()
        if body:
            match = _MULTIPART_BOUNDARY_RE.search(content_type or "")
            boundary = match.group(1).encode("ascii", "ignore") if match else b""
            digest.update(body.replace(boundary, b"") if boundary else body)
            return "body:" + digest.hexdigest()
    except Exception as exc:
        print("[DEDUPE] file identity omitted error_type=", type(exc).__name__)
    return None


def _dedupe_key(flow: http.HTTPFlow, event: dict, identity: str) -> str:
    parts = (event.get("targetLlm"), event.get("clientId"), get_client_ip(flow) or "", identity,
             event.get("finalAction"), event.get("reasonCode") or "")
    return hashlib.sha256("\x1f".join(str(part) for part in parts).encode("utf-8", "surrogatepass")).hexdigest()


def _dedupe_check(key: str, arrived_at: float, now: float, request_id=None):
    """None = 기록한다(창 안에 같은 열쇠 없음). (억제 수, 간격 초, 기록한 requestId) = 같은 업로드의 재시도라 기록하지 않는다."""
    with _dedupe_lock:
        while _dedupe_seen:
            oldest = next(iter(_dedupe_seen.values()))
            if len(_dedupe_seen) <= DEDUPE_MAX_KEYS and now - oldest[0] <= DEDUPE_RETENTION_SECONDS:
                break
            _dedupe_seen.popitem(last=False)
        entry = _dedupe_seen.get(key)
        if entry is not None and arrived_at - entry[0] <= DEDUPE_WINDOW_SECONDS:
            gap = arrived_at - entry[0]
            entry[0] = max(entry[0], now)
            entry[1] += 1
            _dedupe_seen.move_to_end(key)
            return entry[1], gap, entry[2]
        _dedupe_seen[key] = [now, 0, request_id]
        _dedupe_seen.move_to_end(key)
        if len(_dedupe_seen) > DEDUPE_MAX_KEYS:
            _dedupe_seen.popitem(last=False)
        return None


def _attempt_arrived_at(flow: http.HTTPFlow, now: float) -> float:
    """이번 시도가 도착한 시각(perf_counter 기준) = 요청 첫 바이트(mitmproxy timestamp_start, 벽시계를 지금과의 차로 바꿈).
    request 훅은 본문을 다 받은 뒤 불리므로 훅 시작 시각(gw_metrics_started_at)은 업로드 시간만큼 늦다.
    첫 바이트 시각이 없거나 믿을 수 없으면(미래 = 시계가 뒤로 감, 보관 시간보다 오래 전 = 앞 기록이 이미 지워졌을 수 있음)
    request 훅 시작 시각을 쓴다(그것도 없으면 지금)."""
    started = flow.metadata.get("gw_metrics_started_at")
    fallback = started if isinstance(started, float) else now
    first_byte = getattr(flow.request, "timestamp_start", None)
    if isinstance(first_byte, (int, float)) and not isinstance(first_byte, bool):
        age = time.time() - first_byte
        if 0.0 <= age <= DEDUPE_RETENTION_SECONDS:
            return min(fallback, now - age)
    return fallback


def _suppress_duplicate_record(flow: http.HTTPFlow, event: dict, identity: str) -> bool:
    """같은 업로드의 재시도 기록이면 True(이벤트·알림을 보내지 않음). 확인이 실패하면 보낸다."""
    try:
        now = time.perf_counter()
        arrived_at = _attempt_arrived_at(flow, now)
        key = _dedupe_key(flow, event, identity)
        seen = _dedupe_check(key, arrived_at, now, event.get("requestId"))
    except Exception as exc:
        print("[DEDUPE] check failed; record kept error_type=", type(exc).__name__)
        return False
    if seen is None:
        return False
    count, gap, recorded_request_id = seen
    print(f"[DEDUPE] duplicate record suppressed service={event.get('targetLlm')} action={event.get('finalAction')} "
          f"reason={event.get('reasonCode') or '-'} requestId={event.get('requestId')} "
          f"recordedRequestId={recorded_request_id or '-'} key={key[:12]} "
          f"gap_ms={gap * 1000:.0f} window_s={DEDUPE_WINDOW_SECONDS:.3f} suppressed={count}")
    return True


def _send_deferred_dashboard_event(flow: http.HTTPFlow):
    events = flow.metadata.pop("gw_dashboard_events", [])
    identities = flow.metadata.pop("gw_dedupe_identity", None) or {}
    legacy_event = flow.metadata.pop("gw_dashboard_event", None)
    if legacy_event:
        events.append(legacy_event)
    if not events:
        return
    settings = _flow_settings(flow)
    for event in events:
        identity = identities.get(id(event))
        if identity and _suppress_duplicate_record(flow, event, identity):
            continue    # 차단은 이 시도에도 했다. 대시보드 이벤트·확장 알림만 처음 시도 한 번
        latency_ms = flow.metadata.get("gw_latency_ms")
        if latency_ms is not None:
            # Dashboard Security Event v1의 허용 필드. Traffic 화면은 이 값을
            # 집계해 평균/P95 지연시간을 계산한다.
            # 2026-10-10: Gateway 자체 처리 시간(요청 시작 → request() 처리 끝). LLM 응답 시간은 빠진다
            event["processingMs"] = latency_ms
        finalize_metadata(event, latency_ms, getattr(flow.response, 'status_code', None))
        delivery = _gateway_event_outbox.enqueue(event)
        if delivery.status not in {"queued", "duplicate"}:
            print(f"[DASHBOARD] durable enqueue status={delivery.status}")
        # Administrator-only text proof must not be broadcast to client extensions.
        extension_event = {k: v for k, v in event.items() if k not in {"maskEvidence", "gateway"}}
        extension_event["notify"] = _extension_notify(event, settings)
        review_url = request_review_url(event.get("requestId"))
        if review_url:
            extension_event["reviewUrl"] = review_url
            extension_event["reviewLabel"] = "처리 근거 보기 · 관리자 로그인 필요"
        send_extension_event(extension_event)


def _report_uninspected(flow: http.HTTPFlow, result, context: dict) -> bool:
    """검사하지 않은 파일의 판정 집행. 차단 → 403 + 이벤트(True). 허용 → 원본 전달, 이벤트는 보류(False).

    허용 이벤트는 바로 남기지 않는다: 같은 요청의 다른 파일·칸이 뒤에서 막으면 이 파일도 나가지 않으므로.
    request() 처리가 끝났을 때 실제로 전달되는 경우(flow.response 없음)에만 기록한다(_flush_pending_uninspected).
    enforce를 거치지 않아 앞서 정한 요청 지표(gw_action 등)를 덮지 않는다.
    """
    if getattr(result, "blocked", False):
        return enforce_detection_result(flow, result, context)
    print("\n===== GATEWAY ACTION =====")
    print("→ PASS (file not inspected; recorded) reason_code=", getattr(result, "reason_code", None))
    flow.metadata.setdefault("gw_pending_uninspected", []).append((result, context))
    if not flow.metadata.get("gw_action"):
        flow.metadata["gw_action"] = "allow"
    return False


def _flush_pending_uninspected(flow: http.HTTPFlow) -> None:
    pending = flow.metadata.pop("gw_pending_uninspected", None)
    if not pending:
        return
    if flow.response is not None:   # 차단·503으로 끝남: 이 파일들은 나가지 않았다
        print("[FILE] uninspected allow events dropped; request ended with status=", flow.response.status_code)
        return
    for result, context in pending:
        report_dashboard_event(flow, result, context, enforcement_applied=False, http_status=None)


# PDFium(pypdfium2)은 스레드 안전하지 않다(2026-10-11, 파일 업로드 작업 스레드): 두 스레드가 동시에 부르면 프로세스가
# 비정상 종료될 수 있어 PDFium을 쓰는 분류는 한 번에 하나만. PDFium을 부르는 곳은 gateway_file_classes._classify의
# PDF 갈래(b"%PDF-" in data[:1024])뿐이고, 여기서는 그보다 넓게(본문 어디든 %PDF-) 잠근다.
# 다른 곳에서 PDFium을 부르게 되면 그 호출도 이 잠금 안에서 해야 한다. 2단계 PDF 가림(_mask_pdf_upload)은 돌려받은 PDF를
# _classify_upload로 다시 분류하므로 이 잠금을 그대로 쓴다(Gateway는 PDF를 지우지 않음 - 지우기는 Detection).
_pdfium_lock = threading.Lock()


def _classify_upload(flow: http.HTTPFlow, data: bytes | None, *, filename=None, mime=None) -> FileClass:
    """업로드 바이트 분류(inspect_upload). 상한: 바이트 GATEWAY_FILE_MAX_BYTES, 글자·추출 시간은 파일 대기 시간 기준.
    data None = 파일 part를 찾지 못함(검사할 수 없음). 로그에는 형식·크기만(이름·내용 없음).
    PDF일 수 있으면 _pdfium_lock 안에서 분류한다. 시간 한도는 잠금을 얻은 뒤부터 잰다(앞 PDF를 기다린 시간 때문에
    '너무 큼'으로 막지 않게). 기다린 시간은 따로 로그에 남긴다."""
    settings = _flow_settings(flow)
    started = time.monotonic()
    if data is None:
        file_class = FileClass("unknown", "other", "no_file_part")
    else:
        # bytes가 아니면(inspect_upload가 bytes()로 바꿔 읽는다) 안전한 쪽으로 잠근다
        needs_pdfium = not isinstance(data, (bytes, bytearray)) or b"%PDF-" in data
        guard = _pdfium_lock if needs_pdfium else contextlib.nullcontext()
        with guard:
            locked = time.monotonic()
            if guard is _pdfium_lock and locked - started >= 0.001:
                print(f"[FILE] pdf lock wait_ms={(locked - started) * 1000:.1f}")
            started = locked
            file_class = inspect_upload(data, filename=filename, mime=mime, max_bytes=GATEWAY_FILE_MAX_BYTES,
                                        max_chars=file_max_chars(settings.file.seconds),
                                        deadline=started + settings.file.seconds)
    print(f"[FILE] kind={file_class.kind} group={file_class.group or 'inspectable'} detail={file_class.detail} "
          f"bytes={len(data or b'')} chars={len(file_class.text or '')} "
          f"ms={(time.monotonic() - started) * 1000:.1f}")
    if file_class.kind == "pdf" and file_class.pages:
        print(f"[FILE] pdf pages={file_class.pages} text_pages={file_class.text_pages} "
              f"image_only_pages={file_class.image_only_pages}")
    return file_class



# ============================================================
# 파일 처리 실패 기록 (2026-10-10)
# Claude 파일 경로와 ChatGPT 웹 파일 경로(처리 중 예상 못한 오류·칸 무응답 차단)가 쓴다.
# ============================================================

class FileInspectionUnavailable(ValueError):
    """칸별 검사 중 Detection 무응답(파일 처리 '차단'). 원본을 보내지 않고, 기록용으로 그 칸의 판정을 들고 있다."""

    def __init__(self, detection_result):
        super().__init__("file segment inspection unavailable")
        self.detection_result = detection_result


def _file_failure_result(request_id: str, cause=None) -> DetectionResult:
    """파일 처리 실패를 대시보드에 남길 판정. Detection 판정이 아니라 Gateway가 원본을 막은 것이다.

    cause는 실패 직전 판정이다.
      Detection 무응답        -> 검사 불가(FAIL_CLOSED) + 그 오류 코드(detection_timeout 등)
      개인정보 유형별 건수 있음 -> 그 건수를 그대로 싣는다(reason_code file_masking_failed)
      그 밖                   -> 탐지 근거 없음(reason_code file_masking_failed)
    대시보드 Security Event v1 검사(validation.mjs)를 통과하는 값만 쓴다.
    """
    values = {}
    reason_code = "file_masking_failed"
    if getattr(cause, "degraded", False):
        values = {"degraded": True, "degraded_reason": getattr(cause, "degraded_reason", None),
                  "failure_policy": "FAIL_CLOSED"}
        cause_code = getattr(cause, "reason_code", None)
        if isinstance(cause_code, str) and re.fullmatch(r"detection_[a-z_]{1,80}", cause_code):
            reason_code = cause_code
    elif str(getattr(cause, "decided_by", "") or "").lower() == "signature":
        pii_hits = _merged_pii_hits([cause])
        if pii_hits:
            values = {"decided_by": "signature", "pii_hits": pii_hits}
    return DetectionResult(blocked=True, reason="gateway file processing failed", request_id=request_id,
                           action="block", reason_code=reason_code, http_status=503, **values)


def fail_closed_file_upload(
    flow: http.HTTPFlow,
    dashboard_context: dict | None,
    request_id: str | None = None,
    *,
    cause=None,
) -> bool:
    """파일 처리 실패: 원본을 보내지 않는다(기존 503 file_masking_failed 응답 그대로) + 대시보드 BLOCK 기록.

    칸 검사 무응답·가림 재현 실패·multipart 교체 실패·예상 못한 오류가 이전에는 이벤트 없는 503이라
    대시보드에 보이지 않았다. 응답은 바꾸지 않고 이벤트만 더한다. 항상 True(응답을 직접 만듦).
    응답을 먼저 정한다: 기록 단계에서 예외가 나도 원본은 나가지 않는다(예외가 request() 밖으로 나가면
    mitmproxy는 기록만 하고 원본을 보낸다).
    """
    flow.response = http.Response.make(503, b'{"error":{"type":"file_masking_failed"}}',
                                       {'Content-Type': 'application/json'})
    flow.metadata["gw_action"] = "block"
    try:
        result = _file_failure_result(request_id or str(uuid.uuid4()), cause)
        flow.metadata["gw_request_id"] = result.request_id
        flow.metadata["gw_detection_result"] = result
        flow.metadata["gw_detector"] = _result_detector(result)
        flow.metadata["gw_degraded"] = result.degraded

        print("\n===== GATEWAY ACTION =====")
        print("→ BLOCK (file processing failed)")
        print("Request ID:", result.request_id)
        print("Reason code:", result.reason_code)
        report_dashboard_event(
            flow,
            result,
            dashboard_context,
            enforcement_applied=True,
            http_status=503,
        )
    except Exception as exc:
        print("[DASHBOARD] file failure event omitted error_type=", type(exc).__name__)
    return True


class _CellInspection:
    """DOCX·XLSX 칸 검사(ChatGPT·Claude 같은 규칙, 2026-10-10). sanitize_uploaded_file이 칸마다 detect를 부른다.

    - 차단 판정 → BlockedBySignaturePolicy(403 + 이벤트)
    - Detection 무응답 + 파일 처리 '차단'(FAIL_CLOSED) → FileInspectionUnavailable(503 + 이벤트)
    - Detection 무응답 + '원본 전송'(FAIL_OPEN) → 그 칸은 가리지 않고 넘긴다. 첫 무응답 판정을 unanswered에 남긴다
      (남은 칸은 _inspect가 호출 없이 같은 실패로 돌려준다 → 요청당 대기 1회)
    """

    def __init__(self, flow: http.HTTPFlow):
        self.flow = flow
        self.results = []        # 응답이 온 칸 판정
        self.unanswered = None   # 첫 무응답(FAIL_OPEN) 판정

    def detect(self, segment, child_request_id, source):
        detection = _inspect(self.flow, "file", request_id=child_request_id, normalized_text=segment.text,
                             source=source)
        if getattr(detection, "degraded", False):
            if _claude_action(detection) in _FILE_BLOCK_ACTIONS:
                raise FileInspectionUnavailable(detection)
            if self.unanswered is None:
                self.unanswered = detection
            return detection   # allow_degraded → 칸은 그대로(gateway_runtime이 allow로 바꿔 적용)
        if _claude_action(detection) in _FILE_BLOCK_ACTIONS:
            raise BlockedBySignaturePolicy(detection)
        self.results.append(detection)
        return detection

    @property
    def masked(self):
        return [d for d in self.results if _claude_action(d) in _FILE_MASK_ACTIONS]

    def degraded_report(self, request_id: str, changed: bool):
        """칸 일부 무응답(원본 전송)의 기록용 판정. 앞 칸을 가렸으면 MASK, 아니면 ALLOW - 둘 다 degraded(FAIL_OPEN)."""
        cause = self.unanswered
        if not changed or not self.masked:
            return failure_result(cause.degraded_reason, "FAIL_OPEN", request_id)
        report = copy.copy(self.masked[0])
        report.pii_hits = _merged_pii_hits(self.masked) or report.pii_hits
        report.degraded = True
        report.degraded_reason = cause.degraded_reason
        report.failure_policy = "FAIL_OPEN"
        report.reason_code = cause.reason_code
        return report



# ============================================================
# PDF 가림 (2026-10-11, masking-design/SPEC.md §3.5·§3.6, 2단계)
# 가림 판정이 난 PDF는 Detection `POST /redact/pdf`가 지워(검은 상자, 글자 넣지 않음) 돌려준다. Gateway는 돌려받은 PDF를
# 검사 때와 같은 추출기(_classify_upload, pypdfium2)로 다시 뽑아 한 번 더 검사하고, 남은 것이 하나라도 있거나 검사가
# 돌지 않으면 보내지 않는다. 어떤 실패에서도 원본은 나가지 않는다(처리 설정이 '원본 전송'이어도).
#   차단(쪽 판정)                                   → 403 그 쪽 판정 그대로(BlockedBySignaturePolicy)
#   지원 안 함·잔존·재분류 실패·재검사 적중·글자 상한·차단 → 403 mask_application_failed + 이벤트(MaskRefused)
#   시간 초과·연결·HTTP 오류·busy·timeout·error·무응답    → 503 file_masking_failed + 이벤트(FileInspectionUnavailable)
# 검토용 사본은 만들지 않는다(메모리만, review_id 없음 → X-LLMSGW-Review-* 머리도 없음).
# ============================================================

# MaskRefused(가린 파일을 보낼 수 없음, PDF·HWP·HWPX 공통, 명세 §0.3)는 signature_detector.contracts에 있다(HWP·HWPX 가림
# 모듈도 같은 예외를 낸다, 2026-10-11 3단계). 403 mask_application_failed + 이벤트로 끝낸다
# (enforce_detection_result(..., mask_callback=lambda _: False), 기록용 판정 = 파일 전체가 가림이면 그 판정, 아니면 가림 칸 판정).


# 재검사에서 세지 않는 적중: 정책상 허용(allow·allow_notify)과 가림 표시 글자 자체([KR_RRN]·[PASSWORD] 등, Detection
# redact_pdf.REPLACEMENT_TOKEN과 같은 모양). 그 밖(mask_partial·mask_full·block_notify·모르는 값)은 모두 남은 적중
_VERIFY_UNCOUNTED_ACTIONS = {"allow", "allow_notify"}
_REPLACEMENT_TOKEN_RE = re.compile(r"\[[A-Z][A-Z0-9_]*\]")


def _verify_variant_id(request_id: str, tag: str) -> str:
    """'{rid}#verify' → '{rid}#{tag}#verify'(보고서·전후 보기가 '#verify'로 끝나는 기록을 빼는 규칙을 그대로 타게)."""
    if request_id.endswith("#verify"):
        return f"{request_id[:-len('#verify')]}#{tag}#verify"
    return f"{request_id}#{tag}"


def _residual_types(findings, text: str, replaced, index_map=None) -> list[str]:
    """남은 적중의 유형 목록(값 없음). replaced = 원문 좌표 구간 목록(그 안에 완전히 든 적중은 셈하지 않음, HWP·HWPX용).
    index_map이 있으면 text는 공백을 뺀 글자이고 index_map[i] = 원문 위치. 좌표가 틀린 적중은 남은 것으로 센다."""
    left = []
    for finding in findings:
        if not isinstance(finding, dict):
            left.append("unknown")
            continue
        if finding.get("policy_action") in _VERIFY_UNCOUNTED_ACTIONS:
            continue
        entity = str(finding.get("entity_type") or "unknown")[:64]
        start, end = finding.get("start"), finding.get("end")
        if type(start) is not int or type(end) is not int or not 0 <= start < end <= len(text):
            left.append(entity)
            continue
        if _REPLACEMENT_TOKEN_RE.fullmatch(text[start:end]):
            continue
        if index_map is not None:
            start, end = index_map[start], index_map[end - 1] + 1
        if any(a <= start and end <= b for a, b in replaced):
            continue
        left.append(entity)
    return left


def _verify_masked(flow: http.HTTPFlow, text: str, *, request_id: str, source: dict, replaced=()) -> None:
    """보낼 가린 파일의 마지막 검사(명세 §3.5.1). 통과하면 None, 아니면 예외. _inspect를 거치지 않는다
    (글자 상한 지름길·무응답 넘기기·FAIL_OPEN이 '검사 없는 통과'를 돌려주므로).

    두 번 검사한다: 뽑은 글자 그대로(request_id) + 공백을 모두 뺀 글자('{rid}#nows#verify', 공백·줄바꿈으로 쪼개진 값).
    매번 성공 조건을 모두 채워야 한다: 무응답 아님 · signature_status ok · 차단 판정 아님 · findings가 목록 · 남은 적중 0.
      - 같은 요청에 앞선 무응답 / 이번 무응답(degraded) → FileInspectionUnavailable(503)
      - 글자 상한(파일 대기 시간 기준) 초과·그 밖 실패 → MaskRefused(403)
    """
    rule = _flow_settings(flow).timeout_for("file")
    keep = [index for index, char in enumerate(text) if not char.isspace()]
    squashed = "".join(text[index] for index in keep)
    passes = [(text, None, request_id)]
    if squashed and squashed != text:
        passes.append((squashed, keep, _verify_variant_id(request_id, "nows")))
    for variant, index_map, variant_id in passes:
        down = flow.metadata.get("gw_detection_down")
        if down:
            print("→ Verify skipped; earlier call failed:", down[0])
            raise FileInspectionUnavailable(failure_result(down[0], "FAIL_CLOSED", variant_id))
        if len(variant) > file_max_chars(rule.seconds):
            raise MaskRefused("verify_too_large")
        flow.metadata["gw_inspected"] = True
        checked = check_detection(DETECTION_BASE_URL, request_id=variant_id, normalized_text=variant,
                                  source={**source, "signature_original_offsets": True},
                                  timeout=rule.seconds, failure_policy="FAIL_CLOSED")
        if getattr(checked, "degraded", False):
            if getattr(checked, "degraded_reason", None) in _DETECTION_DOWN_REASONS:
                flow.metadata["gw_detection_down"] = (checked.degraded_reason, rule.seconds)
            raise FileInspectionUnavailable(checked)
        if getattr(checked, "signature_status", None) != "ok":
            raise MaskRefused("verify_signature_incomplete")
        if _claude_action(checked) in _FILE_BLOCK_ACTIONS:
            raise MaskRefused("verify_blocked")
        findings = getattr(checked, "findings", None)
        if not isinstance(findings, list):
            raise MaskRefused("verify_findings_missing")
        left = _residual_types(findings, variant, replaced, index_map)
        if left:
            print(f"[PDF-MASK] verify residual pass={'nows' if index_map is not None else 'text'} "
                  f"types={dict(Counter(left))}")
            raise MaskRefused("verify_residual")


def _mask_pdf_upload(flow: http.HTTPFlow, result, *, prefix: str, data: bytes, filename, endpoint: str,
                     service: str):
    """가림 판정 PDF 하나 → (보낼 바이트, RedactResult). 명세 §3.5 순서:
    ① 같은 요청에 앞선 무응답 → FileInspectionUnavailable ② Detection /redact/pdf(총시한 = 파일 대기 시간)
    ③ 상태별 처리(§3.6) ④ 보낼 바이트(masked = 결과, clean = 원본)를 _classify_upload로 다시 분류 - pdf·검사 가능·글자 있음이
    아니면 MaskRefused(스캔·암호·20MiB 초과·글자 0 포함) ⑤ _verify_masked(replaced=(), 글자를 넣지 않으므로 무시 구간 없음).
    prefix: Detection 쪽 기록 번호 앞부분(ChatGPT '{rid}', Claude '{rid}#file:{i}'). 작업 스레드에서 돈다."""
    request_id = getattr(result, "request_id", None) or prefix
    down = flow.metadata.get("gw_detection_down")
    if down:
        print("→ PDF mask skipped; earlier call failed:", down[0])
        raise FileInspectionUnavailable(failure_result(down[0], "FAIL_CLOSED", request_id))
    seconds = _flow_settings(flow).timeout_for("file").seconds
    flow.metadata["gw_inspected"] = True
    redacted = redact_pdf_detection(DETECTION_BASE_URL, prefix, data,
                                    {"service": service, "endpoint": endpoint, "filename": filename}, seconds)
    print(f"[PDF-MASK] redact status={redacted.status} reason={redacted.reason} pages={redacted.pages} "
          f"text_pages={redacted.text_pages} masked_pages={redacted.masked_pages} pii_hits={redacted.pii_hits} "
          f"residual={redacted.residual} bytes_in={len(data)} bytes_out={len(redacted.pdf or b'')} "
          f"ms={redacted.elapsed_ms}")
    status = redacted.status
    if status == "blocked" and redacted.decision is not None:
        raise BlockedBySignaturePolicy(redacted.decision)   # 쪽 판정(그 쪽 decisions 행 번호) 그대로 403
    if status in ("unsupported", "residual"):
        raise MaskRefused(f"redact_{status}")
    if status == "masked" and redacted.pdf:
        outgoing = redacted.pdf
    elif status == "clean":
        outgoing = data
    else:
        # busy·timeout·error(호출 실패 포함): 연결 안 됨·시간 초과면 그 무응답으로, 아니면 파일 전체 판정으로 기록
        reason = redacted.reason if status == "error" and redacted.reason in _DETECTION_DOWN_REASONS else None
        raise FileInspectionUnavailable(failure_result(reason, "FAIL_CLOSED", request_id) if reason else result)
    file_class = _classify_upload(flow, outgoing, filename=filename, mime="application/pdf")
    if file_class.kind != "pdf" or file_class.group or not (file_class.text or "").strip():
        raise MaskRefused(f"reclassify_{file_class.group or file_class.kind}")
    _verify_masked(flow, file_class.text.strip(), request_id=f"{prefix}#verify",
                   source={"service": service, "endpoint": endpoint.split("?", 1)[0], "content_type": "file",
                           "filename": filename, "format": "pdf"},
                   replaced=())
    return outgoing, redacted


def _mask_chatgpt_pdf_upload(flow: http.HTTPFlow, result, context: dict, *, request_id: str, file_bytes: bytes,
                             filename, endpoint: str, apply_sanitized) -> bool:
    """ChatGPT 웹 PDF(원본 PUT·multipart) 가림 집행. 가린 PDF를 apply_sanitized로 본문에 넣고(원본 PUT은 크기 표 →
    완료 요청 file_size 고침, multipart는 경계 확인) MASK 기록(파일 전체 판정 사본, 유형별 건수만 Detection 응답 값).
    응답을 직접 만들었으면(차단·503) True."""
    started = time.monotonic()
    try:
        outgoing, redacted = _mask_pdf_upload(flow, result, prefix=request_id, data=file_bytes, filename=filename,
                                              endpoint=endpoint, service="ChatGPT")
        if outgoing == file_bytes:
            # 파일 전체 검사는 가리라고 했는데 Detection은 지울 것이 없다고 했다(clean, 재검사도 통과): 가림을 보일 수
            # 없으니 보내지 않는다(DOCX '칸 가림 재현 실패'와 같은 503, 여기서는 기록도 남김)
            raise ValueError("whole-file mask not reproduced by the PDF redaction")
        apply_sanitized(FileMaskOutcome(outgoing, None, None, (time.monotonic() - started) * 1000))
        record = copy.copy(result)   # 파일 전체 판정 원본은 그대로 두고 기록용 사본만 바꾼다
        if redacted.pii_hits:
            record.pii_hits = redacted.pii_hits
    except BlockedBySignaturePolicy as exc:
        print("[PDF-MASK] ChatGPT PDF blocked by a page decision")
        return enforce_detection_result(flow, exc.detection_result, context)
    except MaskRefused as exc:
        print("[PDF-MASK] ChatGPT PDF mask refused reason=", exc.reason)
        return enforce_detection_result(flow, result, context, mask_callback=lambda _: False)
    except FileInspectionUnavailable as exc:
        print("[PDF-MASK] ChatGPT PDF mask unavailable; upload blocked")
        return fail_closed_file_upload(flow, context, request_id, cause=exc.detection_result)
    except Exception as exc:
        print("[PDF-MASK] ChatGPT PDF mask failed error_type=", type(exc).__name__)
        return fail_closed_file_upload(flow, context, request_id, cause=result)
    print(f"[PDF-MASK] ChatGPT PDF masked bytes={len(file_bytes)}->{len(outgoing)}")
    return enforce_detection_result(flow, record, context, mask_callback=lambda _: True)



# ============================================================
# Claude Web 파일 업로드 Detection
# ============================================================

def _is_claude_multipart_file_upload(flow: http.HTTPFlow) -> bool:
    req = flow.request

    if req.host != "claude.ai" or req.method.upper() != "POST":
        return False

    content_type = req.headers.get("content-type", "").lower()
    if not content_type.startswith("multipart/form-data"):
        return False

    # Claude Web의 업로드 API 경로는 변경될 수 있으므로
    # 특정 endpoint 하나에 고정하지 않고 실제 multipart part에
    # filename이 있는지를 기준으로 파일 업로드를 식별한다.
    return True


# multipart 본문에 파일 part 머리가 보이는지(해석이 실패했거나 part를 하나도 못 찾았을 때). 내용은 읽지 않는다.
# 이름이 빈 칸(filename="", 파일을 고르지 않은 칸)은 파일 part가 아니다(10/11, 해석기도 일부러 건너뜀).
# 따옴표 이름·맨 이름·RFC 5987(filename*=UTF-8''…) 모두
_MULTIPART_FILENAME_RE = re.compile(rb'(?im)^content-disposition:[^\r\n]*\bfilename\*?=\s*(?:"[^"\r\n]+"|[^\s;"]+)')


def extract_claude_file_upload_info(flow: http.HTTPFlow, errors: list | None = None) -> list[dict]:
    """
    Claude Web multipart 요청에서 파일 part를 모두 추출한다.

    로그에는 파일 원문을 출력하지 않는다. filename/body size만 기록하고
    실제 파일 bytes는 Detection 직전 메모리에서만 처리한다.
    errors: 주면 해석 예외 이름을 담는다(2026-10-10, 해석 실패를 "파일 없음"과 구분).
    """
    req = flow.request
    body = req.raw_content or b""
    content_type = req.headers.get("content-type", "")

    if not body or not content_type.lower().startswith("multipart/form-data"):
        return []

    files = []

    try:
        raw_message = (
            b"Content-Type: "
            + content_type.encode("utf-8", errors="ignore")
            + b"\r\n\r\n"
            + body
        )
        message = BytesParser(policy=default).parsebytes(raw_message)

        for part in message.iter_parts():
            filename = part.get_filename()
            if not filename:
                continue

            file_bytes = part.get_payload(decode=True) or b""
            files.append({
                "filename": filename,
                "file_bytes": file_bytes,
                "body_size": len(file_bytes),
                "part_name": part.get_param(
                    "name",
                    header="Content-Disposition",
                ),
                "content_type": part.get_content_type(),
            })

    except Exception as e:
        print("❌ Claude multipart parsing failed:", type(e).__name__)
        if errors is not None:
            errors.append(type(e).__name__)

    return files


def create_claude_file_event(
    flow: http.HTTPFlow,
    files: list[dict],
    *,
    texts: list[str] | None = None,
    request_id: str | None = None,
) -> dict:
    """texts가 있으면(2026-10-10, inspect_upload가 바이트로 정한 형식의 글자) 그 글자를 쓴다. 없으면 이름 확장자로 추출."""
    contents = []

    for index, item in enumerate(files):
        filename = item.get("filename")
        file_bytes = item.get("file_bytes") or b""
        if texts is not None:
            extracted_text = texts[index]
        else:
            extracted_text = extract_file_text(filename, file_bytes) if file_bytes else None

        content = {
            "type": "file",
            "filename": filename,
            "body_size": len(file_bytes),
        }

        if extracted_text is not None:
            content["content"] = extracted_text

        contents.append(content)

    return {
        "request_id": request_id or str(uuid.uuid4()),
        "timestamp": datetime.now(timezone.utc).isoformat(),
        "client": {"ip": get_client_ip(flow)},
        "service": "Claude",
        "request": {
            "method": flow.request.method,
            "endpoint": flow.request.path.split("?", 1)[0],
        },
        "contents": contents,
        "metadata": {
            "content_count": len(contents),
            "body_size": len(flow.request.raw_content or b""),
        },
    }


def detect_claude_file_upload(flow: http.HTTPFlow, identity: dict) -> bool:
    """
    Claude Web 파일 업로드: 파일마다 분류 → 검사할 수 없는 파일 처리 → 검사 가능한 파일 글자를 이어 1회 검사
    → (DOCX/XLSX) 칸별 가림 → 집행.

    2026-10-10 (ChatGPT 파일 경로와 같게):
    - 가려야 하는데 가릴 수 없는 파일(TXT 등)이 하나라도 있으면 요청 전체를
      차단 + 이벤트(mask_application_failed, 403). 이전에는 이벤트 없는 503이었다.
    - PDF(2026-10-11): Detection /redact/pdf가 지운 PDF를 다시 검사해 바꿔 넣는다(_mask_pdf_upload).
      보낼 수 없으면(MaskRefused) 403 mask_application_failed + 이벤트, 가림 호출 실패는 503 + 이벤트
    - 처리 실패(칸 검사 무응답·가림 재현 실패·multipart 교체 실패·예상 못한 오류)는
      503 그대로(원본을 보내지 않음) + 이벤트. 이전에는 이벤트가 없었다.
    2026-10-10 (Gateway 처리 설정):
    - 형식은 바이트로 정한다(이름·MIME 아님, inspect_upload). 검사할 수 없는 파일(이미지·스캔 PDF·그 밖 형식·
      상한 초과)은 설정에 따라 차단(Detection 전에 요청 전체 403 + 그 파일 이벤트 1건) 또는 허용(전달 + 기록)
    - 글자 없는 문서는 통과가 아니라 그 밖 형식(숨은 곳을 못 본 것과 구분할 수 없음)
    - 해석에 실패했는데 본문에 파일 part 머리가 보이면 그 밖 형식으로 처리(이전: 파일 0개로 보고 통과)
    """
    if not _is_claude_multipart_file_upload(flow):
        return False

    errors = []
    files = extract_claude_file_upload_info(flow, errors)

    print("\n===== CLAUDE FILE UPLOAD DETECTED =====")
    print("Host:", flow.request.host)
    print("Method:", flow.request.method)
    print("Path:", flow.request.path.split("?", 1)[0])
    print("File count:", len(files))

    context = {"identity": identity, "targetLlm": "Claude", "contentType": "file"}
    if not files:
        if _MULTIPART_FILENAME_RE.search(flow.request.raw_content or b""):
            print("→ File part header found but not parsed; treated as uninspectable errors=", errors)
            context["fileIdentity"] = upload_identity(body=flow.request.raw_content,
                                                      content_type=flow.request.headers.get("content-type", ""))
            return _report_uninspected(flow, uninspected_result(flow, "other", str(uuid.uuid4())), context)
        print("→ No file part found; PASS THROUGH")
        return False

    # 같은 업로드 재시도 기록 한 번만(2026-10-11): 요청 전체 판정의 신원 = 모든 파일의 바이트 + 이름
    context["fileIdentity"] = upload_identity(files=[(item.get("filename"), item.get("file_bytes")) for item in files])
    progress = {}   # 실패 기록용: 요청 번호와 파일 전체 판정
    try:
        return _guard_claude_file_upload(flow, files, context, progress)
    except BlockedBySignaturePolicy as exc:
        return enforce_detection_result(flow, exc.detection_result, context)
    except FileInspectionUnavailable as exc:
        print('[SIGNATURE] Claude segment inspection unavailable; upload blocked')
        return fail_closed_file_upload(flow, context, progress.get("request_id"), cause=exc.detection_result)
    except MaskRefused as exc:
        # 가린 PDF·HWP·HWPX를 보낼 수 없음(2026-10-11, 명세 §3.6·§4.1): 403 mask_application_failed + 이벤트.
        # 기록 = 파일 전체가 가림이면 그 판정, 아니면 가림 칸 판정
        print('[FILE-MASK] Claude mask refused reason=', exc.reason)
        record = _mask_record(progress.get("result"), progress.get("cells"))
        if record is None:
            return fail_closed_file_upload(flow, context, progress.get("request_id"), cause=progress.get("result"))
        return enforce_detection_result(flow, record, context, mask_callback=lambda _: False)
    except Exception as exc:
        print('[SIGNATURE] Claude multipart masking failed error_type=', type(exc).__name__)
        return fail_closed_file_upload(flow, context, progress.get("request_id"), cause=progress.get("result"))


def _guard_claude_file_upload(flow: http.HTTPFlow, files: list[dict], context: dict, progress: dict) -> bool:
    """파일마다 분류 → 검사 불가 처리 → 검사 가능한 파일 글자 1회 검사 → (DOCX/XLSX) 칸별 가림 → 집행.
    처리 실패는 예외로 올린다(detect_claude_file_upload가 503 + 기록)."""
    request_id = str(uuid.uuid4())
    progress["request_id"] = request_id
    classes = [_classify_upload(flow, item.get("file_bytes") or b"", filename=item.get("filename"),
                                mime=item.get("content_type")) for item in files]

    def file_request_id(index):
        return request_id if len(files) == 1 else f"{request_id}#upload:{index}"

    def file_context(index):
        # 파일 하나의 기록은 그 파일의 신원(같은 업로드 재시도 기록 한 번만, 2026-10-11)
        if len(files) == 1:
            return context
        item = files[index]
        return {**context, "fileIdentity": upload_identity(files=[(item.get("filename"), item.get("file_bytes"))])}

    uninspected = [(index, uninspected_result(flow, fc.group, file_request_id(index)))
                   for index, fc in enumerate(classes) if fc.group]
    for index, result in uninspected:
        if result.blocked:
            # 검사할 수 없는 파일을 막는 설정: Detection 전에 요청 전체를 막는다. 그 파일 이벤트 1건만
            print("[FILE] Claude upload has a file that cannot be inspected; upload blocked index=", index)
            return enforce_detection_result(flow, result, file_context(index))
    for index, result in uninspected:
        _report_uninspected(flow, result, file_context(index))   # 허용: 기록은 실제로 전달될 때만

    inspectable = [index for index, fc in enumerate(classes) if not fc.group and fc.text]
    if not inspectable:
        print("→ No checkable file text; PASS THROUGH")
        return False

    event = create_claude_file_event(flow, [files[index] for index in inspectable],
                                     texts=[classes[index].text for index in inspectable], request_id=request_id)
    print_gateway_event(event)
    send_to_next_module(event)

    normalized_text = extract_file_checkable_text(event)

    if not normalized_text:
        print("→ No checkable file text extracted; PASS THROUGH")
        return False

    print("Extracted file text length:", len(normalized_text))
    print("→ Sending Claude file text to Detection Module")

    source = {
        "service": "Claude",
        "endpoint": flow.request.path.split("?", 1)[0],
        "content_type": "file",
        "filename": ", ".join(
            str(item.get("filename"))
            for item in event.get("contents", [])
            if item.get("filename")
        ),
    }
    result = _inspect(flow, "file", request_id=request_id, normalized_text=normalized_text, source=source)

    progress["result"] = result

    print("\n===== CLAUDE FILE DETECTION RESULT =====")
    print_detection_summary(result)

    if _claude_action(result) in _FILE_BLOCK_ACTIONS:
        return enforce_detection_result(flow, result, context)
    if getattr(result, "reason_code", None) == "file_too_large":
        # 상한 초과 + 원본 전송 설정: 검사·칸 가림 없이 전달 + 기록
        return _report_uninspected(flow, result, context)

    requested_mask = _claude_action(result) in _FILE_MASK_ACTIONS
    if requested_mask and any(classes[index].kind not in _FILE_MASKABLE_KINDS for index in inspectable):
        # 가릴 수 없는 파일(TXT 등)이 있다 - 요청 전체를 차단 + 이벤트.
        # ChatGPT와 같은 규칙(guard_chatgpt_file_upload). 일부만 가린 요청을 보내지 않는다.
        print('[SIGNATURE] Claude file mask cannot be applied to every file; upload blocked')
        return enforce_detection_result(flow, result, context, mask_callback=lambda _: False)

    cells = _CellInspection(flow)
    progress["cells"] = cells
    endpoint = flow.request.path.split('?', 1)[0]
    # PDF(2026-10-11, 명세 §3.5): Detection이 지운 PDF를 다시 검사해 바꿔 넣는다. 칸 검사보다 먼저(실패하면 칸 호출 없이 끝남).
    # 지울 것이 없는 PDF(clean)는 재검사로 원본을 확인한 뒤 원본 그대로. 실패는 예외(detect_claude_file_upload가 403·503)
    pdf_outgoing = {}
    if requested_mask:
        for file_index in inspectable:
            if classes[file_index].kind == "pdf":
                item = files[file_index]
                pdf_outgoing[file_index], _ = _mask_pdf_upload(
                    flow, result, prefix=f"{request_id}#file:{file_index}", data=item['file_bytes'],
                    filename=item.get('filename'), endpoint=endpoint, service="Claude")
    replacements = []
    for file_index, (item, file_class) in enumerate(zip(files, classes)):
        original = item['file_bytes']
        if file_index in pdf_outgoing:
            replacements.append((original, pdf_outgoing[file_index]))
            continue
        if file_class.group or not file_class.text or file_class.kind not in _SEGMENT_SANITIZED_EXTENSIONS:
            replacements.append((original, original))
            continue

        def detect_segment(segment, child_request_id, filename=item['filename']):
            source = {'service': 'Claude', 'endpoint': endpoint, 'content_type': 'file', 'filename': filename,
                      'context_label': canonical_label(segment.context), 'signature_original_offsets': True}
            if getattr(segment, 'block_order', None) is not None:
                source['block_order'] = segment.block_order   # HWP·HWPX(2026-10-11): 기록 화면 순서 복원용
            return cells.detect(segment, child_request_id, source)

        # Review stays in memory here; do not add to unexpired original-review storage.
        # 가림은 바이트로 정한 형식(이름을 바꾼 DOCX도 DOCX로 가린다)
        if file_class.kind in _HWP_KINDS:
            # HWP·HWPX(2026-10-11, 명세 §4): 칸 가림 → 가린 파일을 다시 읽어 재검사(실패는 예외 → 403·503)
            artifacts = sanitize_uploaded_file(f"upload.{file_class.kind}", original,
                                               f"{request_id}#file:{file_index}", detect_segment,
                                               allow_hwp=True, max_decompressed=GATEWAY_FILE_MAX_BYTES)
            if artifacts.sanitized_bytes != original:
                _verify_hwp_artifacts(flow, artifacts, cells, request_id=f"{request_id}#file:{file_index}",
                                      kind=file_class.kind, service="Claude", endpoint=endpoint,
                                      filename=item.get('filename'))
        else:
            artifacts = sanitize_uploaded_file(f"upload.{file_class.kind}", original,
                                               f"{request_id}#file:{file_index}", detect_segment)
        replacements.append((original, artifacts.sanitized_bytes))
    changed = any(original != sanitized for original, sanitized in replacements)
    if requested_mask and cells.unanswered is not None:
        # 파일 전체 판정이 가림인데 칸 일부를 검사하지 못했다: 처리 설정과 무관하게 원본을 보내지 않는다(503 + 기록)
        raise FileInspectionUnavailable(cells.unanswered)
    if requested_mask and not changed:
        if any(classes[index].kind in _HWP_KINDS for index in inspectable):
            # HWP·HWPX(명세 §4.4): 파일 전체에서만 잡힌 값 - 403 + 이벤트(기록 = 파일 전체 판정)
            raise MaskRefused('mask_not_reproduced')
        raise ValueError('whole-file mask not reproduced by segments')
    if changed:
        rewritten = replace_claude_files(flow.request.raw_content or b'',
                                         flow.request.headers.get('content-type', ''), replacements)
        flow.request.raw_content = rewritten
        flow.request.headers['content-length'] = str(len(rewritten))
        print('[SIGNATURE] Claude multipart sanitized files=', len(files))
        if not requested_mask:
            result = cells.masked[0]
    if cells.unanswered is not None:
        # 파일 전체는 통과, 칸 일부 무응답(원본 전송 설정): 앞 칸 가림은 적용하고 전달. 기록은 degraded(FAIL_OPEN)
        result = cells.degraded_report(request_id, changed)
    return enforce_detection_result(flow, result, context, mask_callback=lambda _: changed)



_FILE_REF_RE = re.compile(r"^[A-Za-z0-9_-]{1,100}$")
_UUID_RE = re.compile(r"^[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{12}$")
CHATGPT_UPLOADED_REFS: "OrderedDict[str, float]" = OrderedDict()


def canonical_file_ref(value) -> str | None:
    """ChatGPT 파일 번호를 한 가지 모양으로. 업로드 주소의 00000000-ddcc-… 와 질문의 file_00000000ddcc… 가 같아진다."""
    if not isinstance(value, str):
        return None
    ref = value.strip()
    for prefix in ("sediment://", "file-service://"):
        if ref.startswith(prefix):
            ref = ref[len(prefix):]
    for prefix in ("file_", "file-"):
        if ref.startswith(prefix):
            ref = ref[len(prefix):]
            break
    if _UUID_RE.match(ref):
        ref = ref.replace("-", "")
    if re.fullmatch(r"[0-9a-fA-F]{32}", ref):
        ref = ref.lower()
    return ref if _FILE_REF_RE.match(ref) else None


def chatgpt_attachment_refs(body) -> list[str]:
    """질문 본문에 붙은 파일 번호들(attachments[].id, asset_pointer). 내용은 읽지 않는다."""
    refs: list[str] = []

    def add(value):
        ref = canonical_file_ref(value)
        if ref and ref not in refs:
            refs.append(ref)

    def walk(value):
        if isinstance(value, dict):
            attachments = value.get("attachments")
            if isinstance(attachments, list):
                for item in attachments:
                    if isinstance(item, dict):
                        add(item.get("id"))
            if isinstance(value.get("asset_pointer"), str):
                add(value["asset_pointer"])
            for child in value.values():
                walk(child)
        elif isinstance(value, list):
            for child in value:
                walk(child)

    walk(body or {})
    return refs[:20]


def _remember_uploaded_ref(ref: str | None) -> None:
    if not ref:
        return
    with _upload_state_lock:   # 작업 스레드 여럿이 함께 부른다: 넣기·순서 바꾸기·지우기를 한 번에
        CHATGPT_UPLOADED_REFS[ref] = time.time()
        CHATGPT_UPLOADED_REFS.move_to_end(ref)
        while len(CHATGPT_UPLOADED_REFS) > 1000:
            CHATGPT_UPLOADED_REFS.popitem(last=False)


def _docx_block_orders(file_bytes: bytes) -> dict[str, int]:
    """본문 문단(pN)·표(tN)가 본문에서 몇 번째 덩어리인지 - 칸 번호(body:pN, body:tN:…)와 같은 순번 규칙."""
    try:
        from docx import Document

        body = Document(BytesIO(file_bytes)).element.body
    except Exception:
        return {}
    orders: dict[str, int] = {}
    paragraphs = tables = 0
    for index, child in enumerate(body.iterchildren()):
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            orders[f"p{paragraphs}"] = index
            paragraphs += 1
        elif tag == "tbl":
            orders[f"t{tables}"] = index
            tables += 1
    return orders


def _segment_block_order(orders: dict[str, int], segment_id: str) -> int | None:
    parts = str(segment_id).split(":")
    if len(parts) >= 5 and parts[2] == "docx" and parts[3] == "body":
        return orders.get(parts[4])
    return None


def _merged_pii_hits(results) -> str | None:
    """실제로 가린 칸들의 유형별 건수 합계 "KR_RRN:3,EMAIL_ADDRESS:6"."""
    counts: dict[str, int] = {}
    for result in results:
        for part in str(getattr(result, "pii_hits", None) or "").split(","):
            name, _, number = part.partition(":")
            number = number.strip()
            # isdigit()만으로는 '²' 같은 글자도 통과해 int()에서 예외가 난다 - ASCII 숫자만 센다
            if name.strip() and number.isascii() and number.isdigit():
                counts[name.strip()] = counts.get(name.strip(), 0) + int(number)
    return ",".join(f"{name}:{count}" for name, count in sorted(counts.items())) or None


_FILE_MASK_ACTIONS = {"mask", "mask_partial", "mask_full"}
_FILE_BLOCK_ACTIONS = {"block", "block_notify"}
_HWP_KINDS = {"hwp", "hwpx"}   # 2026-10-11 3단계(명세 §4): 칸 가림 + 가린 파일 재검사(_verify_masked)
_SEGMENT_SANITIZED_EXTENSIONS = {"docx", "xlsx"} | _HWP_KINDS
_FILE_MASKABLE_KINDS = _SEGMENT_SANITIZED_EXTENSIONS | {"pdf"}   # PDF는 Detection /redact/pdf(2026-10-11)


def _mask_record(whole, cells):
    """가림을 보낼 수 없을 때(MaskRefused) 403 기록용 판정(명세 §4.1): 파일 전체가 가림이면 그 판정, 아니면 가림 칸 판정.
    둘 다 없으면 None(부르는 쪽이 503 + 기록). 파일 전체 '통과' 판정으로 집행하면 mask_callback과 무관하게 통과하므로 쓰지 않는다."""
    if whole is not None and _claude_action(whole) in _FILE_MASK_ACTIONS:
        return whole
    masked = cells.masked if cells is not None else []
    return masked[0] if masked else None


def _verify_hwp_artifacts(flow: http.HTTPFlow, artifacts, cells, *, request_id: str, kind: str, service: str,
                          endpoint: str, filename) -> None:
    """가린 HWP·HWPX의 마지막 검사(명세 §4.4·§4.5). 통과하면 None.
    - 칸 일부 무응답(원본 전송 설정)인데 바꾼 칸이 있음 → FileInspectionUnavailable(503, 명세 §4.5 넷째 줄)
    - 재검사 입력이 없음 → MaskRefused / 재검사(_verify_masked, 바꾼 구간 안에 완전히 든 적중만 무시) 실패 → 403·503"""
    if cells.unanswered is not None:
        raise FileInspectionUnavailable(cells.unanswered)
    verify = getattr(artifacts, "verify", None)
    if verify is None:
        raise MaskRefused(f"{kind}_verify_missing")
    _verify_masked(flow, verify.text, request_id=f"{request_id}#verify",
                   source={"service": service, "endpoint": endpoint.split("?", 1)[0], "content_type": "file",
                           "filename": filename, "format": kind},
                   replaced=verify.spans)


def guard_chatgpt_file_upload(
    flow: http.HTTPFlow,
    identity: dict,
    *,
    request_id: str,
    normalized_text: str,
    filename: str | None,
    file_bytes: bytes,
    endpoint: str,
    apply_sanitized,
    file_ref: str | None = None,
    kind: str | None = None,
    file_identity: str | None = None,
) -> bool:
    """ChatGPT 웹 파일 업로드: 파일 전체 검사 → 칸별 가림 → 집행·대시보드 기록.

    Claude 파일 경로(detect_claude_file_upload)와 같은 순서다(2026-10-08).
    이전에는 칸별 가림만 하고 결과를 대시보드에 보내지 않았고(이벤트 0건),
    칸별 가림을 지원하지 않는 PDF·TXT는 내용과 무관하게 503이었다.

    - 파일 전체가 차단 대상(기밀 문서 일치 등)이면 차단 + 이벤트
    - DOCX/XLSX: 칸별 가림. 한 칸이라도 차단 대상이면 차단 + 이벤트, 가렸으면 마스킹 이벤트
    - PDF(2026-10-11): 가려야 하면 Detection /redact/pdf가 지운 PDF를 다시 검사해 보낸다(_mask_chatgpt_pdf_upload)
    - 그 밖의 형식: 가릴 것이 없으면 원본 그대로 통과, 가려야 하면 가릴 수 없으므로 차단 + 이벤트
    apply_sanitized(outcome)는 가린 bytes를 요청에 적용한다(실제로 바뀐 경우에만 호출).
    응답을 직접 만들었으면(차단·503) True.

    2026-10-10 (Gateway 처리 설정)
    - 형식은 kind(inspect_upload가 바이트로 정함). 없으면 이전처럼 파일 이름 확장자
    - Detection 호출은 파일 대기 시간·처리 설정(_inspect). 상한 초과(file_too_large)면 칸 검사 없음
    - 칸 무응답: 차단 설정 → 503 + 이벤트. 원본 전송 설정 → 남은 칸은 호출 없이 넘기고 앞 칸 가림은 적용,
      기록은 degraded. 파일 전체가 가림이면 설정과 무관하게 503 + 이벤트(_CellInspection)
    2026-10-11: file_identity(upload_identity)가 있으면 같은 업로드 재시도의 기록·알림은 한 번만
    """
    context = {"identity": identity, "targetLlm": "ChatGPT", "contentType": "file"}
    if file_ref:
        context["fileId"] = file_ref
        _remember_uploaded_ref(file_ref)
    if file_identity:
        context["fileIdentity"] = file_identity
    result = _inspect(
        flow,
        "file",
        request_id=request_id,
        normalized_text=normalized_text,
        source={"service": "ChatGPT", "endpoint": endpoint,
                "content_type": "file", "filename": filename},
    )
    print("\n===== CHATGPT FILE DETECTION RESULT =====")
    print_detection_summary(result)
    action = getattr(result, "action", None)
    if action in _FILE_BLOCK_ACTIONS:
        return enforce_detection_result(flow, result, context)
    if getattr(result, "reason_code", None) == "file_too_large":
        # 상한 초과 + 원본 전송 설정: 검사·칸 가림 없이 전달 + 기록
        return _report_uninspected(flow, result, context)

    if kind:
        extension = kind
    else:
        extension = filename.rsplit(".", 1)[-1].lower() if filename and "." in filename else ""
    if extension == "pdf" and action in _FILE_MASK_ACTIONS:
        # 2026-10-11(명세 §3.5): Detection이 지운 PDF를 다시 검사해 보낸다. 실패는 모두 403·503(원본 안 보냄)
        return _mask_chatgpt_pdf_upload(flow, result, context, request_id=request_id, file_bytes=file_bytes,
                                        filename=filename, endpoint=endpoint, apply_sanitized=apply_sanitized)
    if extension not in _SEGMENT_SANITIZED_EXTENSIONS:
        # 칸별 가림을 지원하지 않는 형식: 가려야 하면 enforce가 차단으로 바꾼다(원본을 보내지 않음)
        return enforce_detection_result(flow, result, context, mask_callback=lambda _: False)

    cells = _CellInspection(flow)
    block_orders = _docx_block_orders(file_bytes) if extension == "docx" else {}
    hwp = extension in _HWP_KINDS

    def detect_segment(segment, child_request_id):
        source = {"service": "ChatGPT", "endpoint": endpoint,
                  "content_type": "file", "filename": filename,
                  "context_label": canonical_label(segment.context),
                  "signature_original_offsets": True}
        block_order = getattr(segment, "block_order", None)   # HWP·HWPX는 칸이 들고 온다(2026-10-11)
        if block_order is None:
            block_order = _segment_block_order(block_orders, segment.segment_id)
        if block_order is not None:
            source["block_order"] = block_order   # 기록 화면에서 문단·표 앞뒤를 되살리는 데만 쓴다
        return cells.detect(segment, child_request_id, source)

    def refuse(reason):
        # 가린 HWP·HWPX를 보낼 수 없음(명세 §4.1): 403 mask_application_failed + 이벤트
        print("[HWP-MASK] ChatGPT mask refused reason=", reason)
        record = _mask_record(result, cells)
        if record is None:
            return fail_closed_file_upload(flow, context, request_id, cause=result)
        return enforce_detection_result(flow, record, context, mask_callback=lambda _: False)

    try:
        # 가림은 바이트로 정한 형식으로(이름을 바꾼 DOCX도 DOCX로 가린다)
        if hwp:
            # HWP·HWPX(2026-10-11): 검토용 사본을 남기지 않고, 가린 결과는 아래에서 다시 검사한다
            outcome = sanitize_for_upload(f"upload.{extension}", file_bytes, request_id, detect_segment,
                                          keep_review=False, allow_hwp=True, max_decompressed=GATEWAY_FILE_MAX_BYTES)
        else:
            outcome = sanitize_for_upload(f"upload.{extension}", file_bytes, request_id, detect_segment)
    except BlockedBySignaturePolicy as exc:
        return enforce_detection_result(flow, exc.detection_result, context)
    except FileInspectionUnavailable as exc:
        print("[SIGNATURE] ChatGPT segment inspection unavailable; upload blocked")
        return fail_closed_file_upload(flow, context, request_id, cause=exc.detection_result)
    except MaskRefused as exc:
        return refuse(exc.reason)
    except Exception as exc:
        print("[SIGNATURE] ChatGPT file masking failed error_type=", type(exc).__name__, "reason=", str(exc))
        flow.response = http.Response.make(503, b'{"error":{"type":"file_masking_failed"}}',
                                           {"Content-Type": "application/json"})
        return True

    changed = outcome.sanitized_bytes != file_bytes
    if action in _FILE_MASK_ACTIONS and cells.unanswered is not None:
        # 파일 전체 판정이 가림인데 칸 일부를 검사하지 못했다: 처리 설정과 무관하게 원본을 보내지 않는다
        print("[SIGNATURE] ChatGPT file mask cannot be reproduced; some segments unanswered")
        return fail_closed_file_upload(flow, context, request_id, cause=cells.unanswered)
    if action in _FILE_MASK_ACTIONS and not changed:
        if hwp:
            # HWP·HWPX(명세 §4.4): 파일 전체에서만 잡힌 값(이름표가 옆 칸이 아닌 값 등) - 403 + 이벤트
            return refuse("mask_not_reproduced")
        # 파일 전체 검사는 가리라고 했는데 칸별 가림이 바꾼 것이 없다 - 원본을 보내지 않는다(Claude 경로와 같음)
        print("[SIGNATURE] ChatGPT file mask was not reproduced by segments; upload blocked")
        flow.response = http.Response.make(503, b'{"error":{"type":"file_masking_failed"}}',
                                           {"Content-Type": "application/json"})
        return True
    if hwp and changed:
        try:
            _verify_hwp_artifacts(flow, outcome, cells, request_id=request_id, kind=extension, service="ChatGPT",
                                  endpoint=endpoint, filename=filename)
        except MaskRefused as exc:
            return refuse(exc.reason)
        except FileInspectionUnavailable as exc:
            print("[HWP-MASK] ChatGPT masked file re-check unavailable; upload blocked")
            return fail_closed_file_upload(flow, context, request_id, cause=exc.detection_result)
        except Exception as exc:
            print("[HWP-MASK] ChatGPT masked file re-check failed error_type=", type(exc).__name__)
            return fail_closed_file_upload(flow, context, request_id, cause=result)
    if changed:
        apply_sanitized(outcome)
        masked_segments = cells.masked
        if action not in _FILE_MASK_ACTIONS:
            # 파일 전체로는 통과였지만 칸 단위 문맥으로 가린 경우 - 가린 칸의 판정으로 기록한다
            result = masked_segments[0]
        # 대시보드 유형별 건수 = 실제로 가린 칸들의 합계(계좌번호처럼 칸 이름표를 봐야 잡히는 것 포함)
        merged = _merged_pii_hits(masked_segments)
        if merged:
            result = copy.copy(result)   # 칸 판정 원본은 그대로 두고 기록용 사본만 바꾼다
            result.pii_hits = merged
    if cells.unanswered is not None:
        # 파일 전체는 통과, 칸 일부 무응답(원본 전송 설정): 앞 칸 가림은 적용하고 전달. 기록은 degraded(FAIL_OPEN)
        result = cells.degraded_report(request_id, changed)
    return enforce_detection_result(flow, result, context, mask_callback=lambda _: changed)


def _guard_chatgpt_web_upload(
    flow: http.HTTPFlow,
    identity: dict,
    *,
    file_bytes: bytes | None,
    filename: str | None,
    mime: str | None,
    endpoint: str,
    file_ref: str | None,
    apply_sanitized,
    make_event,
) -> bool:
    """ChatGPT 웹 파일(원본 PUT·multipart) 2026-10-10: 바이트로 분류 → 검사할 수 없으면 설정에 따라 차단·허용
    → 검사 가능하면 그 글자로 guard_chatgpt_file_upload. 처리 중 예상 못한 오류는 503 + 이벤트(원본 안 보냄).
    make_event(file_class, request_id)는 다음 모듈용 이벤트를 만든다(글자는 file_class.text)."""
    context = {"identity": identity, "targetLlm": "ChatGPT", "contentType": "file"}
    if file_ref:
        context["fileId"] = file_ref
    request_id = str(uuid.uuid4())
    try:
        # 같은 업로드 재시도 기록 한 번만(2026-10-11): 원본 PUT은 주소의 파일 번호, multipart는 바이트 + 이름
        context["fileIdentity"] = upload_identity(
            file_ref=file_ref, files=[(filename, file_bytes)] if file_bytes is not None else None,
            body=flow.request.raw_content, content_type=flow.request.headers.get("content-type", ""))
        file_class = _classify_upload(flow, file_bytes, filename=filename, mime=mime)
        if file_class.group:
            if file_ref:
                _remember_uploaded_ref(file_ref)
            return _report_uninspected(flow, uninspected_result(flow, file_class.group, request_id), context)
        if not file_class.text:
            print("→ Empty file; PASS THROUGH")
            return False

        event = make_event(file_class, request_id)
        print_gateway_event(event)
        send_to_next_module(event)
        normalized_text = extract_file_checkable_text(event)
        if not normalized_text:   # inspect_upload가 글자 없는 파일을 other로 돌려주므로 오지 않는다
            return _report_uninspected(flow, uninspected_result(flow, "other", request_id), context)

        resolved_filename = (event.get("contents") or [{}])[0].get("filename")
        print("\n===== FILE TEXT EXTRACTION SUCCESS =====")
        print("Kind:", file_class.kind)
        print("Extracted text length:", len(normalized_text))
        print("→ Sending extracted file text to Detection Module")
        return guard_chatgpt_file_upload(
            flow, identity,
            request_id=request_id, normalized_text=normalized_text,
            filename=resolved_filename, file_bytes=file_bytes,
            endpoint=endpoint, apply_sanitized=apply_sanitized,
            file_ref=file_ref, kind=file_class.kind, file_identity=context["fileIdentity"],
        )
    except Exception as exc:
        # 예상하지 못한 오류에서도 원본을 보내지 않는다 + 기록(이전: request() 밖으로 나가 원본 전달되던 곳 포함)
        print("[FILE] ChatGPT upload processing failed error_type=", type(exc).__name__)
        return fail_closed_file_upload(flow, context, request_id)



# ============================================================
# Claude Web Request 사용자 프롬프트 추출
# ============================================================
def extract_claude_prompt(flow: http.HTTPFlow) -> str | None:
    """
    Claude Web의 실제 대화 전송 요청에서 사용자 prompt를 추출한다.

    현재 확인된 Claude Web completion 형태:
      POST /api/organizations/{orgId}/chat_conversations/{convId}/completion
      Content-Type: application/json

    일반적인 body의 prompt 필드를 우선 사용하고,
    구조가 달라진 경우 messages 계열의 user text를 보조적으로 찾는다.
    인증정보/쿠키/전체 body는 로그에 출력하지 않는다.
    """
    req = flow.request

    if req.host != "claude.ai" or req.method.upper() != "POST":
        return None

    path = req.path.split("?", 1)[0]

    # Claude Web 실제 completion endpoint
    if not re.match(
        r"^/api/organizations/[^/]+/chat_conversations/[^/]+/completion$",
        path,
    ):
        return None

    content_type = req.headers.get("content-type", "").lower()
    if "application/json" not in content_type:
        return None

    try:
        body = json.loads(req.get_text(strict=False) or "")
    except (json.JSONDecodeError, TypeError, UnicodeDecodeError):
        print("❌ Claude completion JSON parsing failed")
        return None

    if not isinstance(body, dict):
        return None

    # 1순위: Claude Web completion의 prompt
    prompt = body.get("prompt")
    if isinstance(prompt, str) and prompt.strip():
        return prompt.strip()

    # 2순위: messages 구조가 사용되는 경우
    messages = body.get("messages")
    candidates = []

    def walk_messages(value):
        if isinstance(value, dict):
            role = value.get("role")
            content = value.get("content")

            if role in {"user", "human"}:
                if isinstance(content, str) and content.strip():
                    candidates.append(content.strip())
                elif isinstance(content, list):
                    for item in content:
                        if isinstance(item, dict):
                            item_type = item.get("type")
                            if item_type in {"text", "input_text"}:
                                item_text = item.get("text")
                                if isinstance(item_text, str) and item_text.strip():
                                    candidates.append(item_text.strip())

            for child in value.values():
                walk_messages(child)

        elif isinstance(value, list):
            for child in value:
                walk_messages(child)

    if isinstance(messages, (dict, list)):
        walk_messages(messages)

    if candidates:
        return candidates[-1]

    return None


def extract_claude_attachments(flow: http.HTTPFlow) -> list[dict]:
    """
    Claude Web completion 본문의 attachments[] 항목에서 검사할 글자를 뽑는다(2026-10-10).

    브라우저가 직접 읽은 작은 텍스트 파일·긴 붙여넣기는 multipart 업로드로 가지 않고
    extracted_content 칸에 글자로 실려 간다. 이전에는 prompt만 검사해서 이 글자는 검사 없이 나갔다.
    파일 이름(file_name)도 그대로 Claude에 가므로 따로 검사한다.
    prompt와 합치지 않는다(event_adapter.py 2026-09-23 원칙: 항목마다 한 번 검사).
    내용·파일 이름은 로그에 남기지 않는다.

    반환: [{"index": attachments[] 안 위치, "filename": file_name 또는 None,
            "text": 정규화한 extracted_content 또는 None, "name_text": 정규화한 file_name 또는 None}]
    글자도 이름도 없는 항목은 뺀다.
    """
    req = flow.request

    if req.host != "claude.ai" or req.method.upper() != "POST":
        return []

    path = req.path.split("?", 1)[0]
    if not re.match(
        r"^/api/organizations/[^/]+/chat_conversations/[^/]+/completion$",
        path,
    ):
        return []

    content_type = req.headers.get("content-type", "").lower()
    if "application/json" not in content_type:
        return []

    try:
        body = json.loads(req.get_text(strict=False) or "")
    except (json.JSONDecodeError, TypeError, UnicodeDecodeError):
        return []   # 오류 로그는 extract_claude_prompt가 이미 남긴다

    attachments = body.get("attachments") if isinstance(body, dict) else None
    if not isinstance(attachments, list):
        return []

    parts = []
    for index, item in enumerate(attachments):
        if not isinstance(item, dict):
            continue
        content = item.get("extracted_content")
        text = normalize_text(content) if isinstance(content, str) and content.strip() else ""
        filename = item.get("file_name")
        has_name = isinstance(filename, str) and bool(filename.strip())
        name_text = normalize_text(filename) if has_name else ""
        if not text and not name_text:  # 비었거나 보이지 않는 문자뿐인 항목
            continue
        parts.append({
            "index": index,
            "filename": filename if has_name else None,
            "text": text or None,
            "name_text": name_text or None,
        })
    return parts


def _apply_claude_attachment_mask(
    flow: http.HTTPFlow,
    index: int,
    normalized_text: str,
    result,
    field: str = "extracted_content",
) -> bool:
    """attachments[index]의 한 칸(field: extracted_content 또는 file_name)만 가린 글자로 바꾼다.

    _apply_json_mask를 쓰지 않는다: 글자가 같은 다른 칸(prompt 등)까지 바꾸고, 가림 근거(maskEvidence)를
    남기는데 대시보드는 그 근거를 질문(prompt) 이벤트에서만 받는다(validation.mjs). 파일 이벤트에 실리면
    이벤트 전체가 거부된다.
    """
    replacement = getattr(result, "masked_text", None)
    if replacement is None:
        replacement = _masked_text(normalized_text, getattr(result, "mask_spans", None) or [])
    if replacement is None:
        return False
    try:
        body = json.loads(flow.request.get_text(strict=False) or "")
    except (json.JSONDecodeError, TypeError, UnicodeDecodeError):
        return False
    attachments = body.get("attachments") if isinstance(body, dict) else None
    if not isinstance(attachments, list) or not 0 <= index < len(attachments):
        return False
    item = attachments[index]
    value = item.get(field) if isinstance(item, dict) else None
    # 검사한 글자가 들어 있는 칸인지 확인한 뒤에만 바꾼다
    if not isinstance(value, str) or normalize_text(value) != normalized_text:
        return False
    item[field] = replacement
    flow.request.set_text(json.dumps(body, ensure_ascii=False, separators=(",", ":")))
    return True


# 여러 칸을 따로 검사한 요청 하나의 지표 값: 가장 강한 조치
_REQUEST_ACTION_RANK = {"allow": 0, "allow_notify": 1, "allow_degraded": 2, "mask": 3, "block": 4}


def _claude_action(result) -> str:
    """판정의 조치. 옛 클라이언트 결과(action 없음)는 enforce_detection_result와 같이 blocked로 정한다."""
    return getattr(result, "action", None) or ("block" if getattr(result, "blocked", False) else "allow")


def _request_level_action(result) -> str:
    action = _claude_action(result)
    if action in _FILE_BLOCK_ACTIONS:
        return "block"
    if action in _FILE_MASK_ACTIONS:
        return "mask"   # enforce_detection_result가 가림 집행 후 남기는 값과 같다
    return action


def _strongest_result(results):
    """가장 강한 조치의 판정(차단 > 가림 > 무응답 통과 > 알림 통과 > 통과). 같으면 앞 것(질문이 먼저)."""
    return max(results, key=lambda result: _REQUEST_ACTION_RANK.get(_request_level_action(result), 0))


def _attachment_report_result(results):
    """첨부 항목 하나(글자·이름 칸)의 기록용 판정: 강한 쪽. 두 칸 모두 가렸으면 유형별 건수를 합친 사본."""
    report = _strongest_result(results)
    masked = [result for result in results if _claude_action(result) in _FILE_MASK_ACTIONS]
    if len(masked) > 1:
        merged = _merged_pii_hits(masked)
        if merged:
            report = copy.copy(report)   # 칸 판정 원본은 그대로 두고 기록용 사본만 바꾼다
            report.pii_hits = merged
    return report


def _record_request_outcome(flow: http.HTTPFlow, results) -> None:
    """요청 전체 지표 값(response()가 _finalize_gateway_metric에 넘기는 gw_action·gw_detector·gw_degraded)을 정한다.

    enforce_detection_result는 부를 때마다 이 값을 덮어써 마지막 칸의 판정만 남는다 - 질문을 가렸어도 뒤에
    깨끗한 첨부가 있으면 ALLOW, 질문 검사가 무응답이었어도 degraded=False로 집계됐다.
    조치·탐지기는 가장 강한 판정에서, degraded는 한 칸이라도 무응답이면 True.
    """
    strongest = _strongest_result(results)
    flow.metadata["gw_request_id"] = getattr(strongest, "request_id", None)
    flow.metadata["gw_detection_result"] = strongest
    flow.metadata["gw_action"] = _request_level_action(strongest)
    flow.metadata["gw_detector"] = _result_detector(strongest)
    flow.metadata["gw_degraded"] = any(bool(getattr(result, "degraded", False)) for result in results)


def detect_claude_attachments(
    flow: http.HTTPFlow,
    attachments: list[dict],
    base_request_id: str,
) -> list[tuple[dict, object]]:
    """
    본문 안 파일 글자·이름(attachments[])을 칸마다 Detection에 보낸다(2026-10-10). 집행·기록은 detect_claude_web이 한다.

    - 칸: 항목마다 글자(extracted_content)와 이름(file_name). 요청 번호 "<prompt 요청 번호>#attachment:<위치>",
      이름 칸은 끝에 ":name". '#' 앞으로 같은 요청의 판정을 되묶는다
    - 차단 판정이 나오면 거기서 멈춘다. 요청 전체가 막히므로 남은 칸은 보내지 않는다
    - 파일 규칙(대기 시간·처리 설정·글자 상한)으로 검사한다(_inspect 'file').
      Detection 무응답 넘기기는 _inspect가 한다: 연결 안 됨이나 같거나 짧은 대기의 시간 초과가 앞에 있으면
      다시 부르지 않고 파일 규칙으로 판정을 새로 만든다(이벤트는 칸마다). 질문 0.5초 시간 초과 뒤에는
      파일 대기 시간(기본 2초)으로 다시 부른다
    반환: [(칸, 판정)] 검사한 순서대로. 칸 = {"field", "index", "text"}
    """
    endpoint = flow.request.path.split("?", 1)[0]

    print("\n===== CLAUDE ATTACHMENT TEXT DETECTED =====")
    print("Attachment count:", len(attachments))
    print("Endpoint:", endpoint)

    checked = []
    for item in attachments:
        for field, text, suffix in (("extracted_content", item["text"], ""), ("file_name", item["name_text"], ":name")):
            if not text:
                continue
            request_id = f"{base_request_id}#attachment:{item['index']}{suffix}"
            print("Attachment index:", item["index"], "field:", field, "text length:", len(text))

            result = _inspect(
                flow,
                "file",
                request_id=request_id,
                normalized_text=text,
                source={
                    "service": "Claude",
                    "endpoint": endpoint,
                    "content_type": "file",
                    "filename": item["filename"],
                },
            )
            print_detection_summary(result)

            checked.append(({"field": field, "index": item["index"], "text": text}, result))
            if _claude_action(result) in _FILE_BLOCK_ACTIONS:
                return checked

    return checked


def detect_claude_web(flow: http.HTTPFlow, identity: dict) -> bool:
    """
    Claude Web 사용자 prompt를 Detection Module에 전달하고
    기존 Gateway BLOCK / MASK / PASS 정책을 적용한다.

    2026-10-10: 본문 attachments[]의 글자(extracted_content: 브라우저가 읽은 파일 글자·긴 붙여넣기)와
    파일 이름(file_name)도 prompt와 따로, 칸마다 검사한다. prompt가 비어 있어도 검사한다.
    집행되지 않은 결과가 대시보드에 남지 않게 나눠 처리한다.
      1. 검사: 질문 → 첨부 칸 순서. 차단이 나오면 멈추고 그 차단 하나만 집행·기록(요청 전체 403)
      2. 가림: 가릴 칸을 모두 가린다. 하나라도 못 가리면 그 칸만 차단(mask_application_failed)으로 기록
      3. 기록: 질문 1건 + 첨부 항목마다 1건(글자·이름 중 강한 판정). 질문 이벤트의 attachedFileIds와
         파일 이벤트의 fileId("<prompt 요청 번호>-attachment-<위치>")로 대시보드가 한 요청으로 묶는다
    요청 전체 지표(gw_action 등)는 마지막에 가장 강한 판정으로 정한다.
    첨부 칸이 없는 요청은 이전과 같다(Detection 1회, 이벤트 1건).
    """
    prompt = extract_claude_prompt(flow)
    # 첨부 칸은 prompt를 가리기 전 원본 본문에서 읽는다
    attachments = extract_claude_attachments(flow)

    if not prompt and not attachments:
        print("→ Claude user prompt not found")
        return False

    # 1. 검사
    checked = []   # [(칸, 판정)] 검사한 순서대로
    prompt_result = None
    if prompt:
        print("\n===== CLAUDE PROMPT DETECTED =====")
        print("Prompt length:", len(prompt))
        print("Endpoint:", flow.request.path.split("?", 1)[0])

        prompt_result = detect_web_prompt(
            flow,
            "Claude",
            prompt,
        )
        checked.append(({"field": "prompt"}, prompt_result))

    base_request_id = getattr(prompt_result, "request_id", None) or str(uuid.uuid4())
    if attachments and (prompt_result is None or _claude_action(prompt_result) not in _FILE_BLOCK_ACTIONS):
        checked += detect_claude_attachments(
            flow,
            attachments,
            base_request_id,
        )
    if not checked:
        return False

    # 질문과 첨부를 잇는 번호. 대시보드 파일 번호 모양(^[A-Za-z0-9_-]{1,100}$)일 때만 싣는다 - 아니면 이벤트 전체가 거부된다
    file_ids = {}
    for item in attachments:
        file_id = f"{base_request_id}-attachment-{item['index']}"
        if _FILE_REF_RE.fullmatch(file_id):
            file_ids[item["index"]] = file_id

    def dashboard_context_for(part, *, link_attachments=False):
        if part["field"] == "prompt":
            context = {"identity": identity, "targetLlm": "Claude", "contentType": "prompt"}
            if link_attachments and file_ids:
                context["attachedFileIds"] = list(file_ids.values())
            return context
        context = {"identity": identity, "targetLlm": "Claude", "contentType": "file"}
        if part["index"] in file_ids:
            context["fileId"] = file_ids[part["index"]]
        return context

    part, result = checked[-1]
    if _claude_action(result) in _FILE_BLOCK_ACTIONS:
        # 요청 전체를 막는다. 앞 칸의 통과·가림은 집행되지 않았으므로 기록하지 않는다
        return enforce_detection_result(flow, result, dashboard_context_for(part))

    # 2. 가림. 첨부 칸 먼저, 질문은 마지막: _apply_json_mask는 본문에서 질문과 글자가 같은 칸을 모두 바꾸므로,
    #    먼저 하면 같은 글자의 첨부 칸이 검사한 글자 확인(_apply_claude_attachment_mask)에서 어긋난다
    masked = [(part, result) for part, result in checked if _claude_action(result) in _FILE_MASK_ACTIONS]
    for part, result in sorted(masked, key=lambda pair: pair[0]["field"] == "prompt"):
        try:
            if part["field"] == "prompt":
                applied = _apply_json_mask(flow, prompt, result)
            else:
                applied = _apply_claude_attachment_mask(flow, part["index"], part["text"], result, part["field"])
        except Exception as exc:
            print("MASK application failed:", type(exc).__name__)
            applied = False
        if not applied:
            # 못 가린 칸이 있으면 요청 전체를 막는다(enforce가 mask_application_failed 차단으로 바꾼다). 그 차단 하나만 기록
            return enforce_detection_result(flow, result, dashboard_context_for(part), mask_callback=lambda _: False)

    # 3. 기록. 가림은 2에서 이미 적용했다(mask_callback은 적용됨만 알린다)
    if prompt_result is not None:
        enforce_detection_result(
            flow,
            prompt_result,
            dashboard_context=dashboard_context_for({"field": "prompt"}, link_attachments=True),
            mask_callback=lambda _: True,
        )
    for item in attachments:
        results = [result for part, result in checked if part["field"] != "prompt" and part["index"] == item["index"]]
        enforce_detection_result(
            flow,
            _attachment_report_result(results),
            dashboard_context=dashboard_context_for({"field": "file", "index": item["index"]}),
            mask_callback=lambda _: True,
        )
    _record_request_outcome(flow, [result for _, result in checked])
    return False



def extract_gemini_prompt(flow: http.HTTPFlow) -> str | None:
    """
    Gemini Web StreamGenerate 요청의 f.req 내부에서
    사용자 입력 텍스트를 추출한다.

    실제 구조:
    form-data:
        f.req = [
            null,
            "[[... , \"USER_PROMPT\", ...]]"
        ]
    """

    req = flow.request

    if req.method.upper() != "POST":
        return None

    path = req.path.split("?", 1)[0]

    if "StreamGenerate" not in path:
        return None

    try:
        form = req.urlencoded_form

        raw = form.get("f.req")

        if not raw:
            return None

        outer = json.loads(raw)

        if not isinstance(outer, list):
            return None

        # Gemini는 두 번째 element가 JSON 문자열
        if len(outer) < 2 or not isinstance(outer[1], str):
            return None

        nested = json.loads(outer[1])

        if not isinstance(nested, list):
            return None

        candidates = []

        def walk(obj):
            if isinstance(obj, str):
                value = obj.strip()

                # prompt 후보 필터
                if (
                    len(value) >= 2
                    and not value.startswith("!")
                    and not value.startswith("http")
                    and value not in ("ko", "en")
                ):
                    candidates.append(value)

            elif isinstance(obj, list):
                for item in obj:
                    walk(item)

            elif isinstance(obj, dict):
                for value in obj.values():
                    walk(value)

        walk(nested)

        if not candidates:
            return None

        # 가장 긴 문자열을 사용자 prompt 후보로 사용
        return max(candidates, key=len)

    except Exception as e:
        print("Gemini prompt extraction failed:", e)
        return None


def detect_gemini_web(flow: http.HTTPFlow, identity: dict):
    prompt = extract_gemini_prompt(flow)

    if not prompt:
        print("→ Gemini user prompt not found")
        return False

    print("\n===== GEMINI PROMPT DETECTED =====")
    print("Prompt length:", len(prompt))

    result = detect_web_prompt(
        flow,
        "Gemini",
        prompt,
    )

    return enforce_detection_result(
        flow,
        result,
        dashboard_context={
            "identity": identity,
            "targetLlm": "Gemini",
            "contentType": "prompt",
        },
        mask_callback=lambda detection_result: _apply_gemini_mask(flow, prompt, detection_result),
    )


# ============================================================
# Gateway Traffic / Performance Metrics
# ============================================================
# 프롬프트/파일 원문은 저장하지 않으며 프로세스 메모리에서만 집계한다.

METRICS_WINDOW_SECONDS = 60
METRICS_LATENCY_MAX = 5000
_gateway_metrics_lock = threading.RLock()
_gateway_client_inflight = Counter()
_gateway_started_requests = deque()
_gateway_completed_requests = deque()
_gateway_latencies_ms = deque(maxlen=METRICS_LATENCY_MAX)
_gateway_metrics = {
    "total_requests": 0,
    "completed_requests": 0,
    "error_requests": 0,
    "timeout_requests": 0,
    "degraded_requests": 0,
    # 2026-10-15 지표 확장 칸(대시보드가 metrics_ext를 알린 뒤에만 전송, CONTRACT.md). 셈은 늘 한다(요청마다 찍는 지표 로그에도 남김)
    "inspected_requests": 0,          # 끝난 요청 중 Gateway가 검사·판정한 것(연결 IP 표 '검사 요청'과 같은 기준)
    "inspected_timeout_requests": 0,  # 그중 연결 시간 초과로 끝난 것
    "client_cancelled_requests": 0,   # error_requests 중 브라우저가 취소·끊은 것(HTTP/2 CANCEL, HTTP/1 클라이언트 끊김)
    "action_counts": Counter(),
    "llm_counts": Counter(),
    "detector_counts": Counter(),
}
# 확장 칸 협상: 옛 대시보드는 모르는 칸이 오면 지표 전체를 거부하므로 응답에 metrics_ext가 온 뒤에만 싣는다
_metrics_ext = MetricsExtNegotiator()
# 브라우저(클라이언트) 취소로 보는 mitmproxy 오류 문구(12.2.3). 'stream reset by client (CANCEL)' = HTTP/2 스트림 취소
# (답변 스트리밍·폴링을 브라우저가 멈춤), 'Client disconnected.' = HTTP/1 클라이언트가 응답을 다 받기 전에 끊음.
# INTERNAL_ERROR 리셋·'peer closed connection'(어느 쪽이 닫았는지 문구로 알 수 없음)·시간 초과는 연결 오류로 둔다
_CLIENT_CANCEL_ERRORS = frozenset({"stream reset by client (CANCEL)", "Client disconnected."})
# 연결 중인 IP 표(2026-10-12): 지표 전송기의 "clients" 칸에만 싣는다(요청마다 찍는 지표 로그에는 없음, IP 목록 로그 없음)
_client_tracker = ClientTracker()


def _gateway_tracked(flow: http.HTTPFlow) -> bool:
    host = (flow.request.host or "").lower()
    return bool(service_for(host)) or host == OPENAI_HOST


def _gateway_service_name(flow: http.HTTPFlow) -> str:
    """llm_distribution에 쓰는 LLM 이름(LLM 사이트 HTTP 요청 전체). 연결 IP 표 targets는 _inspected_target_llm(10/15)."""
    return service_for(flow.request.host) or (
        "OpenAI" if flow.request.host == OPENAI_HOST else "Unknown"
    )


def _prune_metric_window(queue, cutoff):
    while queue and queue[0] < cutoff:
        queue.popleft()


def _start_gateway_metric(flow: http.HTTPFlow):
    if not _gateway_tracked(flow):
        return
    now = time.time()
    client_ip = get_client_ip(flow)
    with _gateway_metrics_lock:
        _prune_metric_window(_gateway_started_requests, now - METRICS_WINDOW_SECONDS)
        _prune_metric_window(_gateway_completed_requests, now - METRICS_WINDOW_SECONDS)
        _gateway_metrics["total_requests"] += 1
        _gateway_started_requests.append(now)
        if client_ip:
            _gateway_client_inflight[client_ip] += 1
    # 처리 중·대기 연결만(요청 수는 검사 때). 식별 헤더는 아직 지우기 전이라 여기서 읽는다(by_client 묶음, 10/15)
    flow.metadata["gw_metrics_client_id"] = flow.request.headers.get("x-llmsgw-client-id")
    flow.metadata["gw_client_ticket"] = _client_tracker.request_started(
        client_ip, flow.metadata["gw_metrics_client_id"], _flow_conn_id(flow))
    flow.metadata["gw_metrics_started_at"] = time.perf_counter()
    flow.metadata["gw_metrics_client_ip"] = client_ip
    flow.metadata["gw_metrics_finalized"] = False
    _emit_gateway_metrics()


def _gateway_p95(values):
    if not values:
        return 0.0
    ordered = sorted(values)
    # Same nearest-rank 95th percentile as Dashboard, including small samples.
    index = max(0, (len(ordered) * 95 + 99) // 100 - 1)
    return round(float(ordered[index]), 2)


def _flow_conn_id(flow: http.HTTPFlow):
    """mitmproxy 클라이언트 연결 id(연결 IP 표 by_client 묶음: 연결 ↔ 확장 ID). 없으면 None."""
    return getattr(getattr(flow, "client_conn", None), "id", None)


def _flow_client_id(flow: http.HTTPFlow):
    """이 요청의 확장 프로그램 식별 헤더 값: 읽어 둔 identity → 지표 시작 때 읽은 값 → 아직 남은 헤더. 형식 검사는 표가 한다."""
    identity = flow.metadata.get("llmsgw_identity")
    if isinstance(identity, dict) and identity.get("client_id"):
        return identity.get("client_id")
    return flow.metadata.get("gw_metrics_client_id") or flow.request.headers.get("x-llmsgw-client-id")


# 확장 칸(2026-10-15, metrics_ext 1). 정의는 docs/handoff/gateway-metrics-fix-2026-10-15/CONTRACT.md
_METRICS_EXT_COUNT_FIELDS = ("latency_samples", "inspected_requests", "degraded_requests", "timeout_requests",
                             "inspected_timeout_requests", "client_cancelled_requests", "connection_error_requests")


def _gateway_metrics_snapshot():
    """10/14까지의 지표 칸 그대로(이벤트 메타데이터·지표 전송 기본). 확장 칸은 _gateway_metrics_ext_snapshot."""
    return _gateway_metrics_ext_snapshot()[0]


def _gateway_metrics_ext_snapshot():
    """(옛 칸 dict, 확장 셈 dict) - 같은 잠금 안에서 한 번에 읽어 두 값이 서로 맞는다.
    확장 셈: latency_samples = 평균·P95를 낸 처리 시간 표본 수(최근 METRICS_LATENCY_MAX건), inspected_requests = 끝난 요청 중
    검사·판정한 것(degraded_rate·timeout_rate의 검사 요청 기준 분모), degraded_requests·timeout_requests = 두 비율의 분자(옛 정의 그대로),
    inspected_timeout_requests = 검사 요청 중 시간 초과, client_cancelled_requests = error_requests 중 브라우저 취소,
    connection_error_requests = error_requests - client_cancelled_requests."""
    now = time.time()
    with _gateway_metrics_lock:
        _prune_metric_window(_gateway_started_requests, now - METRICS_WINDOW_SECONDS)
        _prune_metric_window(_gateway_completed_requests, now - METRICS_WINDOW_SECONDS)
        latencies = list(_gateway_latencies_ms)
        total = max(1, _gateway_metrics["total_requests"])
        counts = {
            "latency_samples": len(latencies),
            "inspected_requests": _gateway_metrics["inspected_requests"],
            "degraded_requests": _gateway_metrics["degraded_requests"],
            "timeout_requests": _gateway_metrics["timeout_requests"],
            "inspected_timeout_requests": _gateway_metrics["inspected_timeout_requests"],
            "client_cancelled_requests": _gateway_metrics["client_cancelled_requests"],
            "connection_error_requests": _gateway_metrics["error_requests"] - _gateway_metrics["client_cancelled_requests"],
        }
        return {
            "active_clients": sum(
                1 for count in _gateway_client_inflight.values() if count > 0
            ),
            "concurrent_requests": sum(_gateway_client_inflight.values()),
            "requests_per_sec": round(
                len(_gateway_started_requests) / METRICS_WINDOW_SECONDS, 3
            ),
            "avg_latency_ms": (
                round(sum(latencies) / len(latencies), 2) if latencies else 0.0
            ),
            "p95_latency_ms": _gateway_p95(latencies),
            "timeout_rate": round(_gateway_metrics["timeout_requests"] / total, 4),
            "degraded_rate": round(_gateway_metrics["degraded_requests"] / total, 4),
            "total_requests": _gateway_metrics["total_requests"],
            "completed_requests": _gateway_metrics["completed_requests"],
            "error_requests": _gateway_metrics["error_requests"],
            "action_distribution": dict(_gateway_metrics["action_counts"]),
            "llm_distribution": dict(_gateway_metrics["llm_counts"]),
            "detector_distribution": dict(_gateway_metrics["detector_counts"]),
            "window_seconds": METRICS_WINDOW_SECONDS,
            "event_delivery": _gateway_event_outbox.snapshot(),
        }, counts


from gateway_metrics_delivery import MetricsPublisher


def _publisher_snapshot():
    """지표 전송기 전용: 지표 + (대시보드가 설정 칸을 준 뒤에만) 설정 적용 보고 + 연결 중인 IP 표(clients, 10/12)
    + 확장 프로그램 중계 접속 현황(relay, 10/12 시스템 상태).
    요청마다 찍는 지표 로그·이벤트 메타데이터(_gateway_metrics_snapshot)에는 넣지 않는다.
    2026-10-15: 대시보드 응답에 metrics_ext가 온 뒤에만 확장 칸(metrics_ext·표본 수·검사 요청 수·브라우저 취소 수 등,
    clients 항목마다 by_client)을 더 싣고, 이때는 표본이 0건이면 avg_latency_ms·p95_latency_ms를 0.0 대신 null로 보낸다.
    옛 대시보드(모르는 칸이면 지표 전체 400, 평균·P95가 숫자가 아니면 422)에는 10/14 칸 그대로 보낸다."""
    extended = _metrics_ext.include()
    metrics, counts = _gateway_metrics_ext_snapshot()
    if extended:
        metrics["metrics_ext"] = METRICS_EXT_VERSION
        metrics.update(counts)
        if counts["latency_samples"] == 0:
            metrics["avg_latency_ms"] = None
            metrics["p95_latency_ms"] = None
    return {**metrics, **GATEWAY_SETTINGS.report_field(), **_access_field(), **_clients_field(extended),
            **_relay_field()}


def _access_field() -> dict:
    """접속 차단 적용 보고(access_applied, 10/12). 대시보드에서 gateway_access 칸을 받은 뒤(또는 사본을 읽은 뒤)만.
    보고를 못 만들어도 지표는 보낸다(칸만 뺀다)."""
    try:
        return GATEWAY_ACCESS.report_field()
    except Exception as exc:
        print(f"[METRICS] access_applied omitted error_type={type(exc).__name__}")
        return {}


def _clients_field(by_client: bool = False) -> dict:
    """표를 못 만들어도 지표는 보낸다(칸만 뺀다). 값은 남기지 않는다. by_client: 확장 칸 협상 뒤 IP 항목마다 확장 ID 묶음."""
    try:
        return {"clients": _client_tracker.snapshot(by_client=by_client)}
    except Exception as exc:
        print(f"[METRICS] clients omitted error_type={type(exc).__name__}")
        return {}


def _relay_field() -> dict:
    """중계(127.0.0.1:8766 /status)를 0.5초 안에 한 번 묻는다. 지표 전송기 스레드에서만 부른다(요청마다 부르지 않는다).
    실패해도 지표는 보낸다: {"relay": {"ok": False, "connected": 0, "registered": 0}}."""
    try:
        return {"relay": fetch_relay_status()}
    except Exception as exc:
        print(f"[METRICS] relay status failed error_type={type(exc).__name__}", flush=True)
        return {"relay": {"ok": False, "connected": 0, "registered": 0}}


def client_connected(client):
    """mitmproxy 훅: 프록시로 TCP 연결이 들어옴(연결 하나에 요청 여러 개). 연결 중인 IP 표에만 쓴다."""
    _client_tracker.connected(_peer_ip(client), getattr(client, "id", None))


def client_disconnected(client):
    _client_tracker.disconnected(_peer_ip(client), getattr(client, "id", None))
    _access_connection_ids.forget(getattr(client, "id", None))


def _peer_ip(client):
    peername = getattr(client, "peername", None)
    return peername[0] if peername else None


def _each(*handlers):
    """지표 응답 하나를 처리 설정·접속 차단에 각각 넘긴다. 한쪽 예외가 다른 쪽을 막지 않는다."""
    def deliver(value):
        for handler in handlers:
            try:
                handler(value)
            except Exception as exc:
                print(f"[DASHBOARD] metrics response handling failed error_type={type(exc).__name__}")
    return deliver


_gateway_metrics_publisher = MetricsPublisher(
    _publisher_snapshot,
    on_payload=_each(GATEWAY_SETTINGS.offer, GATEWAY_ACCESS.offer, _metrics_ext.offer),
    on_status=_each(GATEWAY_SETTINGS.on_metrics_status, GATEWAY_ACCESS.on_metrics_status, _metrics_ext.on_metrics_status),
)
_gateway_event_outbox = EventOutbox()


def running():
    """Keep aggregate telemetry fresh even when no user is sending requests."""
    _gateway_event_outbox.start()
    _gateway_metrics_publisher.start()


def done():
    _gateway_metrics_publisher.stop()
    _gateway_event_outbox.stop()
    _shutdown_file_executor()


def _emit_gateway_metrics():
    """Notify one bounded worker; never do network work on request hooks."""
    _gateway_metrics_publisher.notify()


# 설정을 적용하면 다음 지표를 바로 보내 "반영됨" 보고를 앞당긴다(최소 간격 1초는 전송기가 지킨다)
GATEWAY_SETTINGS.set_on_change(_emit_gateway_metrics)
GATEWAY_ACCESS.set_on_change(_emit_gateway_metrics)


def _finalize_gateway_metric(
    flow: http.HTTPFlow,
    *,
    action=None,
    detector=None,
    degraded=False,
    error=False,
    timeout=False,
    client_cancelled=False,
):
    if (
        not flow.metadata.get("gw_metrics_started_at")
        or flow.metadata.get("gw_metrics_finalized")
    ):
        return
    flow.metadata["gw_metrics_finalized"] = True
    round_trip_ms = round(
        (time.perf_counter() - flow.metadata["gw_metrics_started_at"]) * 1000,
        2,
    )
    # 2026-10-10: 지연 시간 = Gateway 자체 처리 시간(요청 시작 → request() 처리 끝: upstream 전달 또는 Gateway 응답).
    # 이전에는 응답이 끝날 때까지 재서 LLM 답변 시간(예: 17초)이 들어갔다. 집계(완료 수 등)는 지금처럼 응답 때 한다.
    # 응답까지의 시간은 로그에만 남긴다(대시보드 지표·이벤트 형식은 모르는 칸을 거부한다)
    latency_ms = flow.metadata.get("gw_processing_ms")
    if not isinstance(latency_ms, (int, float)) or isinstance(latency_ms, bool):
        latency_ms = round_trip_ms
    service = _gateway_service_name(flow)
    client_ip = flow.metadata.get("gw_metrics_client_ip")
    # 평균·p95 표본은 Gateway가 검사·판정·응답을 한 요청만(10/11, request()가 정함). 집계(요청 수·완료 수 등)는 모든 요청.
    # request()를 거치지 않은 flow(표시 없음)는 이전처럼 표본에 넣는다
    latency_sample = flow.metadata.get("gw_latency_sample", True)
    # 검사 요청(10/15 확장 칸 inspected_requests): 연결 IP 표 '검사 요청'과 같은 기준(Detection을 불렀거나 판정을 정함).
    # 검사 응답 없음(degraded)은 늘 검사 요청이다 → degraded_requests ≤ inspected_requests
    inspected = bool(flow.metadata.get("gw_inspected") or flow.metadata.get("gw_action") or degraded)
    with _gateway_metrics_lock:
        _gateway_metrics["completed_requests"] += 1
        _gateway_completed_requests.append(time.time())
        if latency_sample:
            _gateway_latencies_ms.append(latency_ms)
        if error:
            _gateway_metrics["error_requests"] += 1
            if client_cancelled:
                _gateway_metrics["client_cancelled_requests"] += 1
        if timeout:
            _gateway_metrics["timeout_requests"] += 1
        if inspected:
            _gateway_metrics["inspected_requests"] += 1
            if timeout:
                _gateway_metrics["inspected_timeout_requests"] += 1
        if degraded:
            _gateway_metrics["degraded_requests"] += 1
        if action:
            _gateway_metrics["action_counts"][str(action).upper()] += 1
        _gateway_metrics["llm_counts"][service] += 1
        if detector:
            _gateway_metrics["detector_counts"][str(detector)] += 1
        if client_ip:
            current = _gateway_client_inflight.get(client_ip, 0) - 1
            if current > 0:
                _gateway_client_inflight[client_ip] = current
            else:
                _gateway_client_inflight.pop(client_ip, None)
    _client_tracker.request_finished(flow.metadata.get("gw_client_ticket"))
    flow.metadata["gw_latency_ms"] = latency_ms
    flow.metadata["gw_round_trip_ms"] = round_trip_ms
    metrics, counts = _gateway_metrics_ext_snapshot()
    flow.metadata["gw_metrics_snapshot"] = metrics
    _emit_gateway_metrics()
    # 파일 업로드 작업 스레드 대기(10/11): gateway_ms(지표 지연)에는 넣지 않고 따로 남긴다(작업 스레드가 다 찼을 때 줄 선 시간)
    queue_wait_ms = flow.metadata.get("gw_file_queue_wait_ms")
    print("\n===== GATEWAY TRAFFIC METRICS =====")
    print(f"[METRICS] gateway_ms={latency_ms} round_trip_ms={round_trip_ms} "
          f"latency_sample={'yes' if latency_sample else 'no'}"
          + (f" file_queue_wait_ms={queue_wait_ms}" if queue_wait_ms is not None else ""))
    # 10/15: 옛 칸 뒤에 확장 셈(표본 수·검사 요청 수 등)도 붙여 남긴다(값은 수만, IP·식별자 없음). 옛 칸 값은 그대로(0.0 포함)
    print(json.dumps({**metrics, **counts}, ensure_ascii=False))


def _result_detector(result):
    for name in ("detector", "detector_name", "decided_by"):
        value = getattr(result, name, None)
        if value:
            return value
    return None


def _result_dashboard_metadata(result):
    names = (
        "severity", "match_status", "coverage", "similarity", "confidence",
        "doc_id", "doc_span", "input_span", "reason_code", "policy_version",
    )
    return {
        name: getattr(result, name)
        for name in names
        if getattr(result, name, None) is not None
    }


def _gateway_error_is_timeout(flow: http.HTTPFlow) -> bool:
    message = str(getattr(flow, "error", "") or "").lower()
    return "timeout" in message or "timed out" in message


def _gateway_error_is_client_cancel(flow: http.HTTPFlow) -> bool:
    """브라우저가 요청·답변 스트림을 취소하거나 끊은 것(2026-10-15, 확장 칸 client_cancelled_requests).
    mitmproxy 오류 문구로 가른다(_CLIENT_CANCEL_ERRORS). 문구 끝 공백·대소문자 차이는 보지 않는다(정확히 같은 문구만)."""
    error = getattr(flow, "error", None)
    message = getattr(error, "msg", None)
    if not isinstance(message, str):
        message = str(error or "")
    return message.strip() in _CLIENT_CANCEL_ERRORS


# ============================================================
# Claude 새 화면 rpc 요청 구조 기록 (2026-10-10)
# claude.ai 새 화면은 질문을 POST …/claudeai-rpc/…ConversationService/PerformAction 등으로 보내는데
# 본문 형식을 모른다(지금은 검사하지 않고 그대로 나간다). 형식을 정하기 위해 구조만 남긴다:
# Content-Type·본문 길이·JSON 모양(키 이름, 값의 종류, 문자열 길이, 배열 길이). 값은 남기지 않는다.
# 전달에는 영향이 없다(읽기만, 예외는 삼킨다). GATEWAY_LOG_CLAUDE_RPC_SHAPE=0이면 끈다(기본 켬).
# 10/11 실측: PerformAction은 application/proto(Connect unary, protobuf 89~1212바이트), StreamTimeline은
# application/connect+proto(5바이트 머리 프레임). JSON 모양으로는 non_json만 남아서, 이 둘은 gateway_proto로 걷고
# 잎마다 구조만 남긴다(아래 "protobuf 구조 기록").
# 10/11 저녁: 질문 칸(PerformAction 2.3)이 정해져 그 칸은 검사한다(아래 "Claude 새 화면 질문 검사"). 구조 기록은 그대로
# 두고(다른 동작 찾기용) 줄에 prompt_paths(검사 대상 칸 경로, 값 없음)를 더한다. 실제 검사 결과는 검사 뒤
# '[CLAUDE-RPC-PROMPT] inspected=N of M …' 한 줄(10/11 밤 검토 4: 구조 기록은 검사보다 먼저라 검사 결과가 아님).
# 10/11 밤 검토 3: unary PerformAction의 질문 칸(2.3)은 하위로 풀지 않고 {path 2.3, kind prompt, len, marker만} 한 항목.
# ============================================================

CLAUDE_RPC_SHAPE_PER_HOUR = 10          # 메서드마다 시간당 기록 수
CLAUDE_RPC_SHAPE_MAX_BODY = 1024 * 1024  # 이보다 큰 본문은 모양을 풀지 않는다
_CLAUDE_RPC_SHAPE_MAX_DEPTH = 8
_CLAUDE_RPC_SHAPE_MAX_NODES = 400
_CLAUDE_RPC_SHAPE_MAX_KEYS = 40
_CLAUDE_RPC_SHAPE_LIST_SAMPLES = 2
# 필드 이름 모양(10/11): lowerCamel(conversationUuid·sha256Hash)·snake_case(parent_message_uuid)만.
# 점·붙임표(파일 이름·이름), 밑줄 뒤 숫자·대문자 연속(msg_01AbCd 같은 번호)은 값으로 본다
_CLAUDE_RPC_FIELD_RE = re.compile(r"[a-z][a-z0-9]*(?:[A-Z][a-z0-9]+|_[a-z][a-z0-9]*)*")
_CLAUDE_RPC_FIELD_MAX = 40
_CLAUDE_RPC_SAFE_NAME_RE = re.compile(r"^[A-Za-z0-9_.]{1,100}$")
_claude_rpc_shape_counts: dict = {}
_claude_rpc_shape_window = [0.0]
_claude_rpc_shape_lock = threading.Lock()


def _claude_rpc_shape_enabled() -> bool:
    return (os.getenv("GATEWAY_LOG_CLAUDE_RPC_SHAPE") or "1").strip().lower() not in {"0", "false", "off", "no"}


def _claude_rpc_field_name(key) -> bool:
    return isinstance(key, str) and len(key) <= _CLAUDE_RPC_FIELD_MAX and bool(_CLAUDE_RPC_FIELD_RE.fullmatch(key))


def _json_shape(value, depth=0, budget=None):
    """값 없이 구조만: 객체 {키: 모양}, 배열 ["<list:N>", 앞 2개 모양], 문자열 "str:길이", 수 "num", 참거짓 "bool", null.

    키도 값일 수 있다(파일 이름·사람 이름·메시지 번호로 된 지도). 10/11: 키 하나라도 필드 이름 모양이 아니거나
    키가 _CLAUDE_RPC_SHAPE_MAX_KEYS개를 넘으면 지도로 보고 키 이름 없이 {"<map:키 수>": 첫 값의 모양}만 남긴다
    (같은 지도의 로마자 이름 키처럼 필드 모양을 한 키도 함께 숨긴다)."""
    if budget is None:
        budget = [_CLAUDE_RPC_SHAPE_MAX_NODES]
    if budget[0] <= 0:
        return "…"
    budget[0] -= 1
    if isinstance(value, dict):
        if depth >= _CLAUDE_RPC_SHAPE_MAX_DEPTH:
            return f"<object:{len(value)}>"
        if len(value) > _CLAUDE_RPC_SHAPE_MAX_KEYS or not all(_claude_rpc_field_name(key) for key in value):
            first = next(iter(value.values()))
            return {f"<map:{len(value)}>": _json_shape(first, depth + 1, budget)}
        return {key: _json_shape(child, depth + 1, budget) for key, child in value.items()}
    if isinstance(value, list):
        if depth >= _CLAUDE_RPC_SHAPE_MAX_DEPTH:
            return f"<list:{len(value)}>"
        return [f"<list:{len(value)}>"] + [_json_shape(child, depth + 1, budget)
                                           for child in value[:_CLAUDE_RPC_SHAPE_LIST_SAMPLES]]
    if isinstance(value, str):
        return f"str:{len(value)}"
    if isinstance(value, bool):
        return "bool"
    if isinstance(value, (int, float)):
        return "num"
    if value is None:
        return "null"
    return type(value).__name__


def _claude_rpc_body_shape(body: bytes):
    """(형식, 모양). JSON 본문 또는 Connect 봉투(1바이트 표시 + 4바이트 길이 + JSON)만 푼다. 그 밖은 모양 없음."""
    if len(body) > CLAUDE_RPC_SHAPE_MAX_BODY:
        return "too_large", None
    candidates = [("json", body)]
    if len(body) >= 5 and body[0] in (0, 1, 2):
        (length,) = struct.unpack(">I", body[1:5])
        candidates.append((f"envelope:flags={body[0]}:frame={length}", body[5:5 + length]))
    for kind, data in candidates:
        if data.lstrip()[:1] not in (b"{", b"["):
            continue
        try:
            return kind, _json_shape(json.loads(data.decode("utf-8")))
        except (ValueError, UnicodeDecodeError, RecursionError):
            continue
    return "non_json", None


# protobuf 구조 기록(10/11). 잎 하나 = {"path": 필드 번호 경로 "1.3.2", "wire", "len": 바이트 수, "kind", "flags": [참인 표시]}
# (프레임 본문이면 "frame": 번호). 값·지도 키·부분 문자열은 남기지 않는다(필드 번호는 스키마 정보라 남긴다).
# 한계(10/11 검토): 글자가 메시지로 잘못 읽히면 틀 바이트가 경로 번호·len으로 남는다(잘못 읽힌 단계마다 태그 글자 1자 + 길이 값.
#   예 '* 김철수 …' → 경로 끝 5·len 32 → '* '). 경로를 줄이면 질문 칸 찾기(marker_paths)가 틀려져 그대로 둔다
#   - 10/11 밤: 질문 칸으로 정해진 unary PerformAction 2.3은 예외. values_at이 준 자리(loc)와 그 하위는 풀지 않고
#     {"path": "2.3", "wire": 2, "len": 원본 바이트 수, "kind": "prompt"} 한 항목(표시는 marker만). 구조 걷기(엄격)가 못 읽은
#     본문이어도 values_at(넉넉)이 찾은 칸은 끝에 붙인다(표지 찾기). 다른 칸·다른 동작(15 등)·다른 메서드는 그대로
# 압축 해제 상한은 본문 합계(gateway_proto). 구조 기록은 leaves()(검사용 후보 없음)만 쓴다
# 표시(글자 잎만): has_whitespace · has_hangul · has_latin_word(앞뒤가 영문·숫자가 아닌 로마자 3자 이상) ·
#   digit_run_6plus(숫자 6자리 이상 연속) · has_at_sign · looks_uuid(값 전체가 UUID) ·
#   looks_base64(값 전체가 16자 이상 base64/base64url, 대문자·소문자·숫자 모두 있음) ·
#   marker(시험 표지 CLAUDE_RPC_MARKER가 들어 있음: 가짜 질문 'PGTEST 주민번호 900101-1234568'을 보내 질문 칸을 찾는다)
# 바이트 잎은 marker만. 시험 표지가 원본 바이트에 있으면 메서드 시간당 상한을 넘어도 기록한다(따로 시간당 CLAUDE_RPC_MARKER_PER_HOUR건)
# 줄 길이 상한 CLAUDE_RPC_LINE_MAX(잎 CLAUDE_RPC_PROTO_MAX_LEAVES개까지, 넘으면 반씩 줄이고 "truncated": true)
CLAUDE_RPC_MARKER = "PGTEST"
CLAUDE_RPC_MARKER_PER_HOUR = 20
CLAUDE_RPC_PROTO_MAX_LEAVES = 80
CLAUDE_RPC_LINE_MAX = 8000
_CLAUDE_RPC_PROTO_TYPES = {      # Content-Type → 프레임(봉투) 본문인지
    "application/proto": False,
    "application/x-protobuf": False,
    "application/protobuf": False,
    "application/connect+proto": True,
}
_CLAUDE_RPC_FLAG_TESTS = (
    ("has_whitespace", re.compile(r"\s").search),
    ("has_hangul", re.compile("[ᄀ-ᇿ㄰-㆏ꥠ-꥿가-힣ힰ-퟿]").search),
    ("has_latin_word", re.compile(r"(?<![A-Za-z0-9])[A-Za-z]{3,}(?![A-Za-z0-9])").search),
    ("digit_run_6plus", re.compile(r"[0-9]{6,}").search),
    ("has_at_sign", lambda text: "@" in text),
    ("looks_uuid", re.compile(r"[0-9A-Fa-f]{8}-[0-9A-Fa-f]{4}-[0-9A-Fa-f]{4}-[0-9A-Fa-f]{4}-[0-9A-Fa-f]{12}").fullmatch),
    ("looks_base64", lambda text: len(text) >= 16
        and re.fullmatch(r"[A-Za-z0-9+/]+={0,2}|[A-Za-z0-9_-]+={0,2}", text) is not None
        and re.search(r"[A-Z]", text) is not None and re.search(r"[a-z]", text) is not None
        and re.search(r"[0-9]", text) is not None),
    ("marker", lambda text: CLAUDE_RPC_MARKER in text),
)
_claude_rpc_marker_count = [0]


def _claude_rpc_leaf_flags(kind: str, value) -> list[str]:
    """참인 표시 이름만(값은 돌려주지 않는다)."""
    if kind == "bytes" and isinstance(value, bytes):
        return ["marker"] if CLAUDE_RPC_MARKER.encode() in value else []
    if kind != "text" or not isinstance(value, str):
        return []
    return [name for name, test in _CLAUDE_RPC_FLAG_TESTS if test(value)]


def _claude_rpc_prompt_entry(value: bytes) -> dict:
    """구조 기록의 질문 칸 한 항목: 경로·길이·marker만(질문 내용에서 나온 다른 표시·하위 경로 없음)."""
    entry = {"path": CLAUDE_RPC_PROMPT_PATH, "wire": gateway_proto.WIRE_LEN, "len": len(value), "kind": "prompt"}
    if CLAUDE_RPC_MARKER.encode() in value:
        entry["flags"] = ["marker"]
    return entry


def _claude_rpc_proto_structure(req, body: bytes, enveloped: bool, prompt=None) -> dict:
    """protobuf 본문의 구조 칸(format·leaves_total·marker_paths·prompt_paths·leaves). 값은 담지 않는다.
    prompt: unary PerformAction이면 _claude_rpc_prompt_parts 결과(질문 칸 자리·값), 아니면 None.
    prompt_paths: 질문으로 검사할 칸의 경로(값 없음. 칸마다 하나, 10개까지). 검사 결과가 아니다(결과는 검사 뒤 한 줄)."""
    fmt = "connect+proto" if enveloped else "proto"
    prompt = prompt if isinstance(prompt, dict) and prompt.get("target") else {}
    prompt_paths = [CLAUDE_RPC_PROMPT_PATH for _ in prompt.get("texts") or []][:10]
    prompt_values = {tuple(loc): value for loc, value in prompt.get("values") or []}
    depth = len(CLAUDE_RPC_PROMPT_PATH.split("."))
    if len(body) > CLAUDE_RPC_SHAPE_MAX_BODY:
        return {"format": "too_large", "prompt_paths": prompt_paths}
    if enveloped:
        data, encoding = body, req.headers.get("connect-content-encoding", "")
    else:
        # Connect unary 압축은 Content-Encoding. 상한 안에서 직접 푼다(mitmproxy 해제는 크기 상한이 없다)
        data, encoding = gateway_proto.decompress(body, req.headers.get("content-encoding", "")), None
        if data is None:
            return {"format": fmt, "leaves_total": 1, "marker_paths": [], "prompt_paths": prompt_paths,
                    "leaves": [{"path": "", "wire": None, "len": len(body), "kind": "unparsed",
                                "reason": "content_encoding"}]}
    leaves, markers, shown = [], [], set()

    def add_prompt(loc):
        shown.add(loc)
        entry = _claude_rpc_prompt_entry(prompt_values[loc])
        if "flags" in entry:
            markers.append(entry["path"])
        leaves.append(entry)

    for record, value in gateway_proto.leaves(data, enveloped=enveloped, encoding=encoding):
        owner = tuple(record["loc"][:depth]) if prompt_values and len(record["loc"]) >= depth else None
        if owner in prompt_values:          # 질문 칸(또는 그 하위로 읽힌 잎): 한 항목만
            if owner not in shown:
                add_prompt(owner)
            continue
        entry = {"path": record["path"], "wire": record["wire"], "len": record["length"], "kind": record["kind"]}
        if "frame" in record:
            entry["frame"] = record["frame"]
        if record.get("reason") in gateway_proto.REASONS:
            entry["reason"] = record["reason"]
        flags = _claude_rpc_leaf_flags(record["kind"], value)
        if flags:
            entry["flags"] = flags
        if "marker" in flags:
            markers.append(f"{record['frame']}:{record['path']}" if "frame" in record else record["path"])
        leaves.append(entry)
    for loc in sorted(prompt_values):       # 구조 걷기(엄격)가 못 읽은 본문에서 values_at(넉넉)이 찾은 질문 칸
        if loc not in shown:
            add_prompt(loc)
    return {"format": fmt, "leaves_total": len(leaves), "marker_paths": markers[:10], "prompt_paths": prompt_paths,
            "leaves": leaves[:CLAUDE_RPC_PROTO_MAX_LEAVES]}


def _claude_rpc_line(record: dict) -> str:
    """protobuf 구조 기록 한 줄. CLAUDE_RPC_LINE_MAX를 넘으면 잎을 반씩 줄인다(JSON은 늘 온전)."""
    line = json.dumps(record, ensure_ascii=True, separators=(",", ":"))
    leaves = record.get("leaves")
    if isinstance(leaves, list) and record.get("leaves_total", 0) > len(leaves):
        record = {**record, "truncated": True}
        line = json.dumps(record, ensure_ascii=True, separators=(",", ":"))
    while len(line) > CLAUDE_RPC_LINE_MAX and isinstance(leaves, list) and leaves:
        leaves = leaves[:len(leaves) // 2]
        record = {**record, "leaves": leaves, "truncated": True}
        line = json.dumps(record, ensure_ascii=True, separators=(",", ":"))
    return line


def _claude_rpc_rate_ok(method: str, marker_hint: bool) -> bool:
    """메서드마다 시간당 CLAUDE_RPC_SHAPE_PER_HOUR건(메서드 50종까지). 시험 표지가 보이면 따로 시간당 CLAUDE_RPC_MARKER_PER_HOUR건 더."""
    now = time.time()
    with _claude_rpc_shape_lock:
        if now - _claude_rpc_shape_window[0] >= 3600:
            _claude_rpc_shape_window[0] = now
            _claude_rpc_shape_counts.clear()
            _claude_rpc_marker_count[0] = 0
        count = _claude_rpc_shape_counts.get(method, 0)
        if (method in _claude_rpc_shape_counts or len(_claude_rpc_shape_counts) < 50) \
                and count < CLAUDE_RPC_SHAPE_PER_HOUR:
            _claude_rpc_shape_counts[method] = count + 1
            return True
        if marker_hint and _claude_rpc_marker_count[0] < CLAUDE_RPC_MARKER_PER_HOUR:
            _claude_rpc_marker_count[0] += 1
            return True
        return False


def _log_claude_rpc_shape(flow: http.HTTPFlow) -> None:
    """claude.ai …/claudeai-rpc/… POST의 구조만 기록한다(메서드마다 시간당 CLAUDE_RPC_SHAPE_PER_HOUR건). 값은 남기지 않는다.
    JSON(봉투 포함)은 키·값 종류 모양, protobuf(application/proto·connect+proto)는 잎마다 경로·종류·길이·표시."""
    try:
        req = flow.request
        if req.method.upper() != "POST" or not _claude_rpc_shape_enabled():
            return
        path = req.path.split("?", 1)[0]
        marker = "/claudeai-rpc"
        if marker not in path:
            return
        segments = [part for part in path.split(marker, 1)[1].split("/") if part]
        method = "/".join(part if _CLAUDE_RPC_SAFE_NAME_RE.fullmatch(part) else "<id>" for part in segments[-2:])
        body = req.raw_content or b""
        content_type = (req.headers.get("content-type", "") or "").split(";", 1)[0].strip()[:100]
        enveloped = _CLAUDE_RPC_PROTO_TYPES.get(content_type.lower())
        marker_hint = enveloped is not None and CLAUDE_RPC_MARKER.encode() in body
        if not _claude_rpc_rate_ok(method, marker_hint):
            return
        record = {
            "method": method or "<none>",
            "content_type": content_type,
            "content_encoding": (req.headers.get("content-encoding", "") or "")[:40],
            "bytes": len(body),
        }
        if enveloped is None:
            kind, shape = _claude_rpc_body_shape(body)
            record.update(format=kind, shape=shape)
            line = json.dumps(record, ensure_ascii=True, separators=(",", ":"))
        else:
            if enveloped:
                record["connect_content_encoding"] = (req.headers.get("connect-content-encoding", "") or "")[:40]
            record.update(_claude_rpc_proto_structure(req, body, enveloped, _claude_rpc_prompt_target(flow)))
            line = _claude_rpc_line(record)
        print("[CLAUDE-RPC] " + line)
    except Exception as exc:
        print("[CLAUDE-RPC] shape log failed error_type=", type(exc).__name__)


# ============================================================
# Claude 새 화면 질문 검사 (2026-10-11)
# 10/11 05:08 UTC 운영 실측(구조만): 새 화면은 질문을 POST …/claudeai-rpc/anthropic.bard.api.v1alpha.ConversationService/
# PerformAction, application/proto(Connect unary, protobuf 약 1 KB)로 보낸다. 가짜 질문 'PGTEST 주민번호 900101-1234568'이
# 잎 2.3(글자 42바이트) 하나에만 있었다. 같은 요청의 다른 잎: 1.1.1 글자(21)·1.2 UUID·1.14 글자(5)·1.15 바이트(32, 모름:
# 무결성 표지일 수 있음)·2.1/2.2 UUID·2.7.2 글자(15, 모델로 보임)·2.8.1.8.1.1 로마자 글자 12개(기능 목록으로 보임)·
# 2.12 글자(10)·2.13 글자(5)·숫자들. 같은 화면의 다른 PerformAction(976바이트)은 2 대신 15 아래에 비슷한 구조, 질문 칸 없음.
# 사용자 결정(10/11): 질문을 검사하고 개인정보면 막는다(이 경로는 아직 가리지 않음).
#   - 대상: claude.ai POST …ConversationService/PerformAction, Content-Type application/proto(Content-Encoding gzip·deflate는
#     gateway_proto로 상한 안에서 품). 잎 2.3의 원본 바이트(gateway_proto.values_at: 글자·메시지 판정 없이)를 UTF-8로 읽어
#     정규화한 글자. 2.3이 여러 개면 하나씩 검사(차단이 나오면 멈춤). 2.3이 없거나 비면 질문 전송이 아님 → 그대로 전달
#   - 본문을 읽지 못하면(압축 해제 실패·모르는 압축·protobuf 아님·4 MiB 초과·상한) 오늘처럼 전달 + 구조 기록, 기록 없음.
#     '[CLAUDE-RPC-PROMPT] not inspected reason=<낱말> bytes=N' 한 줄만 남긴다
#   - 10/11 밤 검토 1: values_at은 표준 protobuf 파서만큼 넉넉하게 읽는다(예약 번호·모르는 그룹·최소 길이가 아닌 varint).
#     동작 메시지(2)가 그래도 메시지가 아니면 사유(not_protobuf): 다른 2에서 찾은 질문은 검사하고 'partly unreadable' 한 줄,
#     찾은 질문이 없으면 위 'not inspected' 한 줄. PerformAction에서 못 읽음을 실패 닫힘(403·503)으로 바꿀지는 사용자 결정
#   - 10/11 밤 검토 2: 검사 중 예외(Detection 응답 깨짐 등)는 질문 처리 설정의 실패 판정(OTHER)으로 집행한다
#     (예외가 request 훅 밖으로 나가면 mitmproxy가 원래 요청을 보낸다)
#   - 10/11 밤 검토 4: 검사가 끝나면 '[CLAUDE-RPC-PROMPT] inspected=N of M degraded=D action=<gw_action> status=<403|503|forwarded>'
#     한 줄(값 없음). N = 실제 판정을 받은 칸(무응답·실패 판정 뺌), M = 검사 대상 칸, D = 무응답·실패 판정 칸
#   - 검사: 다른 질문과 같은 길(detect_web_prompt → _inspect 'prompt': 질문 대기 시간·처리 설정, service Claude,
#     content_type prompt, endpoint 경로). 글자 상한 없음(다른 질문과 같다)
#   - 판정: allow → 바이트 그대로 전달 / mask → 403 mask_application_failed(가림 기본 끔) / block → Detection 사유로 403 /
#     Detection 무응답 → 질문 처리 설정(원본 전송 = 전달 + degraded 기록, 차단 = 503)
#   - 차단 응답: Connect unary 오류(Claude 화면이 문구를 보여 준다). HTTP 상태(403·503)는 다른 차단과 같고, 본문
#     {"code": "permission_denied"|"unavailable", "message": 판정의 사용자 문구}, Content-Type application/json,
#     X-LLMSGW-Review-URL은 다른 차단과 같게
#   - 기록: 요청마다 이벤트 1건(가장 강한 판정, 질문 칸 여러 개가 모두 가림이면 유형별 건수 합침) + 확장 알림(기존 길).
#     차단은 같은 업로드 재시도처럼 기록 한 번(열쇠의 신원 = 정규화한 질문 글자 sha256. 창·보관은 같은 값)
#   - 가림(GATEWAY_CLAUDE_RPC_MASK=1일 때만, 기본 0): 가림 판정 칸만 replace_text로 바꾸고 다시 인코딩(같은 압축) + MASK 기록.
#     1.15(32바이트)가 본문 무결성 표지라면 Claude가 바뀐 본문을 거부할 수 있다 → 켜기 전에 실제 화면 시험 필요.
#     가림 근거(maskEvidence)는 싣지 않는다(대시보드가 JSON_TEXT·GEMINI_FORM_TEXT 표현만 받는다)
#   - application/connect+proto·JSON PerformAction, 2 말고 다른 동작 번호(15 등)의 글자는 아직 검사하지 않는다(구조 기록으로 찾음)
# ============================================================

CLAUDE_RPC_PROMPT_PATH = "2.3"
CLAUDE_RPC_MASK_BLOCK_MESSAGE = "개인정보가 들어 있어 보안 정책에 따라 전송을 차단했습니다. 개인정보를 지우고 다시 보내 주세요."
CLAUDE_RPC_BLOCK_MESSAGE = "보안 정책에 따라 외부 전송을 차단했습니다."
_CLAUDE_RPC_PERFORM_ACTION_RE = re.compile(r"/claudeai-rpc/(?:[A-Za-z0-9_]+\.)*ConversationService/PerformAction$")
_CONNECT_ERROR_CODES = {403: "permission_denied", 429: "resource_exhausted", 503: "unavailable", 504: "deadline_exceeded"}


def _claude_rpc_mask_enabled() -> bool:
    return (os.getenv("GATEWAY_CLAUDE_RPC_MASK") or "0").strip() == "1"


def _find_claude_rpc_prompt(flow: http.HTTPFlow) -> dict:
    """{"target": False} = 질문 검사 대상 요청이 아님. 대상이면 {"target": True, "reason": 읽지 못한 사유 또는 None,
    "partial": 일부 동작 메시지를 못 읽은 사유 또는 None(찾은 질문은 검사), "values": [(loc, 원본 바이트)],
    "texts": [(loc, 정규화한 글자)](빈 글자 뺌), "data": 압축 푼 본문, "encoding": 압축 방식}.
    예외를 내보내지 않는다(실패 = 읽지 못함)."""
    req = flow.request
    try:
        host = (req.host or "").lower()
        if req.method.upper() != "POST" or not (host == "claude.ai" or host.endswith(".claude.ai")):
            return {"target": False}
        if not _CLAUDE_RPC_PERFORM_ACTION_RE.search(req.path.split("?", 1)[0]):
            return {"target": False}
        content_type = (req.headers.get("content-type", "") or "").split(";", 1)[0].strip().lower()
        if _CLAUDE_RPC_PROTO_TYPES.get(content_type) is not False:    # unary protobuf만(connect+proto·JSON은 아님)
            return {"target": False}
    except Exception:
        return {"target": False}
    found = {"target": True, "reason": None, "partial": None, "values": [], "texts": [], "data": None, "encoding": ""}
    try:
        raw = req.raw_content or b""
        found["bytes"] = len(raw)
        encoding = (req.headers.get("content-encoding", "") or "").strip().lower()
        found["encoding"] = encoding
        data = gateway_proto.decompress(raw, encoding)
        if data is None:
            found["reason"] = "content_encoding"
            return found
        values, reason = gateway_proto.values_at(data, CLAUDE_RPC_PROMPT_PATH)
        if reason and not values:
            found["reason"] = reason
            return found
        found["partial"] = reason            # 다른 동작 메시지(2)를 못 읽음: 찾은 질문은 그대로 검사
        found["data"], found["values"] = data, values
        for loc, value in values:
            text = normalize_text(value.decode("utf-8", errors="replace"))
            if text:
                found["texts"].append((loc, text))
    except Exception as exc:
        print("[CLAUDE-RPC-PROMPT] read failed error_type=", type(exc).__name__)
        found.update(reason="error", partial=None, values=[], texts=[], data=None)
    return found


def _claude_rpc_prompt_parts(flow: http.HTTPFlow) -> dict:
    """요청마다 한 번만 읽는다(구조 기록과 검사가 같이 쓴다). 검사가 끝나면 지운다."""
    found = flow.metadata.get("gw_claude_rpc_prompt")
    if found is None:
        found = _find_claude_rpc_prompt(flow)
        flow.metadata["gw_claude_rpc_prompt"] = found
    return found


def _claude_rpc_prompt_target(flow: http.HTTPFlow):
    """구조 기록용: 질문 검사 대상 요청이면 읽은 질문 칸(_claude_rpc_prompt_parts), 아니면 None.
    줄의 prompt_paths(검사 대상 칸)와 질문 칸 한 항목 묶기에 쓴다. 검사 결과는 검사 뒤 따로 한 줄."""
    try:
        found = _claude_rpc_prompt_parts(flow)
    except Exception:
        return None
    if not found.get("target") or found.get("reason"):
        return None
    return found


def _claude_rpc_prompt_identity(texts) -> str:
    """재시도 기록 한 번의 신원: 정규화한 질문 글자들(순서대로, 길이 붙여) sha256. 메모리 안에서만 쓴다."""
    digest = hashlib.sha256()
    for _, text in texts:
        data = text.encode("utf-8", "surrogatepass")
        digest.update(len(data).to_bytes(8, "big") + data)
    return "prompt:" + digest.hexdigest()


def _apply_claude_rpc_mask(flow: http.HTTPFlow, found: dict, checked) -> bool:
    """가림 판정 칸(2.3)만 가린 글자로 바꾸고 같은 압축으로 다시 싣는다. 바꾼 본문을 다시 읽어 확인한 뒤에만 True.
    못 하면 원래 본문 그대로 두고 False(→ 차단)."""
    data = found.get("data")
    if data is None:
        return False
    expected = {}
    for loc, text, result in checked:
        if _claude_action(result) not in _FILE_MASK_ACTIONS:
            continue
        replacement = getattr(result, "masked_text", None)
        if replacement is None:
            replacement = _masked_text(text, getattr(result, "mask_spans", None) or [])
        if not isinstance(replacement, str):
            return False
        rebuilt = gateway_proto.replace_text(data, loc, replacement)
        if rebuilt is None:
            return False
        data = rebuilt
        expected[loc] = replacement.encode("utf-8")
    if not expected:
        return False
    values, reason = gateway_proto.values_at(data, CLAUDE_RPC_PROMPT_PATH)
    if reason or [loc for loc, _ in values] != [loc for loc, _ in found["values"]]:
        return False
    if any(loc in expected and value != expected[loc] for loc, value in values):
        return False
    req = flow.request
    original_raw, original_length = req.raw_content, req.headers.get("content-length")
    req.content = data     # mitmproxy가 Content-Encoding(gzip·deflate)대로 다시 압축하고 content-length를 고친다
    if gateway_proto.decompress(req.raw_content or b"", found.get("encoding")) != data:
        req.raw_content = original_raw
        if original_length is not None:
            req.headers["content-length"] = original_length
        return False
    return True


def _claude_rpc_connect_error(flow: http.HTTPFlow) -> None:
    """enforce_detection_result가 만든 차단 응답(JSON error)을 Connect unary 오류로 바꾼다(Claude 화면이 message를 보여 준다).
    상태 코드·X-LLMSGW-Review-URL·기록(이미 남김)은 그대로."""
    response = flow.response
    message = None
    try:
        error = json.loads(response.get_text(strict=False) or "{}").get("error")
        if isinstance(error, dict) and isinstance(error.get("message"), str) and error["message"].strip():
            message = error["message"]
    except (ValueError, AttributeError, TypeError):
        pass
    status = response.status_code
    headers = {"Content-Type": "application/json"}
    review_url = response.headers.get("X-LLMSGW-Review-URL")
    if review_url:
        headers["X-LLMSGW-Review-URL"] = review_url
    body = {"code": _CONNECT_ERROR_CODES.get(status, "permission_denied"), "message": message or CLAUDE_RPC_BLOCK_MESSAGE}
    flow.response = http.Response.make(status, json.dumps(body, ensure_ascii=False).encode("utf-8"), headers)


def _claude_rpc_inspection_failure(flow: http.HTTPFlow):
    """검사 중 예외의 판정: Detection 실패(OTHER)와 같고 질문 규칙의 실패 처리를 따른다(설정을 못 읽으면 차단 쪽)."""
    try:
        policy = _flow_settings(flow).timeout_for("prompt").failure_policy
    except Exception:
        policy = "FAIL_CLOSED"
    return failure_result("OTHER", policy if policy in ("FAIL_OPEN", "FAIL_CLOSED") else "FAIL_CLOSED", str(uuid.uuid4()))


def inspect_claude_rpc_prompt(flow: http.HTTPFlow, identity: dict) -> bool:
    """claude.ai 새 화면 질문(PerformAction 2.3) 검사·집행(위 설명). True = 질문을 검사함(차단이면 flow.response가 있음).
    False = 대상이 아니거나 질문 칸이 없거나 본문을 읽지 못함 → 오늘처럼 그대로."""
    found = _claude_rpc_prompt_parts(flow)
    try:
        if not found.get("target"):
            return False
        if found.get("reason"):
            print(f"[CLAUDE-RPC-PROMPT] not inspected reason={found['reason']} bytes={found.get('bytes', 0)}")
            return False
        if found.get("partial"):
            print(f"[CLAUDE-RPC-PROMPT] partly unreadable reason={found['partial']} bytes={found.get('bytes', 0)}")
        texts = found["texts"]
        if not texts:
            return False
        print("\n===== CLAUDE RPC PROMPT DETECTED =====")
        print("Prompt fields:", len(texts), "path:", CLAUDE_RPC_PROMPT_PATH)
        print("Endpoint:", flow.request.path.split("?", 1)[0])

        checked = []    # [(loc, 글자, 판정)]
        for loc, text in texts:
            try:
                result = detect_web_prompt(flow, "Claude", text)
            except Exception as exc:
                # 검사 중 예외(10/11 밤 검토 2): request 훅 밖으로 나가면 mitmproxy가 원래 요청을 보낸다 →
                # Detection 실패와 같이 질문 처리 설정(차단 = 503, 원본 전송 = 전달 + degraded 기록)
                print("[CLAUDE-RPC-PROMPT] inspection failed error_type=", type(exc).__name__)
                result = _claude_rpc_inspection_failure(flow)
            checked.append((loc, text, result))
            if _claude_action(result) in _FILE_BLOCK_ACTIONS:
                break     # 요청 전체를 막는다. 남은 칸은 보내지 않는다
        results = [result for _, _, result in checked]
        context = {"identity": identity, "targetLlm": "Claude", "contentType": "prompt"}
        action = _request_level_action(_strongest_result(results))
        if action == "block":
            context["fileIdentity"] = _claude_rpc_prompt_identity(texts)
            enforce_detection_result(flow, _strongest_result(results), context)
        elif action == "mask":
            report = _attachment_report_result(results)
            applied = False
            if _claude_rpc_mask_enabled():
                try:
                    applied = _apply_claude_rpc_mask(flow, found, checked)
                except Exception as exc:
                    print("MASK application failed:", type(exc).__name__)
                    applied = False
            if applied:
                print("[CLAUDE-RPC-PROMPT] masked fields=", sum(
                    _claude_action(result) in _FILE_MASK_ACTIONS for result in results))
                enforce_detection_result(flow, report, context, mask_callback=lambda _: True)
            else:
                # 이 형식은 아직 가리지 않는다(또는 가림을 확인하지 못함): 원본을 보내지 않고 막는다
                report.action, report.blocked, report.reason_code = "block", True, "mask_application_failed"
                report.user_message, report.http_status = CLAUDE_RPC_MASK_BLOCK_MESSAGE, 403
                context["fileIdentity"] = _claude_rpc_prompt_identity(texts)
                enforce_detection_result(flow, report, context)
        else:
            enforce_detection_result(flow, _strongest_result(results), context)
        flow.metadata["gw_degraded"] = any(bool(getattr(result, "degraded", False)) for result in results)
        if flow.response is not None:
            _claude_rpc_connect_error(flow)
        degraded = sum(bool(getattr(result, "degraded", False)) for result in results)
        print(f"[CLAUDE-RPC-PROMPT] inspected={len(results) - degraded} of {len(texts)} degraded={degraded} "
              f"action={flow.metadata.get('gw_action') or 'allow'} "
              f"status={flow.response.status_code if flow.response is not None else 'forwarded'}")
        return True
    finally:
        flow.metadata.pop("gw_claude_rpc_prompt", None)


# ============================================================
# 파일 업로드 작업 스레드 (2026-10-11, T0 A안)
# 이전에는 request 훅이 동기라 업로드 하나를 검사하는 동안(추출·칸 검사·Detection 대기) 이벤트 루프가 통째로 멈춰
# 다른 사용자의 질문이 그만큼 기다렸다(T0 로컬 실측: XLSX 1,448칸·DOCX 1,000칸 업로드 중 질문 p95 +15~25초,
# 새 연결 12~50% 실패). 이제 파일 업로드 요청만 전용 작업 스레드에서 지금 처리(_request_sync)를 그대로 돌리고
# 질문·그 밖 요청은 지금처럼 루프에서 바로 처리한다(T0 시제품: 업로드 중 질문 p95 +0~3ms).
# - 전용 묶음(GATEWAY_FILE_WORKERS개)을 쓴다. asyncio.to_thread·mitmproxy @concurrent가 쓰는 기본 스레드 묶음은
#   mitmproxy가 upstream 주소 찾기에도 써서 업로드가 그것을 채우면 새 연결까지 막힐 수 있다
# - 작업 스레드 수 = 동시에 검사하는 업로드 수 상한. 넘는 업로드는 줄을 서고(그동안에도 질문은 막히지 않음),
#   줄 선 시간은 gw_file_queue_wait_ms로 따로 남긴다(지표 지연 gw_processing_ms·processingMs에는 넣지 않음)
# - 작업 스레드 안의 예상 못한 예외도 원본을 보내지 않는다(503 + 대시보드 BLOCK 기록). 루프(질문)는 지금 그대로
# - 공유 상태 잠금: _upload_state_lock(파일 위쪽), PDFium: _pdfium_lock(_classify_upload)
# ============================================================

GATEWAY_FILE_WORKERS = 4                 # T0: 업로드 4건 동시(PDF 2·XLSX·DOCX)에도 질문 p95 그대로. 운영 EC2는 2 vCPU
FILE_OFFLOAD_MIN_BYTES = 256 * 1024      # 업로드 경로인데 위 갈래에 안 맞는 큰 본문(모르는 업로드 모양)도 작업 스레드로
_GEMINI_UPLOAD_HOST = "push.clients6.google.com"
_file_executor: ThreadPoolExecutor | None = None   # 처음 쓸 때 만든다. done()이 닫으면 다음 요청 때 다시 만든다
_file_executor_lock = threading.Lock()


def _submit_file_work(fn, *args):
    """전용 작업 스레드 묶음에 넘긴다. 만들기·넘기기를 한 잠금 안에서 해 done()이 막 닫은 묶음에 넘기지 않는다."""
    global _file_executor
    with _file_executor_lock:
        if _file_executor is None:
            _file_executor = ThreadPoolExecutor(max_workers=GATEWAY_FILE_WORKERS, thread_name_prefix="gw-file")
        return _file_executor.submit(fn, *args)


def _shutdown_file_executor() -> None:
    """done()에서: 묶음을 닫는다. 이미 받은 업로드는 끝까지 처리한다(취소하면 그 요청이 응답 없이 멈춘다).
    시험·스크립트 다시 읽기처럼 같은 모듈을 계속 쓰면 다음 업로드 때 새 묶음을 만든다."""
    global _file_executor
    with _file_executor_lock:
        pool, _file_executor = _file_executor, None
    if pool is not None:
        pool.shutdown(wait=False)


def _is_llm_upload_path(host: str, path: str) -> bool:
    """LLM 서비스의 업로드 경로(큰 본문 기준에만 쓴다). host 소문자, path는 '?' 앞."""
    service = service_for(host)
    return (host == _GEMINI_UPLOAD_HOST
            or (host.endswith("oaiusercontent.com") and path.startswith("/files/"))
            or (service == "ChatGPT" and path.startswith("/backend-api/files"))
            or (host == OPENAI_HOST and path.startswith(("/v1/files", "/v1/uploads")))
            # Claude 업로드: /api/{org}/upload, …/wiggle/upload-file, …/convert_document
            or (service == "Claude" and ("/upload" in path or "convert_document" in path)))


def is_file_upload_request(flow: http.HTTPFlow) -> bool:
    """작업 스레드로 보낼 파일 업로드 요청인지. _handle_request의 업로드 갈래(Gemini resumable·ChatGPT 원본 PUT·
    ChatGPT multipart·Claude multipart·OpenAI /v1/files)와 같은 조건이되 대소문자를 너그럽게 본다(더 넓게 = 작업 스레드로
    더 보낸다. 질문 경로는 들어오지 않는다) + 업로드 경로로 가는 FILE_OFFLOAD_MIN_BYTES 이상 본문."""
    req = flow.request
    host = (req.host or "").lower()
    method = (req.method or "").upper()
    path = req.path.split("?", 1)[0]
    multipart = req.headers.get("content-type", "").lower().startswith("multipart/form-data")
    if host == _GEMINI_UPLOAD_HOST and method == "POST" and path == "/upload/":
        return True                                                    # Gemini resumable(시작·파일 바이트)
    if method == "PUT" and host.endswith("oaiusercontent.com") and path.startswith("/files/") and path.endswith("/raw"):
        return True                                                    # ChatGPT 원본 PUT
    if method == "POST" and multipart and (
            (service_for(host) == "ChatGPT" and path.startswith("/backend-api/files/"))   # ChatGPT multipart
            or host == "claude.ai"                                                      # Claude multipart(경로 무관)
            or (host == OPENAI_HOST and path == OPENAI_FILES_ENDPOINT)):                # OpenAI /v1/files
        return True
    return _is_llm_upload_path(host, path) and len(req.raw_content or b"") >= FILE_OFFLOAD_MIN_BYTES


def _file_target_llm(flow: http.HTTPFlow) -> str:
    host = (flow.request.host or "").lower()
    service = service_for(host) or ""
    if host == _GEMINI_UPLOAD_HOST or service == "Gemini":
        return "Gemini"
    return "Claude" if service.startswith("Claude") else "ChatGPT"   # OpenAI API도 ChatGPT(지금 OpenAI 경로 기록과 같음)


def _fail_closed_file_request(flow: http.HTTPFlow, exc: BaseException) -> None:
    """파일 업로드 처리 중 예상 못한 예외(작업 스레드 안, 또는 작업 스레드에 넘기지 못함): 원본을 보내지 않는다.
    응답을 아직 정하지 않았으면 503 file_masking_failed + 대시보드 BLOCK 기록(fail_closed_file_upload: 응답을 먼저 정하고
    기록 실패는 삼킨다). 이미 차단·오류 응답을 정했으면 그대로 둔다(그 응답이 나가고 원본은 나가지 않는다)."""
    print("[FILE-WORKER] upload processing failed; original not forwarded error_type=", type(exc).__name__)
    if flow.response is not None:
        return
    try:
        context = {"identity": flow.metadata.get("llmsgw_identity") or {}, "targetLlm": _file_target_llm(flow),
                   "contentType": "file"}
        fail_closed_file_upload(flow, context)
    except Exception as record_exc:
        print("[FILE-WORKER] failure record omitted error_type=", type(record_exc).__name__)
    if flow.response is None:   # 기록 단계 예외로 응답을 못 정했어도 원본은 보내지 않는다
        flow.response = http.Response.make(503, b'{"error":{"type":"file_masking_failed"}}',
                                           {"Content-Type": "application/json"})


def _run_file_request(flow: http.HTTPFlow) -> None:
    """작업 스레드 본체: 줄 선 시간을 따로 남기고(지표 지연에는 넣지 않음) 지금 처리를 실패 시 차단으로 돌린다."""
    queued = flow.metadata.pop("gw_file_queued_at", None)
    if isinstance(queued, float):
        wait_ms = round((time.perf_counter() - queued) * 1000, 2)
        flow.metadata["gw_file_queue_wait_ms"] = wait_ms
        print(f"[FILE-WORKER] start queue_wait_ms={wait_ms} thread={threading.current_thread().name}")
    _request_sync(flow, fail_closed=True)


# ============================================================
# 접속 차단 (2026-10-12)
# 관리자가 대시보드에서 막은 확장 프로그램 ID(x-llmsgw-client-id), '확장 ID 없는 접속 차단'(기본 꺼짐)을 요청 맨 앞에서 집행한다.
#   - 예외: llm-gateway.local(health·접속 상태·내 기록 중계·검토 파일, 실제 목적지 기준), 접속 IP 127.0.0.0/8·::1
#   - 차단 안내(10/12 사유 안내): 403 페이지·JSON에 사유 이름표·user_memo(관리자가 보이기로 한 경우만)
#   - 접속 상태 GET http://llm-gateway.local/__llmsgw/access-status(_answer_access_status): 확장 프로그램이 묻는 이 브라우저의
#     차단 여부·사유. 차단 중에도 답한다. 확장 ID는 로그에 남기지 않는다
#   - 확장 ID: 식별 머리(지우기 전 값). 확장 프로그램은 최상위 페이지 요청(main_frame)에는 머리를 붙이지 않으므로
#     머리가 없으면 같은 연결의 앞 요청에서 본 ID를 쓴다(_access_connection_ids, 연결이 끊기면 지움)
#   - '확장 ID 없는 접속 차단'은 최상위 페이지(GET·HEAD)·CORS 사전 확인에는 쓰지 않는다(머리를 붙일 수 없는 요청, 내용 없음)
#   - 막은 요청: Detection·본문 해석·파일 작업 스레드 없음, Gateway 지표(요청 수·평균·p95)에 넣지 않음(denied_5m으로 따로)
#   - 보안 이벤트: (규칙 번호 또는 missing, 확장 ID 또는 IP)마다 10분에 1건. 확장 프로그램 알림은 그 이벤트와 함께(기존 중계)
#     10/16: 검사 대상 요청(질문·파일 업로드, _access_request_kind)을 막았을 때만 남기고 입력 유형도 그 종류대로.
#     배경 요청 차단은 이벤트 없이 denied_5m('최근 5분 막은 요청')·[ACCESS] 로그(event=skipped kind=other)로만 센다
#   - 판단 중 예외(코드 결함)는 막지 않고 통과 + 경고 로그(분당 1번)
# ============================================================

_access_connection_ids = ConnectionClientIds()
_access_failure_logged_at = 0.0


def _access_target_llm(flow: http.HTTPFlow) -> str | None:
    """대시보드 이벤트의 targetLlm. LLM 밖 목적지는 None(이벤트 없이 막기만 한다)."""
    host = (flow.request.host or "").lower()
    if host == _GEMINI_UPLOAD_HOST:
        return "Gemini"
    if host == OPENAI_HOST:
        return "ChatGPT"
    service = service_for(host) or ""
    if service.startswith("ChatGPT"):
        return "ChatGPT"
    if service.startswith("Claude"):
        return "Claude"
    return "Gemini" if service == "Gemini" else None


def _enforce_access(flow: http.HTTPFlow) -> bool:
    """True = 막음(403 응답을 정함). 판단 중 예외는 통과(fail-open)."""
    global _access_failure_logged_at
    try:
        req = flow.request
        connection = getattr(flow.client_conn, "id", None)
        header_id = normalize_client_id(req.headers.get(ACCESS_CLIENT_ID_HEADER))
        if header_id:   # 목록이 없을 때도 적어 둔다(차단이 걸린 뒤 같은 연결의 화면 이동에 쓴다)
            _access_connection_ids.remember(connection, header_id)
        access = GATEWAY_ACCESS.current()
        if not access.enforcing:
            return False
        if (req.host or "").lower() in ACCESS_EXEMPT_HOSTS or access_exempt_peer(get_client_ip(flow)):
            return False
        decision = access_decide(
            access, header_id or _access_connection_ids.get(connection),
            navigation=top_level_navigation(req.method, req.headers),
            preflight=cors_preflight(req.method, req.headers),
            now=GATEWAY_ACCESS.now(),
        )
    except Exception as exc:
        now = time.monotonic()
        if now - _access_failure_logged_at >= 60:
            _access_failure_logged_at = now
            print(f"[ACCESS] WARNING check failed; request passed error_type={type(exc).__name__}")
        return False
    if decision is None:
        return False
    _deny_access(flow, decision)
    return True


def _deny_access(flow: http.HTTPFlow, decision) -> None:
    request_id = str(uuid.uuid4())
    navigate = str(flow.request.headers.get("sec-fetch-mode", "")).lower() == "navigate"
    try:
        flow.response = access_block_response(decision, request_id, navigate=navigate)
    except Exception as exc:
        print(f"[ACCESS] block page failed error_type={type(exc).__name__}")
        flow.response = http.Response.make(403, json.dumps({"error": {
            "type": "access_blocked", "reason_code": decision.reason_code, "request_id": request_id}}).encode("utf-8"),
            {"Content-Type": "application/json; charset=utf-8", "Cache-Control": "no-store"})
    flow.metadata["gw_request_id"] = request_id
    flow.metadata["gw_access_blocked"] = decision.reason_code
    GATEWAY_ACCESS.note_denied()
    try:
        client_ip = normalize_ip(get_client_ip(flow))
        if decision.client_id:
            _client_tracker.note_client_id(client_ip, decision.client_id, _flow_conn_id(flow))   # 연결 IP 표에서 '차단 중'을 볼 수 있게
        # 10/16: 보안 이벤트는 검사 대상 요청(질문·파일 업로드)을 막았을 때만. 배경 요청(화면 로딩·폴링·원격 측정·
        # 웹소켓·Gemini batchexecute·Codex 등)은 denied_5m·이 로그 줄로만 센다. 10분 대표 기록 칸도 검사 대상 요청만 쓴다
        # (배경 요청이 칸을 먼저 차지해 뒤따른 질문 차단 기록이 빠지지 않게)
        kind = _access_request_kind(flow)
        target = _access_target_llm(flow) if kind else None
        key = (decision.rule.id if decision.rule else "missing", decision.client_id or client_ip or "-")
        record = bool(target) and GATEWAY_ACCESS.allow_event(key)
        outcome = "recorded" if record else "suppressed" if target else "skipped"
        print(f"[ACCESS] denied reason={decision.reason_code} rule={decision.rule.number if decision.rule else '-'} "
              f"navigate={'yes' if navigate else 'no'} event={outcome} kind={kind or 'other'} "
              f"requestId={request_id}")
        if record:
            _report_access_event(flow, decision, request_id, target=target, content_type=kind)
    except Exception as exc:
        print(f"[ACCESS] event omitted error_type={type(exc).__name__}")


def _answer_access_status(flow: http.HTTPFlow) -> None:
    """GET /__llmsgw/access-status: 이 확장 ID가 지금 막히는지·사유(gateway_access.access_status).
    확장 ID = 확장 프로그램이 직접 붙인 X-LLMSGW-Client-Id(식별 머리와 같은 이름) → 없거나 틀리면 같은 연결에서 본 ID.
    판단 실패는 503(확장 프로그램은 '확인 불가'로 보인다, 차단으로 보이지 않게). 요청마다 로그를 남기지 않는다."""
    req = flow.request
    if req.method != "GET":
        flow.response = access_status_response({"error": {"code": "METHOD_NOT_ALLOWED"}}, 405, {"Allow": "GET"})
        return
    try:
        client_id = (normalize_client_id(req.headers.get(ACCESS_CLIENT_ID_HEADER))
                     or _access_connection_ids.get(getattr(flow.client_conn, "id", None)))
        payload = access_status(GATEWAY_ACCESS.current(), client_id,
                                exempt=access_exempt_peer(get_client_ip(flow)), now=GATEWAY_ACCESS.now())
        flow.response = access_status_response(payload)
    except Exception as exc:
        print(f"[ACCESS] status failed error_type={type(exc).__name__}")
        flow.response = access_status_response({"error": {"code": "STATUS_UNAVAILABLE"}}, 503)


def _access_request_kind(flow: http.HTTPFlow) -> str | None:
    """접속 차단 이벤트를 남길 요청인지와 그 입력 유형(2026-10-16). 막지 않았다면 Gateway가 검사했을 요청만:
      'file'   = 파일 업로드(is_file_upload_request와 같은 조건) 또는 ChatGPT 파일 예약(업로드의 첫 요청, 막히면 바이트 PUT이 없다)
      'prompt' = 질문 전송: ChatGPT /backend-api/f/conversation, Claude completion·새 화면 PerformAction 질문 칸,
                 Gemini StreamGenerate, OpenAI /v1/responses (모두 POST, _handle_request의 질문 갈래와 같은 경로)
      None     = 그 밖(배경 요청). 이벤트 없이 denied_5m·[ACCESS] 로그로만 센다
    Claude PerformAction은 같은 경로에 질문이 아닌 동작도 있어 질문 칸이 있는지만 본다(_find_claude_rpc_prompt, 예외 없음).
    예외는 None(이벤트 없음, 막는 것은 그대로)."""
    try:
        req = flow.request
        if is_file_upload_request(flow) or _is_chatgpt_file_reservation(flow):
            return "file"
        if (req.method or "").upper() != "POST":
            return None
        host = (req.host or "").lower()
        path = _url_path(req.path)
        service = service_for(host)
        if service == "ChatGPT" and path == "/backend-api/f/conversation":
            return "prompt"
        if host == OPENAI_HOST and path == OPENAI_RESPONSES_ENDPOINT:
            return "prompt"
        if service == "Gemini" and "StreamGenerate" in path:
            return "prompt"
        if host == "claude.ai" and _CLAUDE_COMPLETION_PATH_RE.match(path):
            return "prompt"
        if _CLAUDE_RPC_PERFORM_ACTION_RE.search(path):
            found = _find_claude_rpc_prompt(flow)
            if found.get("target") and found.get("texts"):
                return "prompt"
    except Exception as exc:
        print(f"[ACCESS] request kind check failed error_type={type(exc).__name__}")
    return None


_CLAUDE_COMPLETION_PATH_RE = re.compile(r"^/api/organizations/[^/]+/chat_conversations/[^/]+/completion$")


def _report_access_event(flow: http.HTTPFlow, decision, request_id: str, *, target: str, content_type: str) -> None:
    """기존 대시보드 이벤트 경로(report_dashboard_event → response 훅에서 보관함·확장 프로그램 중계).
    target·content_type은 _deny_access가 정한 값(검사 대상 요청만 여기 온다, 10/16)."""
    identity = read_and_strip_identity(flow)   # 이 요청은 나가지 않는다. 열람 키 지문·tenant를 이벤트에 쓴다
    identity.pop("client_id", None)
    if decision.client_id:
        identity["client_id"] = decision.client_id
    result = AccessBlockResult(
        request_id=request_id, reason_code=decision.reason_code,
        access_rule_id=decision.rule.id if decision.reason_code == ACCESS_REASON_BLOCKED and decision.rule else None,
    )
    context = {"identity": identity, "targetLlm": target, "contentType": content_type}
    report_dashboard_event(flow, result, context, enforcement_applied=True, http_status=403)


# ============================================================
# mitmproxy Request Hook
# ============================================================

async def request(flow: http.HTTPFlow):
    """mitmproxy request 훅(2026-10-11부터 async, 운영 mitmproxy 12.2.3은 async 훅을 기다린다).
    파일 업로드 요청(is_file_upload_request)은 전용 작업 스레드에서 지금 처리를 돌리고 끝날 때까지 기다린다.
    그동안 이벤트 루프는 다른 요청을 처리한다. 그 밖 요청(질문 등)은 지금처럼 루프에서 바로 처리한다.
    직접 부르는 시험은 코루틴을 끝까지 돌려야 한다(tests/gw_support.run_hook).
    2026-10-12: 맨 앞에서 접속 차단을 본다(본문 해석·Detection·파일 작업 스레드 전). 막으면 403으로 끝낸다."""
    if _enforce_access(flow):
        return
    if not is_file_upload_request(flow):
        _request_sync(flow)
        return
    flow.metadata["gw_file_queued_at"] = time.perf_counter()
    try:
        await asyncio.wrap_future(_submit_file_work(_run_file_request, flow), loop=asyncio.get_running_loop())
    except Exception as exc:
        # 작업 스레드에 넘기지 못함(종료 중 등) 또는 작업 스레드 안의 실패 처리까지 실패: 원본을 보내지 않는다
        _fail_closed_file_request(flow, exc)


def _request_sync(
    flow: http.HTTPFlow,
    *,
    fail_closed: bool = False,
):
    """2026-10-10: 처리 본체(_handle_request)를 감싸 Gateway 자체 처리 시간을 잰다.
    처리 끝 = upstream으로 전달하기 직전 또는 Gateway가 응답(차단·오류)을 정한 때. LLM 응답 시간은 빠진다.
    보류한 '검사하지 않은 파일' 허용 기록도 여기서 정리한다(실제로 전달될 때만 남김).
    2026-10-11: fail_closed=True(파일 업로드 작업 스레드)면 처리 중 예외도 원본을 보내지 않는다(503 + 기록, 보류한
    허용 기록은 버림). 루프에서 도는 요청(질문 등)은 지금처럼 예외를 mitmproxy에 넘긴다."""
    try:
        try:
            _handle_request(flow)
        except Exception as exc:
            if not fail_closed:
                raise
            _fail_closed_file_request(flow, exc)
    finally:
        try:
            _flush_pending_uninspected(flow)
        except Exception as exc:
            print("[DASHBOARD] uninspected file event omitted error_type=", type(exc).__name__)
        started = flow.metadata.get("gw_metrics_started_at")
        if started is not None and "gw_processing_ms" not in flow.metadata:
            flow.metadata["gw_processing_ms"] = round((time.perf_counter() - started) * 1000, 2)
            flow.metadata["gw_latency_sample"] = _gateway_did_work(flow)
        _count_inspected_request(flow)


def _count_inspected_request(flow: http.HTTPFlow) -> None:
    """연결 중인 IP 표 '최근 5분 검사 요청'·LLM별 수(2026-10-15). request() 처리가 끝날 때(_request_sync) 요청 하나에 한 번.

    검사 요청 = Detection을 부른 요청(gw_inspected: 질문·파일 칸 검사, 대기 초과·연결 불가로 응답이 없었던 것·글자 수 초과로
    건너뛴 것 포함) 또는 Gateway가 판정을 정한 요청(gw_action: 차단·마스킹·허용, 검사 못 한 파일의 차단·허용, 파일 처리 실패 차단,
    Gemini 업로드 검사). 한 요청 안에서 Detection을 여러 번 불러도(질문 + 첨부, 파일 칸 여러 개) 1건.
    세지 않는 것: 판정 없이 그냥 통과한 요청(페이지 로딩·자산·폴링·상태 확인), llm-gateway.local(상태 확인·중계),
    접속 차단으로 막은 요청(request() 맨 앞에서 끝나 여기까지 오지 않음, denied_5m으로 따로 셈).
    같은 업로드를 ChatGPT가 다시 보내면(차단 뒤 재시도) 시도마다 검사하므로 시도마다 1건(대시보드 이벤트는 한 번만 남는다).
    실패해도 요청 처리에는 영향 없음(값은 남기지 않는다)."""
    metadata = flow.metadata
    if metadata.get("gw_inspected_counted") or not (metadata.get("gw_inspected") or metadata.get("gw_action")):
        return
    metadata["gw_inspected_counted"] = True
    try:
        _client_tracker.request_inspected(get_client_ip(flow), _inspected_target_llm(flow),
                                          _flow_client_id(flow), _flow_conn_id(flow))
    except Exception as exc:
        print(f"[METRICS] inspected request count omitted error_type={type(exc).__name__}")


def _inspected_target_llm(flow: http.HTTPFlow) -> str:
    """검사 요청의 LLM 이름: 보안 이벤트 targetLlm과 같게(ChatGPT·Claude·Gemini, OpenAI API는 ChatGPT,
    Gemini 업로드 호스트는 Gemini). 그 밖은 _gateway_service_name."""
    return _access_target_llm(flow) or _gateway_service_name(flow)


def _gateway_did_work(flow: http.HTTPFlow) -> bool:
    """지연 평균·p95 표본인지(10/11): Gateway가 검사(Detection 호출·건너뜀 포함)·판정(gw_action)·자체 응답(차단·오류)을 한 요청.
    정적 파일·폴링처럼 그냥 통과한 요청(0.0x ms)까지 넣으면 가장 가까운 순위 p95가 약 0이 되어 검사 비용이 보이지 않는다."""
    metadata = flow.metadata
    return bool(metadata.get("gw_action") or metadata.get("gw_inspected") or flow.response is not None)


def _handle_request(
    flow: http.HTTPFlow,
):

    req = flow.request
     # Chrome Extension Gateway health check
    host = (req.pretty_host or "").lower()

    if host == "llm-gateway.local":
        if req.path.startswith("/__llmsgw/health"):
            flow.response = http.Response.make(
                200,
                json.dumps({
                    "status": "ok",
                    "service": "llm-security-gateway",
                    "version": "dev",
                }).encode("utf-8"),
                {
                    "Content-Type": "application/json; charset=utf-8",
                    "Cache-Control": "no-store",
                },
            )
            return
        # 접속 상태(10/12): 내 기록 중계보다 앞에서 직접 답한다(그 모듈의 _gateway_own_route에 넣을 필요 없음)
        if req.path.split("?", 1)[0] == ACCESS_STATUS_PATH:
            _answer_access_status(flow)
            return
        # 내 요청 기록 중계(2026-10-11): 목적지만 대시보드 내부 주소로 바꾸고 전달은 mitmproxy가 비동기로 한다
        if my_requests_relay.handle_request(flow):
            return
        review_prefix = "/__llmsgw/review/"
        if req.method == "GET" and req.path.split("?", 1)[0].startswith(review_prefix):
            review_id = req.path.split("?", 1)[0][len(review_prefix):]
            artifact = load_review_artifact(review_id)
            if artifact is None:
                flow.response = http.Response.make(404, b'{"error":"review_not_found"}',
                                                   {"Content-Type": "application/json", "Cache-Control": "no-store"})
                return
            name, payload = artifact
            extension = name.rsplit(".", 1)[-1]
            mime = {"docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}.get(
                        extension, "application/octet-stream")
            flow.response = http.Response.make(
                200, payload, {"Content-Type": mime,
                "Content-Disposition": f'attachment; filename="review-{review_id}.{extension}"',
                "Cache-Control": "no-store"})
            return

    _start_gateway_metric(flow)
    # 요청 하나는 한 판만: 도중에 대시보드 설정이 바뀌어도 이 요청은 시작 때 설정으로 끝까지 처리한다
    flow.metadata["gw_settings"] = GATEWAY_SETTINGS.current()

    content_type = req.headers.get(
        "content-type",
        "",
    )

    # ========================================================
    # ChatGPT 파일 업로드 경로 확인용 DEBUG
    #
    # 현재 웹 파일 업로드가 /backend-api/files/로 잡히지 않아
    # 실제 파일 바이트가 어느 POST 요청으로 들어오는지 확인한다.
    # 파일 내용 자체는 출력하지 않고 URL / Content-Type / 크기만 기록한다.
    # ========================================================
    if req.host == "chatgpt.com" and req.method == "POST":
        raw_body = req.raw_content or b""
        print("\n===== CHATGPT POST DEBUG =====")
        print("Path:", _url_path(req.path))
        print("Content-Type:", content_type)
        print("Body size:", len(raw_body))

    # ========================================================
    # LLM Web Service 식별
    #
    # ChatGPT 웹서비스는 사용자 프롬프트가 포함된
    # /backend-api/f/conversation/prepare 요청을 별도로 파싱.
    # 그 외 LLM 웹서비스 요청은 현재 식별 후 통과.
    # ========================================================

    llm_service = service_for(req.host)
    identity = {}

    # Gemini 파일 업로드는 gemini.google.com이 아니라
    # push.clients6.google.com의 resumable upload로 들어오므로
    # LLM service 식별보다 먼저 처리한다.
    if req.host == "push.clients6.google.com":
        if handle_gemini_resumable_upload(flow, {}):
            return

    if llm_service:
        identity = read_and_strip_identity(flow)
        flow.metadata["llmsgw_identity"] = identity
        _client_tracker.note_client_id(get_client_ip(flow), identity.get("client_id"), _flow_conn_id(flow))

        # Gemini 웹 요청 구조만 확인한다.
        # 현재 Detection / BLOCK / PASS 로직은 변경하지 않는다.
        if llm_service == "Gemini":
            debug_gemini_request_response(flow, "REQUEST")
            debug_gemini_flow_identity(flow)
            debug_gemini_assistant_path(flow)
            debug_gemini_request_structure(flow)
            debug_gemini_f_req_structure(flow)

            # Gemini 사용자 프롬프트 추출 및 Detection Module 전달
            blocked = detect_gemini_web(flow, identity)
            if blocked:
                return

        # ====================================================
        # Claude Web 사용자 프롬프트 Detection
        #
        # Claude Web의 completion 요청은
        # /api/organizations/{orgId}/chat_conversations/{convId}/completion
        # 형태의 JSON POST로 처리한다.
        # ====================================================
        if llm_service == "Claude":
            # ------------------------------------------------
            # Claude Web 파일 업로드 Detection
            # ------------------------------------------------
            # 실제 Claude Web의 파일 업로드 endpoint는 변경될 수 있으므로
            # 특정 path에 고정하지 않고 multipart/form-data + filename part를
            # 기준으로 파일 업로드를 식별한다.
            if _is_claude_multipart_file_upload(flow):
                blocked = detect_claude_file_upload(flow, identity)
                if blocked:
                    return

                # 파일 업로드 요청은 여기서 처리 종료.
                # 파일 원문을 추출하지 못한 경우에도 upstream으로 정상 전달한다.
                print("→ Claude Web file upload: PASS THROUGH")
                return

            # 새 화면 rpc: 구조 기록(값 없음, 다른 동작 찾기용) → PerformAction 질문 칸(2.3) 검사(10/11)
            _log_claude_rpc_shape(flow)
            if inspect_claude_rpc_prompt(flow, identity):
                print("→ Claude rpc prompt:", "BLOCK" if flow.response is not None else "PASS THROUGH")
                return

            # ------------------------------------------------
            # Claude Web 사용자 프롬프트 Detection
            # ------------------------------------------------
            blocked = detect_claude_web(flow, identity)
            if blocked:
                return

            # Claude completion 요청이면 여기서 처리를 끝낸다.
            # 다른 Claude 부수 요청까지 프롬프트 검사 대상으로 삼지 않는다.
            claude_path = req.path.split("?", 1)[0]
            if re.match(
                r"^/api/organizations/[^/]+/chat_conversations/[^/]+/completion$",
                claude_path,
            ):
                print("→ Claude Web prompt: PASS THROUGH")
                return

        print(
            "\n===== LLM WEB TRAFFIC DETECTED ====="
        )

        print(
            "Service:",
            llm_service,
        )

        print(
            "Host:",
            req.host,
        )

        print(
            "Method:",
            req.method,
        )

        print(
            "Path:",
            _url_path(req.path),
        )

        # =================================================
        # ChatGPT Web 실제 파일 바이트 업로드
        # oaiusercontent.com /files/{file_id}/raw
        # =================================================
        raw_path = req.path.split("?", 1)[0]
        is_chatgpt_raw_upload = (
            req.method == "PUT"
            and req.host.endswith("oaiusercontent.com")
            and raw_path.startswith("/files/")
            and raw_path.endswith("/raw")
        )

        if is_chatgpt_raw_upload:
            file_id = _extract_chatgpt_raw_file_id(req.path)
            filename = _chatgpt_file_name(file_id)   # 예약 응답에서 이은 이름(2026-10-12). 없으면 번호로 기록

            print("\n===== CHATGPT RAW FILE UPLOAD DETECTED =====")
            print("Host:", req.host)
            print("Method:", req.method)
            print("Path:", raw_path)
            print("File ID:", file_id)
            print("Content-Type:", req.headers.get("content-type", ""))
            print("Filename:", f"(resolved, {len(filename)} chars)" if filename else "(metadata not resolved)")
            print("Body size:", len(req.raw_content or b""))

            if file_id:
                original_bytes = req.raw_content or b""

                def apply_raw_upload(outcome):
                    req.raw_content = outcome.sanitized_bytes
                    req.headers["content-length"] = str(len(req.raw_content))
                    with _upload_state_lock:   # 작업 스레드가 넣고 루프(완료 요청)가 훑는다
                        CHATGPT_SANITIZED_FILE_SIZE[file_id] = len(req.raw_content)
                    flow.metadata["llmsgw_review_id"] = outcome.review_id
                    flow.metadata['llmsgw_put_size_diagnostic'] = {
                        'original_bytes': len(original_bytes),
                        'forwarded_bytes': len(req.raw_content),
                        'size_changed': len(original_bytes) != len(req.raw_content),
                        'content_changed': True,
                    }
                    print('[SIGNATURE] raw PUT sizes=', flow.metadata['llmsgw_put_size_diagnostic'])
                    print(f"[SIGNATURE] sanitized raw upload review_id={outcome.review_id} elapsed_ms={outcome.elapsed_ms:.1f}")

                # 2026-10-10: 바이트로 분류(inspect_upload) → 검사할 수 없는 파일은 설정에 따라 차단·허용,
                # 검사 가능하면 그 글자로 검사. 처리 중 예상 못한 오류는 503 + 이벤트(_guard_chatgpt_web_upload)
                _guard_chatgpt_web_upload(
                    flow, identity,
                    file_bytes=original_bytes, filename=filename,
                    mime=req.headers.get("content-type", ""), endpoint=raw_path,
                    file_ref=canonical_file_ref(file_id), apply_sanitized=apply_raw_upload,
                    make_event=lambda file_class, request_id: create_chatgpt_raw_file_event(
                        flow, file_id, filename, file_class=file_class, request_id=request_id),
                )
            else:
                print("→ File ID was not resolved; upload blocked")
                flow.response = http.Response.make(503, b'{"error":{"type":"file_masking_failed"}}',
                                                   {"Content-Type": "application/json"})

            return

        is_openai_api = (
            req.host == OPENAI_HOST
            and req.path in (
                OPENAI_RESPONSES_ENDPOINT,
                OPENAI_FILES_ENDPOINT,
            )
        )

        if not is_openai_api:
            # =================================================
            # ChatGPT Web 파일 업로드
            #
            # 실제 ChatGPT 웹에서는 OpenAI API의 /v1/files가 아니라
            # /backend-api/files/... 경로로 파일 업로드가 발생할 수 있다.
            # multipart/form-data인 경우 기존 파일 파서를 재사용하여
            # TXT / PDF / DOCX / XLSX의 텍스트를 추출하고 Detection에 전달한다.
            # =================================================

            web_path = req.path.split("?", 1)[0]
            if req.method == 'POST' and web_path in {'/backend-api/files', '/backend-api/files/'}:
                print('[SIGNATURE] reservation sizes=', request_size_summary(
                    _json_body(flow) or {}, {}))
            if llm_service == "ChatGPT":
                _note_chatgpt_reservation(flow)   # 파일 이름: 응답 훅이 upload_url 번호와 잇는다(2026-10-12)
            is_chatgpt_file_upload = (
                llm_service == "ChatGPT"
                and web_path.startswith("/backend-api/files/")
            )

            if is_chatgpt_file_upload:
                if req.method == "POST" and content_type.startswith("application/json"):
                    _remember_chatgpt_file_metadata(flow)
                    if "claim_and_finish" in web_path or web_path.endswith("/process_upload_stream"):
                        claim_file_id = _extract_chatgpt_claim_file_id(web_path)
                        completion_body, sizes = _json_body(flow) or {}, _sanitized_sizes()
                        with _upload_state_lock:
                            relationship = _upload_identity_diagnostic.compare(
                                get_client_ip(flow), completion_body, sizes, claim_file_id)
                        print('[SIGNATURE] completion identity relationship=', relationship)
                        print('[SIGNATURE] completion size linkage=', request_size_summary(
                            completion_body, sizes, claim_file_id))
                        rewritten_sizes = _rewrite_completion_file_size(flow, claim_file_id)
                        if rewritten_sizes:
                            print("[SIGNATURE] completion file_size rewritten=", rewritten_sizes)
                        print('[SIGNATURE] completion forwarded sizes=', request_size_summary(
                            _json_body(flow) or {}, {}))

                print("\n===== CHATGPT WEB FILE UPLOAD DETECTED =====")
                print("Host:", req.host)
                print("Method:", req.method)
                print("Path:", _url_path(req.path))
                print("Content-Type:", content_type)

                if (
                    req.method == "POST"
                    and content_type.startswith("multipart/form-data")
                ):
                    file_info = extract_file_upload_info(flow)
                    multipart_name = _chatgpt_multipart_file_name(web_path, file_info["filename"])

                    def apply_multipart_upload(outcome):
                        req.raw_content = replace_multipart_file(req.raw_content or b"", content_type, outcome.sanitized_bytes)
                        req.headers["content-length"] = str(len(req.raw_content))
                        flow.metadata["llmsgw_review_id"] = outcome.review_id
                        print(f"[SIGNATURE] sanitized ChatGPT multipart review_id={outcome.review_id} elapsed_ms={outcome.elapsed_ms:.1f}")

                    def make_multipart_event(file_class, request_id):
                        event = create_file_event(flow, file_class=file_class, request_id=request_id)
                        event["service"] = "ChatGPT"
                        if multipart_name != file_info["filename"] and event.get("contents"):
                            event["contents"][0]["filename"] = multipart_name   # 기록 이름(2026-10-12)
                        return event

                    # 2026-10-10: 원본 PUT과 같은 규칙. name="file" part가 없으면 검사할 수 없는 파일(그 밖 형식)
                    _guard_chatgpt_web_upload(
                        flow, identity,
                        file_bytes=file_info["file_bytes"], filename=multipart_name,
                        mime=None, endpoint=_url_path(req.path), file_ref=None,
                        apply_sanitized=apply_multipart_upload, make_event=make_multipart_event,
                    )
                    return

                # claim_and_finish 같은 웹 파일 API가 multipart가 아닐 경우
                # 해당 요청 자체에는 파일 원문이 없을 수 있으므로,
                # 일단 통과시키되 로그로 실제 Content-Type을 확인한다.
                print("→ ChatGPT file endpoint is not multipart/form-data")
                print("→ File bytes are not available in this request")
                print("→ PASS THROUGH")

                return

            # =================================================
            # 실제 ChatGPT 메시지 전송 요청만 사용자 프롬프트 추출
            # =================================================
            prompt = None

            if (
                llm_service == "ChatGPT"
                and web_path == "/backend-api/f/conversation"
            ):
                prompt = extract_chatgpt_prompt(flow)

            if prompt:
                result = detect_web_prompt(
                    flow,
                    llm_service,
                    prompt,
                )

                attached_refs = chatgpt_attachment_refs(_json_body(flow))
                if attached_refs:
                    # 번호만 남긴다. 같은 Gateway에서 올라간 파일인지(known)는 연결 확인용
                    with _upload_state_lock:
                        known_uploads = sum(ref in CHATGPT_UPLOADED_REFS for ref in attached_refs)
                    print("[ATTACH] conversation attachment refs=", len(attached_refs),
                          "known_uploads=", known_uploads)

                # Detection 결과에 따라 실제 BLOCK / PASS 집행
                if enforce_detection_result(
                    flow,
                    result,
                    dashboard_context={
                        "identity": identity,
                        "targetLlm": "ChatGPT",
                        "contentType": "prompt",
                        "attachedFileIds": attached_refs,
                    },
                    mask_callback=lambda detection_result: _apply_json_mask(
                        flow,
                        normalize_text(prompt),
                        detection_result,
                    ),
                ):
                    return
            else:
                print("→ User prompt not found in this request")
                print("→ LLM Web Traffic: PASS THROUGH")

            return

    # OpenAI API는 아래의 기존 상세 파싱 로직으로 처리
    # (나머지 LLM 웹서비스는 현재 여기서 통과)
    # ========================================================
    # 1. OpenAI File Upload
    # ========================================================

    if (
        req.host == OPENAI_HOST
        and req.method == "POST"
        and req.path == OPENAI_FILES_ENDPOINT
        and content_type.startswith(
            "multipart/form-data"
        )
    ):

        print(
            "\n===== OPENAI FILE UPLOAD DETECTED ====="
        )

        print(
            "Host:",
            req.host,
        )

        print(
            "Method:",
            req.method,
        )

        print(
            "Path:",
            _url_path(req.path),
        )

        print(
            "Content-Type:",
            content_type,
        )

        # ----------------------------------------------------
        # File Event 생성
        # ----------------------------------------------------

        event = create_file_event(
            flow
        )

        print_gateway_event(
            event
        )

        # ----------------------------------------------------
        # 다음 모듈 전달
        # ----------------------------------------------------

        send_to_next_module(
            event
        )

        # ----------------------------------------------------
        # 파일에서 추출한 텍스트를 Detection Module에 직접 전달
        # ----------------------------------------------------

        normalized_text = extract_file_checkable_text(event)

        if normalized_text:
            filename = (
                (event.get("contents") or [{}])[0].get("filename")
            )

            print("\n===== FILE TEXT EXTRACTION SUCCESS =====")
            print("Filename:", filename)
            print("Extracted text length:", len(normalized_text))
            print("→ Sending extracted file text to Detection Module")

            file_info = extract_file_upload_info(flow)
            def detect_file_segment(segment, child_request_id):
                return _inspect(
                    flow, "file", request_id=child_request_id,
                    normalized_text=segment.text,
                    source={"service": "OpenAI", "endpoint": _url_path(req.path),
                            "content_type": "file", "filename": filename,
                            "context_label": canonical_label(segment.context),
                            "signature_original_offsets": True},
                )
            try:
                outcome = sanitize_for_upload(filename, file_info["file_bytes"], event["request_id"], detect_file_segment)
                req.raw_content = replace_multipart_file(req.raw_content or b"", content_type, outcome.sanitized_bytes)
                req.headers["content-length"] = str(len(req.raw_content))
                flow.metadata["llmsgw_review_id"] = outcome.review_id
                print(f"[SIGNATURE] sanitized multipart upload review_id={outcome.review_id} elapsed_ms={outcome.elapsed_ms:.1f}")
            except Exception as exc:
                print("[SIGNATURE] multipart file masking failed error_type=", type(exc).__name__)
                flow.response = http.Response.make(503, b'{"error":{"type":"file_masking_failed"}}',
                                                   {"Content-Type": "application/json"})
                return
        else:
            print("\n===== FILE TEXT EXTRACTION FAILED =====")
            print("→ No checkable text extracted from uploaded file")
            print("→ File upload blocked because safe masking cannot be guaranteed")
            flow.response = http.Response.make(503, b'{"error":{"type":"file_masking_failed"}}',
                                               {"Content-Type": "application/json"})

        return

    # ========================================================
    # 2. OpenAI Responses API
    # ========================================================

    if not (
        req.host == OPENAI_HOST
        and req.method == "POST"
        and req.path == OPENAI_RESPONSES_ENDPOINT
        and content_type.startswith(
            "application/json"
        )
    ):

        return

    print(
        "\n===== OPENAI REQUEST DETECTED ====="
    )

    print(
        "Host:",
        req.host,
    )

    print(
        "Method:",
        req.method,
    )

    print(
        "Path:",
        _url_path(req.path),
    )

    print(
        "Content-Type:",
        content_type,
    )

    # --------------------------------------------------------
    # JSON Body 파싱
    # --------------------------------------------------------

    try:

        body = json.loads(
            req.get_text()
        )

    except (
        json.JSONDecodeError,
        TypeError,
    ):

        print(
            "❌ JSON parsing failed"
        )

        return

    # --------------------------------------------------------
    # Gateway Event 생성
    # --------------------------------------------------------

    event = create_responses_event(
        flow,
        body,
    )

    print_gateway_event(
        event
    )

    # --------------------------------------------------------
    # 다음 모듈 전달
    # --------------------------------------------------------

    send_to_next_module(
        event
    )

    # --------------------------------------------------------
    # 항목별로 따로 Detection에 보낸다 (2026-09-23)
    #
    # 프롬프트와 파일을 이어붙여 한 번에 보내면 창 해시가 어느 쪽에도 없는 창을 만들고,
    # 시그니처는 경계를 넘는 문맥 검사 때문에 없던 적중을 만들거나 있던 적중을 잃는다
    # (event_adapter 모듈 docstring에 실측 근거). 항목마다 한 번씩 부른다.
    #
    # 출처별 처리(파일은 차단, 프롬프트는 마스킹 등)는 아직 정하지 않았다. 지금은 옛 동작을
    # 그대로 유지한다 - 한 항목이라도 차단이면 요청 전체를 차단한다. 검사 순서도 항목 순서
    # 그대로다. 정책이 정해지면 이 루프 안에서 분기한다.
    # --------------------------------------------------------

    for part in checkable_parts(event):

        # 항목 종류(prompt·file·image, Gateway가 만든 값)의 규칙. 모르는 종류는 파일 규칙(_inspect)
        result = _inspect(
            flow,
            part["kind"],
            request_id=part_request_id(event["request_id"], part["index"]),
            normalized_text=part["text"],
            source={
                "service": event.get("service"),
                "content_type": part["kind"],
                "filename": part["filename"],
            },
        )

        print()
        print(
            f"===== DETECTION RESULT "
            f"[{part['kind']} #{part['index']}"
            f"{' ' + part['filename'] if part['filename'] else ''}] ====="
        )
        print_detection_summary(result)

        # Detection 결과에 따라 실제 BLOCK / PASS 집행
        if enforce_detection_result(
            flow,
            result,
            dashboard_context={
                "identity": identity,
                "targetLlm": "ChatGPT",
                "contentType": part["kind"],
            },
            mask_callback=lambda detection_result, part=part: _apply_json_mask(
                flow,
                part["text"],
                detection_result,
                body,
            ),
        ):
            return

def debug_gemini_streamgenerate_response(flow: http.HTTPFlow):
    """
    Gemini StreamGenerate 응답의 상태/헤더/크기만 안전하게 기록한다.
    응답 body 내용, 쿠키, 인증값, query string은 출력하지 않는다.
    """
    try:
        req = flow.request
        if req.host != "gemini.google.com" or req.method.upper() != "POST":
            return

        path = req.path.split("?", 1)[0]
        if "StreamGenerate" not in path:
            return

        resp = flow.response
        print("\n===== GEMINI STREAMGENERATE RESPONSE TRACE =====")
        print("Flow ID:", getattr(flow, "id", None))
        print("Status:", resp.status_code if resp is not None else "NONE")

        if resp is None:
            print("Response: NONE")
            print("===== END GEMINI STREAMGENERATE RESPONSE TRACE =====\n")
            return

        ctype = resp.headers.get("content-type", "")
        content_length = resp.headers.get("content-length", "")
        transfer_encoding = resp.headers.get("transfer-encoding", "")

        print("Content-Type:", ctype.split(";", 1)[0].strip())
        print("Content-Length header:", content_length or "<none>")
        print("Transfer-Encoding:", transfer_encoding or "<none>")
        print("Response body size:", len(resp.raw_content or b""))
        print("===== END GEMINI STREAMGENERATE RESPONSE TRACE =====\n")
    except Exception as exc:
        print("\n===== GEMINI STREAMGENERATE RESPONSE TRACE ERROR =====")
        print("Error type:", type(exc).__name__)
        print("===== END GEMINI STREAMGENERATE RESPONSE TRACE ERROR =====\n")


def response(flow: http.HTTPFlow):
    # ChatGPT 파일 이름(2026-10-12): 예약 응답이 브라우저에 가기 전에 upload_url 번호와 이름을 잇는다(예외는 안에서 삼킴)
    _link_chatgpt_reservation_name(flow)
    my_requests_relay.handle_response(flow)
    if 'llmsgw_put_size_diagnostic' in flow.metadata and flow.response is not None:
        print('[SIGNATURE] raw PUT response status=', flow.response.status_code)
    review_id = flow.metadata.get("llmsgw_review_id")
    if (review_id and flow.response is not None
            and os.getenv("SIGNATURE_REVIEW_HEADERS", "1") == "1"
            and os.getenv("SIGNATURE_PRIVATE_REVIEW_DOWNLOAD", "0") == "1"):
        flow.response.headers["X-LLMSGW-Review-ID"] = review_id
        flow.response.headers["X-LLMSGW-Review-URL"] = f"http://llm-gateway.local/__llmsgw/review/{review_id}"
    request_path = flow.request.path.split("?", 1)[0]
    if (flow.request.host == 'chatgpt.com' and flow.request.method == 'POST'
            and request_path.startswith('/backend-api/files') and flow.response is not None
            and 'claim_and_finish' not in request_path and not request_path.endswith('/process_upload_stream')):
        creation_text = flow.response.get_text(strict=False) or ''
        with _upload_state_lock:
            relationship = _upload_identity_diagnostic.remember(get_client_ip(flow), creation_text)
        print('[SIGNATURE] creation identity relationship=', relationship)
    if ("claim_and_finish" in request_path or request_path == "/backend-api/files/process_upload_stream") and flow.response is not None:
        print("[SIGNATURE] completion response status=", flow.response.status_code,
              "summary=", completion_summary(flow.response.get_text(strict=False) or ""))
        print('[SIGNATURE] completion response availability=', {
            'body_available': flow.response.raw_content is not None,
            'body_bytes': len(flow.response.raw_content or b''),
            'streaming': bool(flow.response.stream),
        })
    debug_gemini_request_response(flow, "RESPONSE")
    debug_gemini_streamgenerate_response(flow)
    _finalize_gateway_metric(
        flow,
        action=flow.metadata.get("gw_action"),
        detector=flow.metadata.get("gw_detector"),
        degraded=flow.metadata.get("gw_degraded", False),
        error=False,
        timeout=False,
    )
    _send_deferred_dashboard_event(flow)


def error(flow: http.HTTPFlow):
    """Gemini flow의 실제 mitmproxy-level connection error만 안전하게 기록한다."""
    my_requests_relay.handle_error(flow)
    _finalize_gateway_metric(
        flow,
        action=flow.metadata.get("gw_action"),
        detector=flow.metadata.get("gw_detector"),
        degraded=flow.metadata.get("gw_degraded", False),
        error=True,
        timeout=_gateway_error_is_timeout(flow),
        client_cancelled=_gateway_error_is_client_cancel(flow),
    )
    _send_deferred_dashboard_event(flow)
    try:
        req = flow.request
        if req.host != "gemini.google.com" or req.method.upper() != "POST":
            return

        path = req.path.split("?", 1)[0]
        if "StreamGenerate" not in path:
            return

        err = getattr(flow, "error", None)
        print("\n===== GEMINI STREAMGENERATE ERROR =====")
        print("Flow ID:", getattr(flow, "id", None))
        if err is None:
            print("Error: <none>")
        else:
            print("Error type:", type(err).__name__)
            print("Error message:", str(err)[:500])
        print("===== END GEMINI STREAMGENERATE ERROR =====\n")
    except Exception as exc:
        print("\n===== GEMINI STREAMGENERATE ERROR TRACE FAILED =====")
        print("Error type:", type(exc).__name__)
        print("===== END GEMINI STREAMGENERATE ERROR TRACE FAILED =====\n")
