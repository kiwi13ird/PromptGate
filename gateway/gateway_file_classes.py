"""업로드 파일 분류·글자 추출 (2026-10-10, ChatGPT 웹·Claude 경로 전용).

판단은 바이트만으로 한다. 파일 이름·MIME은 클라이언트가 정하는 값이라 로그용으로만 받는다.
분류 결과(kind)로 글자를 뽑는다(확장자로 다시 고르지 않는다).

group
  None          검사 가능(text를 Detection에 보냄). 0바이트는 보낼 내용이 없어 통과(text "")
  'image'       PNG·JPEG·GIF·WEBP 구조를 끝까지 확인한 것, 글자 없이 그림만 든 DOCX·XLSX
  'scanned_pdf' 글자가 하나도 없고 글자 없는 쪽에 그림·도형이 있는 PDF
  'other'       그 밖(PPTX·RTF·압축·암호·손상·글자 없는 문서·예상 못한 오류, 닫힌 목록 밖·암호·배포용 HWP·HWPX 등)
  'too_large'   바이트·압축 해제 크기·글자 수·추출 시간 상한 초과

inspect_upload는 어떤 예외도 밖으로 내지 않는다: 예외 → group 'other', detail 'error:<예외 이름>'.
Gemini·OpenAI API 경로(gateway.extract_file_text, pypdf)는 이 모듈을 쓰지 않는다.

HWP 5.0·HWPX(2026-10-11, masking-design/SPEC.md §2.3): Detection 등록과 같은 파일 signature_detector/document/hwpdoc.py
(md5 동일)로 글자를 뽑는다. hwpdoc·olefile은 HWP를 만났을 때만 불러온다 - olefile이 없으면 HWP는 'other'
(error:ImportError, 기본 차단)이고 다른 형식·애드온 기동에는 영향이 없다. 'FileHeader' 스트림 이름이 없는 CFB는 olefile을
부르지 않는다. HWP·HWPX 밖에 다른 리더가 읽는 바이트가 있거나(<kind>_embedded_container) HWPX와 DOCX·XLSX 부품이 한
패키지에 섞이면(ambiguous_package) 'other'.
"""

from __future__ import annotations

import html
import io
import json
import re
import struct
import time
import zipfile
from dataclasses import dataclass

_WS_RE = re.compile(r"\s+")
# 글자 파일로 보지 않는 제어 문자. 허용: 탭·줄바꿈·VT·FF·CR, 그리고 터미널 제어 BEL·BS·ESC(10/11: 색을 넣은
# 빌드·CI 로그가 막히지 않게. Detection 정규화가 이 문자들을 지운다). NUL(UTF-16 글자·바이너리)과 나머지 C0은 거부
_BAD_CONTROL_RE = re.compile(r"[\x00-\x06\x0e-\x1a\x1c-\x1f]")
_CFB = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_PNG = b"\x89PNG\r\n\x1a\n"
_ODT_MIMETYPE = "application/vnd.oasis.opendocument.text"
_HWPX_MIMETYPE = "application/hwp+zip"


@dataclass(frozen=True)
class FileClass:
    kind: str                 # pdf docx xlsx odt text png jpeg gif webp cfb hwp hwpx rtf zip pptx unknown empty
    group: str | None         # None image scanned_pdf other too_large
    detail: str               # 로그용 짧은 이름(파일 이름·내용 없음)
    text: str | None = None   # 검사 가능하면 Detection에 보낼 문자열
    image_only_pages: int = 0
    pages: int = 0
    text_pages: int = 0


class _Limit(Exception):
    def __init__(self, kind, detail):
        super().__init__(detail)
        self.kind = kind
        self.detail = detail


def _nonspace(text: str) -> int:
    return len(_WS_RE.sub("", text))


class _Budget:
    """추출 도중의 시간·글자 한도(10/11). XML 요소·행을 하나 지날 때마다 시계를 보고(tick),
    모은 글자는 공백을 뺀 수로 센다(add). 넘으면 _Limit(kind, 'time'|'chars')."""

    __slots__ = ("kind", "deadline", "clock", "max_chars", "chars")

    def __init__(self, kind, *, deadline, clock, max_chars):
        self.kind, self.deadline, self.clock, self.max_chars, self.chars = kind, deadline, clock, max_chars, 0

    def tick(self):
        if self.clock() > self.deadline:
            raise _Limit(self.kind, "time")

    def add(self, text: str):
        if text:
            self.chars += _nonspace(text)
            if self.chars > self.max_chars:
                raise _Limit(self.kind, "chars")


def inspect_upload(data, *, filename=None, mime=None, max_bytes: int, max_chars: int, deadline: float,
                   clock=time.monotonic) -> FileClass:
    """업로드 바이트 → FileClass. filename·mime은 판단에 쓰지 않는다(로그용)."""
    del filename, mime
    try:
        return _classify(bytes(data or b""), max_bytes=max_bytes, max_chars=max_chars,
                         deadline=deadline, clock=clock)
    except _Limit as limit:
        return FileClass(limit.kind, "too_large", limit.detail)
    except Exception as exc:   # MemoryError·RecursionError 포함
        return FileClass("unknown", "other", f"error:{type(exc).__name__}")


def _classify(data: bytes, *, max_bytes, max_chars, deadline, clock) -> FileClass:
    if not data:
        return FileClass("empty", None, "empty", text="")
    image = _image_class(data)            # 바이트 상한보다 먼저(구조 확인은 수 ms)
    if image is not None:
        return image
    if len(data) > max_bytes:
        return FileClass("unknown", "too_large", "bytes")
    if clock() > deadline:
        return FileClass("unknown", "too_large", "time")
    if b"%PDF-" in data[:1024]:
        return extract_pdf(data, max_chars=max_chars, deadline=deadline, clock=clock)
    if data.startswith(b"PK\x03\x04"):
        result = _zip_class(data, max_bytes=max_bytes, max_chars=max_chars, deadline=deadline, clock=clock)
    elif data.startswith(_CFB):
        # HWP 5.0이면 글자를 뽑는다. DOC·XLS·PPT·암호 걸린 OOXML·깨진 CFB는 지금처럼 cfb(other)
        result = _cfb_class(data, max_bytes=max_bytes, max_chars=max_chars, deadline=deadline, clock=clock)
    elif data.startswith(b"{\\rtf"):
        return FileClass("rtf", "other", "rtf")       # 한글이 \'b0\'a1·\uN?로 들어가 원문 일치가 안 나온다
    else:
        result = _text_class(data, max_chars=max_chars)
    if result.group is None and clock() > deadline:
        return FileClass(result.kind, "too_large", "time")
    return result


# ------------------------------------------------------------------ 이미지 구조

def _png_end(d: bytes):
    pos = 8
    while pos + 12 <= len(d):
        (length,) = struct.unpack(">I", d[pos:pos + 4])
        chunk_type = d[pos + 4:pos + 8]
        pos += 12 + length
        if pos > len(d):
            return None
        if chunk_type == b"IEND":
            return pos
    return None


def _jpeg_end(d: bytes, start: int = 0):
    """d[start:]에서 시작하는 JPEG(FF D8)의 끝 표지(EOI) 다음 위치(d 기준). 구조가 맞지 않으면 None."""
    pos, n = start + 2, len(d)
    while pos + 2 <= n:
        if d[pos] != 0xFF:
            return None
        marker = d[pos + 1]
        if marker == 0xFF:                 # 채움 바이트
            pos += 1
            continue
        if marker == 0xD9:                 # EOI
            return pos + 2
        if 0xD0 <= marker <= 0xD7 or marker == 0x01:
            pos += 2
            continue
        if pos + 4 > n:
            return None
        (length,) = struct.unpack(">H", d[pos + 2:pos + 4])
        if length < 2:
            return None
        pos += 2 + length
        if marker == 0xDA:                 # SOS: 압축 데이터는 FF 위치만 건너뛴다(FF00·RSTn은 데이터)
            while True:
                pos = d.find(b"\xff", pos)
                if pos < 0 or pos + 1 >= n:
                    return None
                following = d[pos + 1]
                if following != 0x00 and not 0xD0 <= following <= 0xD7 and following != 0xFF:
                    break
                pos += 1 if following == 0xFF else 2
    return None


def _gif_end(d: bytes):
    if len(d) < 13:
        return None
    flags = d[10]
    pos = 13
    if flags & 0x80:
        pos += 3 * (2 ** ((flags & 7) + 1))

    def skip_blocks(p):
        while p < len(d):
            size = d[p]
            p += 1
            if size == 0:
                return p
            p += size
        return None

    while pos is not None and pos < len(d):
        block = d[pos]
        if block == 0x3B:
            return pos + 1
        if block == 0x21:
            pos = skip_blocks(pos + 2)
        elif block == 0x2C:
            if pos + 10 > len(d):
                return None
            local_flags = d[pos + 9]
            pos += 10
            if local_flags & 0x80:
                pos += 3 * (2 ** ((local_flags & 7) + 1))
            pos = skip_blocks(pos + 1)
        else:
            return None
    return None


def _webp_end(d: bytes):
    if len(d) < 12:
        return None
    (size,) = struct.unpack("<I", d[4:8])
    return 8 + size if 8 + size <= len(d) else None


def _image_class(d: bytes) -> FileClass | None:
    """PNG·JPEG·GIF·WEBP만. 끝까지 맞고 뒤가 비었거나 0x00뿐 → image, 뒤에 데이터 → other(<형식>_trailer),
    머리만 맞고 구조가 깨짐 → other(broken_image). 이미지 머리가 아니면 None.

    여러 장 JPEG(10/11, MPF: Ultra HDR·이득 지도·깊이 지도 사진): 끝 표지 바로 뒤에 다음 JPEG(FF D8 FF)가 오면
    그것도 끝까지 따라간다. 모든 바이트가 온전한 JPEG들(+ 끝의 0x00)일 때만 image. 모션 사진(뒤에 MP4)·
    JPEG + PDF처럼 JPEG가 아닌 덧붙이기는 그대로 jpeg_trailer(명세 §14-15 기본값)."""
    if d.startswith(_PNG):
        kind, end = "png", _png_end(d)
    elif d.startswith(b"\xff\xd8"):
        kind, end = "jpeg", _jpeg_end(d)
        while end is not None and d.startswith(b"\xff\xd8\xff", end):
            following = _jpeg_end(d, end)
            if following is None:          # 다음 JPEG가 깨짐 → 그 앞까지만 그림, 나머지는 덧붙이기
                break
            end = following
    elif d[:6] in (b"GIF87a", b"GIF89a"):
        kind, end = "gif", _gif_end(d)
    elif d[:4] == b"RIFF" and d[8:12] == b"WEBP":
        kind, end = "webp", _webp_end(d)
    else:
        return None
    if end is None:
        return FileClass(kind, "other", "broken_image")
    if d[end:].strip(b"\x00"):
        return FileClass(kind, "other", f"{kind}_trailer")
    return FileClass(kind, "image", kind)


# ------------------------------------------------------------------ 글자 파일

def _decode_text(data: bytes):
    for encoding in ("utf-8-sig", "cp949"):
        try:
            text = data.decode(encoding)
        except UnicodeDecodeError:
            continue
        return None if _BAD_CONTROL_RE.search(text) else text
    return None


class _JsonPairs(list):
    """JSON 객체의 (키, 값) 쌍 전부(json.loads object_pairs_hook). dict로 받으면 같은 키가 여러 번 나올 때
    마지막 값만 남는데 원본 바이트(앞 값 포함)는 그대로 전달되므로, 쌍을 모두 남겨 모두 검사한다(10/11)."""


def _json_lines(value, out: list) -> None:
    """키·문자열 값·숫자를 문서 순서대로(\\uXXXX는 json이 이미 풀었다). 같은 키가 여러 번 나오면 모두."""
    if isinstance(value, _JsonPairs):
        for key, child in value:
            out.append(str(key))
            _json_lines(child, out)
    elif isinstance(value, list):
        for child in value:
            _json_lines(child, out)
    elif isinstance(value, str):
        out.append(value)
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        out.append(str(value))


def _text_class(data: bytes, *, max_chars) -> FileClass:
    text = _decode_text(data)
    if text is None:
        return FileClass("unknown", "other", "unknown")
    head = text.lstrip()[:1]
    detail = "text"
    if head in ("{", "["):
        try:
            value = json.loads(text, object_pairs_hook=_JsonPairs)
        except ValueError:
            value = None
        else:
            lines: list = []
            _json_lines(value, lines)
            text, detail = "\n".join(lines), "json"
    elif head == "<":                          # HTML·XML·SVG: &#xAC00; 같은 문자 참조를 푼다
        text, detail = html.unescape(text), "markup"
    if not text.strip():
        return FileClass("text", "other", "empty")
    if len(text) > max_chars and _nonspace(text) > max_chars:
        return FileClass("text", "too_large", "chars")
    return FileClass("text", None, detail, text=text)


# ------------------------------------------------------------------ ZIP 계열

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"


def _tag(namespace, name):
    return f"{{{namespace}}}{name}"


_W_P, _W_T, _W_TBL, _W_TR, _W_TC = (_tag(_W, name) for name in ("p", "t", "tbl", "tr", "tc"))
_W_TXBX = _tag(_W, "txbxContent")
_W_SDT, _W_SDT_CONTENT = _tag(_W, "sdt"), _tag(_W, "sdtContent")
_W_CUSTOM_XML, _W_SMART_TAG = _tag(_W, "customXml"), _tag(_W, "smartTag")
_MC_ALT, _MC_CHOICE, _MC_FALLBACK = (_tag(_MC, name) for name in ("AlternateContent", "Choice", "Fallback"))


def _xml(data: bytes):
    from lxml import etree
    parser = etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=False, recover=False)
    return etree.fromstring(data, parser)


def _flatten(element, budget):
    """블록 칸을 감싸는 껍데기(내용 제어·사용자 XML·대체 콘텐츠)를 벗긴다. mc:Fallback(VML 중복)은 버린다."""
    for child in element:
        budget.tick()
        if child.tag in (_W_SDT,):
            for content in child.iterchildren(_W_SDT_CONTENT):
                yield from _flatten(content, budget)
        elif child.tag in (_W_CUSTOM_XML, _W_SMART_TAG):
            yield from _flatten(child, budget)
        elif child.tag == _MC_ALT:
            for choice in child.iterchildren(_MC_CHOICE):
                yield from _flatten(choice, budget)
                break
        else:
            yield child


def _paragraph(paragraph, budget):
    """이 문단에 직접 속한 w:t 글자와 글상자(w:txbxContent)를 문서 순서대로, 문단 아래를 한 번만 훑어 모은다.

    안쪽 문단(w:p, 글상자 안 문단 등)은 그 문단 몫이라 내려가지 않고, mc:Fallback(VML 대체본, 같은 글자)은 버린다.
    글상자 안이라도 문단 밖에 있는 글자는 이 문단 몫이다. 10/11: 이전에는 w:t·글상자마다 조상을 거슬러 올라가
    '깊이 × w:t 수'만큼 걸렸다(수 KB DOCX로 7~18초). 지금은 요소 수에 비례하고 요소마다 시간 한도를 본다.
    반환 (글자, 글상자 목록)."""
    texts, boxes = [], []
    stack = list(paragraph)
    stack.reverse()
    while stack:
        node = stack.pop()
        budget.tick()
        tag = node.tag
        if tag == _W_P or tag == _MC_FALLBACK:
            continue
        if tag == _W_T:
            texts.append(node.text or "")
        elif tag == _W_TXBX:
            boxes.append(node)
        children = list(node)
        if children:
            children.reverse()
            stack.extend(children)
    return "".join(texts), boxes


def _block_lines(container, out: list, budget) -> None:
    """문단·표를 문서 순서대로. 표는 행마다 칸 글자를 ' | '로(Detection 등록 preprocessing/pipeline.py와 같음)."""
    for child in _flatten(container, budget):
        if child.tag == _W_P:
            text, boxes = _paragraph(child, budget)
            if text.strip():
                budget.add(text)
                out.append(text)
            for box in boxes:
                _block_lines(box, out, budget)
        elif child.tag == _W_TBL:
            for row in _flatten(child, budget):
                if row.tag != _W_TR:
                    continue
                cells = []
                nested: list = []
                for cell in _flatten(row, budget):
                    if cell.tag != _W_TC:
                        continue
                    lines: list = []
                    for item in _flatten(cell, budget):
                        if item.tag == _W_P:
                            text, boxes = _paragraph(item, budget)
                            if text.strip():
                                budget.add(text)
                                lines.append(text)
                            for box in boxes:
                                _block_lines(box, nested, budget)
                        elif item.tag == _W_TBL:            # 표 안 표: 이 행 뒤에 따로 줄로
                            _block_lines([item], nested, budget)
                    cells.append("\n".join(lines))
                if any(cell.strip() for cell in cells):
                    out.append(" | ".join(cells))
                out.extend(nested)


def _docx_text(zf: zipfile.ZipFile, names: set, budget) -> str:
    """머리글·바닥글(구역별 파트) → 본문(문서 순서, 글상자·표 안 표·내용 제어 포함) → 각주·미주·메모.
    budget: 추출 도중 시간·글자 한도(넘으면 _Limit('docx', 'time'|'chars'))."""
    out: list = []
    for name in sorted(n for n in names if re.fullmatch(r"word/(?:header|footer)\d*\.xml", n)):
        budget.tick()
        _block_lines(_xml(zf.read(name)), out, budget)
    budget.tick()
    document = _xml(zf.read("word/document.xml"))
    body = document.find(_tag(_W, "body"))
    if body is not None:
        _block_lines(body, out, budget)
    for name in ("word/footnotes.xml", "word/endnotes.xml", "word/comments.xml"):
        if name in names:
            budget.tick()
            for note in _xml(zf.read(name)):     # w:footnote·w:endnote·w:comment
                _block_lines(note, out, budget)
    return "\n".join(out)


def _xlsx_text(data: bytes, *, budget) -> tuple[str, bool]:
    """지금 방식([Sheet: 이름] + 셀 값 ' | '). 반환 (글자, 셀 값이 있는지). 시트 이름 줄은 비었는지 판단에 쓰지 않는다.
    budget: 행마다 시간 한도(10/11), 셀 글자 수 한도(넘으면 _Limit('xlsx', 'time'|'chars'))."""
    from openpyxl import load_workbook
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        lines, has_cells = [], False
        for sheet in workbook.worksheets:
            lines.append(f"[Sheet: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                budget.tick()
                values = [str(value) for value in row if value is not None]
                if values:
                    line = " | ".join(values)
                    if line.strip():
                        has_cells = True
                    lines.append(line)
                    budget.add(line)
        return "\n".join(lines), has_cells
    finally:
        workbook.close()


def _odt_text(zf: zipfile.ZipFile) -> str:
    """지금 방식(content.xml 태그 제거)."""
    with zf.open("content.xml") as handle:
        xml = handle.read().decode("utf-8", errors="ignore")
    return html.unescape(" ".join(re.sub(r"<[^>]+>", " ", xml).split()))


def _zip_class(data: bytes, *, max_bytes, max_chars, deadline, clock) -> FileClass:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        infos = zf.infolist()
        # 압축을 푼 선언 크기 합. 실제 해제는 zipfile이 선언 크기까지만 읽고 CRC로 확인한다(어긋나면 예외 → other)
        if sum(info.file_size for info in infos) > max_bytes:
            return FileClass("zip", "too_large", "zip_bytes")
        names = {info.filename for info in infos}
        mimetype = None
        if "mimetype" in names:
            with zf.open("mimetype") as handle:
                mimetype = handle.read(128).decode("ascii", errors="replace").strip()
        # HWPX 표지(hwpdoc.sniff와 같은 두 표지 중 하나)를 DOCX·XLSX 부품보다 먼저 본다(10/11 검토 지적 3): DOCX 부품만 검사하면
        # 같은 패키지의 HWPX 섹션 글자가 검사 없이 나간다. 두 형식이 섞인 패키지는 어느 리더가 무엇을 읽을지 몰라 other
        if mimetype == _HWPX_MIMETYPE or "Contents/content.hpf" in names:
            if "word/document.xml" in names or "xl/workbook.xml" in names:
                return FileClass("zip", "other", "ambiguous_package")
            if b"%PDF-" in data or not _zip_layout_exact(data, infos, deadline=deadline, clock=clock):
                return FileClass("hwpx", "other", "hwpx_embedded_container")
            return _hwp_class("hwpx", data, max_bytes=max_bytes, max_chars=max_chars, deadline=deadline, clock=clock)
        if "word/document.xml" in names:
            text = _docx_text(zf, names, _Budget("docx", deadline=deadline, clock=clock, max_chars=max_chars))
            if not text.strip():
                media = any(name.startswith("word/media/") for name in names)
                return FileClass("docx", "image" if media else "other", "docx_media_only" if media else "empty")
            if len(text) > max_chars and _nonspace(text) > max_chars:
                return FileClass("docx", "too_large", "chars")
            return FileClass("docx", None, "docx", text=text)
        if "xl/workbook.xml" in names:
            text, has_cells = _xlsx_text(data, budget=_Budget("xlsx", deadline=deadline, clock=clock,
                                                              max_chars=max_chars))
            if not has_cells:
                media = any(name.startswith("xl/media/") for name in names)
                return FileClass("xlsx", "image" if media else "other", "xlsx_media_only" if media else "empty")
            return FileClass("xlsx", None, "xlsx", text=text)
        if mimetype == _ODT_MIMETYPE and "content.xml" in names:
            text = _odt_text(zf)
            if not text.strip():
                return FileClass("odt", "other", "empty")
            if len(text) > max_chars and _nonspace(text) > max_chars:
                return FileClass("odt", "too_large", "chars")
            return FileClass("odt", None, "odt", text=text)
        if any(name.startswith("ppt/") for name in names):
            return FileClass("pptx", "other", "pptx")
        return FileClass("zip", "other", "zip")


# ------------------------------------------------------------------ HWP 5.0·HWPX (hwpdoc, 2026-10-11)

# 미리보기 대조에 남기는 글자 = Detection 창 해시 압축(windowhash/hashing.py compact)과 같은 글자(한글·영문·숫자, 소문자)
_COMPACT_KEEP_RE = re.compile(r"[가-힣A-Za-z0-9]")
# 한컴 미리보기 글자(PrvText)의 줄 경계와 표 칸 경계(칸 글자는 <…>로 들어간다)
_PREVIEW_SPLIT_RE = re.compile(r"[\r\n<>]+")
# CFB 디렉터리의 'FileHeader' 스트림 이름(UTF-16LE, 소문자). olefile은 이름을 대소문자 없이 찾으므로 bytes.lower()(ASCII만
# 바꿈 - re.I 바이트 패턴과 같은 규칙)로 찾는다. 20MiB에서 re.I 검색 약 135ms, lower()+찾기 약 10ms(10/11 로컬)
_FILEHEADER_NAME = "fileheader".encode("utf-16-le")
_ZIP_END = b"PK\x05\x06"
_HWP_IMPORT_WARNED: list = []


def _hwpdoc():
    """공통 추출 모듈(HWP·HWPX를 만났을 때만 불러옴). olefile은 hwpdoc이 CFB를 읽을 때 다시 함수 안에서 불러온다."""
    from signature_detector.document import hwpdoc
    return hwpdoc


def _hwp_unavailable(kind: str, exc: ImportError) -> FileClass:
    """hwpdoc·olefile·lxml을 불러오지 못함 → other(기본 차단). 운영자가 알아볼 수 있게 프로세스마다 한 번 남긴다."""
    if not _HWP_IMPORT_WARNED:
        _HWP_IMPORT_WARNED.append(True)
        print(f"[FILE] HWP/HWPX inspection unavailable (missing module: {getattr(exc, 'name', None) or 'unknown'}); "
              "treated as other")
    return FileClass(kind, "other", "error:ImportError")


def _cfb_class(data: bytes, *, max_bytes, max_chars, deadline, clock) -> FileClass:
    """CFB 머리: HWP 5.0이면 글자를 뽑고, 아니면(DOC·XLS·PPT·암호 걸린 OOXML·깨진 CFB) 기준판처럼 cfb.

      ① 'FileHeader' 스트림 이름(UTF-16, 대소문자 무시 - olefile 이름 찾기와 같음)이 바이트 어디에도 없으면 HWP가 아니다
         (hwpdoc.sniff·read 모두 그 스트림이 있어야 HWP) → olefile을 부르지 않고 cfb. 10/11 검토 지적 2: 디렉터리 항목
         15만 8천 개 19.4MiB 조작 파일을 olefile로 두 번 열어 2.2~2.9초 동안 동기 훅이 멈췄다(기준판 0ms). olefile이 없는
         운영 venv에서도 DOC·XLS는 기준판과 같은 cfb(ImportError 경고 없음)
      ② 이름이 있으면 kind는 hwp. olefile로 한 번만 열고(read), 깨짐(HwpCorrupt)이면 다시 열지 않고 hwp_corrupt(other)
      ③ 다른 리더가 읽는 바이트 → hwp_embedded_container(other). 10/11 검토 지적 1: 깨끗한 HWP 뒤에 붙이거나 BinData 그림
         스트림 안에 넣은 개인정보 DOCX·HWPX·PDF를 ZIP·PDF 리더가 읽는데 hwpdoc은 HWP만 읽어 통과시켰다(기준판은 cfb other)
         - ZIP 끝 레코드(PK 05 06)·PDF 머리(%PDF-)가 바이트 어디에든 있음(덧붙임·BinData 안 모두)
         - 파일이 마지막 할당 섹터(FAT)에서 끝나지 않음: 머리 없는 PDF를 덧붙여도 pypdf·PyMuPDF가 읽는다(10/11 실측)
         실제 HWP·HWT 807개(팀 19·한컴 서식 포함): 두 표지 0건, 모두 섹터 격자에 맞고 마지막 할당 섹터에서 끝남"""
    if _FILEHEADER_NAME not in data.lower():
        return FileClass("cfb", "other", "cfb")
    if _ZIP_END in data or b"%PDF-" in data:
        return FileClass("hwp", "other", "hwp_embedded_container")
    end = _cfb_end(data)
    if end is None:
        return FileClass("hwp", "other", "hwp_corrupt")
    if end != len(data):
        return FileClass("hwp", "other", "hwp_embedded_container")
    return _hwp_class("hwp", data, max_bytes=max_bytes, max_chars=max_chars, deadline=deadline, clock=clock,
                      unavailable_kind="cfb")


def _cfb_end(data: bytes) -> int | None:
    """CFB 마지막 할당 섹터(FAT 칸이 FREESECT가 아닌 마지막 번호)의 끝 위치. 섹터 n은 파일 (n + 1) × 섹터 크기에서 시작.
    FAT 섹터 목록 = 머리 DIFAT 109칸 + DIFAT 사슬(MS-CFB). 머리 칸·사슬이 파일 밖을 가리키면 None(깨짐)."""
    if len(data) < 512:
        return None
    shift = struct.unpack_from("<H", data, 0x1E)[0]
    if shift not in (9, 12):
        return None
    size = 1 << shift
    sectors = len(data) // size - 1
    n_fat = struct.unpack_from("<I", data, 0x2C)[0]
    difat_next, n_difat = struct.unpack_from("<II", data, 0x44)
    if sectors < 1 or n_fat > sectors or n_difat > sectors:
        return None
    fat_sectors = list(struct.unpack_from("<109I", data, 0x4C))
    for _ in range(n_difat):
        if difat_next >= sectors:
            return None
        chain = struct.unpack_from(f"<{size // 4}I", data, size * (difat_next + 1))
        fat_sectors += chain[:-1]
        difat_next = chain[-1]
    if any(sector >= sectors for sector in fat_sectors[:n_fat]):
        return None
    fat = b"".join(data[size * (sector + 1):size * (sector + 2)] for sector in fat_sectors[:n_fat])
    used = (len(fat.rstrip(b"\xff")) + 3) // 4          # 마지막 할당 섹터 번호 + 1 (뒤쪽 FREESECT 칸을 뗌)
    return size * (used + 1)


def _zip_layout_exact(data: bytes, infos, *, deadline, clock) -> bool:
    """ZIP 바이트 전체 = 지역 항목들(머리 + 압축 데이터)이 0부터 빈틈없이 → 중앙 디렉터리 → 끝 레코드(주석 없음, 파일 끝).

    10/11 검토 지적 1: zipfile·hwpdoc은 끝 레코드와 중앙 디렉터리로 항목을 찾는다. 그 밖 바이트(뒤에 붙인 PDF·ZIP, 앞에 붙인
    ZIP, 항목 사이 빈틈, ZIP 주석)와 지역 머리만 다른 항목(이름·크기)은 앞에서부터 읽는 리더(libarchive·Java ZipInputStream
    등)나 PDF 리더가 읽는데 우리는 보지 않는다. 중앙 디렉터리 항목의 지역 머리가 이름·방식·플래그·CRC·크기까지 같아야 하고,
    데이터 기술자(플래그 bit 3)·ZIP64·여러 디스크는 받지 않는다. 실제 HWPX·HWTX 24개(한컴 서식 20 포함) 모두 이 규칙을 지킴."""
    end = len(data) - 22
    if end < 0 or data.rfind(_ZIP_END) != end or data[-2:] != b"\x00\x00":
        return False
    disk, cd_disk, count_disk, count, cd_size, cd_offset = struct.unpack_from("<HHHHII", data, end + 4)
    if disk or cd_disk or count_disk != count or count != len(infos) or cd_offset + cd_size != end:
        return False
    pos = 0
    for index, info in enumerate(sorted(infos, key=lambda item: item.header_offset)):
        if index % 1024 == 0 and clock() > deadline:          # 항목 수만 개짜리 조작 ZIP도 시간 한도 안에서
            raise _Limit("hwpx", "time")
        if info.header_offset != pos or pos + 30 > cd_offset or data[pos:pos + 4] != b"PK\x03\x04":
            return False
        flags, method, _, _, crc, compressed, size, name_len, extra_len = struct.unpack_from("<HHHHIIIHH", data, pos + 6)
        name = info.orig_filename.encode("utf-8" if info.flag_bits & 0x800 else "cp437")
        if (flags & 0x08 or (flags, method, crc, compressed, size) != (info.flag_bits, info.compress_type, info.CRC,
                                                                       info.compress_size, info.file_size)
                or data[pos + 30:pos + 30 + name_len] != name):
            return False
        pos += 30 + name_len + extra_len + compressed
    return pos == cd_offset


def _compact(text: str) -> str:
    return "".join(_COMPACT_KEEP_RE.findall(text)).lower()


def _preview_mismatch(preview: str, body: str, tick) -> bool:
    """미리보기 글자(PrvText)에 본문에 없는 글자가 있는지. 미리보기는 검사 글자에 넣지 않으므로(명세 §2.1) 본문과 맞을 때만
    통과시킨다. 둘 다 Detection 압축 글자로, 본문은 레코드 순서(HwpDoc.paragraph_text - 한컴 미리보기 순서와 같음).
      ① 미리보기 전체가 본문의 부분열(명세 §2.3)
      ② 미리보기의 줄·표 칸 조각 하나하나가 본문 안에 이어진 글자로 있음. ①만으로는 본문에 없는 줄(옛 판의 개인정보 등)도
         그 숫자·글자가 본문 곳곳에 흩어져 있기만 하면 통과한다
    실측(10/11): 팀 HWP 19개·한컴 견본·서식 275개(미리보기 있는 것, 조각 7,551개) 모두 통과, 본문만 바꾼 합성 HWP 3개 모두 불일치.
    """
    compact_preview = _compact(preview)
    if not compact_preview:
        return False
    compact_body = _compact(body)
    if len(compact_preview) > len(compact_body):
        return True
    pos = 0
    for ch in compact_preview:
        pos = compact_body.find(ch, pos)
        if pos < 0:
            return True
        pos += 1
    for piece in set(_compact(piece) for piece in _PREVIEW_SPLIT_RE.split(preview)):
        tick()                                 # 조작한 미리보기(조각 수만 개)가 시간 한도를 넘지 않게
        if piece and piece not in compact_body:
            return True
    return False


_PREVIEW_LINE_RE = re.compile(r"[\r\n]+")
_PREVIEW_CELL_RE = re.compile(r"[<>]+")


def preview_extra_text(preview: str, body: str, tick=lambda: None) -> str:
    """미리보기가 본문과 다르면(_preview_mismatch) 검사할 미리보기 글자, 같으면 ''(2026-10-14).

    줄마다 한 줄, 표 칸(<…>)은 ' | '로 잇는다(lines()의 표 행과 같은 모양). 이전에는 불일치 파일을 검사 불가(other,
    <kind>_preview_mismatch)로 막았다. 미리보기를 검사 글자에 넣으면 막을 이유가 없다: 가려야 하면 HWP·HWPX 가림이
    미리보기 글자·그림을 모두 지우고(hwp_mask·hwpx_mask), 통과면 검사한 글자 그대로 나간다.
    본문만 바꾸고 미리보기를 다시 만들지 않은 파일(바이트 편집·한컴 밖 도구로 만든 HWP)이 여기에 든다."""
    if not _preview_mismatch(preview, body, tick):
        return ""
    lines = []
    for line in _PREVIEW_LINE_RE.split(preview):
        cells = [cell.strip() for cell in _PREVIEW_CELL_RE.split(line)]
        cells = [cell for cell in cells if cell and cell != "\x00"]
        if cells:
            lines.append(" | ".join(cells))
    return "\n".join(lines)


def _hwp_class(kind: str, data: bytes, *, max_bytes, max_chars, deadline, clock, unavailable_kind=None) -> FileClass:
    """HWP·HWPX 바이트 → FileClass(명세 §2.3). 검사 글자 = '\\n'.join(HwpDoc.lines()) - 본문 + 부가 글자(수식·필드 명령·
    요약 정보·개체 설명문·HWPX 메타·바탕쪽). Detection 등록 글자(본문만)는 이 글자 압축열의 앞부분이다.
    부르는 쪽(_cfb_class·_zip_class)이 먼저 다른 리더가 읽는 숨은 바이트(<kind>_embedded_container)를 거른다.

      닫힌 목록 밖 개체·암호·배포용·DRM·문서 이력 → other(사유 코드 그대로: hwp_unsupported_object, hwp_distribution …)
      깨짐 → other(<kind>_corrupt) / 압축 해제·구조 수 상한 → too_large(<kind>_decompressed)
      시간·글자 한도(tick·add) → too_large(time·chars)
      hwpdoc·olefile·lxml 없음 → other(error:ImportError, kind는 unavailable_kind가 있으면 그것)
      본문 글자 없음 → other(<kind>_media_only: 그림 BinData 있음, empty) - 그림만 든 문서를 image(기본 허용)로 보내지 않음
      미리보기 불일치 → 검사 가능(<kind>_preview_extra): 검사 글자 끝에 미리보기 글자를 붙인다(2026-10-14,
                       preview_extra_text. 이전: other(<kind>_preview_mismatch)로 기본 차단)
    """
    try:
        hwpdoc = _hwpdoc()
    except ImportError as exc:
        return _hwp_unavailable(unavailable_kind or kind, exc)
    budget = _Budget(kind, deadline=deadline, clock=clock, max_chars=max_chars)
    try:
        doc = hwpdoc.read(data, max_decompressed=max_bytes, tick=budget.tick, add=budget.add)
    except ImportError as exc:
        return _hwp_unavailable(unavailable_kind or kind, exc)
    except hwpdoc.HwpUnsupported as exc:
        return FileClass(kind, "other", exc.kind)          # 짧은 사유 코드만(exc.detail은 남기지 않음)
    except hwpdoc.HwpTooLarge:
        return FileClass(kind, "too_large", f"{kind}_decompressed")
    except hwpdoc.HwpCorrupt:
        return FileClass(kind, "other", f"{kind}_corrupt")
    if not any(p.text.strip() for p in doc.paragraphs):
        return FileClass(kind, "other", f"{kind}_media_only" if doc.media else "empty")
    budget.tick()
    extra = preview_extra_text(doc.preview_text(), doc.paragraph_text(), budget.tick)
    text = "\n".join(doc.lines())
    if extra:
        text += "\n" + extra          # 본문 글자 뒤(Detection 등록 글자 = 본문이 앞부분인 규칙 유지)
    if len(text) > max_chars and _nonspace(text) > max_chars:
        return FileClass(kind, "too_large", "chars")
    return FileClass(kind, None, f"{kind}_preview_extra" if extra else kind, text=text)


# ------------------------------------------------------------------ PDF (pypdfium2)

def _annotation_strings(page, raw):
    """쪽의 주석 문자열(Contents, 양식 칸 값 V)과 첨부 파일 주석 여부."""
    import ctypes

    texts, has_file = [], False
    count = raw.FPDFPage_GetAnnotCount(page.raw)
    for index in range(count):
        annot = raw.FPDFPage_GetAnnot(page.raw, index)
        if not annot:
            continue
        try:
            if raw.FPDFAnnot_GetSubtype(annot) == raw.FPDF_ANNOT_FILEATTACHMENT:
                has_file = True
            for key in (b"Contents", b"V"):
                size = raw.FPDFAnnot_GetStringValue(annot, key, None, 0)
                if size <= 2:
                    continue
                buffer = ctypes.create_string_buffer(size)
                raw.FPDFAnnot_GetStringValue(annot, key, ctypes.cast(buffer, ctypes.POINTER(raw.FPDF_WCHAR)), size)
                value = buffer.raw[:size - 2].decode("utf-16-le", errors="ignore")
                if value.strip():
                    texts.append(value)
        finally:
            raw.FPDFPage_CloseAnnot(annot)
    return "\n".join(texts), has_file


def extract_pdf(data: bytes, *, max_chars: int, deadline: float, clock=time.monotonic) -> FileClass:
    """쪽 글자(get_text_range: 쪽 밖 글자까지) + 주석 문자열. 한글 CMap(UniKS-UTF16-H 등)도 읽는다.

    첨부 파일·XFA → other. 전체 글자 0자이고 글자 없는 쪽에 그림·도형이 있으면 scanned_pdf, 없으면 other(empty).
    시간 한도(파일 대기 시간)는 쪽마다, 글자 상한은 공백을 뺀 글자 수로 확인한다.
    """
    import pypdfium2 as pdfium
    import pypdfium2.raw as raw

    visual = [raw.FPDF_PAGEOBJ_IMAGE, raw.FPDF_PAGEOBJ_PATH]
    try:
        pdf = pdfium.PdfDocument(data)          # 암호 필요·손상 → 예외
    except Exception as exc:
        return FileClass("pdf", "other", f"pdf_unreadable:{type(exc).__name__}")
    try:
        if pdf.count_attachments():
            return FileClass("pdf", "other", "pdf_attachment")
        if pdf.get_formtype() in (raw.FORMTYPE_XFA_FULL, raw.FORMTYPE_XFA_FOREGROUND):
            return FileClass("pdf", "other", "pdf_xfa")
        texts, nonspace, zero_visual, text_pages = [], 0, 0, 0
        pages = len(pdf)
        for index in range(pages):
            if clock() > deadline:
                return FileClass("pdf", "too_large", "pdf_time", pages=pages)
            page = pdf[index]
            try:
                textpage = page.get_textpage()
                try:
                    text = textpage.get_text_range()
                finally:
                    textpage.close()
                annot_text, has_file = _annotation_strings(page, raw)
                if has_file:
                    return FileClass("pdf", "other", "pdf_attachment", pages=pages)
                page_text = text + ("\n" + annot_text if annot_text else "")
                count = _nonspace(page_text)
                if count == 0:
                    # 글자 0자 쪽에서만 그림·도형을 본다(글자 쪽까지 보면 Word PDF에서 +50%)
                    if any(True for _ in page.get_objects(filter=visual, max_depth=15)):
                        zero_visual += 1
                else:
                    text_pages += 1
            finally:
                page.close()
            texts.append(page_text.replace("\r\n", "\n"))
            nonspace += count
            if nonspace > max_chars:
                return FileClass("pdf", "too_large", "pdf_chars", pages=pages)
    finally:
        pdf.close()
    if nonspace == 0:
        if zero_visual:
            return FileClass("pdf", "scanned_pdf", "pdf_scanned", image_only_pages=zero_visual, pages=pages)
        return FileClass("pdf", "other", "empty", pages=pages)
    return FileClass("pdf", None, "pdf", text="\n".join(t for t in texts if t.strip()),
                     image_only_pages=zero_visual, pages=pages, text_pages=text_pages)
