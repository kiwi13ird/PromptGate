import tempfile
import unittest
from pathlib import Path

from docx import Document
from presidio_analyzer import RecognizerResult

from signature_detector.document.docx_sanitizer import sanitize_docx


class StubAnalyzer:
    values = {
        "employee@example.com": "EMAIL_ADDRESS",
        "010-1234-5678": "KR_PHONE_NUMBER",
        "1999.03.14": "KR_DATE_OF_BIRTH",
    }

    def analyze(self, text, language, entities):
        results = []
        for value, entity_type in self.values.items():
            start = text.find(value)
            if start >= 0:
                results.append(RecognizerResult(entity_type, start, start + len(value), 0.9))
        return results


class DocxSanitizerTest(unittest.TestCase):
    def test_masks_body_table_header_and_footer(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            input_path = Path(temp_dir) / "input.docx"
            output_path = Path(temp_dir) / "nested" / "output.docx"
            doc = Document()
            doc.add_paragraph("이메일: employee@example.com")
            table = doc.add_table(rows=1, cols=2)
            table.rows[0].cells[0].text = "생년월일"
            table.rows[0].cells[1].text = "1999.03.14"
            header_table = doc.add_table(rows=2, cols=2)
            header_table.rows[0].cells[0].text = "성명"
            header_table.rows[0].cells[1].text = "생년월일"
            header_table.rows[1].cells[0].text = "테스트"
            header_table.rows[1].cells[1].text = "1999.03.14"
            doc.sections[0].header.paragraphs[0].text = "010-1234-5678"
            doc.sections[0].footer.paragraphs[0].text = "employee@example.com"
            doc.save(input_path)

            sanitize_docx(input_path, output_path, StubAnalyzer(), list(StubAnalyzer.values.values()))
            sanitized = Document(output_path)
            self.assertIn("[EMAIL_ADDRESS]", sanitized.paragraphs[0].text)
            self.assertEqual("[DATE_OF_BIRTH]", sanitized.tables[0].cell(0, 1).text)
            self.assertEqual("[DATE_OF_BIRTH]", sanitized.tables[1].cell(1, 1).text)
            self.assertIn("[PHONE_NUMBER]", sanitized.sections[0].header.paragraphs[0].text)
            self.assertIn("[EMAIL_ADDRESS]", sanitized.sections[0].footer.paragraphs[0].text)


if __name__ == "__main__":
    unittest.main()
