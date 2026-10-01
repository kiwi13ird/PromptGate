from docx import Document


def extract_docx(path):
    doc = Document(path)
    texts = []

    # 일반 문단
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()

        if text:
            texts.append(text)

    # 표
    for table in doc.tables:
        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]

            if cells:
                texts.append(" | ".join(cells))

    return texts