"""DOCX/XLSX를 원본 위치에 되돌릴 수 있는 segment로 다룬다."""

from __future__ import annotations

from copy import deepcopy
from dataclasses import dataclass
from difflib import SequenceMatcher
from io import BytesIO
from time import perf_counter

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.text.run import Run
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from signature_detector.contracts import CheckableSegment, SegmentFinding, SegmentResult, validate_segment_result
from signature_detector.policy import PolicyAction


@dataclass(frozen=True)
class SanitizedArtifacts:
    sanitized_bytes: bytes
    review_bytes: bytes
    segment_count: int
    elapsed_ms: float


def _docx_paragraphs(document):
    """지원 영역의 paragraph를 안정적인 순번과 함께 반환한다."""

    seen = set()

    def emit_paragraphs(paragraphs, prefix):
        for index, paragraph in enumerate(paragraphs):
            key = id(paragraph._p)
            if key not in seen:
                seen.add(key)
                yield f"{prefix}:p{index}", paragraph

    def emit_tables(tables, prefix):
        for table_index, table in enumerate(tables):
            for row_index, row in enumerate(table.rows):
                for cell_index, cell in enumerate(row.cells):
                    cell_prefix = f"{prefix}:t{table_index}:r{row_index}:c{cell_index}"
                    yield from emit_paragraphs(cell.paragraphs, cell_prefix)
                    yield from emit_tables(cell.tables, cell_prefix)

    yield from emit_paragraphs(document.paragraphs, "body")
    yield from emit_tables(document.tables, "body")
    for section_index, section in enumerate(document.sections):
        yield from emit_paragraphs(section.header.paragraphs, f"header{section_index}")
        yield from emit_tables(section.header.tables, f"header{section_index}")
        yield from emit_paragraphs(section.footer.paragraphs, f"footer{section_index}")
        yield from emit_tables(section.footer.tables, f"footer{section_index}")


def extract_docx_segments(file_bytes: bytes) -> list[CheckableSegment]:
    document = Document(BytesIO(file_bytes))
    return [
        CheckableSegment(f"f:0:docx:{locator}", "file", paragraph.text)
        for locator, paragraph in _docx_paragraphs(document)
        if paragraph.text.strip()
    ]


def _style_for_offset(runs, offset):
    cursor = 0
    for run in runs:
        end = cursor + len(run.text)
        if offset < end or (offset == end and end == len("".join(r.text for r in runs))):
            return deepcopy(run._r.rPr) if run._r.rPr is not None else None
        cursor = end
    return None


def _rewrite_docx_paragraph(paragraph, result: SegmentResult, review: bool) -> None:
    original = paragraph.text
    if result.masked_text is None:
        raise ValueError("DOCX mask result requires masked_text")
    pieces = []
    for tag, original_start, original_end, masked_start, masked_end in SequenceMatcher(
        None, original, result.masked_text, autojunk=False
    ).get_opcodes():
        if tag == "equal":
            pieces.append((original[original_start:original_end], original_start, False))
        else:
            # Detection의 offset은 정규화 문자열 기준일 수 있으므로 원문에 적용하지 않는다.
            # 완성된 masked_text와 원문 diff만 사용한다.
            pieces.append((result.masked_text[masked_start:masked_end], original_start, True))

    old_runs = list(paragraph.runs)
    for run in old_runs:
        paragraph._p.remove(run._r)
    for text, style_offset, is_mask in pieces:
        if not text:
            continue
        run = paragraph.add_run(text)
        rpr = _style_for_offset(old_runs, style_offset)
        if rpr is not None:
            run._r.insert(0, rpr)
        if review and is_mask:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _docx_result_map(segments, results):
    segment_by_id = {segment.segment_id: segment for segment in segments}
    result_by_id = {result.segment_id: result for result in results}
    if set(result_by_id) != set(segment_by_id) or len(result_by_id) != len(results):
        raise ValueError("DOCX Detection results are missing, duplicated, or unexpected")
    for segment_id, result in result_by_id.items():
        # 파일 재구성에는 정규화 offset 대신 완성된 masked_text를 사용한다.
        if result.segment_id != segment_id:
            raise ValueError("Detection returned a different DOCX segment_id")
        if result.action in {PolicyAction.MASK_PARTIAL, PolicyAction.MASK_FULL}:
            if result.masked_text is None or result.masked_text == segment_by_id[segment_id].text:
                raise ValueError("DOCX mask action requires changed masked_text")
        if result.action is PolicyAction.BLOCK_NOTIFY:
            raise PermissionError("DOCX blocked by Signature policy")
    return result_by_id


def sanitize_docx_bytes(file_bytes: bytes, results: list[SegmentResult]) -> SanitizedArtifacts:
    started = perf_counter()
    segments = extract_docx_segments(file_bytes)
    result_by_id = _docx_result_map(segments, results)
    outputs = []
    for review in (False, True):
        document = Document(BytesIO(file_bytes))
        for locator, paragraph in _docx_paragraphs(document):
            segment_id = f"f:0:docx:{locator}"
            result = result_by_id.get(segment_id)
            if result and result.action in {PolicyAction.MASK_PARTIAL, PolicyAction.MASK_FULL}:
                _rewrite_docx_paragraph(paragraph, result, review)
        stream = BytesIO()
        document.save(stream)
        outputs.append(stream.getvalue())
    return SanitizedArtifacts(outputs[0], outputs[1], len(segments), (perf_counter() - started) * 1000)


def extract_xlsx_segments(file_bytes: bytes) -> list[CheckableSegment]:
    workbook = load_workbook(BytesIO(file_bytes), data_only=False)
    segments = []
    for sheet_index, worksheet in enumerate(workbook.worksheets):
        for row in worksheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.strip():
                    segments.append(
                        CheckableSegment(
                            f"f:0:xlsx:s{sheet_index}:r{cell.row - 1}:c{cell.column - 1}",
                            "file",
                            cell.value,
                        )
                    )
    return segments


def sanitize_xlsx_bytes(file_bytes: bytes, results: list[SegmentResult]) -> SanitizedArtifacts:
    started = perf_counter()
    segments = extract_xlsx_segments(file_bytes)
    segment_by_id = {segment.segment_id: segment for segment in segments}
    result_by_id = {result.segment_id: result for result in results}
    if set(result_by_id) != set(segment_by_id) or len(result_by_id) != len(results):
        raise ValueError("XLSX Detection results are missing, duplicated, or unexpected")
    for segment_id, result in result_by_id.items():
        # cell 전체 문자열을 교체하므로 정규화 offset을 원본 cell에 적용하지 않는다.
        if result.segment_id != segment_id:
            raise ValueError("Detection returned a different XLSX segment_id")
        if result.action in {PolicyAction.MASK_PARTIAL, PolicyAction.MASK_FULL}:
            if result.masked_text is None or result.masked_text == segment_by_id[segment_id].text:
                raise ValueError("XLSX mask action requires changed masked_text")
        if result.action is PolicyAction.BLOCK_NOTIFY:
            raise PermissionError("XLSX blocked by Signature policy")

    outputs = []
    review_fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
    for review in (False, True):
        workbook = load_workbook(BytesIO(file_bytes), data_only=False)
        for segment_id, result in result_by_id.items():
            if result.action not in {PolicyAction.MASK_PARTIAL, PolicyAction.MASK_FULL}:
                continue
            _, _, _, sheet_token, row_token, column_token = segment_id.split(":")
            worksheet = workbook.worksheets[int(sheet_token[1:])]
            cell = worksheet.cell(row=int(row_token[1:]) + 1, column=int(column_token[1:]) + 1)
            cell.value = result.masked_text
            if review:
                cell.fill = review_fill
        stream = BytesIO()
        workbook.save(stream)
        outputs.append(stream.getvalue())
    return SanitizedArtifacts(outputs[0], outputs[1], len(segments), (perf_counter() - started) * 1000)


def result_from_masking(segment: CheckableSegment, masking_result) -> SegmentResult:
    return SegmentResult(
        segment.segment_id,
        masking_result.action,
        masking_result.masked_text,
        tuple(
            SegmentFinding(
                finding.entity_type,
                finding.start,
                finding.end,
                finding.replacement,
                finding.action,
                finding.score,
            )
            for finding in masking_result.findings
        ),
    )
