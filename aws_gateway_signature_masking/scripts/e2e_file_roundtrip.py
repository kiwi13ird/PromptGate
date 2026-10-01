#!/usr/bin/env python3
"""Exercise actual Gateway file hooks against Detection and a local upstream."""

from __future__ import annotations

import json
import os
import sys
import tempfile
import threading
import types
import urllib.request
from email.parser import BytesParser
from email.policy import default
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from io import BytesIO
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook

from e2e_prompt_roundtrip import (
    FakeFlow,
    FakeResponse,
    Headers,
    NoRedis,
    install_mitmproxy_stub,
    install_requests_stub,
    read_json,
    send_json,
    configure_v2_policy_if_present,
)


ROOT = Path(__file__).resolve().parents[1]
CANDIDATE = Path(os.getenv("E2E_GATEWAY_DIR", ROOT / "worktrees" / "candidate_8081_signature_masking"))
DETECTION = Path(os.getenv("E2E_DETECTION_DIR", ROOT.parent / "detection"))


class UploadRequest:
    def __init__(self, host, method, path, content_type, body):
        self.host = host
        self.pretty_host = host
        self.method = method
        self.path = path
        self.headers = Headers({"content-type": content_type, "content-length": str(len(body))})
        self.raw_content = body

    def get_text(self):
        return self.raw_content.decode("utf-8")


def docx_bytes(text):
    document = Document()
    document.add_paragraph(text)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def xlsx_bytes(text):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet["A1"] = text
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def extracted_text(filename, payload):
    if filename.endswith(".docx"):
        return "\n".join(p.text for p in Document(BytesIO(payload)).paragraphs)
    workbook = load_workbook(BytesIO(payload), data_only=False)
    return "\n".join(
        str(cell.value) for worksheet in workbook.worksheets for row in worksheet.iter_rows()
        for cell in row if cell.value is not None
    )


def multipart(filename, payload, boundary="signature-e2e"):
    body = (
        f"--{boundary}\r\nContent-Disposition: form-data; name=\"purpose\"\r\n\r\nassistants\r\n"
        f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"{filename}\"\r\n"
        "Content-Type: application/octet-stream\r\n\r\n"
    ).encode() + payload + f"\r\n--{boundary}--\r\n".encode()
    return f"multipart/form-data; boundary={boundary}", body


def multipart_file(content_type, body):
    message = BytesParser(policy=default).parsebytes(
        b"Content-Type: " + content_type.encode() + b"\r\n\r\n" + body
    )
    for part in message.iter_parts():
        if part.get_param("name", header="Content-Disposition") == "file":
            return part.get_filename(), part.get_payload(decode=True)
    raise AssertionError("upstream did not receive a file part")


def main() -> int:
    detection_port, upstream_port = 18001, 18081
    os.environ["DETECTION_BASE_URL"] = f"http://127.0.0.1:{detection_port}"
    os.environ["SIGNATURE_POLICY_JSON"] = json.dumps({"default_action": "mask_full"})
    review_dir = tempfile.TemporaryDirectory(prefix="signature-review-")
    os.environ["SIGNATURE_REVIEW_DIR"] = review_dir.name
    install_mitmproxy_stub()
    install_requests_stub()
    sys.path.insert(0, str(DETECTION))
    sys.path.insert(0, str(CANDIDATE))
    from detection_service import pipeline
    import gateway
    policy_tmp = tempfile.TemporaryDirectory(prefix="signature-policy-")
    configure_v2_policy_if_present(pipeline, Path(policy_tmp.name))

    gateway.send_to_next_module = lambda event: None
    received = []

    class DetectionHandler(BaseHTTPRequestHandler):
        def do_POST(self):
            request = read_json(self)
            response = pipeline.run(
                NoRedis(), request["normalized_text"], request_id=request["request_id"],
                source=request.get("source"), decision_log=None,
            )
            send_json(self, 200, response.model_dump())

        def log_message(self, *_):
            pass

    class UpstreamHandler(BaseHTTPRequestHandler):
        def do_POST(self):
            size = int(self.headers.get("content-length", "0"))
            body = self.rfile.read(size)
            received.append((self.headers.get("content-type"), body))
            send_json(self, 200, {"id": "mock-file"})

        def do_PUT(self):
            size = int(self.headers.get("content-length", "0"))
            received.append((self.headers.get("content-type"), self.rfile.read(size)))
            send_json(self, 200, {"id": "mock-raw"})

        def log_message(self, *_):
            pass

    servers = [
        ThreadingHTTPServer(("127.0.0.1", detection_port), DetectionHandler),
        ThreadingHTTPServer(("127.0.0.1", upstream_port), UpstreamHandler),
    ]
    for server in servers:
        threading.Thread(target=server.serve_forever, daemon=True).start()

    secret = "010-1234-5678"
    cases = [("sample.docx", docx_bytes(f"담당자 연락처 {secret}")),
             ("sample.xlsx", xlsx_bytes(f"담당자 연락처 {secret}"))]
    results = []
    try:
        for filename, original in cases:
            content_type, body = multipart(filename, original)
            flow = FakeFlow(UploadRequest("api.openai.com", "POST", "/v1/files", content_type, body))
            gateway.request(flow)
            if flow.response is not None:
                raise AssertionError(f"Gateway blocked supported file: {flow.response.raw_content!r}")
            if flow.request.headers.get("content-length") != str(len(flow.request.raw_content)):
                raise AssertionError("Gateway did not refresh multipart Content-Length")
            request = urllib.request.Request(
                f"http://127.0.0.1:{upstream_port}/v1/files", data=flow.request.raw_content,
                headers={"Content-Type": content_type}, method="POST",
            )
            with urllib.request.urlopen(request, timeout=5) as response:
                response.read()
            upstream_name, upstream_file = multipart_file(*received[-1])
            upstream_text = extracted_text(upstream_name, upstream_file)
            review_id = flow.metadata["llmsgw_review_id"]
            review_matches = list(Path(review_dir.name).glob(f"{review_id}.*"))
            review_text = extracted_text(filename, review_matches[0].read_bytes())
            passed = secret not in upstream_text and "[KR_PHONE_NUMBER]" in upstream_text and secret not in review_text
            results.append((f"OpenAI multipart {filename}", passed, upstream_text, review_matches[0]))

        raw_name, raw_original = cases[0]
        file_id = "file-signature-e2e"
        gateway.CHATGPT_FILE_META[file_id] = raw_name
        raw_flow = FakeFlow(UploadRequest(
            "uploads.oaiusercontent.com", "PUT", f"/files/{file_id}/raw",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document", raw_original,
        ))
        gateway.request(raw_flow)
        if raw_flow.response is not None:
            raise AssertionError(f"Gateway blocked supported raw file: {raw_flow.response.raw_content!r}")
        if raw_flow.request.headers.get("content-length") != str(len(raw_flow.request.raw_content)):
            raise AssertionError("Gateway did not refresh raw PUT Content-Length")
        request = urllib.request.Request(
            f"http://127.0.0.1:{upstream_port}/files/{file_id}/raw", data=raw_flow.request.raw_content,
            headers={"Content-Type": raw_flow.request.headers.get("content-type")}, method="PUT",
        )
        with urllib.request.urlopen(request, timeout=5) as response:
            response.read()
        raw_text = extracted_text(raw_name, received[-1][1])
        results.append(("ChatGPT raw PUT DOCX", secret not in raw_text and "[KR_PHONE_NUMBER]" in raw_text,
                        raw_text, Path(review_dir.name)))

        unsupported_type, unsupported_body = multipart("unknown.bin", b"sensitive opaque bytes")
        blocked_flow = FakeFlow(UploadRequest("api.openai.com", "POST", "/v1/files", unsupported_type, unsupported_body))
        gateway.request(blocked_flow)
        blocked = blocked_flow.response is not None and blocked_flow.response.status_code == 503
        results.append(("unsupported file fail-closed", blocked, "503" if blocked else "not blocked", None))
    finally:
        for server in reversed(servers):
            server.shutdown()
            server.server_close()
        policy_tmp.cleanup()

    print("\n=== E2E FILE ROUND TRIP ===")
    for name, passed, text, review_path in results:
        print(f"{name}: {'PASS' if passed else 'FAIL'}")
        print("  upstream/review text:", text)
        if review_path:
            print("  review artifact:", review_path)
    review_dir.cleanup()
    return 0 if all(item[1] for item in results) else 1


if __name__ == "__main__":
    raise SystemExit(main())
